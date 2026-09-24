"""验证确认稿的材料识别、意见写入和发布入口。"""
from pathlib import Path
import pytest

from openpyxl import Workbook, load_workbook
from pypdf import PdfWriter

from risk_audit.checks.registry import CheckContext
from risk_audit.models import Entity, FileRecord, Finding, ParsedSheet
from risk_audit.util import sha256_file


def make_file(tmp_path, name, sheet_types):
    """构造材料记录；tmp_path 为测试目录，name 为文件名，sheet_types 为表类别。"""
    path = tmp_path / name
    book = Workbook()
    book.active.title = '业务表'
    book.active.append(['控制措施编号', '审核意见'])
    book.active.append(['M1', '人工保留意见'])
    book.save(path)
    sheets = [ParsedSheet('业务表', st, [1], {'measure_id': 1}, {'审核意见': 2}, 2, 2, []) for st in sheet_types]
    return FileRecord(path, Path(name), sha256_file(path), 'xlsx', '205H', [], False, '06', 'default', 'matrix', sheets)


def test_explanation_pdf_and_application_names(tmp_path):
    """PDF 说明及申请应被识别；tmp_path 为测试目录。"""
    from risk_audit.inventory_v180 import scan_package_v180
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.add_metadata({'/Title': '信通公司主体说明'})
    writer.write(tmp_path / '信通公司附件.pdf')
    entities = {'205H': Entity('205H', '信通公司')}
    records = scan_package_v180(tmp_path, entities, {'信通公司': '205H'})
    assert len(records) == 1
    assert records[0].material_type == 'explanation'
    assert records[0].true_format == 'pdf'
    assert records[0].entity_code == '205H'


@pytest.mark.parametrize('suffix,header_row', [('.xlsx', 1), ('.xls', 1), ('.xlsx', 60)])
def test_excel_candidate_without_material_keyword_is_not_lost(tmp_path, suffix, header_row):
    """参数为目录、允许扩展名及表头行；业务标题或晚表头不能被扫描静默漏掉。"""
    from risk_audit.inventory_v180 import scan_package_v180
    from risk_audit.readers.confirmed_v180 import parse_workbook_v180
    import json
    book = Workbook()
    book.active.title = '风控矩阵'
    for column, value in enumerate(['控制措施编号', '控制措施', '责任主体'], 1):
        book.active.cell(header_row, column, value)
    for column, value in enumerate(['M1', '核对报表', '信通公司-财务部'], 1):
        book.active.cell(header_row + 1, column, value)
    path = tmp_path / ('06信通公司报送' + suffix)
    book.save(path)
    files = scan_package_v180(tmp_path, {'205H': Entity('205H', '信通公司')}, {'信通公司': '205H'})
    assert len(files) == 1
    aliases = json.loads((Path(__file__).parents[1] / 'rulepacks/releases/1.8.0/field_aliases.json').read_text())
    assert parse_workbook_v180(files[0], path, aliases)[0].sheet_type == 'matrix'


def test_combined_matrix_and_lists_count_both_categories(tmp_path):
    """合并工作簿应同时满足矩阵与三清单类；tmp_path 为测试目录。"""
    from risk_audit.checks.materials_v180 import required_documents_v3
    file = make_file(tmp_path, '信通公司06矩阵及三清单.xlsx', ['matrix', 'position_duty', 'incompatible_position', 'system_rule'])
    ctx = CheckContext([], [file], [], {'205H': Entity('205H', '信通公司')}, {}, {}, ['205H'], [{'business_code': '06', 'variant_id': 'default', 'required': True}])
    assert not required_documents_v3(ctx, {'material_types': ['matrix', 'three_lists']})


def test_same_hash_duplicate_not_multiple_versions(tmp_path):
    """重复内容不算多个版本；tmp_path 为测试目录。"""
    from risk_audit.checks.materials_v180 import required_documents_v3
    first = make_file(tmp_path, '信通公司06矩阵.xlsx', ['matrix'])
    other = make_file(tmp_path, '信通公司06矩阵副本.xlsx', ['matrix'])
    other.sha256 = first.sha256
    ctx = CheckContext([], [first, other], [], {}, {}, {}, ['205H'], [{'business_code': '06', 'required': True}])
    assert not required_documents_v3(ctx, {'material_types': ['matrix']})


def test_uploaded_unmatched_three_lists_is_not_reported_as_missing(tmp_path):
    """已上传但未识别业务表的三清单不能误报缺件；tmp_path 为测试目录。"""
    from risk_audit.checks.materials_v180 import required_documents_v3
    path = tmp_path / '信通公司06三清单.xlsx'
    file = FileRecord(path, Path(path.name), 'hash', 'xlsx', '205H', [], False,
                      '06', 'default', 'three_lists',
                      parse_errors=['未识别到正式业务工作表'])
    context = CheckContext([], [file], [], {'205H': Entity('205H', '信通公司')}, {}, {},
                           ['205H'], [{'business_code': '06', 'variant_id': 'default', 'required': True}])

    issues = required_documents_v3(context, {'material_types': ['three_lists']})

    assert len(issues) == 1
    assert issues[0]['evidence']['unavailable_reason'] == '未识别到正式业务工作表'
    assert 'missing_material' not in issues[0]['evidence']


def test_missing_column_is_written_as_review_under_v180():
    """明确缺列需写待核实意见，不应因旧分流策略隐藏；无参数。"""
    from risk_audit.issue_routing import route_issue
    from risk_audit.models import Record
    row = Record('position_duty', '205H', '06', 'default', 'a.xlsx', '岗位', 2, {}, 'r')
    issue = {'record': row, 'kind': 'review', 'evidence': {'issue_type': 'field_unavailable', 'field': 'person_names'}}
    result = route_issue(issue, {'check_id': 'person_required'}, {'version': 2, 'technical_issues': 'internal', 'unfinished_is_pass': False})
    assert result['kind'] == 'review'
    assert result['evidence']['publication_channel'] == 'unit'


@pytest.mark.parametrize(
    ('issue_type', 'field', 'check_id'),
    [
        ('separation_field_unavailable', 'role', 'handler_reviewer_overlap_0916'),
        ('responsibility_unavailable', 'duty', 'broad_responsibility_pattern'),
    ],
)
def test_position_duty_missing_fields_stay_on_source_row(issue_type, field, check_id):
    """岗位清单缺字段意见必须写回原行；参数为问题类型、字段和检查编号。"""
    from risk_audit.issue_routing import route_issue
    from risk_audit.models import Record

    row = Record('position_duty', '20L5', '07', 'default', '三清单.xlsx', 'Sheet2', 4, {}, 'row-4')
    issue = {
        'record': row,
        'kind': 'review',
        'evidence': {'issue_type': issue_type, 'field': field},
    }

    result = route_issue(
        issue,
        {'check_id': check_id},
        {'version': 2, 'technical_issues': 'internal', 'unfinished_is_pass': False},
    )

    assert result['kind'] == 'review'
    assert result['record'] is row
    assert result['location_policy'] == 'row'
    assert result['evidence']['publication_channel'] == 'unit'


@pytest.mark.parametrize('has_record', [True, False])
def test_missing_matrix_fields_are_public_material_review(has_record):
    """has_record 为是否有明细；空矩阵或有明细都应提示资料级缺列待核实。"""
    from risk_audit.issue_routing import route_issue
    from risk_audit.models import Record
    row = Record('matrix', '205H', '06', 'default', 'a.xlsx', '矩阵', 2, {}, 'r')
    issue = {'record': row if has_record else None, 'kind': 'review', 'evidence': {'issue_type': 'schema_field_mapping', 'missing_fields': ['control_objective'], 'unavailable_reason': '未能读取控制目标字段，请核实是否缺列。'}}
    result = route_issue(issue, {'check_id': 'matrix_schema'}, {'version': 2, 'technical_issues': 'internal', 'unfinished_is_pass': False})
    assert result['kind'] == 'review'
    assert result['evidence']['publication_channel'] == 'unit'
    assert result['location_policy'] == 'material'


def test_reused_audit_column_rebuilds_current_opinion_only(tmp_path):
    """复用审核意见列时只保留本轮意见；tmp_path 为测试目录。"""
    from risk_audit.writer import write_outputs
    file = make_file(tmp_path, '信通公司06矩阵.xlsx', ['matrix'])
    finding = Finding('k', 'id', 'r', 'c', 'R05', 1, 'violation', '205H', '信通公司', '06', 'default', file.relative_path.as_posix(), '业务表', 2, '请补充姓名。', {}, 'row')
    output = tmp_path / 'output'
    metadata = tmp_path / 'metadata'
    write_outputs([file], [finding], output, metadata_dir=metadata)
    book = load_workbook(output / file.relative_path)
    assert book.active.max_column == 2
    assert book.active['B2'].value == '请补充姓名。'
    # 本轮没有意见时必须清空旧内容，不能把任何历史审核结果带入新副本。
    write_outputs([file], [], output, metadata_dir=metadata)
    assert load_workbook(output / file.relative_path).active['B2'].value is None


def test_rerun_uses_corrected_source_business_data(tmp_path):
    """重跑使用整改后的业务数据并清空旧意见；tmp_path 为隔离目录。"""
    from risk_audit.writer import write_outputs
    file = make_file(tmp_path, '信通公司06矩阵.xlsx', ['matrix'])
    output = tmp_path / 'output'
    metadata = tmp_path / 'metadata'
    write_outputs([file], [], output, metadata_dir=metadata)
    source = load_workbook(file.source)
    source.active['A2'] = '整改后措施'
    source.save(file.source)
    file.sha256 = sha256_file(file.source)
    write_outputs([file], [], output, metadata_dir=metadata)
    result = load_workbook(output / file.relative_path)
    assert result.active['A2'].value == '整改后措施'
    assert result.active['B2'].value is None


def test_rerun_clears_source_and_output_old_opinions(tmp_path):
    """源件和旧副本意见均不带入重跑结果；tmp_path 为隔离目录。"""
    from risk_audit.writer import write_outputs
    file = make_file(tmp_path, '信通公司06矩阵.xlsx', ['matrix'])
    output = tmp_path / 'output'
    metadata = tmp_path / 'metadata'
    write_outputs([file], [], output, metadata_dir=metadata)
    source = load_workbook(file.source)
    source.active['B2'] = '本次源材料新增人工意见'
    source.save(file.source)
    file.sha256 = sha256_file(file.source)
    warnings, _ = write_outputs([file], [], output, metadata_dir=metadata)
    value = load_workbook(output / file.relative_path).active['B2'].value
    assert value is None
    assert not any(warning['type'] == 'human_opinion_sources_differ' for warning in warnings)


def test_material_opinion_keeps_rule_number_without_material_marker(tmp_path):
    """整套材料意见保留规则编号且不增加范围标签；tmp_path 为测试目录。"""
    from risk_audit.writer import write_outputs
    file = make_file(tmp_path, '信通公司06矩阵.xlsx', ['matrix'])
    finding = Finding('k', 'id', 'r', 'c', 'R01', 1, 'violation', '205H', '信通公司', '06', 'default', '', '', None, '【第1条】缺少三清单。', {}, 'material')
    output = tmp_path / 'output'
    write_outputs([file], [finding], output, metadata_dir=tmp_path / 'metadata')
    assert load_workbook(output / file.relative_path).active['B2'].value == '【第1条】缺少三清单。'


def test_repeated_opinion_in_other_row_cannot_change_rule_order(tmp_path):
    """tmp_path 为目录；其它行重复第3条提示时，本行仍须第3条先于第14条。"""
    from risk_audit.writer import write_outputs
    file = make_file(tmp_path, '信通公司06矩阵.xlsx', ['matrix'])
    first = Finding('k1', 'id1', 'r3', 'c3', 'R03', 1, 'violation', '205H', '信通公司', '06', 'default', file.relative_path.as_posix(), '业务表', 2, '【第3条】请补充不适用原因。', {}, 'row')
    second = Finding('k2', 'id2', 'r14', 'c14', 'R14', 1, 'violation', '205H', '信通公司', '06', 'default', file.relative_path.as_posix(), '业务表', 2, '【第14条】请核实适用性匹配情况。', {}, 'row')
    from dataclasses import replace
    third = replace(first, row=3, finding_key='k3', finding_id='id3')
    output = tmp_path / 'output'
    write_outputs([file], [first, second, third], output, metadata_dir=tmp_path / 'metadata')
    value = load_workbook(output / file.relative_path).active['B2'].value
    assert value.index('【第3条】') < value.index('【第14条】')


def test_reused_opinion_column_keeps_right_side_business_cells_readable(tmp_path):
    """tmp_path 为隔离目录；复用中间意见列不能缩小工作表范围并隐藏右侧业务数据。"""
    from risk_audit.writer import write_outputs
    file = make_file(tmp_path, '信通公司06矩阵.xlsx', ['matrix'])
    book = load_workbook(file.source)
    book.active['C1'] = '其它业务字段'
    book.active['C2'] = '现行业务数据'
    book.save(file.source)
    file.sha256 = sha256_file(file.source)
    output = tmp_path / 'output'
    write_outputs([file], [], output, metadata_dir=tmp_path / 'metadata')
    result = load_workbook(output / file.relative_path, read_only=True)
    assert list(result.active.values)[1][-1] == '现行业务数据'


def test_no_writable_workbook_retains_pending_material_opinions(tmp_path):
    """tmp_path 为输出目录；整个包无表格时意见须保留待回填，不能丢失。"""
    import json
    from risk_audit.writer import write_outputs
    finding = Finding('k', 'id', 'r', 'c', 'R01', 1, 'violation', '205H', '信通公司', '06', 'default', '', '', None, '缺少三清单。', {}, 'material')
    metadata = tmp_path / 'metadata'
    warnings, _ = write_outputs([], [finding], tmp_path, metadata_dir=metadata)
    pending = json.loads((metadata / '待回填资料级意见.json').read_text())
    assert pending[0]['entity_name'] == '信通公司'
    assert pending[0]['message'] == '缺少三清单。'
    assert any(warning['type'] == 'no_writable_sheet' for warning in warnings)


def test_missing_unit_material_is_carried_only_by_its_reporting_group(tmp_path):
    """tmp_path 为隔离目录；所属主体缺件可落所属报送包，不得落到其它市公司的材料。"""
    import json
    from risk_audit.writer import write_outputs
    file = make_file(tmp_path, '信通公司06矩阵.xlsx', ['matrix'])
    entities = {'205H': Entity('205H', '信通公司', '信通公司'), 'child': Entity('child', '长沙县公司', '长沙公司')}
    finding = Finding('k', 'id', 'r', 'c', 'R01', 1, 'violation', 'child', '长沙县公司', '06', 'default', '', '', None, '缺少三清单。', {}, 'material')
    output = tmp_path / 'output'
    metadata = tmp_path / 'metadata'
    warnings, _ = write_outputs([file], [finding], output, entities=entities, metadata_dir=metadata)
    assert load_workbook(output / file.relative_path).active['B2'].value is None
    assert json.loads((metadata / '待回填资料级意见.json').read_text())[0]['entity_code'] == 'child'
    assert any(warning['type'] == 'no_writable_sheet' for warning in warnings)
    entities['205H'] = Entity('205H', '长沙公司本部', '长沙公司')
    second = tmp_path / 'same_group'
    write_outputs([file], [finding], second, entities=entities, metadata_dir=tmp_path / 'second_metadata')
    assert '长沙县公司' in load_workbook(second / file.relative_path).active['B2'].value
