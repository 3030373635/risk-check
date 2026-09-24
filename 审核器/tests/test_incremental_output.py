"""验证主体完成即输出、批次元数据累计以及实时日志。"""
import json
import logging
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from risk_audit import runner
from risk_audit.models import Finding


ROOT = Path(__file__).resolve().parents[2]


def create_batch(tmp_path):
    """创建两个主体的真实材料；tmp_path 为隔离测试目录，返回输入目录及主体名册。"""
    input_root = tmp_path / 'input'
    entities_path = tmp_path / 'entities.xlsx'
    entities = Workbook()
    entities.active.append(['单位名称', '单位代码', '上级单位', '对应主业单位', '是否存续'])
    for code, name in [('A001', '测试主体甲有限公司'), ('B001', '测试主体乙有限公司')]:
        entities.active.append([name, code, '测试上级公司', '', '是'])
        directory = input_root / code
        directory.mkdir(parents=True)
        book = Workbook()
        book.active.title = '风控矩阵'
        book.active.append(['控制措施编号', '控制措施', '控制系统', '控制载体', '责任主体', '是否适用及原因'])
        book.active.append(['M1', '核对报表', '/', '/', name + '-财务部-核算专责', '适用'])
        book.save(directory / f'06{name}风控矩阵.xlsx')
    entities.save(entities_path)
    return input_root, entities_path


def execute_batch(tmp_path, input_root, entities_path, **options):
    """执行真实审核；前三个参数为隔离目录、输入和名册，options 为审核选项。"""
    from run_audit import configure_soffice
    configure_soffice()
    return runner.audit(input_root, tmp_path / 'output', ROOT / '审核器/rulepacks/releases/1.8.0',
                        entities_path, ROOT, tmp_path / 'runs', run_id='incremental', **options)


def test_first_entity_is_written_before_second_starts(tmp_path, monkeypatch, capsys):
    """参数为测试目录、替换工具和终端捕获器；验证副本、文件日志及静默终端。"""
    input_root, entities_path = create_batch(tmp_path)
    original_engine = runner.run_engine
    original_parse = runner.parse_files
    original_baselines = runner.load_baselines
    baseline_calls = []
    parsed_entities = []

    def observe_baselines(*args, **kwargs):
        """args/kwargs 为公共基准参数；公共基准只加载一次，且扫描后立即记录主体。"""
        assert not parsed_entities
        log = (tmp_path / 'runs/incremental/audit.log').read_text()
        assert '识别到2个主体' in log and 'A001' in log and 'B001' in log
        baseline_calls.append(True)
        return original_baselines(*args, **kwargs)

    def observe_parse(files, *args, **kwargs):
        """files 为当前主体材料，args/kwargs 为读取参数；第二主体解析前首份副本须已写出。"""
        assert len(baseline_calls) == 1
        codes = {file.entity_code for file in files}
        assert len(codes) == 1, '报送材料必须按主体逐组解析'
        code = next(iter(codes))
        if code == 'B001':
            assert list((tmp_path / 'output/A001').glob('*.xlsx'))
            assert (tmp_path / 'runs/incremental/entities/0001/result.json').exists()
        parsed_entities.append(code)
        return original_parse(files, *args, **kwargs)

    def observe_engine(pack, *args):
        """在第二个主体开始时检查第一份副本；pack 为主体规则包，args 为引擎参数。"""
        if pack['submission_scope']['entity_codes'] == ['B001']:
            first_path = next((tmp_path / 'output/A001').glob('*.xlsx'), None)
            assert first_path is not None, '第一主体完成后必须立即输出'
            book = load_workbook(first_path)
            assert '【' in book.active.cell(2, book.active.max_column).value
            book.close()
            records = json.loads((tmp_path / 'runs/incremental/entities/0001/result.json').read_text())
            assert records['entity_code'] == 'A001'
            assert records['write_completed'] is True
            assert '主体输出完成' in (tmp_path / 'runs/incremental/audit.log').read_text()
        return original_engine(pack, *args)

    monkeypatch.setattr(runner, 'run_engine', observe_engine)
    monkeypatch.setattr(runner, 'parse_files', observe_parse)
    monkeypatch.setattr(runner, 'load_baselines', observe_baselines)
    result = execute_batch(tmp_path, input_root, entities_path)
    assert result['write_completed'] is True
    ownership = json.loads((Path(result['audit_metadata_dir']) / 'ownership.json').read_text())
    assert {Path(path).parts[0] for path in ownership} == {'A001', 'B001'}
    assert len(result['entity_results']) == 2
    assert all(item['write_completed'] for item in result['entity_results'])
    assert parsed_entities == ['A001', 'B001']
    assert len(baseline_calls) == 1
    captured = capsys.readouterr()
    assert captured.out == ''
    assert captured.err == ''


def test_later_failure_keeps_first_output_and_logs_traceback(tmp_path, monkeypatch):
    """参数为隔离目录和替换工具；后续主体异常不能丢失已完成副本，日志须及时关闭。"""
    input_root, entities_path = create_batch(tmp_path)
    original_engine = runner.run_engine
    logger = logging.getLogger('risk_audit')
    previous_handlers = list(logger.handlers)

    def fail_second(pack, *args):
        """模拟第二主体引擎异常；pack 为规则包，args 为其余引擎参数。"""
        if pack['submission_scope']['entity_codes'] == ['B001']:
            raise RuntimeError('测试第二主体异常')
        return original_engine(pack, *args)

    monkeypatch.setattr(runner, 'run_engine', fail_second)
    with pytest.raises(RuntimeError, match='测试第二主体异常'):
        execute_batch(tmp_path, input_root, entities_path)
    assert list((tmp_path / 'output/A001').glob('*.xlsx'))
    log_text = (tmp_path / 'runs/incremental/audit.log').read_text()
    assert 'B001' in log_text
    assert 'Traceback' in log_text
    assert '审核异常终止' in log_text
    assert logger.handlers == previous_handlers


def use_small_baseline(input_root, monkeypatch):
    """input_root 为真实测试材料目录，monkeypatch 为替换工具；使用单份真实矩阵避免无关基准转换。"""
    from risk_audit.util import sha256_file
    original = runner.load_baselines
    source = next((input_root / 'A001').glob('*.xlsx'))
    registry = {'entries': [{'path': str(source.relative_to(input_root)), 'business_code': '06',
                             'sha256': sha256_file(source)}]}

    def load_small(project_root, baseline_registry, aliases, work_dir, **options):
        """前两项为生产基准参数，aliases/work_dir 为读取配置及目录，options 为隐藏表等选项。"""
        return original(input_root, registry, aliases, work_dir, **options)

    monkeypatch.setattr(runner, 'load_baselines', load_small)


def test_second_entity_parse_failure_keeps_first_output(tmp_path, monkeypatch):
    """tmp_path 为测试目录，monkeypatch 为替换工具；后续主体读取异常仍保留第一主体副本与结果。"""
    input_root, entities_path = create_batch(tmp_path)
    use_small_baseline(input_root, monkeypatch)
    original_parse = runner.parse_files

    def fail_second(files, *args, **kwargs):
        """files 为本主体材料，args/kwargs 为读取配置；模拟第二主体读取器不可恢复异常。"""
        if any(file.entity_code == 'B001' for file in files):
            assert list((tmp_path / 'output/A001').glob('*.xlsx'))
            raise RuntimeError('第二主体解析异常')
        return original_parse(files, *args, **kwargs)

    monkeypatch.setattr(runner, 'parse_files', fail_second)
    with pytest.raises(RuntimeError, match='第二主体解析异常'):
        execute_batch(tmp_path, input_root, entities_path)
    assert list((tmp_path / 'output/A001').glob('*.xlsx'))
    assert (tmp_path / 'runs/incremental/entities/0001/result.json').is_file()
    assert not (tmp_path / 'output/B001').exists()


@pytest.mark.parametrize('manual', [False, True])
def test_trial_parses_each_entity_and_saves_final_scope(tmp_path, monkeypatch, manual):
    """tmp_path 为测试目录，monkeypatch 为替换工具，manual 控制是否手动覆盖范围。"""
    input_root, entities_path = create_batch(tmp_path)
    use_small_baseline(input_root, monkeypatch)
    options = {'write': False}
    if manual:
        scope_file = tmp_path / 'scope.json'
        scopes = {code: {'entity_codes': [code], 'businesses': [
            {'business_code': '06', 'variant_id': 'default', 'required': True}]} for code in ['A001', 'B001']}
        scope_file.write_text(json.dumps(scopes))
        options['scope_file'] = scope_file
    result = execute_batch(tmp_path, input_root, entities_path, **options)
    assert not (tmp_path / 'output').exists()
    assert result['parsed_files'] == 2 and result['records'] == 2
    snapshot = json.loads((tmp_path / 'runs/incremental/snapshot.json').read_text())
    assert set(snapshot['submission_scopes']) == {'A001', 'B001'}
    assert snapshot['scope_source']['mode'] == ('manual_override' if manual else 'uploaded_materials')
    inventory = json.loads((tmp_path / 'runs/incremental/inventory.json').read_text())
    assert [file['sheets'][0]['records'] for file in inventory] == [1, 1]


def test_legacy_conversion_files_are_isolated_by_entity(tmp_path, monkeypatch):
    """tmp_path 为测试目录，monkeypatch 为替换工具；同名旧格式文件的转换副本不能相互覆盖。"""
    from risk_audit.readers import excel
    input_root, entities_path = create_batch(tmp_path)
    use_small_baseline(input_root, monkeypatch)
    converted = {}
    for code in ['A001', 'B001']:
        (input_root / code / '06风控矩阵.xls').write_bytes(bytes.fromhex('D0CF11E0A1B11AE1'))

    def convert_fixture(source, directory):
        """source 为旧格式输入，directory 为转换目录；以真实工作簿模拟外部转换器的输出。"""
        directory.mkdir(parents=True, exist_ok=True)
        path = directory / '06风控矩阵.xlsx'
        book = Workbook()
        book.active.title = '风控矩阵'
        book.active.append(['控制措施编号', '控制措施', '责任主体'])
        book.active.append(['M2', source.parent.name, '财务部'])
        book.save(path)
        converted[source.parent.name] = path
        return path, {}

    monkeypatch.setattr(excel, 'convert_xls', convert_fixture)
    execute_batch(tmp_path, input_root, entities_path)
    assert converted['A001'] != converted['B001']
    for code, path in converted.items():
        book = load_workbook(path)
        assert book.active['B2'].value == code
        book.close()
        assert (tmp_path / 'output' / code / '06风控矩阵.xlsx').is_file()


def test_trial_records_each_entity_without_output(tmp_path):
    """tmp_path 为隔离目录；试跑记录各主体结果和日志，但不生成审核副本。"""
    input_root, entities_path = create_batch(tmp_path)
    result = execute_batch(tmp_path, input_root, entities_path, write=False)
    run_directory = Path(result['run_dir'])
    assert not (tmp_path / 'output').exists()
    assert not (run_directory / '_risk_audit').exists()
    trial_reports = [
        '本次审核范围.md', '本次审核范围.json', '主体代码待确认.md', 'entity_alerts.json',
        '隐藏工作表处理提示.md', 'hidden_sheet_alerts.json', '集中复核事项.md', '复核事项.csv',
        'review_tasks.json', 'review_tasks_summary.json', '责任表述集中确认.csv',
        'responsibility_patterns.json', 'responsibility_patterns_summary.json', 'internal_diagnostics.json',
        '内部处理事项.csv', 'capability_diagnostics.json', '程序判定能力明细.csv',
    ]
    assert all((run_directory / name).is_file() for name in trial_reports)
    assert Path(result['submission_scope']['report']).parent == run_directory
    assert Path(result['hidden_sheets']['report']).parent == run_directory
    assert Path(result['review_tasks']['review_task_file']).parent == run_directory
    assert Path(result['internal_diagnostics']['report']).parent == run_directory
    assert Path(result['entity_code_alerts']['report']).parent == run_directory
    assert len(result['entity_results']) == 2
    assert all(item['write_completed'] is False for item in result['entity_results'])
    assert Path(result['log_file']).is_file()


def test_output_state_accumulates_pending_and_unparsed_files(tmp_path):
    """tmp_path 为隔离目录；后续主体写入不得清掉前面的待回填及未审核清单。"""
    from risk_audit.writer import OutputState, write_outputs
    from risk_audit.models import FileRecord
    from risk_audit.util import sha256_file
    state = OutputState()
    metadata = tmp_path / 'runs/_risk_audit'
    for code in ['A001', 'B001']:
        source = tmp_path / f'{code}.xlsx'
        source.write_bytes(b'unreadable workbook')
        file = FileRecord(source, Path(source.name), sha256_file(source), 'xlsx', code,
                          [], False, '06', 'default', 'matrix')
        finding = Finding(code, code, 'r', 'c', 'R01', 1, 'violation', code, code, '06',
                          'default', '', '', None, '缺少三清单。', {}, 'material')
        write_outputs([file], [finding], tmp_path / 'output', output_state=state,
                      metadata_dir=metadata)
    pending = json.loads((metadata / '待回填资料级意见.json').read_text())
    assert [item['entity_code'] for item in pending] == ['A001', 'B001']
    unparsed = json.loads((metadata / '未审核文件.json').read_text())
    assert [item['file'] for item in unparsed] == ['A001.xlsx', 'B001.xlsx']


def test_incremental_rerun_preserves_old_ownership_for_later_entities(tmp_path):
    """tmp_path 为隔离目录；上一轮后续主体的程序意见不可因增量归属覆盖而变成人工意见。"""
    from dataclasses import replace
    from test_delivery_v180 import make_file
    from risk_audit.writer import OutputState, write_outputs
    first = make_file(tmp_path, '甲公司06矩阵.xlsx', ['matrix'])
    second = make_file(tmp_path, '乙公司06矩阵.xlsx', ['matrix'])
    finding = Finding('a', 'a', 'r', 'c', 'R05', 1, 'violation', '205H', '测试主体', '06',
                      'default', str(first.relative_path), '业务表', 2, '上轮程序意见', {}, 'row')
    output = tmp_path / 'output'
    metadata = tmp_path / 'runs/_risk_audit'
    write_outputs([first, second], [finding, replace(finding, finding_key='b', finding_id='b',
                                                   file_path=str(second.relative_path))], output,
                  metadata_dir=metadata)
    state = OutputState()
    for file in [first, second]:
        write_outputs([file], [], output, output_state=state, metadata_dir=metadata)
        # 第一主体写出后，第二主体的旧程序归属必须仍在，保证中断后可安全重跑。
        assert str(second.relative_path) in json.loads((metadata / 'ownership.json').read_text())
        book = load_workbook(output / file.relative_path)
        assert book.active['B2'].value == '人工保留意见'
        book.close()
    assert len(json.loads((metadata / 'ownership.json').read_text())) == 2


def test_batch_path_collision_is_rejected_before_first_output(tmp_path):
    """tmp_path 为隔离目录；格式转换目标冲突必须在任何主体输出之前被整批拒绝并记录日志。"""
    input_root, entities_path = create_batch(tmp_path)
    original = next((input_root / 'A001').glob('*.xlsx'))
    from risk_audit.configuration.loader import load_pack
    pack = load_pack(ROOT / '审核器/rulepacks/releases/1.8.0')
    legacy_path = next(ROOT / item['path'] for item in pack['baseline_registry']['entries']
                       if Path(item['path']).suffix.lower() == '.xls')
    original.with_suffix('.xls').write_bytes(legacy_path.read_bytes())
    with pytest.raises(ValueError, match='输出路径重复'):
        execute_batch(tmp_path, input_root, entities_path)
    assert not (tmp_path / 'output').exists()
    assert '输出路径重复' in (tmp_path / 'runs/incremental/audit.log').read_text()


def test_cli_failure_keeps_json_separate_from_logs(tmp_path, capsys):
    """参数为隔离目录和终端捕获器；命令异常写文件，stdout 仍为可解析 JSON。"""
    from risk_audit.cli import main
    input_root, entities_path = create_batch(tmp_path)
    code = main(['trial', '--input', str(input_root), '--pack', str(tmp_path / 'missing-pack'),
                 '--entities', str(entities_path), '--project-root', str(ROOT),
                 '--runs-root', str(tmp_path / 'runs'), '--run-id', 'cli-error'])
    captured = capsys.readouterr()
    assert code == 2
    assert json.loads(captured.out)['ok'] is False
    assert captured.err == ''
    log_text = (tmp_path / 'runs/cli-error/audit.log').read_text()
    assert '审核异常终止' in log_text
    assert 'Traceback' in log_text
