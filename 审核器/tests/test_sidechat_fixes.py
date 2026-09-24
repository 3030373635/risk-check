from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace
import pytest
from openpyxl import Workbook

from risk_audit.checks.registry import CheckContext, build_registry, schema_contains
from risk_audit.checks.coverage import responsibility_coverage
from risk_audit.checks.applicability import applicability_values
from risk_audit.configuration.loader import load_pack
from risk_audit.configuration.validator import ConfigError, validate_pack
from risk_audit.models import FieldValue, Record, FileRecord, ParsedSheet, Entity
from risk_audit.readers.excel import parse_workbook
from risk_audit.engine import run_engine
from risk_audit.opinion_text import advice

ROOT=Path(__file__).resolve().parents[2]
DRAFT=ROOT/'审核器/rulepacks/drafts/sidechat-111'
CP={'generic_responsibilities':['各部门','相关人员','各级单位'],'confirmed_mappings':{},'confirmed_aliases':{}}
AP={'placeholders':['','/','无','待定','-','--']}


def rec(kind='matrix',row=3,entity='205H',business='06',variant='default',**values):
    data={'measure_id':'M1','applicability':'是','responsibility':'信通-各部门-相关人员',
          'department':'财务部','position':'核算专责','duty':'对报表的准确性负主体责任',**values}
    fields={k:FieldValue(v,v,f'A{row}') for k,v in data.items()}
    return Record(kind,entity,business,variant,kind+'.xlsx',kind,row,fields,f'{kind}-{entity}-{business}-{row}')


def ctx(matrices=(),duties=()):
    return CheckContext(list(matrices),[],[*matrices,*duties],{'205H':Entity('205H','信通')},{},{},['205H'],[{'business_code':'06','variant_id':'default'}])


def test_generic_matrix_with_related_duty_is_only_a_coverage_limit():
    issues=responsibility_coverage(ctx([rec()],[rec('position_duty')]),CP)
    assert len(issues)==1 and issues[0]['kind']=='limitation'
    assert issues[0]['evidence']['issue_type']=='coverage_mapping_unconfirmed'


def test_generic_does_not_hide_actual_absence_of_any_duty_record():
    issues=responsibility_coverage(ctx([rec()],[]),CP)
    assert len(issues)==1 and issues[0]['kind']=='violation'
    assert issues[0]['evidence']['issue_type']=='duty_record_missing'


def test_confirmed_generic_mapping_checks_every_required_pair():
    source=rec(); d=rec('position_duty')
    p={**CP,'confirmed_mappings':{source.value('responsibility'):[{'department':'财务部','position':'核算专责'}]}}
    assert responsibility_coverage(ctx([source],[d]),p)==[]
    p['confirmed_mappings'][source.value('responsibility')].append({'department':'财务部','position':'稽核专责'})
    issues=responsibility_coverage(ctx([source],[d]),p)
    assert issues[0]['evidence']['issue_type']=='specific_responsibility_missing'


def test_generic_item_does_not_suppress_check_of_a_specific_item():
    source=rec(responsibility='信通-各部门-相关人员、信通-财务部-稽核专责')
    duties=[rec('position_duty'),rec('position_duty',row=4,measure_id='M2',position='稽核专责')]
    issues=responsibility_coverage(ctx([source],duties),CP)
    assert {i['kind'] for i in issues}=={'review','limitation'}
    concrete=next(i for i in issues if i['kind']=='review')
    assert concrete['evidence']['missing_responsibilities']==['信通-财务部-稽核专责']
    assert '稽核专责' in advice('matrix_duty_coverage','review',concrete['evidence'])


@pytest.mark.parametrize('kw',[{'entity':'OTHER'},{'business':'07'},{'variant':'other'},{'measure_id':'M2'}])
def test_other_scopes_do_not_satisfy_duty_coverage(kw):
    issues=responsibility_coverage(ctx([rec()],[rec('position_duty',**kw)]),CP)
    assert issues[0]['evidence']['issue_type']=='duty_record_missing'


def test_unconfirmed_alias_is_not_a_missing_responsibility_finding():
    source=rec(responsibility='信通-财资部-核算专责');d=rec('position_duty')
    assert responsibility_coverage(ctx([source],[d]),CP)[0]['kind']=='limitation'
    assert responsibility_coverage(ctx([source],[d]),{**CP,'confirmed_aliases':{'财资部':'财务部'}})==[]


def test_unknown_applicability_is_not_a_missing_duty_finding():
    assert responsibility_coverage(ctx([rec(applicability='')],[]),CP)[0]['kind']=='limitation'
    assert responsibility_coverage(ctx([rec(applicability='不适用：没有该业务')],[]),CP)==[]


@pytest.mark.parametrize('text,reason,category',[
 ('是','',''),('适用','',''),('不适用','不承担该业务',''),('不适用：不承担该业务','',''),
 ('否；不承担该业务','',''),('不适用','','not_applicable_reason_empty'),
 ('否','/','not_applicable_reason_empty'),('','已写备注','applicability_empty'),
 ('/','','applicability_empty')])
def test_real_applicability_values_matter_not_split_headers(text,reason,category):
    issues=applicability_values(ctx([rec(applicability=text,applicability_reason=reason)]),AP)
    assert ([i['evidence']['issue_type'] for i in issues] == [category]) if category else issues==[]


def test_missing_reason_formula_cache_is_not_a_pass():
    row=rec(applicability='不适用',applicability_reason='')
    row.fields['applicability_reason'].state='formula_no_cache'
    assert applicability_values(ctx([row]),AP)[0]['evidence']['issue_type']=='reason_cache_missing'


def test_split_header_acceptance_does_not_skip_missing_business_fields(tmp_path):
    pack=load_pack(DRAFT);path=tmp_path/'matrix.xlsx';w=Workbook();s=w.active;s.title='风控矩阵'
    s.append(['风控矩阵']);s.append(['控制措施编号','控制措施','控制载体','责任主体','是否适用','修改备注'])
    s.append(['M1','控制','单据','财务部-核算专责','否','无此业务']);w.save(path)
    f=FileRecord(path,Path(path.name),'h','xlsx','205H',[],False,'06','default','matrix')
    f.sheets=parse_workbook(f,path,pack['field_aliases'])
    c=ctx();c.files=[f];c.resources={'field_aliases':pack['field_aliases']}
    p={'sheet_types':['matrix'],'required_fields':[],'check_applicability_tail':True,'accept_split_applicability':True}
    assert schema_contains(c,p)==[]
    issues=schema_contains(c,{**p,'required_fields':['control_system']})
    assert any('control_system' in x['evidence'].get('missing_fields',[]) for x in issues)


def test_unit_paired_reason_takes_precedence_over_an_earlier_modification_note(tmp_path):
    pack=load_pack(DRAFT);path=tmp_path/'matrix.xlsx';w=Workbook();s=w.active;s.title='风控矩阵'
    s.append(['风控矩阵']);s.append(['控制措施编号','控制措施','责任主体','修改备注','信通公司是否适用','信通公司修改备注'])
    s.append(['M1','控制','财务部-核算专责','历史备注','否','无此业务']);w.save(path)
    f=FileRecord(path,Path(path.name),'h','xlsx','205H',[],False,'06','default','matrix')
    parsed=parse_workbook(f,path,pack['field_aliases'])[0]
    assert parsed.columns['applicability_reason']==6
    assert applicability_values(ctx(parsed.records),AP)==[]
    s['F3']='';w.save(path);rows=parse_workbook(f,path,pack['field_aliases'])[0].records
    assert applicability_values(ctx(rows),AP)[0]['evidence']['issue_type']=='not_applicable_reason_empty'


def test_limits_are_reported_in_status_without_generating_a_material_opinion():
    pack=load_pack(DRAFT);pack['rules']=[r for r in pack['rules'] if r['rule_id']=='duties.coverage']
    source=rec();d=rec('position_duty')
    f=FileRecord(Path('test.xlsx'),Path('test.xlsx'),'h','xlsx','205H',[],False,'06','default','matrix',
                 [ParsedSheet('matrix','matrix',[2],{}, {},9,3,[source]),ParsedSheet('position_duty','position_duty',[2],{}, {},9,3,[d])])
    limits=[];findings,statuses=run_engine(pack,build_registry(),[f],{'205H':Entity('205H','信通')},{},limitations=limits)
    assert findings==[] and len(limits)==1
    assert statuses[0].status=='partial' and statuses[0].limitations==1
    assert limits[0]['row']==3 and limits[0]['business_code']=='06'


def test_sidechat_pack_keeps_deleted_text_as_r09_and_does_not_adopt_paused_word_rules():
    pack=load_pack(DRAFT);r9=next(r for r in pack['rules'] if r['display_code']=='R09')
    assert [(c['operator'],c['params']['match_scope']) for c in r9['checks']]==[('deleted_text_reappears','measure')]
    assert not any(r['display_code']=='R13' for r in pack['rules'])
    for name in ['R05.json','R09.json','R10.json']:
        assert (DRAFT/'rules'/name).read_bytes()==(ROOT/'审核器/rulepacks/releases/1.1.0/rules'/name).read_bytes()
    assert not any(c['operator'] in {'reference_exists','set_subset'} for r in pack['rules'] if r['enabled'] for c in r['checks'])


def test_empty_confirmed_mapping_cannot_silently_approve_a_responsibility():
    pack=load_pack(DRAFT);rule=next(r for r in pack['rules'] if r['rule_id']=='duties.coverage')
    rule['checks'][0]['params']['confirmed_mappings']={'各部门':[]}
    with pytest.raises(ConfigError):validate_pack(pack,build_registry())


def test_generic_department_not_reported_missing_just_because_its_words_exist_elsewhere():
    source=rec(responsibility='信通-管理部门-专业管理人员')
    duties=[rec('position_duty'),rec('position_duty',row=4,measure_id='M2',department='管理部门',position='专业管理人员')]
    issues=responsibility_coverage(ctx([source],duties),CP)
    assert len(issues)==1 and issues[0]['kind']=='limitation'


def test_moved_applicability_field_is_not_falsely_reported_as_a_deleted_baseline_column(tmp_path):
    pack=load_pack(DRAFT);path=tmp_path/'matrix.xlsx';w=Workbook();s=w.active;s.title='风控矩阵'
    s.append(['风控矩阵']);s.append(['控制措施编号','控制措施','控制载体','责任主体','其他说明','是否适用','修改备注'])
    s.append(['M1','控制','单据','财务部-核算专责','','是','']);w.save(path)
    f=FileRecord(path,Path(path.name),'h','xlsx','205H',[],False,'06','default','matrix')
    f.sheets=parse_workbook(f,path,pack['field_aliases'])
    c=ctx();c.files=[f];c.resources={'field_aliases':pack['field_aliases']}
    c.baselines={('06','default'):{'fields':['applicability'],'business_header_paths':[{'coordinate':'E2','column':5,'path':[{'name':'是否适用','coordinate':'E2'}]}]}}
    p={'sheet_types':['matrix'],'required_fields':[],'compare_baseline':True,'ignore_fields':['applicability','applicability_reason'],'check_applicability_tail':True,'accept_split_applicability':True}
    assert schema_contains(c,p)==[]


@pytest.mark.parametrize('kw',[{'entity':''},{'business':''},{'variant':'unknown'}])
def test_unresolved_scopes_are_not_merged_into_a_confirmed_coverage_result(kw):
    source=rec(responsibility='信通-财务部-核算专责',**kw)
    duties=[rec('position_duty',**kw)]
    issues=responsibility_coverage(ctx([source],duties),CP)
    assert len(issues)==1 and issues[0]['evidence']['issue_type']=='coverage_scope_unconfirmed'
    assert issues[0]['kind']=='limitation'
