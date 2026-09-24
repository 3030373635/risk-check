"""审核统计表按单位、业务和规则组合汇总的行为验证。"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from risk_audit.audit_statistics import write_audit_statistics
from risk_audit.models import FileRecord


def _file(tmp_path: Path, relative_path: str, material_type: str, business_code: str) -> FileRecord:
    """构建审核统计样例文件；tmp_path 为测试目录，relative_path 为包内路径。"""
    return FileRecord(
        tmp_path / Path(relative_path).name,
        Path(relative_path),
        f'hash-{relative_path}',
        'xlsx',
        '205H',
        [],
        False,
        business_code,
        'default',
        material_type,
    )


def test_statistics_groups_nested_unit_businesses_and_rule_combinations(tmp_path):
    """tmp_path 为隔离目录；单位取业务上级完整路径，备注合并矩阵与三清单意见。"""
    matrix = _file(tmp_path, 'a/b/c/d/风控矩阵.xlsx', 'matrix', '09')
    lists = _file(tmp_path, 'a/b/c/d/三清单.xlsx', 'three_lists', '09')
    system = _file(tmp_path, 'a/b/c/e/风控矩阵.xlsx', 'matrix', '07')
    ownership = {
        str(matrix.relative_path): [
            {'sheet': '岗位清单', 'cell': 'H2', 'program_text': '【第10条】岗位职责请按照国网标准句式编制。'},
            {'sheet': '风控矩阵', 'cell': 'H4', 'program_text': '【第5条】【第3条】待核实：未能读取必需字段。'},
        ],
        str(lists.relative_path): [
            {'sheet': '岗位清单', 'cell': 'I3', 'program_text': '【第10条】岗位职责请按照国网标准句式编制。'},
            {'sheet': '岗位清单', 'cell': 'I5', 'program_text': '【第3条】【第5条】字段对应待核实。'},
        ],
        str(system.relative_path): [
            {'sheet': '系统规则', 'cell': 'G2', 'program_text': '【第7条】请核实与矩阵对应的系统规则是否应该做出修改。'},
        ],
    }
    target = tmp_path / '审核统计表.xlsx'

    write_audit_statistics(target, [matrix, lists, system], ownership, '包名')

    sheet = load_workbook(target)['审核统计表']
    assert list(sheet.values) == [
        ('单位', '备注'),
        ('包名/a/b/c',
         'd：\n【第3条】【第5条】待核实：未能读取必需字段。 2条。\n'
         '【第10条】岗位职责请按照国网标准句式编制。 2条。\n\n'
         'e：\n【第7条】请核实与矩阵对应的系统规则是否应该做出修改。 1条。'),
    ]


def test_statistics_writes_none_for_recognized_business_without_opinions(tmp_path):
    """tmp_path 为隔离目录；已识别但无程序意见的业务仍须输出“无”。"""
    sales = _file(tmp_path, '1.某供电公司/01营销售电/风控矩阵.xlsx', 'matrix', '01')
    trade = _file(tmp_path, '1.某供电公司/02交易与购电/风控矩阵.xlsx', 'matrix', '02')
    ownership = {
        str(sales.relative_path): [
            {'sheet': '风控矩阵', 'cell': 'H2', 'program_text': '【第10条】请修改岗位职责。'},
        ],
    }
    target = tmp_path / '审核统计表.xlsx'

    write_audit_statistics(target, [sales, trade], ownership, '包名')

    assert list(load_workbook(target)['审核统计表'].values) == [
        ('单位', '备注'),
        ('包名/1.某供电公司',
         '01营销售电：\n【第10条】请修改岗位职责。 1条。\n\n'
         '02交易与购电：\n无'),
    ]


def test_statistics_keeps_root_files_with_duplicate_display_codes_separate(tmp_path):
    """tmp_path 为隔离目录；根目录下两个 01 业务必须显示为两个统计分区。"""
    sales = _file(tmp_path, '01风控矩阵-营销售电.xlsx', 'matrix', '01')
    sales.business_id = 'sales_electricity'
    sales.preservation['business_name'] = '营销售电'
    equity = _file(tmp_path, '01风控矩阵-股权管理.xlsx', 'matrix', '01')
    equity.business_id = 'equity_management'
    equity.preservation['business_name'] = '股权（产权）管理'
    target = tmp_path / '审核统计表.xlsx'

    write_audit_statistics(target, [sales, equity], {}, '测试主体有限公司')

    text = load_workbook(target)['审核统计表']['B2'].value
    assert '01 营销售电：\n无' in text
    assert '01 股权（产权）管理：\n无' in text


def test_statistics_formats_header_and_protects_excel_text(tmp_path):
    """tmp_path 为隔离目录；表头样式明显，外部目录文本不能触发 Excel 公式。"""
    file = _file(tmp_path, 'a/=1+1/风控矩阵.xlsx', 'matrix', '09')
    ownership = {str(file.relative_path): [
        {'sheet': '风控矩阵', 'cell': 'H2', 'program_text': '【第10条】请修改。'},
    ]}
    target = tmp_path / '审核统计表.xlsx'

    write_audit_statistics(target, [file], ownership, '=PACKAGE')

    sheet = load_workbook(target, data_only=False)['审核统计表']
    assert sheet['A1'].font.bold is True
    assert sheet['A1'].fill.fill_type == 'solid'
    assert sheet['A1'].alignment.horizontal == 'center'
    assert sheet['B2'].alignment.wrap_text is True
    assert sheet.freeze_panes == 'A2'
    assert sheet.auto_filter.ref == 'A1:B2'
    assert sheet['A2'].data_type == 's' and sheet['A2'].value.startswith("'=PACKAGE/a")
    assert sheet['B2'].data_type == 's' and sheet['B2'].value.startswith("'=1+1：")


def test_statistics_recognizes_business_before_unit_directory(tmp_path):
    """tmp_path 为隔离目录；真实包中业务目录在单位之前时仍按单位合并。"""
    sales = _file(tmp_path, '1.营销售电-ok/1.国网某供电分公司/01风控矩阵.xlsx', 'matrix', '01')
    trade = _file(tmp_path, '2.交易与购电-ok/1.国网某供电分公司/02风控矩阵.xlsx', 'matrix', '02')
    ownership = {
        str(sales.relative_path): [
            {'sheet': '风控矩阵', 'cell': 'H2', 'program_text': '【第10条】请修改岗位职责。'},
        ],
        str(trade.relative_path): [
            {'sheet': '风控矩阵', 'cell': 'H2', 'program_text': '【第7条】请核实系统规则。'},
        ],
    }
    target = tmp_path / '审核统计表.xlsx'

    write_audit_statistics(target, [sales, trade], ownership, '包名')

    assert list(load_workbook(target)['审核统计表'].values) == [
        ('单位', '备注'),
        ('包名/1.国网某供电分公司',
         '1.营销售电-ok：\n【第10条】请修改岗位职责。 1条。\n\n'
         '2.交易与购电-ok：\n【第7条】请核实系统规则。 1条。'),
    ]


def test_statistics_prefers_named_business_over_numbered_center_unit(tmp_path):
    """tmp_path 为隔离目录；同编号业务与中心单位并存时优先识别业务名。"""
    file = _file(tmp_path, '01营销售电/1.某技术技能培训中心/风控矩阵.xlsx', 'matrix', '01')
    ownership = {str(file.relative_path): [
        {'sheet': '风控矩阵', 'cell': 'H2', 'program_text': '【第10条】请修改岗位职责。'},
    ]}
    target = tmp_path / '审核统计表.xlsx'

    write_audit_statistics(target, [file], ownership, '包名')

    assert list(load_workbook(target)['审核统计表'].values) == [
        ('单位', '备注'),
        ('包名/1.某技术技能培训中心', '01营销售电：\n【第10条】请修改岗位职责。 1条。'),
    ]


def test_statistics_warns_and_keeps_header_when_written_file_is_unknown(tmp_path, caplog):
    """tmp_path/caplog 为测试工具；无法归属的写回记录告警而不中断输出。"""
    ownership = {'未知文件.xlsx': [
        {'sheet': '风控矩阵', 'cell': 'H2', 'program_text': '【第1条】缺少业务材料。'},
    ]}
    target = tmp_path / '审核统计表.xlsx'

    write_audit_statistics(target, [], ownership, '包名')

    assert '审核统计表未能定位意见所属业务目录' in caplog.text
    assert list(load_workbook(target)['审核统计表'].values) == [('单位', '备注')]


def test_statistics_splits_actual_program_opinions_from_one_written_cell(tmp_path):
    """tmp_path 为隔离目录；同一审核意见单元格内的独立意见分别按规则组合统计。"""
    file = _file(tmp_path, 'a/b/c/d/风控矩阵.xlsx', 'matrix', '09')
    ownership = {
        str(file.relative_path): [{
            'sheet': '风控矩阵',
            'cell': 'H2',
            'program_text': '【第10条】请修改岗位职责。\n【第3条】【第5条】请核实必需字段。',
        }],
    }
    target = tmp_path / '审核统计表.xlsx'

    write_audit_statistics(target, files=[file], ownership=ownership, package_name='包名')

    assert load_workbook(target)['审核统计表']['B2'].value == (
        'd：\n【第3条】【第5条】请核实必需字段。 1条。\n'
        '【第10条】请修改岗位职责。 1条。'
    )
