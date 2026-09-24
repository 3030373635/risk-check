"""验证主体内按业务及电压变体独立解析、审核并即时输出。"""
import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from risk_audit import runner
from risk_audit.util import sha256_file


ROOT = Path(__file__).resolve().parents[2]


def create_business_batch(tmp_path, units):
    """tmp_path 为测试目录，units 为业务及变体列表；返回输入、名册和各单元源文件。"""
    input_root = tmp_path / 'input'
    entities_path = tmp_path / 'entities.xlsx'
    roster = Workbook()
    roster.active.append(['单位名称', '单位代码', '上级单位', '对应主业单位', '是否存续'])
    roster.active.append(['测试主体甲有限公司', 'A001', '', '', '是'])
    roster.save(entities_path)
    sources = {}
    for business, variant in units:
        label = {'default': '业务', '35-220kv': '35千伏-220千伏电网基建',
                 '500-750kv': '500千伏-750千伏电网基建'}[variant]
        directory = input_root / 'A001' / f'{business} {label}'
        directory.mkdir(parents=True)
        matrix = Workbook()
        matrix.active.title = '风控矩阵'
        matrix.active.append(['控制措施编号', '控制措施', '控制系统', '控制载体', '责任主体', '是否适用及原因'])
        matrix.active.append(['M1', '核对报表', '/', '/', '测试主体甲有限公司-财务部-核算专责', '适用'])
        matrix_path = directory / f'{business}风控矩阵-{label}-测试主体甲有限公司.xlsx'
        matrix.save(matrix_path)
        lists = Workbook()
        lists.active.title = '岗位内控责任清单'
        lists.active.append(['控制措施编号', '控制措施', '部门', '岗位名称', '人员姓名', '人员编号', '岗位职责', '角色'])
        lists.active.append(['M1', '核对报表', '财务部', '核算专责', '张三', 'E1', '对报表负主体责任', '经办'])
        incompatible = lists.create_sheet('不相容岗位清单')
        incompatible.append(['控制措施编号', '部门', '不相容业务角色', '岗位A', '岗位B', '岗位职责'])
        incompatible.append(['M1', '财务部', '经办/审核', '核算专责', '主任', '核对报表'])
        system = lists.create_sheet('系统控制规则清单')
        system.append(['控制措施编号', '规则名称', '规则内容'])
        system.append(['M1', '报表规则', '核对报表'])
        lists_path = directory / f'{business}三清单-{label}-测试主体甲有限公司.xlsx'
        lists.save(lists_path)
        sources[(business, variant)] = [matrix_path, lists_path]
    return input_root, entities_path, sources


def use_business_baseline(input_root, sources, monkeypatch):
    """input_root/sources 为真实测试材料，monkeypatch 为替换工具；读取一份真实基准并记录加载次数。"""
    original = runner.load_baselines
    key = next(iter(sources))
    path = sources[key][0]
    registry = {'entries': [{'path': str(path.relative_to(input_root)), 'business_code': key[0],
                             'variant_id': key[1], 'sha256': sha256_file(path)}]}
    calls = []

    def load_small(project_root, baseline_registry, aliases, work_dir, **options):
        """前两项为生产基准配置，aliases/work_dir 为读取配置及目录，options 为读取选项。"""
        calls.append(True)
        return original(input_root, registry, aliases, work_dir, **options)

    monkeypatch.setattr(runner, 'load_baselines', load_small)
    return calls


@pytest.mark.parametrize('units', [[('06', 'default'), ('07', 'default')],
                                  [('03', '35-220kv'), ('03', '500-750kv')]])
@pytest.mark.parametrize('write', [True, False])
@pytest.mark.parametrize('manual', [False, True])
def test_businesses_are_parsed_audited_and_written_independently(tmp_path, monkeypatch, units, write, manual):
    """tmp_path/monkeypatch 为测试工具，units/write/manual 控制业务组合、输出及手动范围。"""
    input_root, entities_path, sources = create_business_batch(tmp_path, units)
    baseline_calls = use_business_baseline(input_root, sources, monkeypatch)
    original_parse, original_engine = runner.parse_files, runner.run_engine
    parsed, audited = [], []

    def observe_parse(files, *args, **kwargs):
        """files 为待解析单元，args/kwargs 为读取参数；后续业务解析前检查前一单元结果可见。"""
        keys = {(file.entity_code, file.business_code, file.variant_id) for file in files}
        assert len(keys) == 1, '一次只能解析一个主体业务及矩阵类型'
        assert len(baseline_calls) == 1
        if parsed:
            first = tmp_path / 'runs/business/entities/0001/businesses/0001/result.json'
            assert first.is_file(), '前一业务结果须在下一业务解析前落盘'
            for source in sources[parsed[-1]]:
                assert (tmp_path / 'output' / source.relative_to(input_root)).exists() == write
        _, business, variant = next(iter(keys))
        parsed.append((business, variant))
        return original_parse(files, *args, **kwargs)

    def observe_engine(pack, registry, files, entities, baselines, *args):
        """参数为规则、注册表、本单元文件、名册及基准；args 为运行标识和限制收集器。"""
        scope = pack['submission_scope']
        assert len(scope['businesses']) == 1, '审核范围不得包含其他业务或变体'
        business = scope['businesses'][0]
        key = business['business_code'], business['variant_id']
        assert {(file.business_code, file.variant_id) for file in files} == {key}
        assert set(baselines) <= {key}, '本单元只能使用对应公共基准'
        assert {sheet.sheet_type for file in files for sheet in file.sheets} == {
            'matrix', 'position_duty', 'incompatible_position', 'system_rule'}
        audited.append(key)
        return original_engine(pack, registry, files, entities, baselines, *args)

    monkeypatch.setattr(runner, 'parse_files', observe_parse)
    monkeypatch.setattr(runner, 'run_engine', observe_engine)
    scope_file = None
    if manual:
        scope_file = tmp_path / 'scope.json'
        scope_file.write_text(json.dumps({'A001': {'entity_codes': ['A001'], 'businesses': [
            {'business_code': business, 'variant_id': variant, 'required': True} for business, variant in units]}}))
    result = runner.audit(input_root, tmp_path / 'output', ROOT / '审核器/rulepacks/releases/1.8.0',
                          entities_path, ROOT, tmp_path / 'runs', run_id='business', write=write, scope_file=scope_file)
    assert sorted(parsed) == sorted(units) and audited == parsed
    assert len(baseline_calls) == 1
    assert len(result['business_results']) == 2
    assert len(result['entity_results'][0]['business_results']) == 2
    snapshot = json.loads((tmp_path / 'runs/business/snapshot.json').read_text())
    assert len(snapshot['submission_scopes']['A001']['businesses']) == 2


def test_later_business_failure_preserves_finished_business(tmp_path, monkeypatch):
    """tmp_path/monkeypatch 为测试工具；后续业务异常不能丢失同一主体已完成的业务副本。"""
    units = [('06', 'default'), ('07', 'default')]
    input_root, entities_path, sources = create_business_batch(tmp_path, units)
    use_business_baseline(input_root, sources, monkeypatch)
    original = runner.parse_files

    def fail_second(files, *args, **kwargs):
        """files 为当前业务文件，args/kwargs 为读取参数；模拟第二业务不可恢复读取异常。"""
        if any(file.business_code == '07' for file in files):
            raise RuntimeError('第二业务解析异常')
        return original(files, *args, **kwargs)

    monkeypatch.setattr(runner, 'parse_files', fail_second)
    with pytest.raises(RuntimeError, match='第二业务解析异常'):
        runner.audit(input_root, tmp_path / 'output', ROOT / '审核器/rulepacks/releases/1.8.0',
                     entities_path, ROOT, tmp_path / 'runs', run_id='business')
    assert all((tmp_path / 'output' / source.relative_to(input_root)).exists() for source in sources[units[0]])
    progress = json.loads((tmp_path / 'runs/business/entities/0001/result.json').read_text())
    assert progress['completed_businesses'] == 1 and progress['write_completed'] is False


@pytest.mark.parametrize('fail_copy', [False, True])
def test_explanation_is_copied_once_and_failure_keeps_entity_incomplete(tmp_path, monkeypatch, fail_copy):
    """tmp_path/monkeypatch 为测试工具，fail_copy 控制说明复制异常；业务完成不能提前标记主体完成。"""
    from docx import Document
    input_root, entities_path, sources = create_business_batch(tmp_path, [('06', 'default'), ('07', 'default')])
    use_business_baseline(input_root, sources, monkeypatch)
    document = Document()
    document.add_paragraph('报送材料说明')
    source = input_root / 'A001/测试主体甲有限公司说明.docx'
    document.save(source)
    original = runner.write_outputs
    copies = []

    def observe_copy(files, *args, **kwargs):
        """files 为输出文件，args/kwargs 为意见和输出配置；说明须单独原样复制一次。"""
        if any(file.material_type == 'explanation' for file in files):
            assert all(file.material_type == 'explanation' for file in files)
            copies.append(True)
            if fail_copy:
                raise RuntimeError('说明复制异常')
        return original(files, *args, **kwargs)

    monkeypatch.setattr(runner, 'write_outputs', observe_copy)
    if fail_copy:
        with pytest.raises(RuntimeError, match='说明复制异常'):
            runner.audit(input_root, tmp_path / 'output', ROOT / '审核器/rulepacks/releases/1.8.0',
                         entities_path, ROOT, tmp_path / 'runs', run_id='business')
        progress = json.loads((tmp_path / 'runs/business/entities/0001/result.json').read_text())
        assert progress['completed_businesses'] == 2 and progress['write_completed'] is False
    else:
        runner.audit(input_root, tmp_path / 'output', ROOT / '审核器/rulepacks/releases/1.8.0',
                     entities_path, ROOT, tmp_path / 'runs', run_id='business')
        assert sha256_file(tmp_path / 'output' / source.relative_to(input_root)) == sha256_file(source)
    assert len(copies) == 1


@pytest.mark.parametrize('uploaded_business', ['06', '07'])
def test_application_list_does_not_create_missing_business(tmp_path, monkeypatch, uploaded_business):
    """tmp_path/monkeypatch 为测试工具，uploaded_business 为已报业务；辅助清单不得扩展审核范围。"""
    input_root, entities_path, sources = create_business_batch(tmp_path, [(uploaded_business, 'default')])
    use_business_baseline(input_root, sources, monkeypatch)
    attachment = Workbook()
    attachment.active.append(['单位', '设备（资产）管理', '物资（服务）采购与实施'])
    attachment.active.append(['测试主体甲有限公司', '适用', '适用'])
    attachment.save(input_root / 'A001/附件1应用清单.xlsx')
    result = runner.audit(input_root, tmp_path / 'output', ROOT / '审核器/rulepacks/releases/1.8.0',
                          entities_path, ROOT, tmp_path / 'runs', run_id='business')
    assert [item['business_code'] for item in result['business_results']] == [uploaded_business]
    missing_business = '07' if uploaded_business == '06' else '06'
    scope = json.loads((Path(result['audit_metadata_dir']) / '本次审核范围.json').read_text())
    businesses = next(iter(scope['submission_scopes'].values()))['businesses']
    assert [item['business_code'] for item in businesses] == [uploaded_business]
    findings = json.loads((tmp_path / 'runs/business/findings.json').read_text())
    assert not any(item['business_code'] == missing_business for item in findings)
