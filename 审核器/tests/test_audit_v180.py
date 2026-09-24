"""构造 Excel 验证确认稿规则从读取到写回的完整流程。"""
from copy import deepcopy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

from risk_audit.configuration.loader import load_pack
from risk_audit.runner import audit
from risk_audit.util import sha256_file, write_json

ROOT = Path(__file__).resolve().parents[2]


def create_sample(root):
    """创建有意含问题的构造材料；root 为独立测试输入目录，不代表单位实际报送。"""
    root.mkdir(parents=True, exist_ok=True)
    path = root / '06信通公司矩阵及三清单-程序构造样例.xlsx'
    book = Workbook()
    matrix = book.active
    matrix.title = '风控矩阵'
    matrix.append(['控制措施编号', '控制措施', '控制系统', '控制载体', '责任主体', '是否适用及原因'])
    matrix.append(['M2', '省公司措施', '/', '/', '省公司-财务部-主任', '是'])
    deleted = CellRichText([TextBlock(InlineFont(strike=True), '旧审批单'), TextBlock(InlineFont(), '新审批单')])
    matrix.append(['M10', '核对报表', '/', deleted, '国网湖南省电力有限公司信息通信分公司-财务部-核算专责', '适用'])
    matrix.append(['M1', '不开展的业务', '/', '/', '国网湖南省电力有限公司信息通信分公司-财务部-核算专责', '否：本期无此业务'])
    duties = book.create_sheet('岗位内控责任清单')
    duties.append(['控制措施编号', '控制措施', '部门', '岗位名称', '人员姓名', '人员编号', '岗位职责', '角色', '审核意见'])
    duties.append(['M10', '核对报表', '财务部', '合同经办人员', '张三', 'E1', '对旧审批单的准确性负主体责任', '经办', '人工样例意见'])
    duties.append(['M10', '核对报表', '财务部', '核算专责', '张三', 'E1', '对旧审批单的准确性负审核责任', '审核'])
    duties.append(['M1', '不开展的业务', '财务部', '核算专责', '李四', 'E2', '负审核责任', '经办'])
    incompatible = book.create_sheet('不相容岗位清单')
    incompatible.append(['控制措施编号', '部门', '不相容业务角色', '岗位A', '岗位B', '岗位职责'])
    incompatible.append(['M2', '管理部门', '经办/审核', '负责人', '相关人员', '省公司职责'])
    incompatible.append(['M10', '管理部门', '经办/审核', '负责人', '核算专责', '对旧审批单负主体责任'])
    system = book.create_sheet('系统控制规则清单')
    system.append(['控制措施编号', '规则名称', '规则内容'])
    system.append(['M10', '规则2', '示例规则'])
    system.append(['M1', '规则1', '示例规则'])
    book.save(path)
    return path


def test_confirmed_pack_full_audit_preserves_inputs_and_human_opinions(tmp_path):
    """tmp_path 为隔离目录；验证新规则意见、适用范围、原始哈希及人工意见保留。"""
    from run_audit import configure_soffice
    configure_soffice()
    input_root = tmp_path / 'input'
    path = create_sample(input_root)
    before = sha256_file(path)
    scope = tmp_path / 'scope.json'
    write_json(scope, {'205H': {'entity_codes': ['205H'], 'businesses': [{'business_code': '06', 'variant_id': 'default', 'required': True}]}})
    result = audit(input_root, tmp_path / 'output', ROOT / '审核器/rulepacks/releases/1.8.0', ROOT / '审核/会计主体清单20260907.xlsx', ROOT, tmp_path / 'runs', scope_file=scope)
    assert result['write_completed']
    assert result['parsed_files'] == 1
    assert not result['statuses'].get('failed')
    assert sha256_file(path) == before
    output = load_workbook(tmp_path / 'output' / path.name)
    assert '【资料级】' not in output['风控矩阵']['G2'].value
    assert '【第3条】待核实' in output['风控矩阵']['G2'].value
    assert '控制目标' in output['风控矩阵']['G2'].value
    assert output['风控矩阵']['G2'].value.index('【第3条】') < output['风控矩阵']['G2'].value.index('【第14条】')
    assert '【第14条】请核实适用性匹配情况。' in output['风控矩阵']['G2'].value
    assert '人工样例意见' in output['岗位内控责任清单']['I2'].value
    assert '【第13条】' in output['岗位内控责任清单']['I2'].value
    assert '【第8条】' in output['岗位内控责任清单']['I3'].value
    assert '【第11条】经办对应主体责任、审核对应审核责任、审批对应审批责任。' in output['岗位内控责任清单']['I4'].value
    assert output['不相容岗位清单']['G2'].value is None
    assert '【第6条】' in output['不相容岗位清单']['G3'].value
    assert '【第13条】' in output['不相容岗位清单']['G3'].value


def test_configuration_rejects_invalid_responsibility_patterns():
    """无参数；非法适用范围必须在写入任何输出之前拒绝。"""
    import pytest
    from risk_audit.checks.registry import build_registry
    from risk_audit.configuration.validator import ConfigError, validate_pack
    pack = deepcopy(load_pack(ROOT / '审核器/rulepacks/releases/1.8.0'))
    pack['manifest']['status'] = 'draft'
    pack['responsibility_applicability']['restrictions'][0]['patterns'] = ['[']
    with pytest.raises(ConfigError, match='非法'):
        validate_pack(pack, build_registry())


def test_unparsed_visible_workbook_with_hidden_sheet_cannot_report_completion(tmp_path):
    """隐藏辅助表不能掩盖可见表漏读；tmp_path 为实际审核流程的隔离目录。"""
    from run_audit import configure_soffice
    configure_soffice()
    input_root = tmp_path / 'input'
    input_root.mkdir()
    path = input_root / '06信通公司风控矩阵.xlsx'
    book = Workbook()
    book.active.title = '待确认资料'
    book.active.append(['原始资料'])
    book.active.append(['需要人工确认的内容'])
    book.create_sheet('历史参考').sheet_state = 'hidden'
    book.save(path)
    lists = Workbook()
    lists.active.title = '系统控制规则清单'
    lists.active.append(['控制措施编号', '规则名称', '规则内容'])
    lists.active.append(['M1', '权限校验', '禁止越权'])
    duties = lists.create_sheet('岗位内控责任清单')
    duties.append(['控制措施编号', '部门', '岗位名称', '岗位职责'])
    incompatible = lists.create_sheet('不相容岗位清单')
    incompatible.append(['不相容业务角色', '岗位A', '岗位B'])
    lists.save(input_root / '06信通公司三清单.xlsx')
    scope = tmp_path / 'scope.json'
    write_json(scope, {'205H': {'entity_codes': ['205H'], 'businesses': [{'business_code': '06', 'variant_id': 'default', 'required': True}]}})
    result = audit(input_root, tmp_path / 'output', ROOT / '审核器/rulepacks/releases/1.8.0', ROOT / '审核/会计主体清单20260907.xlsx', ROOT, tmp_path / 'runs', scope_file=scope)
    assert result['write_completed'] is False
    assert result['has_check_limits'] is True
    assert result['unparsed_files'] == 1
    assert result['unparsed_file_paths'] == ['06信通公司风控矩阵.xlsx']
    assert sha256_file(tmp_path / 'output' / path.name) == sha256_file(path)
