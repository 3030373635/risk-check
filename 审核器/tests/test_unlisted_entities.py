"""验证清单外主体继续审核，且分组和最终标注可追溯。"""
import json
import shutil
from pathlib import Path

import pytest
from openpyxl import Workbook

from risk_audit import runner
from risk_audit.entity_groups import prepare_entity_groups
from risk_audit.models import Entity, FileRecord
from risk_audit.util import sha256_file
from test_business_incremental import ROOT, create_business_batch, use_business_baseline


@pytest.mark.parametrize('write', [False, True])
@pytest.mark.parametrize('manual', [False, True])
def test_unlisted_entity_is_audited_and_marked(tmp_path, monkeypatch, write, manual):
    """tmp_path/monkeypatch 为测试工具，write/manual 控制副本写入和手动范围。"""
    input_root, roster_path, sources = create_business_batch(tmp_path, [('06', 'default'), ('07', 'default')])
    use_business_baseline(input_root, sources, monkeypatch)
    roster = Workbook()
    roster.active.append(['单位名称', '单位代码', '上级单位', '对应主业单位', '是否存续'])
    roster.active.append(['其他已登记有限公司', 'B001', '', '', '是'])
    roster.save(roster_path)
    hashes = {path: sha256_file(path) for paths in sources.values() for path in paths}
    scope_file = None
    if manual:
        scope_file = tmp_path / 'scope.json'
        scope_file.write_text(json.dumps({'B001': {'entity_codes': ['B001'], 'businesses': []}}))
    result = runner.audit(input_root, tmp_path / 'output', ROOT / '审核器/rulepacks/releases/1.8.0',
                          roster_path, ROOT, tmp_path / 'runs', run_id='unlisted', write=write, scope_file=scope_file)
    entity = next(item for item in result['entity_results'] if item['input_files'])
    assert entity['entity_code'] == ''
    assert entity['entity_registry_status'] == 'not_listed'
    assert entity['entity_registry_message'] == '主体不在会计主体清单中'
    assert entity['entity_name'] == '测试主体甲有限公司'
    assert entity['completed_businesses'] == 2
    assert result['parsed_files'] == 4 and not result['statuses'].get('failed')
    assert result['write_completed'] == write
    assert result['has_check_limits']
    assert result['entity_code_alerts']['files'] == 4
    run_directory = tmp_path / 'runs/unlisted'
    report_directory = Path(result['audit_metadata_dir']) if write else run_directory
    report = (report_directory / '主体代码待确认.md').read_text()
    assert '主体不在会计主体清单中' in report and '已继续审核' in report
    assert '| 测试主体甲有限公司 | 待确认 | 主体不在会计主体清单中 |' in (report_directory / '本次审核范围.md').read_text()
    findings = json.loads((run_directory / 'findings.json').read_text())
    statuses = json.loads((run_directory / 'statuses.json').read_text())
    completeness = [item for item in statuses if item['rule_id'] == 'documents.completeness']
    assert len(completeness) == 4 and all(item['status'] == 'executed' for item in completeness), '清单外主体也要执行资料完整性检查'
    assert not any(item['evidence'].get('missing_material') == '主体差异说明' for item in findings)
    assert all(sha256_file(path) == before for path, before in hashes.items())
    for path in hashes:
        assert (tmp_path / 'output' / path.relative_to(input_root)).exists() == write
    if write:
        assert '主体不在会计主体清单中' in (Path(result['audit_metadata_dir']) / '主体代码待确认.md').read_text()


def test_different_unlisted_entities_remain_separate(tmp_path, monkeypatch):
    """tmp_path/monkeypatch 为测试工具；不同清单外主体不能共享矩阵与三清单。"""
    input_root, roster_path, sources = create_business_batch(tmp_path, [('06', 'default')])
    use_business_baseline(input_root, sources, monkeypatch)
    roster = Workbook()
    roster.active.append(['单位名称', '单位代码', '上级单位', '对应主业单位', '是否存续'])
    roster.active.append(['测试主体甲有限公司', 'A001', '', '', '是'])
    roster.save(roster_path)
    for index, name in enumerate(['清单外主体甲有限公司', '清单外主体乙有限公司']):
        directory = input_root / name
        directory.mkdir()
        source = sources[('06', 'default')][index]
        shutil.copyfile(source, directory / source.name.replace('测试主体甲有限公司', name))
    observed = []
    original = runner.run_engine

    def observe_engine(pack, registry, files, *args):
        """pack/registry/files 为审核上下文，args 为名册及运行参数；验证每次只处理一个主体。"""
        codes = {file.entity_code for file in files}
        assert len(codes) == 1 and None not in codes
        assert pack['submission_scope']['entity_codes'] == list(codes)
        observed.extend(codes)
        return original(pack, registry, files, *args)

    monkeypatch.setattr(runner, 'run_engine', observe_engine)
    result = runner.audit(input_root, tmp_path / 'output', ROOT / '审核器/rulepacks/releases/1.8.0',
                          roster_path, ROOT, tmp_path / 'runs', run_id='separate', write=False)
    assert len(set(observed)) == 3
    assert {item['entity_name'] for item in result['entity_results'] if item['entity_registry_status'] == 'not_listed'} == {'清单外主体甲有限公司', '清单外主体乙有限公司'}
    findings = json.loads((tmp_path / 'runs/separate/findings.json').read_text())
    assert len([item for item in findings if item['check_id'] == 'required_matrix_and_lists']) == 2


@pytest.mark.parametrize('candidate,missing,status', [
    (None, False, 'not_listed'), (None, True, 'missing_code'), ('UNKNOWN', False, 'code_not_listed')])
def test_registry_status_retains_original_code_evidence(tmp_path, candidate, missing, status):
    """tmp_path 为目录，candidate/missing/status 为候选代码、名册缺代码状态和期望分类。"""
    name = '样例服务有限公司'
    path = tmp_path / name / '06风控矩阵.xlsx'
    file = FileRecord(path, path.relative_to(tmp_path), 'hash', 'xlsx', candidate, [], False,
                      '06', 'default', 'matrix', parse_errors=['无法确定会计主体', '独立读取错误'])
    roster = [Entity(None, name, source_row=12)] if missing else []
    entities = {'A001': Entity('A001', '已登记主体有限公司')}
    groups = prepare_entity_groups([file], tmp_path, entities, roster)
    identity = groups[file.entity_code]
    assert identity['entity_registry_status'] == status
    assert identity['candidate_code'] == candidate and identity['entity_code'] == ''
    assert file.parse_errors == ['独立读取错误']
    assert set(entities) == {'A001'}, '本次分组不能修改正式名册'


def test_unconfirmed_short_name_is_not_reported_as_missing_from_roster(tmp_path):
    """tmp_path 为目录；简称无法匹配时必须提示待确认，不能断言清单缺少该单位。"""
    path = tmp_path / '样例公司' / '06风控矩阵.xlsx'
    file = FileRecord(path, path.relative_to(tmp_path), 'hash', 'xlsx', None, [], False, '06', 'default', 'matrix')
    identity = next(iter(prepare_entity_groups([file], tmp_path, {}).values()))
    assert identity['entity_registry_status'] == 'unconfirmed'
    assert '需确认名称或简称' in identity['entity_registry_message']


@pytest.mark.parametrize('subdirectory', ['', 'A001'])
@pytest.mark.parametrize('suffix', ['有限公司', '电力调度控制中心'])
def test_shared_directory_uses_separate_names_for_unlisted_entities(tmp_path, subdirectory, suffix):
    """tmp_path 为输入根，subdirectory 为共享目录；不同主体按名称隔离，同名矩阵与三清单关联。"""
    files = []
    for name, material in [('样例主体甲' + suffix, 'matrix'), ('样例主体甲' + suffix, 'three_lists'),
                           ('样例主体乙' + suffix, 'matrix')]:
        filename = f'06{material}-{name}.xlsx'
        relative = Path(subdirectory) / filename
        files.append(FileRecord(tmp_path / relative, relative, 'hash', 'xlsx',
                                None, [], False, '06', 'default', material))
    groups = prepare_entity_groups(files, tmp_path, {})
    assert len(groups) == 2
    assert files[0].entity_code == files[1].entity_code != files[2].entity_code


@pytest.mark.parametrize('confirmed', [False, True])
@pytest.mark.parametrize('named_file', [False, True])
@pytest.mark.parametrize('contains_parent_name', [False, True])
def test_unlisted_child_does_not_inherit_registered_parent(tmp_path, confirmed, named_file, contains_parent_name):
    """tmp_path 为目录，confirmed 控制扫描器，named_file 控制文件全称，contains_parent_name 控制名称包含父公司。"""
    from risk_audit.inventory import scan_package
    from risk_audit.inventory_v180 import scan_package_v180
    root = tmp_path / '已登记父公司有限公司'
    name = (root.name if contains_parent_name else '') + '清单外子公司有限公司'
    directory = root / name / '06 业务'
    directory.mkdir(parents=True)
    filename = f'06风控矩阵-{name}.xlsx' if named_file else '06风控矩阵.xlsx'
    book = Workbook()
    book.active.append(['控制措施编号', '控制措施', '责任主体'])
    book.save(directory / filename)
    scan = scan_package_v180 if confirmed else scan_package
    files = scan(root, {'P001': Entity('P001', root.name)}, {})
    assert files[0].entity_code is None
    assert not files[0].entity_conflict
    assert f'报送名称:{name}' in files[0].entity_evidence


@pytest.mark.parametrize('confirmed', [False, True])
def test_package_scanner_skips_dot_prefixed_files(tmp_path, confirmed):
    """tmp_path 为输入目录，confirmed 控制扫描器版本；点开头文件不得进入审核清单。"""
    from risk_audit.inventory import scan_package
    from risk_audit.inventory_v180 import scan_package_v180
    root = tmp_path / '测试单位有限公司'
    root.mkdir()
    book = Workbook()
    book.active.append(['控制措施编号', '控制措施', '责任主体'])
    book.save(root / '06风控矩阵.xlsx')
    (root / '.~06风控矩阵.xlsx').write_bytes(b'office-temporary-file')
    scan = scan_package_v180 if confirmed else scan_package

    files = scan(root, {'E001': Entity('E001', root.name)}, {})

    assert [file.relative_path.name for file in files] == ['06风控矩阵.xlsx']


def test_missing_code_parent_does_not_replace_explicit_child_names(tmp_path):
    """tmp_path 为输入目录；缺代码母公司只能匹配本组明确名称，不能覆盖子公司。"""
    files = []
    for name in ['样例主体甲有限公司', '样例主体乙有限公司']:
        relative = Path('缺代码母公司有限公司') / f'06风控矩阵-{name}.xlsx'
        files.append(FileRecord(tmp_path / relative, relative, 'hash', 'xlsx', None, [], False,
                                '06', 'default', 'matrix'))
    groups = prepare_entity_groups(files, tmp_path, {}, [Entity(None, '缺代码母公司有限公司', source_row=12)])
    assert len(groups) == 2
    assert {item['entity_name'] for item in groups.values()} == {'样例主体甲有限公司', '样例主体乙有限公司'}
    assert all(item['entity_registry_status'] == 'not_listed' for item in groups.values())


def test_single_unit_attachment_cannot_replace_explicit_unlisted_names(tmp_path):
    """tmp_path 为测试目录；单单位应用清单不能把清单外甲和丙的材料强行归属乙。"""
    from risk_audit.inventory_v180 import scan_package_v180
    from risk_audit.submission_scope import resolve_submission_scopes
    root = tmp_path / 'input'
    root.mkdir()
    for name, material in [('样例主体甲有限公司', '风控矩阵'), ('样例主体丙有限公司', '三清单')]:
        book = Workbook()
        book.active.append(['控制措施编号', '控制措施', '责任主体'])
        book.save(root / f'06{material}-{name}.xlsx')
    attachment = Workbook()
    attachment.active.append(['单位', '设备（资产）管理'])
    attachment.active.append(['样例主体乙有限公司', '适用'])
    attachment.save(root / '附件1应用清单.xlsx')
    entities = {'B001': Entity('B001', '样例主体乙有限公司')}
    files = scan_package_v180(root, entities, {})
    resolve_submission_scopes(root, files, entities, {}, '第一批')
    assert all(file.entity_code is None for file in files)
    groups = prepare_entity_groups(files, root, entities)
    assert len(groups) == 2 and all(item['entity_registry_status'] == 'not_listed' for item in groups.values())


@pytest.mark.parametrize('confirmed', [False, True])
def test_unlisted_explanation_retains_conflicting_body_evidence(tmp_path, confirmed):
    """tmp_path 为输入目录，confirmed 控制扫描器版本；清单外文件名不能屏蔽正文主体冲突。"""
    from docx import Document
    from risk_audit.inventory import scan_package
    from risk_audit.inventory_v180 import scan_package_v180
    document = Document()
    document.add_paragraph('已登记主体甲有限公司')
    document.add_paragraph('已登记主体乙有限公司')
    document.save(tmp_path / '清单外主体有限公司-主体说明.docx')
    entities = {'A001': Entity('A001', '已登记主体甲有限公司'),
                'B001': Entity('B001', '已登记主体乙有限公司')}
    scan = scan_package_v180 if confirmed else scan_package
    file = scan(tmp_path, entities, {})[0]
    assert file.entity_conflict and file.entity_code is None
    assert '标准全称:已登记主体甲有限公司' in file.entity_evidence
    assert '标准全称:已登记主体乙有限公司' in file.entity_evidence


@pytest.mark.parametrize('confirmed', [False, True])
@pytest.mark.parametrize('filename', [
    '07风控矩阵-计量分公司.xlsx',
    '计量分公司-附件1：07风控矩阵-省公司（审定）8.13.xlsx',
    '07风控矩阵.xlsx',
])
def test_short_name_uses_registered_unit_directory(tmp_path, confirmed, filename):
    """tmp_path 为隔离目录，confirmed 为扫描版本，filename 为实际报送文件命名。"""
    from risk_audit.inventory import scan_package
    from risk_audit.inventory_v180 import scan_package_v180
    name = '湖南朗晟电力产业发展有限公司计量分公司'
    directory = tmp_path / name / '07物资（服务）采购与实施-计量分公司（负责人：某人）'
    directory.mkdir(parents=True)
    book = Workbook()
    book.save(directory / filename)
    scan = scan_package_v180 if confirmed else scan_package
    file = scan(tmp_path, {'209Z': Entity('209Z', name)}, {})[0]
    assert file.entity_code == '209Z'
    assert not file.entity_conflict


@pytest.mark.parametrize('filename', ['01风控矩阵-供服中心.xlsx', '01风控矩阵.xlsx'])
def test_abbreviated_center_uses_full_directory_with_parentheses(tmp_path, filename):
    """tmp_path 为隔离目录，filename 为有短称或无短称材料；括号中的名称不另建主体。"""
    from risk_audit.inventory_v180 import scan_package_v180
    name = '国网湖南省电力有限公司供电服务中心（计量中心）本部'
    directory = tmp_path / name / '01.营销售电'
    directory.mkdir(parents=True)
    book = Workbook()
    book.save(directory / filename)
    file = scan_package_v180(tmp_path, {'20JQ': Entity('20JQ', name)}, {})[0]
    assert file.entity_code == '20JQ'
    assert not file.entity_conflict


def test_incompatible_center_short_name_does_not_use_directory(tmp_path):
    """tmp_path 为隔离目录；目录全称不能覆盖不相容的报送中心名称。"""
    from risk_audit.inventory_v180 import scan_package_v180
    name = '国网湖南省电力有限公司供电服务中心（计量中心）本部'
    directory = tmp_path / name / '01.营销售电'
    directory.mkdir(parents=True)
    book = Workbook()
    book.save(directory / '01风控矩阵-其他服务中心.xlsx')
    file = scan_package_v180(tmp_path, {'20JQ': Entity('20JQ', name)}, {})[0]
    assert file.entity_code is None


def test_directory_full_name_evidence_cannot_be_ignored(tmp_path):
    """tmp_path 为隔离目录；中间目录带其他正式主体全称时必须保留归属冲突。"""
    from risk_audit.inventory_v180 import scan_package_v180
    name_a, name_b = '已登记主体甲有限公司', '已登记主体乙有限公司'
    directory = tmp_path / name_a / (name_b + '报送资料') / '07业务'
    directory.mkdir(parents=True)
    book = Workbook()
    book.save(directory / '07风控矩阵.xlsx')
    file = scan_package_v180(tmp_path, {'A001': Entity('A001', name_a), 'B001': Entity('B001', name_b)}, {})[0]
    assert file.entity_code is None
    assert file.entity_conflict
