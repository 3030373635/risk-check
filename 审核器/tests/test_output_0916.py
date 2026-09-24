"""新版排序和输出的完整业务行为，验证格式、人工意见和幂等。"""
from importlib import import_module
import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.comments import Comment
from openpyxl.styles import Font

from risk_audit.models import Entity, FileRecord, Finding, ParsedSheet
from risk_audit.readers.confirmed_v180 import parse_workbook_v180
from risk_audit.readers.excel import parse_files
from risk_audit.configuration.loader import load_pack
from risk_audit.engine import run_engine
from risk_audit.runner import _has_failures, _remove_legacy_opinion_outputs
from risk_audit.util import sha256_file
from risk_audit.writer import write_outputs

ROOT = Path(__file__).resolve().parents[2]


def module():
    """获取新版输出模块；无参数。"""
    name = 'risk_audit.output_0916'
    from importlib.util import find_spec
    assert find_spec(name), '新版排序输出尚未实现'
    return import_module(name)


def test_legacy_opinion_cleanup_does_not_follow_audit_directory_symlink(tmp_path):
    """tmp_path 为隔离目录；_risk_audit 父链接只删除链接，不得删除外部旧报告。"""
    external = tmp_path / 'external'
    external.mkdir()
    sentinel = external / '意见反馈明细.xlsx'
    sentinel.write_text('外部文件', encoding='utf-8')
    output = tmp_path / 'output'
    output.mkdir()
    audit_link = output / '_risk_audit'
    audit_link.symlink_to(external, target_is_directory=True)

    _remove_legacy_opinion_outputs(output)

    assert not audit_link.exists()
    assert not audit_link.is_symlink()
    assert sentinel.read_text(encoding='utf-8') == '外部文件'


def file_record(path):
    """解析样例材料；path 为测试工作簿。"""
    value = FileRecord(path, Path(path.name), sha256_file(path), 'xlsx', '205H', ['确认简称:信通公司'], False, '09', 'default', 'three_lists')
    value._parser_policy = {'version': 3}
    aliases = load_pack(ROOT / '审核器/rulepacks/releases/1.8.0')['field_aliases']
    value.sheets = parse_workbook_v180(value, path, aliases)
    return value, aliases


def test_sort_moves_only_details_with_merged_values_format_and_human_opinions(tmp_path):
    """tmp_path 为隔离目录；自然数字顺序、合并展开、富文本及人工批注一起随记录移动。"""
    path = tmp_path / '09信通公司三清单.xlsx'
    book = Workbook(); ws = book.active; ws.title = '岗位内控责任清单'
    ws.append(['控制措施编号', '部门', '岗位名称', '人员姓名', '岗位职责', '角色', '审核意见'])
    for number in (12, 12, 2):
        ws.append([f'薪酬业务-{number}.控制点-控制措施05', '财务部', '核算专责', '张三', '对资料真实性负主体责任', '经办'])
    ws.merge_cells('A2:A3'); ws['G2'] = '人工意见'; ws['E2'] = CellRichText([TextBlock(InlineFont(strike=True), '旧文'), '对资料真实性负主体责任'])
    ws['B2'].font = Font(color='FF0000'); ws['B2'].comment = Comment('人工批注', '审核员')
    ws['A6'] = '填表说明：保留'; book.save(path)
    before = sha256_file(path); file, aliases = file_record(path)
    module().sort_duties(file, {}, tmp_path / 'work', aliases)
    output = load_workbook(file._preprocessed_path, rich_text=True)
    assert [output.active.cell(row, 1).value for row in (2, 3, 4)] == ['薪酬业务-2.控制点-控制措施05', '薪酬业务-12.控制点-控制措施05', '薪酬业务-12.控制点-控制措施05']
    assert output.active['G3'].value == '人工意见'
    assert output.active['B3'].comment.text == '人工批注'
    assert output.active['B3'].font.color.rgb == '00FF0000'
    assert output.active['E3'].value[0].font.strike
    assert output.active['A6'].value == '填表说明：保留'
    assert sha256_file(path) == before


def test_sort_moves_leading_data_column_without_header(tmp_path):
    """缺少表头的最左侧部门列也必须随岗位明细整体排序。

    tmp_path 为隔离目录；部门列有业务数据但表头为空，参数无其他要求。
    """
    path = tmp_path / '09信通公司缺部门表头三清单.xlsx'
    book = Workbook(); sheet = book.active; sheet.title = '岗位内控责任清单'
    sheet.append([None, '岗位名称', '人员姓名', '岗位职责', '控制措施编号', '角色'])
    sheet.append(['财务部', '核算专责', '张三', '复核资料', '业务-12.控制点-控制措施01', '审核'])
    sheet.append(['建设部', '工程专责', '李四', '编制资料', '业务-2.控制点-控制措施01', '经办'])
    book.save(path)
    file, aliases = file_record(path)

    module().sort_duties(file, {}, tmp_path / 'work', aliases)

    output = load_workbook(file._preprocessed_path).active
    assert [[output.cell(row, column).value for column in range(1, 7)] for row in (2, 3)] == [
        ['建设部', '工程专责', '李四', '编制资料', '业务-2.控制点-控制措施01', '经办'],
        ['财务部', '核算专责', '张三', '复核资料', '业务-12.控制点-控制措施01', '审核'],
    ]


def test_sort_uses_vertical_region_bounds_for_column_without_header(tmp_path):
    """纵向共表时，下方业务区域的同列表头不得截断上方岗位排序范围。

    tmp_path 为隔离目录；岗位区域 G 列有数据但无表头，下方矩阵的 G 列有表头。
    """
    path = tmp_path / '09信通公司纵向共表三清单.xlsx'
    book = Workbook(); sheet = book.active; sheet.title = '综合填报'
    sheet.append(['岗位内控责任清单'])
    sheet.append(['控制措施编号', '部门', '岗位名称', '人员姓名', '岗位职责', '角色', None])
    sheet.append(['业务-12.控制点-控制措施01', '财务部', '核算专责', '张三', '复核资料', '审核', '补充-12'])
    sheet.append(['业务-2.控制点-控制措施01', '建设部', '工程专责', '李四', '编制资料', '经办', '补充-2'])
    sheet.cell(10, 1, '风控矩阵')
    for column, value in enumerate(['控制措施编号', '控制措施', '责任主体', None, None, None, '控制载体'], 1):
        sheet.cell(11, column, value)
    sheet.cell(12, 1, 'M1'); sheet.cell(12, 2, '复核'); sheet.cell(12, 3, '财务部'); sheet.cell(12, 7, '清单')
    book.save(path)
    file, aliases = file_record(path)

    module().sort_duties(file, {}, tmp_path / 'work', aliases)

    output = load_workbook(file._preprocessed_path).active
    assert [output.cell(row, 7).value for row in (3, 4)] == ['补充-2', '补充-12']
    assert output['G11'].value == '控制载体'
    assert output['G12'].value == '清单'


def test_sort_keeps_unique_hidden_auxiliary_columns_in_place(tmp_path):
    """唯一命名的隐藏辅助列不属于岗位明细，不得随排序移动。

    tmp_path 为隔离目录；H:I 为普通隐藏辅助区，表头不与业务字段重复。
    """
    path = tmp_path / '09信通公司隐藏辅助列三清单.xlsx'
    book = Workbook(); sheet = book.active; sheet.title = '岗位内控责任清单'
    sheet.append(['控制措施编号', '部门', '岗位名称', '人员姓名', '岗位职责', '角色', None, '辅助编号', '辅助结果'])
    sheet.append(['业务-12.控制点-控制措施01', '财务部', '核算专责', '张三', '复核资料', '审核', None, '辅助-12', '=COUNTIF($A:$A,H2)'])
    sheet.append(['业务-2.控制点-控制措施01', '建设部', '工程专责', '李四', '编制资料', '经办', None, '辅助-2', '=COUNTIF($A:$A,H3)'])
    sheet.column_dimensions['H'].hidden = True; sheet.column_dimensions['I'].hidden = True
    book.save(path)
    file, aliases = file_record(path)

    module().sort_duties(file, {}, tmp_path / 'work', aliases)

    output = load_workbook(file._preprocessed_path, data_only=False).active
    assert [output.cell(row, 1).value for row in (2, 3)] == [
        '业务-2.控制点-控制措施01',
        '业务-12.控制点-控制措施01',
    ]
    assert [(output[f'H{row}'].value, output[f'I{row}'].value) for row in (2, 3)] == [
        ('辅助-12', '=COUNTIF($A:$A,H2)'),
        ('辅助-2', '=COUNTIF($A:$A,H3)'),
    ]


def test_sort_moves_only_current_region_audit_column(tmp_path):
    """并排区域的审核意见必须各自跟随所属记录，不得跨区移动。

    tmp_path 为隔离目录；左侧为岗位清单，右侧为矩阵，两个区域都有审核意见。
    """
    path = tmp_path / '09信通公司并排区域三清单.xlsx'
    book = Workbook(); sheet = book.active; sheet.title = '综合填报'
    sheet.cell(1, 1, '岗位内控责任清单'); sheet.cell(1, 9, '风控矩阵')
    for column, value in enumerate(['控制措施编号', '部门', '岗位名称', '人员姓名', '岗位职责', '角色', '审核意见'], 1):
        sheet.cell(2, column, value)
    for column, value in enumerate(['控制措施编号', '控制措施', '责任主体', None, '审核意见'], 9):
        sheet.cell(2, column, value)
    for column, value in enumerate(['业务-12.控制点-控制措施01', '财务部', '核算专责', '张三', '复核资料', '审核', '岗位意见-12'], 1):
        sheet.cell(3, column, value)
    for column, value in enumerate(['业务-2.控制点-控制措施01', '建设部', '工程专责', '李四', '编制资料', '经办', '岗位意见-2'], 1):
        sheet.cell(4, column, value)
    for row, values in ((3, ['M1', '复核矩阵资料', '财务部', None, '矩阵意见-1']),
                        (4, ['M2', '审批矩阵资料', '建设部', None, '矩阵意见-2'])):
        for column, value in enumerate(values, 9):
            sheet.cell(row, column, value)
    book.save(path)
    file, aliases = file_record(path)

    module().sort_duties(file, {}, tmp_path / 'work', aliases)

    output = load_workbook(file._preprocessed_path).active
    assert [output.cell(row, 7).value for row in (3, 4)] == ['岗位意见-2', '岗位意见-12']
    assert [output.cell(row, 9).value for row in (3, 4)] == ['M1', 'M2']
    assert [output.cell(row, 13).value for row in (3, 4)] == ['矩阵意见-1', '矩阵意见-2']


def test_sort_does_not_fallback_to_vertical_region_audit_column(tmp_path):
    """当前岗位区域没有审核列时，不得回退使用下方矩阵的审核列。

    tmp_path 为隔离目录；H:M 为隐藏辅助区，下方矩阵在 M 列设有审核意见。
    """
    path = tmp_path / '09信通公司纵向审核列三清单.xlsx'
    book = Workbook(); sheet = book.active; sheet.title = '综合填报'
    sheet.append(['岗位内控责任清单'])
    sheet.append(['控制措施编号', '部门', '岗位名称', '人员姓名', '岗位职责', '角色', None,
                  '辅助编号', None, None, None, None, '辅助结果'])
    sheet.append(['业务-12.控制点-控制措施01', '财务部', '核算专责', '张三', '复核资料', '审核', None,
                  '辅助-12', None, None, None, None, '结果-12'])
    sheet.append(['业务-2.控制点-控制措施01', '建设部', '工程专责', '李四', '编制资料', '经办', None,
                  '辅助-2', None, None, None, None, '结果-2'])
    sheet.cell(10, 1, '风控矩阵')
    sheet.cell(11, 1, '控制措施编号'); sheet.cell(11, 2, '控制措施'); sheet.cell(11, 3, '责任主体')
    sheet.cell(11, 13, '审核意见')
    sheet.cell(12, 1, 'M1'); sheet.cell(12, 2, '复核'); sheet.cell(12, 3, '财务部'); sheet.cell(12, 13, '矩阵意见')
    sheet.column_dimensions.group('H', 'M', hidden=True)
    book.save(path)
    file, aliases = file_record(path)

    module().sort_duties(file, {}, tmp_path / 'work', aliases)

    output = load_workbook(file._preprocessed_path).active
    assert [output.cell(row, 1).value for row in (3, 4)] == [
        '业务-2.控制点-控制措施01',
        '业务-12.控制点-控制措施01',
    ]
    assert [(output[f'H{row}'].value, output[f'M{row}'].value) for row in (3, 4)] == [
        ('辅助-12', '结果-12'),
        ('辅助-2', '结果-2'),
    ]


def test_sort_respects_side_by_side_region_with_empty_headers(tmp_path):
    """并排区域即使没有可识别表头，也不得被岗位排序跨区移动。

    tmp_path 为隔离目录；右侧矩阵只有明确区域标题和原始数据，业务表头全部为空。
    """
    path = tmp_path / '09信通公司空表头并排区域三清单.xlsx'
    book = Workbook(); sheet = book.active; sheet.title = '综合填报'
    sheet.cell(1, 1, '岗位内控责任清单'); sheet.cell(1, 9, '风控矩阵')
    for column, value in enumerate(['控制措施编号', '部门', '岗位名称', '人员姓名', '岗位职责', '角色'], 1):
        sheet.cell(2, column, value)
    for column, value in enumerate(['业务-12.控制点-控制措施01', '财务部', '核算专责', '张三', '复核资料', '审核'], 1):
        sheet.cell(3, column, value)
    for column, value in enumerate(['业务-2.控制点-控制措施01', '建设部', '工程专责', '李四', '编制资料', '经办'], 1):
        sheet.cell(4, column, value)
    for row, values in ((3, ['矩阵原始-1', '矩阵原始-2', '矩阵原始-3', '矩阵原始-4']),
                        (4, ['矩阵原始-5', '矩阵原始-6', '矩阵原始-7', '矩阵原始-8'])):
        for column, value in enumerate(values, 9):
            sheet.cell(row, column, value)
    book.save(path)
    file, aliases = file_record(path)

    module().sort_duties(file, {}, tmp_path / 'work', aliases)

    output = load_workbook(file._preprocessed_path).active
    assert [output.cell(row, 1).value for row in (3, 4)] == [
        '业务-2.控制点-控制措施01',
        '业务-12.控制点-控制措施01',
    ]
    assert [[output.cell(row, column).value for column in range(9, 13)] for row in (3, 4)] == [
        ['矩阵原始-1', '矩阵原始-2', '矩阵原始-3', '矩阵原始-4'],
        ['矩阵原始-5', '矩阵原始-6', '矩阵原始-7', '矩阵原始-8'],
    ]


def test_sort_clips_oversized_merge_before_styled_blank_tail(tmp_path):
    """验证岗位排序忽略带样式空白尾行中的超长合并范围。

    tmp_path 为隔离目录；这防止 XLS 转换产生带样式空行后，
    Excel/WPS 将末组合并到最大行时中断审核。
    """
    path = tmp_path / '09怀化三清单.xlsx'
    book = Workbook(); ws = book.active; ws.title = '岗位职责清单（怀化）'
    ws.append(['岗位内控责任清单'])
    ws.append(['部门', '岗位名称', '人员姓名', '岗位职责编号', '角色', '岗位职责',
               '控制措施编号', '控制载体', '责任主体', '控制措施内容'])
    ws.append(['财务部', '核算专责', '张三', 1, '经办', '对资料真实性负主体责任',
               '业务-12.控制点-控制措施01', '资料', '财务部门-核算专责', '核对资料'])
    ws.append(['综合服务中心', '离退休事务专责', '李四', 1, '经办', '对资料完整性负主体责任',
               '业务-2.控制点-控制措施01', '清单', '综合部门-事务专责', '复核清单'])
    ws.append([None, None, '李四', 2, '经办', '对资料准确性负主体责任',
               '业务-1.控制点-控制措施01', '清单', '综合部门-事务专责', '审核清单'])
    ws.merge_cells('A4:A5'); ws.merge_cells('B4:B5')
    # LibreOffice 会为 XLS 尾部样式保留空单元格，物理行尾不等于有效内容行尾。
    ws['A1000'].font = Font(bold=True)
    book.save(path)
    file, aliases = file_record(path)

    # 模拟实际报送文件：有效明细只到第 5 行，末组合并却延伸到 Excel 最大行。
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    root = etree.fromstring(parts['xl/worksheets/sheet1.xml'])
    namespaces = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    for merged in root.xpath('//m:mergeCell', namespaces=namespaces):
        if merged.get('ref') in {'A4:A5', 'B4:B5'}:
            merged.set('ref', merged.get('ref').replace(':A5', ':A1048576').replace(':B5', ':B1048576'))
    parts['xl/worksheets/sheet1.xml'] = etree.tostring(root)
    with ZipFile(path, 'w') as archive:
        for name, data in parts.items():
            archive.writestr(name, data)

    module().sort_duties(file, {}, tmp_path / 'work', aliases)

    output = load_workbook(file._preprocessed_path).active
    assert [output.cell(row, 7).value for row in range(3, 6)] == [
        '业务-1.控制点-控制措施01',
        '业务-2.控制点-控制措施01',
        '业务-12.控制点-控制措施01',
    ]
    assert [output.cell(row, 1).value for row in range(3, 6)] == [
        '综合服务中心', '综合服务中心', '财务部',
    ]
    assert list(output.merged_cells.ranges) == []


def test_sort_rejects_oversized_merge_crossing_real_content(tmp_path):
    """验证岗位排序拒绝跨越真实内容行的超长合并范围。

    tmp_path 为隔离目录；明细后的填表说明是其他有效区域，不得按空白尾行裁剪。
    """
    path = tmp_path / '09怀化跨区合并三清单.xlsx'
    book = Workbook(); ws = book.active; ws.title = '岗位职责清单'
    ws.append(['控制措施编号', '部门', '岗位名称', '人员姓名', '岗位职责', '角色'])
    for number in (12, 2, 1):
        ws.append([f'薪酬业务-{number}.控制点-控制措施01', '财务部', '核算专责', '张三',
                   '对资料真实性负主体责任', '经办'])
    ws['A5'] = '填表说明：保留其他有效区域'
    ws.merge_cells('B3:B4'); book.save(path)
    file, aliases = file_record(path)

    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    root = etree.fromstring(parts['xl/worksheets/sheet1.xml'])
    namespaces = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    root.xpath('//m:mergeCell[@ref="B3:B4"]', namespaces=namespaces)[0].set('ref', 'B3:B1048576')
    parts['xl/worksheets/sheet1.xml'] = etree.tostring(root)
    with ZipFile(path, 'w') as archive:
        for name, data in parts.items():
            archive.writestr(name, data)

    with pytest.raises(ValueError, match='岗位排序合并范围跨越有效明细或其他区域'):
        module().sort_duties(file, {}, tmp_path / 'work', aliases)


def test_sort_moves_merged_department_when_fill_department_header_conflicts(tmp_path):
    """tmp_path 为隔离目录；负责填表部门造成字段冲突时，合并部门仍须随整行排序。"""
    path = tmp_path / '03株洲公司三清单.xlsx'
    book = Workbook(); ws = book.active; ws.title = '岗位职责清单'
    ws.append(['岗位内控责任清单'])
    ws.append(['部门', '岗位名称', '岗位人员姓名', '岗位职责编号', '角色', '岗位职责',
               '控制措施编号', '控制载体', '责任主体', '控制措施内容', '负责填表部门'])
    ws.append(['供电公司-财务资产部', '财务专责', '向思源', 1, '审核', '对资料准确性负审核责任。',
               '电网基建业务-2.控制点-控制措施01', '资料', '财务部门-财务专责', '复核资料', '财务部'])
    ws.append([None, '财务主任', '孙岑', 1, '审批', '对资料准确性负审批责任。',
               '电网基建业务-3.控制点-控制措施01', '资料', '财务部门-财务主任', '审批资料', '财务部'])
    ws.append(['供电公司-发展策划部', '发展策划部主任', '朱育兰', 1, '审批', '对投资决策合规性负审批责任。',
               '电网基建业务-1.控制点-控制措施02', '投资清单', '发展部门-发展策划部主任', '审批投资决策', '发展部'])
    ws.append([None, '项目前期管理专责', '曾宪敏', 1, '经办', '对项目前期工作及时性负主体责任。',
               '电网基建业务-4.控制点-控制措施01', '项目前期资料', '发展部门-项目前期管理专责', '办理前期工作', '发展部'])
    ws.merge_cells('A3:A4'); ws.merge_cells('A5:A6'); book.save(path)
    file, aliases = file_record(path)

    module().sort_duties(file, {}, tmp_path / 'work', aliases)

    output = load_workbook(file._preprocessed_path).active
    assert [output.cell(row, 7).value for row in range(3, 7)] == [
        '电网基建业务-1.控制点-控制措施02',
        '电网基建业务-2.控制点-控制措施01',
        '电网基建业务-3.控制点-控制措施01',
        '电网基建业务-4.控制点-控制措施01',
    ]
    assert [output.cell(row, 1).value for row in range(3, 7)] == [
        '供电公司-发展策划部',
        '供电公司-财务资产部',
        '供电公司-财务资产部',
        '供电公司-发展策划部',
    ]
    assert output['C3'].value == '朱育兰'


def test_sort_keeps_hidden_rejected_auxiliary_block_in_place(tmp_path):
    """tmp_path 为隔离目录；隐藏且未选中的重复字段辅助区不随岗位明细排序。"""
    path = tmp_path / '10澄县三清单.xlsx'
    book = Workbook(); ws = book.active; ws.title = '岗位职责清单'
    ws.append(['控制措施编号', '部门', '岗位名称', '人员姓名', '岗位职责', '角色', None, '补充说明', '控制措施编号', '核对结果'])
    measures = ('业务-12.控制点-控制措施01', '业务-2.控制点-控制措施01',
                '业务-1.控制点-控制措施01', '业务-3.控制点-控制措施01')
    for measure, note in zip(measures, ('说明-12', '说明-2', '说明-1', '说明-3')):
        ws.append([measure, '财务部', '核算专责', '张三', '对资料真实性负主体责任', '经办', None, note])
    # 隐藏辅助区的纵向表头跨过首条业务明细，但重复字段未被解析器选为业务列。
    ws.merge_cells('I1:I2')
    ws.column_dimensions['I'].hidden = True; ws.column_dimensions['J'].hidden = True
    ws['I3'] = '辅助-12'; ws['J3'] = '=COUNTIF($A:$A,I3)'
    ws['I4'] = '辅助-2'; ws['J4'] = '=COUNTIF($A:$A,I4)'
    book.save(path)
    file, aliases = file_record(path)

    module().sort_duties(file, {}, tmp_path / 'work', aliases)

    output = load_workbook(file._preprocessed_path, data_only=False).active
    assert [output.cell(row, 1).value for row in (2, 3, 4, 5)] == [
        '业务-1.控制点-控制措施01',
        '业务-2.控制点-控制措施01',
        '业务-3.控制点-控制措施01',
        '业务-12.控制点-控制措施01',
    ]
    assert [output.cell(row, 8).value for row in (2, 3, 4, 5)] == ['说明-1', '说明-2', '说明-3', '说明-12']
    assert {str(cell_range) for cell_range in output.merged_cells.ranges} == {'I1:I2'}
    assert [(output[f'I{row}'].value, output[f'J{row}'].value) for row in (3, 4)] == [
        ('辅助-12', '=COUNTIF($A:$A,I3)'),
        ('辅助-2', '=COUNTIF($A:$A,I4)'),
    ]


def test_sort_keeps_correct_merged_order_unchanged_even_when_baseline_differs(tmp_path):
    """tmp_path 为隔离目录；数字顺序正确时不展开合并、不生成副本，且不服从旧基准顺序。"""
    path = tmp_path / '09信通公司三清单.xlsx'
    book = Workbook(); ws = book.active; ws.title = '岗位内控责任清单'
    ws.append(['控制措施编号', '部门', '岗位名称', '人员姓名', '岗位职责', '角色'])
    for number, department in ((2, '财务部'), (2, '建设部'), (12, '办公室')):
        ws.append([f'薪酬业务-{number}.控制点-控制措施05', department, '核算专责', '张三', '对资料真实性负主体责任', '经办'])
    ws.merge_cells('A2:A3'); ws.merge_cells('C2:C3')
    book.save(path)
    before = sha256_file(path); file, aliases = file_record(path)
    baselines = {('09', 'default'): {'measure_order': [
        '薪酬业务-12.控制点-控制措施05', '薪酬业务-2.控制点-控制措施05',
    ]}}

    module().sort_duties(file, baselines, tmp_path / 'work', aliases)

    assert not hasattr(file, '_preprocessed_path')
    assert sha256_file(path) == before
    assert {str(value) for value in load_workbook(path).active.merged_cells.ranges} == {'A2:A3', 'C2:C3'}


@pytest.mark.parametrize('master_ref', ['G2', 'G5'])
def test_sort_expands_shared_formulas_including_stationary_group_members(tmp_path, master_ref):
    """tmp_path 为隔离目录，master_ref 为共享主公式坐标；排序保留整组公式及绝对引用。"""
    path = tmp_path / '09信通公司三清单.xlsx'
    book = Workbook(); ws = book.active; ws.title = '岗位内控责任清单'
    ws.append(['控制措施编号', '部门', '岗位名称', '人员姓名', '岗位职责', '角色', '拼接'])
    for measure in ('M1', 'M12', 'M2'):
        ws.append([measure, '财务部', '核算专责', '张三', '对资料真实性负主体责任', '经办'])
    for row in range(2, 6):
        ws.cell(row, 7, f'=C{row}&$D{row}&E$2&$F$2')
    ws['B3'].comment = Comment('随记录移动', '审核员')
    book.save(path)
    # 写入与实际文件相同的共享公式结构，主公式可在有效明细之外。
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    root = etree.fromstring(parts['xl/worksheets/sheet1.xml'])
    for cell in root.xpath('//m:c[m:f]', namespaces=ns):
        formula = cell.find('m:f', ns)
        formula.set('t', 'shared'); formula.set('si', '4')
        if cell.get('r') == master_ref:
            formula.set('ref', 'G2:G5')
        else:
            formula.text = None
    parts['xl/worksheets/sheet1.xml'] = etree.tostring(root)
    with ZipFile(path, 'w') as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    before = sha256_file(path); file, aliases = file_record(path)
    module().sort_duties(file, {}, tmp_path / 'work', aliases)
    output_path = Path(file._preprocessed_path)
    output = load_workbook(output_path).active
    assert [output.cell(row, 1).value for row in (2, 3, 4)] == ['M1', 'M2', 'M12']
    assert [output.cell(row, 7).value for row in (2, 3, 4, 5)] == [
        '=C2&$D2&E$2&$F$2', '=C3&$D3&E$2&$F$2',
        '=C4&$D4&E$2&$F$2', '=C5&$D5&E$2&$F$2',
    ]
    assert output['B4'].comment.text == '随记录移动'
    with ZipFile(output_path) as archive:
        root = etree.fromstring(archive.read('xl/worksheets/sheet1.xml'))
        assert not root.xpath('//m:f[@t="shared"]', namespaces=ns)
    assert sha256_file(path) == before
    sorted_hash = sha256_file(output_path)
    module().sort_duties(file, {}, tmp_path / 'work', aliases)
    assert sha256_file(output_path) == sorted_hash


def test_sort_preserves_formula_when_relative_translation_would_leave_sheet(tmp_path):
    """tmp_path 为隔离目录；公式相对引用平移越界时保留原公式并完成排序。"""
    path = tmp_path / '10星美三清单.xlsx'
    book = Workbook(); ws = book.active; ws.title = '岗位职责清单'
    ws.append(['控制措施编号', '部门', '岗位职责编号', '岗位名称', '岗位职责', '角色', '同行公式'])
    ws.append(['员工报账业务-12.控制点-控制措施01', '财务部', 1, '会计', '复核资料', '审核', '=C2&$D2'])
    ws.append(['员工报账业务-13.控制点-控制措施01', '财务部', 2, '会计', '审批资料', '审批', '=C3&$D3'])
    # 该序号公式随记录向上移动时，按单元格位移翻译会产生第 0 行以下的非法引用。
    ws.append(['员工报账业务-1.控制点-控制措施01', '财务部', '=ROW(A1)', '会计', '经办资料', '经办', '=C4&$D4'])
    book.save(path)
    file, aliases = file_record(path)

    module().sort_duties(file, {}, tmp_path / 'work', aliases)

    output = load_workbook(file._preprocessed_path).active
    assert [output.cell(row, 1).value for row in (2, 3, 4)] == [
        '员工报账业务-1.控制点-控制措施01',
        '员工报账业务-12.控制点-控制措施01',
        '员工报账业务-13.控制点-控制措施01',
    ]
    assert output['C2'].value == '=ROW(A1)'
    assert output['G2'].value == '=C2&$D2'


def test_matrix_reference_and_existing_column_precede_single_rightmost_opinion(tmp_path):
    """tmp_path 为隔离目录；基准富文本保留，写入岗位已有编号，重复执行不添列。"""
    path = tmp_path / '09信通公司矩阵.xlsx'
    book = Workbook(); ws = book.active; ws.title = '风控矩阵'
    ws.append(['控制措施编号', '控制措施', '责任主体', '是否适用及原因', '审核意见'])
    ws.append(['业务-2.名称-控制措施05', '核对资料', '财务部', '适用', '人工意见'])
    ws['E2'].comment = Comment('保留意见批注', '审核员'); ws['E2'].hyperlink = 'https://example.com/review'
    duty = book.create_sheet('岗位内控责任清单')
    duty.append(['控制措施编号', '部门', '岗位名称', '人员姓名', '岗位职责', '角色'])
    duty.append(['业务-02.其他名称-控制措施5', '财务部', '核算专责', '张三', '对资料真实性负主体责任', '经办'])
    book.save(path); file, aliases = file_record(path)
    baseline_path = tmp_path / 'baseline.xlsx'
    baseline_book = Workbook(); b = baseline_book.active
    b.append(['控制措施编号', '责任主体']); b.append(['业务-02.旧名称-控制措施5', CellRichText([TextBlock(InlineFont(strike=True), '旧部门'), TextBlock(InlineFont(color='FF0000'), '省公司-财务部')])]); baseline_book.save(baseline_path)
    baseline = {('09', 'default'): {'responsibilities': [{'measure_id': b['A2'].value, 'source_path': str(baseline_path), 'sheet': b.title, 'cell': 'B2'}]}}
    module().enrich_matrices(file, path, baseline, [], [file])
    module().enrich_matrices(file, path, baseline, [], [file])
    output = load_workbook(path, rich_text=True); ws = output.active
    assert [ws.cell(1, col).value for col in range(5, 8)] == ['省公司版本责任主体（核对后删除）', '岗位清单已有的控制措施编号', '审核意见']
    assert ws['E2'].value[0].font.strike and ws['E2'].value[1].font.color.rgb == '00FF0000'
    assert ws['F2'].value == '业务-2.名称-控制措施05'
    assert ws['G2'].value == '人工意见'
    assert ws['G2'].comment.text == '保留意见批注' and ws['G2'].hyperlink.target == 'https://example.com/review'
    assert ws['E2'].comment is None
    assert ws.max_column == 7
    assert ws.auto_filter.ref == 'E1:G2'
    from zipfile import ZipFile
    from lxml import etree
    with ZipFile(path) as archive:
        drawing = etree.fromstring(archive.read('xl/drawings/commentsDrawing1.vml'))
        assert drawing.xpath('//x:ClientData/x:Column/text()', namespaces={'x': 'urn:schemas-microsoft-com:office:excel'}) == ['6']


def test_existing_matrix_output_columns_are_cleared_before_rebuild(tmp_path):
    """tmp_path 为隔离目录；修正版保留三列时，旧数据全部清空并仅写入本轮结果。"""
    path = tmp_path / '09信通公司修正版矩阵.xlsx'
    book = Workbook(); matrix = book.active; matrix.title = '风控矩阵'
    matrix.append([
        '控制措施编号', '控制措施', '责任主体', '是否适用及原因',
        '省公司版本责任主体（核对后删除）', '岗位清单已有的控制措施编号', '审核意见',
    ])
    matrix.append(['业务-2.名称-控制措施05', '核对资料', '财务部', '适用', '旧责任主体', '旧措施编号', '旧审核意见'])
    matrix.append([None, None, None, None, '尾部旧责任主体', '尾部旧措施编号', '尾部旧审核意见'])
    duty = book.create_sheet('岗位内控责任清单')
    duty.append(['控制措施编号', '部门', '岗位名称', '人员姓名', '岗位职责', '角色'])
    duty.append(['业务-2.名称-控制措施05', '财务部', '核算专责', '张三', '核对资料', '经办'])
    book.save(path)
    file, _ = file_record(path)

    write_outputs(
        [file], [], tmp_path / 'output',
        baselines={('09', 'default'): {'responsibilities': []}},
        metadata_dir=tmp_path / 'metadata',
    )

    output = load_workbook(tmp_path / 'output' / path.name).active
    assert [output.cell(1, column).value for column in range(5, 8)] == [
        '省公司版本责任主体（核对后删除）', '岗位清单已有的控制措施编号', '审核意见',
    ]
    assert output['E2'].value is None
    assert output['F2'].value == '业务-2.名称-控制措施05'
    assert output['G2'].value is None
    assert [output.cell(3, column).value for column in range(5, 8)] == [None, None, None]
    assert output.max_column == 7


@pytest.mark.parametrize(('last_header_hidden', 'expected_header_row'), [(True, 2), (False, 3)])
def test_matrix_output_writes_headers_to_last_visible_header_row(
        tmp_path, last_header_hidden, expected_header_row):
    """验证新增列使用最后一个可见表头行。

    tmp_path 为隔离目录，last_header_hidden 控制末表头行是否隐藏，
    expected_header_row 为应写入新增三列标题的行号。
    """
    path = tmp_path / '06信通公司矩阵.xlsx'
    book = Workbook(); sheet = book.active; sheet.title = '设备管理'
    sheet['A1'] = '风控矩阵-设备管理'
    headers = ['业务条线', '控制措施编号', '控制措施', '控制措施分类', None,
               '控制系统', '控制载体', '责任主体', '信通公司是否适用']
    for column, value in enumerate(headers, 1):
        sheet.cell(2, column, value)
    sheet['D3'] = '不相容岗位'; sheet['E3'] = '分级授权'
    sheet.row_dimensions[3].hidden = last_header_hidden
    values = ['设备', 'M01', '核对资料', None, None, '系统', '载体', '部门-岗位', '是']
    for column, value in enumerate(values, 1):
        sheet.cell(4, column, value)
    # 复现模板筛选从末表头行开始的情形，并保留原有筛选条件。
    sheet.auto_filter.ref = 'A3:I4'
    sheet.auto_filter.add_filter_column(2, ['核对资料'])
    book.save(path)

    file = FileRecord(path, Path(path.name), sha256_file(path), 'xlsx', '205H', [], False,
                      '06', 'default', 'matrix')
    file._parser_policy = {'version': 3}
    aliases = load_pack(ROOT / '审核器/rulepacks/releases/1.8.0')['field_aliases']
    file.sheets = parse_workbook_v180(file, path, aliases)
    assert file.sheets[0].header_rows == [2, 3]

    writer_output = tmp_path / 'writer-output'
    write_outputs([file], [], writer_output, metadata_dir=tmp_path / 'writer-metadata')
    writer_sheet = load_workbook(writer_output / path.name).active
    assert writer_sheet.cell(expected_header_row, 10).value == '审核意见'
    assert writer_sheet.cell(5 - expected_header_row, 10).value is None

    write_outputs(
        [file], [], tmp_path / 'output',
        baselines={('06', 'default'): {'responsibilities': []}},
        metadata_dir=tmp_path / 'metadata',
    )
    # 真实重跑会重新解析源件，不复用已被本轮增强更新过的内存坐标。
    rerun_file = FileRecord(path, Path(path.name), sha256_file(path), 'xlsx', '205H', [], False,
                            '06', 'default', 'matrix')
    rerun_file._parser_policy = {'version': 3}
    rerun_file.sheets = parse_workbook_v180(rerun_file, path, aliases)
    write_outputs(
        [rerun_file], [], tmp_path / 'output',
        baselines={('06', 'default'): {'responsibilities': []}},
        metadata_dir=tmp_path / 'metadata',
    )

    output = load_workbook(tmp_path / 'output' / path.name).active
    assert output.row_dimensions[3].hidden is last_header_hidden
    assert [output.cell(expected_header_row, column).value for column in range(10, 13)] == [
        '省公司版本责任主体（核对后删除）', '岗位清单已有的控制措施编号', '审核意见',
    ]
    assert [output.cell(5 - expected_header_row, column).value for column in range(10, 13)] == [None, None, None]
    assert output.max_column == 12
    assert output.auto_filter.ref == f'A{expected_header_row}:L4'
    assert output.auto_filter.filterColumn[0].filters.filter == ['核对资料']


def test_matrix_output_extends_existing_filter_and_keeps_filter_criteria(tmp_path):
    """tmp_path 为隔离目录；原筛选范围扩展到审核意见列，已选条件保持不变。"""
    path = tmp_path / '09信通公司矩阵.xlsx'
    book = Workbook(); ws = book.active; ws.title = '风控矩阵'
    ws.append(['控制措施编号', '控制措施', '责任主体', '是否适用及原因', '审核意见'])
    ws.append(['业务-2.名称-控制措施05', '核对资料', '财务部', '适用', None])
    ws.auto_filter.ref = 'A1:D2'
    ws.auto_filter.add_filter_column(1, ['核对资料'])
    book.save(path)
    file, _ = file_record(path)

    module().enrich_matrices(file, path, {}, [], [file])
    module().enrich_matrices(file, path, {}, [], [file])

    output = load_workbook(path).active
    assert output.auto_filter.ref == 'A1:G2'
    assert len(output.auto_filter.filterColumn) == 1
    assert output.auto_filter.filterColumn[0].colId == 1
    assert output.auto_filter.filterColumn[0].filters.filter == ['核对资料']


def test_matrix_output_places_new_filter_after_sheet_protection(tmp_path):
    """tmp_path 为隔离目录；受保护矩阵的筛选节点顺序必须符合 OOXML 规范。"""
    path = tmp_path / '09信通公司矩阵.xlsx'
    book = Workbook(); ws = book.active; ws.title = '风控矩阵'
    ws.append(['控制措施编号', '控制措施', '责任主体', '是否适用及原因', '审核意见'])
    ws.append(['业务-2.名称-控制措施05', '核对资料', '财务部', '适用', None])
    ws.protection.sheet = True
    book.save(path)
    file, _ = file_record(path)

    module().enrich_matrices(file, path, {}, [], [file])

    with ZipFile(path) as archive:
        root = etree.fromstring(archive.read('xl/worksheets/sheet1.xml'))
    child_names = [etree.QName(child).localname for child in root]
    assert child_names.index('sheetData') < child_names.index('sheetProtection') < child_names.index('autoFilter')


def test_blank_merged_output_area_is_unmerged_and_later_files_continue(tmp_path):
    """tmp_path 为隔离目录；空白合并区域自动取消合并，并继续写出整批文件。"""
    bad_path = tmp_path / '冲突矩阵.xlsx'
    bad_book = Workbook(); bad_sheet = bad_book.active; bad_sheet.title = '风控矩阵'
    bad_sheet.append(['控制措施编号', '控制措施', '责任主体', '是否适用及原因'])
    bad_sheet.append(['业务-1.控制点-控制措施01', '核对资料', '财务部', '适用'])
    # 空白尾部合并横跨程序即将使用的审核意见列，复现合同管理文件的布局冲突。
    bad_sheet.merge_cells('E1:H1')
    bad_book.save(bad_path)
    bad_file, _ = file_record(bad_path)

    good_path = tmp_path / '正常矩阵.xlsx'
    good_book = Workbook(); good_sheet = good_book.active; good_sheet.title = '风控矩阵'
    good_sheet.append(['控制措施编号', '控制措施', '责任主体', '是否适用及原因'])
    good_sheet.append(['业务-2.控制点-控制措施01', '核对资料', '财务部', '适用'])
    good_book.save(good_path)
    good_file, _ = file_record(good_path)

    output = tmp_path / 'output'
    warnings, _ = write_outputs(
        [bad_file, good_file], [], output,
        baselines={('09', 'default'): {'responsibilities': []}},
        metadata_dir=tmp_path / 'metadata',
    )

    skipped = [warning for warning in warnings if warning['type'] == 'output_incompatible_file']
    assert skipped == []
    repaired = load_workbook(output / bad_path.name).active
    assert [repaired.cell(1, column).value for column in range(5, 8)] == [
        '省公司版本责任主体（核对后删除）', '岗位清单已有的控制措施编号', '审核意见',
    ]
    assert 'E1:H1' not in {str(cell_range) for cell_range in repaired.merged_cells.ranges}
    assert load_workbook(output / good_path.name).active['G1'].value == '审核意见'
    report = json.loads((tmp_path / 'metadata/未审核文件.json').read_text(encoding='utf-8'))
    assert report == []
    assert not _has_failures([bad_file, good_file], [], warnings)


@pytest.mark.parametrize('merged_value', ['原表说明', '=SUM(A2:D2)'])
def test_nonempty_merged_output_area_moves_columns_right_until_safe(tmp_path, merged_value):
    """tmp_path 为隔离目录，merged_value 为原表文字或公式；保留非空合并区并将三列输出右移到安全位置。"""
    path = tmp_path / '非空合并矩阵.xlsx'
    book = Workbook(); sheet = book.active; sheet.title = '风控矩阵'
    sheet.append(['控制措施编号', '控制措施', '责任主体', '是否适用及原因'])
    sheet.append(['业务-1.控制点-控制措施01', '核对资料', '财务部', '适用'])
    sheet.merge_cells('E1:H1'); sheet['E1'] = '左侧说明'
    sheet.merge_cells('I1:L1'); sheet['I1'] = merged_value
    sheet.merge_cells('M1:P1')
    book.save(path)
    file, _ = file_record(path)
    assert file.sheets[0].output_column == 13

    output = tmp_path / 'output'
    warnings, _ = write_outputs(
        [file], [], output,
        baselines={('09', 'default'): {'responsibilities': []}},
        metadata_dir=tmp_path / 'metadata',
    )

    assert [warning for warning in warnings if warning['type'] == 'output_incompatible_file'] == []
    result = load_workbook(output / path.name, data_only=False).active
    assert result['E1'].value == '左侧说明'
    assert result['I1'].value == merged_value
    assert {'E1:H1', 'I1:L1'} <= {str(cell_range) for cell_range in result.merged_cells.ranges}
    assert 'M1:P1' not in {str(cell_range) for cell_range in result.merged_cells.ranges}
    assert [result.cell(1, column).value for column in range(13, 16)] == [
        '省公司版本责任主体（核对后删除）', '岗位清单已有的控制措施编号', '审核意见',
    ]


def test_matrix_output_hides_columns_outside_confirmed_keyword_set(tmp_path):
    """tmp_path 为隔离目录；审核副本只显示指定矩阵业务列和程序新增列。"""
    path = tmp_path / '09职工福利矩阵.xlsx'; book = Workbook(); ws = book.active; ws.title = '风控矩阵'
    ws.append(['序号', '关键控制点', '控制措施编号', '控制措施', '经营风险', '责任主体', '是否适用及原因', '审核意见'])
    ws.append([1, '薪酬', '业务-2.名称-控制措施05', '核对资料', '风险', '财务部', '适用', ''])
    book.save(path); file, _ = file_record(path)

    module().enrich_matrices(file, path, {}, [], [file])

    output = load_workbook(path).active
    assert output.column_dimensions['A'].hidden is True
    assert output.column_dimensions['E'].hidden is True
    assert all(output.column_dimensions[column].hidden is not True for column in ('B', 'C', 'D', 'F', 'G'))


@pytest.mark.parametrize('hidden_incompatible', [False, True])
def test_missing_incompatible_sheet_is_still_reported(tmp_path, registry, hidden_incompatible):
    """未匹配可见不相容岗位清单时仍执行第6条；hidden_incompatible 控制是否添加隐藏表。"""
    path = tmp_path / '09三清单-信通公司.xlsx'
    book = Workbook()
    book.active.title = '岗位内控责任清单'
    book.active.append(['控制措施编号', '部门', '岗位名称', '岗位职责'])
    system_rule = book.create_sheet('系统控制规则清单')
    system_rule.append(['控制措施编号', '规则名称', '规则内容'])
    if hidden_incompatible:
        incompatible = book.create_sheet('不相容岗位清单')
        incompatible.append(['不相容业务角色', '岗位A', '岗位B'])
        incompatible.sheet_state = 'hidden'
    book.save(path)
    pack = load_pack(ROOT / '审核器/rulepacks/releases/1.9.7')
    file = FileRecord(path, Path(path.name), 'hash', 'xlsx', '205H', [], False,
                      '09', 'default', 'three_lists')
    file._rules_version = pack['manifest']['version']
    parse_files([file], pack['field_aliases'], tmp_path / 'work')
    assert {sheet.sheet_type for sheet in file.sheets} == {'position_duty', 'system_rule'}

    entities = {'205H': Entity('205H', '信通公司', '省公司')}
    pack['rules'] = [rule for rule in pack['rules'] if rule['rule_id'] == 'lists.incompatible_exists']
    pack['submission_scope'] = {
        'entity_codes': ['205H'],
        'businesses': [{'business_code': '09', 'variant_id': 'default', 'required': True}],
    }

    findings, _ = run_engine(pack, registry, [file], entities, {})

    assert len(findings) == 1
    assert findings[0].check_id == 'incompatible_sheet_exists'
    assert findings[0].message == '【第6条】未找到“不相容岗位清单”，请核实是否漏报。'


def test_new_pack_audit_and_repeat_rebuild_opinions(tmp_path):
    """重跑重建意见且统计表按约定归档；tmp_path 为测试目录。"""
    import json
    from test_audit_v180 import create_sample
    from run_audit import configure_soffice
    from risk_audit.runner import audit
    from risk_audit.util import write_json
    configure_soffice('/Users/67m/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice')
    # 当前主体归属按单位目录全称确定，样例使用名册中的正式名称。
    source = create_sample(tmp_path / '国网湖南省电力有限公司信息通信分公司'); before = sha256_file(source)
    scope = tmp_path / 'scope.json'
    write_json(scope, {'205H': {'entity_codes': ['205H'], 'businesses': [{'business_code': '06', 'variant_id': 'default', 'required': True}]}})
    legacy_root = tmp_path / 'output/_risk_audit'
    legacy_root.mkdir(parents=True)
    (legacy_root / '意见反馈明细.xlsx').write_text('旧版程序报告', encoding='utf-8')
    (legacy_root / '审核情况.xlsx').write_text('旧版程序报告', encoding='utf-8')
    (legacy_root / '意见反馈报告').mkdir()
    (legacy_root / '意见反馈报告/index.html').write_text('旧版程序报告', encoding='utf-8')
    active_pack = ROOT / '审核器/rulepacks/releases/1.9.13'

    result = audit(source.parent, tmp_path / 'output', active_pack, ROOT / '审核/会计主体清单20260907.xlsx', ROOT, tmp_path / 'runs', scope_file=scope)
    assert result['write_completed'] and not result['statuses'].get('failed')
    findings = json.loads((Path(result['run_dir']) / 'findings.json').read_text())
    assert not any(finding['display_code'] in {'R14', 'R13'} for finding in findings)
    assert not any(finding['rule_id'] in {'lists.order', 'lists.order_system', 'lists.order_incompatible'} for finding in findings)
    output = load_workbook(tmp_path / 'output' / source.name, rich_text=True)
    assert list(output['风控矩阵'].values)[0][-3:] == ('省公司版本责任主体（核对后删除）', '岗位清单已有的控制措施编号', '审核意见')
    assert [output['岗位内控责任清单'].cell(row, 1).value for row in (2, 3, 4)] == ['M1', 'M10', 'M10']
    assert '人工样例意见' not in (output['岗位内控责任清单']['I3'].value or '')
    assert 'feedback_report' not in result
    assert 'feedback_html_report' not in result
    assert 'audit_summary_report' not in result
    assert Path(result['audit_statistics_report']) == tmp_path / 'output/审核统计表.xlsx'
    assert Path(result['audit_statistics_report']).is_file()
    assert load_workbook(result['audit_statistics_report']).sheetnames == ['审核统计表']
    assert not (tmp_path / 'output/_risk_audit').exists()
    audit_directory = Path(result['run_dir']) / '_risk_audit'
    assert (audit_directory / 'ownership.json').is_file()
    assert (audit_directory / '本次审核范围.md').is_file()
    assert (audit_directory / '集中复核事项.md').is_file()
    duplicated_reports = [
        '本次审核范围.md',
        '本次审核范围.json',
        '主体代码待确认.md',
        'entity_alerts.json',
        '隐藏工作表处理提示.md',
        'hidden_sheet_alerts.json',
        '集中复核事项.md',
        '复核事项.csv',
        'review_tasks.json',
        'review_tasks_summary.json',
        '责任表述集中确认.csv',
        'responsibility_patterns.json',
        'responsibility_patterns_summary.json',
        'internal_diagnostics.json',
        '内部处理事项.csv',
        'capability_diagnostics.json',
        '程序判定能力明细.csv',
    ]
    assert all((audit_directory / name).is_file() for name in duplicated_reports)
    assert all(not (Path(result['run_dir']) / name).exists() for name in duplicated_reports)
    assert Path(result['submission_scope']['report']).parent == audit_directory
    assert Path(result['hidden_sheets']['report']).parent == audit_directory
    assert Path(result['review_tasks']['review_task_file']).parent == audit_directory
    assert Path(result['internal_diagnostics']['report']).parent == audit_directory
    assert Path(result['entity_code_alerts']['report']).parent == audit_directory
    second = audit(source.parent, tmp_path / 'output', active_pack, ROOT / '审核/会计主体清单20260907.xlsx', ROOT, tmp_path / 'runs', scope_file=scope)
    assert second['write_completed']
    repeated = load_workbook(tmp_path / 'output' / source.name)
    assert '人工样例意见' not in (repeated['岗位内控责任清单']['I3'].value or '')
    assert repeated['风控矩阵'].max_column == output['风控矩阵'].max_column
    assert sha256_file(source) == before


def test_interrupted_run_ownership_replaces_manual_edit_on_retry(tmp_path, monkeypatch):
    """汇总中断后的旧意见在重跑时被替换；参数为测试目录和替换工具。"""
    from test_audit_v180 import create_sample
    from run_audit import configure_soffice
    from risk_audit import runner as audit_runner
    from risk_audit.util import write_json
    configure_soffice('/Users/67m/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice')
    source = create_sample(tmp_path / '国网湖南省电力有限公司信息通信分公司')
    scope = tmp_path / 'scope.json'
    write_json(scope, {'205H': {'entity_codes': ['205H'], 'businesses': [
        {'business_code': '06', 'variant_id': 'default', 'required': True},
    ]}})
    active_pack = ROOT / '审核器/rulepacks/releases/1.9.13'
    original_statistics_writer = audit_runner.write_audit_statistics

    def fail_statistics(*args, **kwargs):
        """args/kwargs 为统计表写入参数；模拟副本及 ownership 完成后的汇总中断。"""
        raise RuntimeError('模拟汇总阶段中断')

    monkeypatch.setattr(audit_runner, 'write_audit_statistics', fail_statistics)
    with pytest.raises(RuntimeError, match='模拟汇总阶段中断'):
        audit_runner.audit(source.parent, tmp_path / 'output', active_pack,
                           ROOT / '审核/会计主体清单20260907.xlsx', ROOT, tmp_path / 'runs', scope_file=scope)

    failed_run = next((tmp_path / 'runs').iterdir())
    assert not (failed_run / 'summary.json').exists()
    ownership = json.loads((failed_run / '_risk_audit/ownership.json').read_text(encoding='utf-8'))
    relative_path, cells = next(iter(ownership.items()))
    owned_cell = cells[0]
    output_path = tmp_path / 'output' / relative_path
    output = load_workbook(output_path)
    output[owned_cell['sheet']][owned_cell['cell']] = '中断后人工意见'
    output.save(output_path)

    monkeypatch.setattr(audit_runner, 'write_audit_statistics', original_statistics_writer)
    result = audit_runner.audit(source.parent, tmp_path / 'output', active_pack,
                                ROOT / '审核/会计主体清单20260907.xlsx', ROOT, tmp_path / 'runs', scope_file=scope)

    repeated = load_workbook(output_path)
    assert '中断后人工意见' not in (repeated[owned_cell['sheet']][owned_cell['cell']].value or '')
    assert Path(result['audit_metadata_dir']).parent == Path(result['run_dir'])


def test_audit_skips_dot_prefixed_temporary_file(tmp_path):
    """tmp_path 为测试目录；Office 点开头临时文件不得中断批次或进入统计。"""
    from test_audit_v180 import create_sample
    from run_audit import configure_soffice
    from risk_audit.runner import audit
    from risk_audit.util import write_json
    configure_soffice('/Users/67m/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice')
    input_root = tmp_path / '国网湖南省电力有限公司信息通信分公司'
    create_sample(input_root)
    # 模拟 Office 临时文件，扫描时应直接忽略。
    (input_root / '.~06风控矩阵.xlsx').write_bytes(b'not-an-excel-workbook')
    scope = tmp_path / 'scope.json'
    write_json(scope, {'205H': {'entity_codes': ['205H'], 'businesses': [
        {'business_code': '06', 'variant_id': 'default', 'required': True},
    ]}})

    result = audit(input_root, tmp_path / 'output', ROOT / '审核器/rulepacks/releases/1.9.13',
                   ROOT / '审核/会计主体清单20260907.xlsx', ROOT, tmp_path / 'runs', scope_file=scope)

    assert result['write_completed'] is True
    assert result['unparsed_files'] == 0
    assert result['unparsed_file_paths'] == []
    report = Path(result['audit_statistics_report'])
    assert report.is_file()
    statistics = load_workbook(report)['审核统计表']
    assert statistics.max_row == 2
    assert statistics['A2'].value == input_root.name
    assert statistics['B2'].value.startswith('业务06：\n')
    assert '【第' in statistics['B2'].value


def test_natural_sort_uses_numbers_before_control_point_names(tmp_path):
    """tmp_path 为测试目录；同控制点的名称差异不能把措施05排到02之前。"""
    path = tmp_path / '09信通公司三清单.xlsx'; book = Workbook(); ws = book.active; ws.title = '岗位内控责任清单'
    ws.append(['控制措施编号', '部门', '岗位名称', '人员姓名', '岗位职责', '角色'])
    for measure in ('业务-2.A名称-控制措施05', '业务-2.Z名称-控制措施02'):
        ws.append([measure, '财务部', '核算专责', '张三', '对资料真实性负主体责任', '经办'])
    book.save(path); file, aliases = file_record(path); module().sort_duties(file, {}, tmp_path / 'work', aliases)
    assert load_workbook(file._preprocessed_path).active['A2'].value == '业务-2.Z名称-控制措施02'


def test_natural_sort_extracts_numbers_from_actual_nonstandard_measure_ids(tmp_path):
    """tmp_path 为测试目录；缺少句点或“控制措施”字样的实际编号仍按控制点和措施数字排序。"""
    path = tmp_path / '08三清单.xlsx'; book = Workbook(); ws = book.active; ws.title = '岗位内控责任清单'
    ws.append(['控制措施编号', '部门', '岗位名称', '人员姓名', '岗位职责', '角色'])
    for measure in ('科技项目管理业务-14A名称-控制措施05', '科技项目管理业务-14Z名称-控制措施02',
                    '设备（资产）管理-15.Z业务02', '设备（资产）管理-15.A业务05'):
        ws.append([measure, '财务部', '核算专责', '张三', '对资料真实性负主体责任', '经办'])
    book.save(path); file, aliases = file_record(path)

    module().sort_duties(file, {}, tmp_path / 'work', aliases)

    output = load_workbook(file._preprocessed_path).active
    values = [output.cell(row, 1).value for row in range(2, 6)]
    assert values == [
        '科技项目管理业务-14Z名称-控制措施02', '科技项目管理业务-14A名称-控制措施05',
        '设备（资产）管理-15.Z业务02', '设备（资产）管理-15.A业务05',
    ]


def test_complete_writeback_replaces_rich_old_opinion(tmp_path):
    """程序意见替换旧富文本且重复运行不累积；tmp_path 为测试目录。"""
    from risk_audit.writer import write_outputs
    path = tmp_path / '09信通公司三清单.xlsx'; book = Workbook(); ws = book.active; ws.title = '岗位内控责任清单'
    ws.append(['控制措施编号', '部门', '岗位名称', '人员姓名', '岗位职责', '角色', '审核意见'])
    ws.append(['M1', '财务部', '核算专责', '张三', '对资料真实性负主体责任', '经办'])
    ws['G2'] = CellRichText([TextBlock(InlineFont(strike=True, color='FF0000'), '旧人工意见'), '现行人工意见']); book.save(path)
    file, _ = file_record(path)
    finding = Finding('r', 's', 'x', 'position_specific', 'R05', 1, 'violation', '205H', '信通公司', '09', 'default', path.name, ws.title, 2, '【第5条】程序意见', {}, 'row')
    for _ in range(2):
        write_outputs([file], [finding], tmp_path / 'output', baselines={}, metadata_dir=tmp_path / 'metadata')
        cell = load_workbook(tmp_path / 'output' / path.name, rich_text=True).active['G2']
        assert str(cell.value) == '【第5条】程序意见'


def test_changed_old_opinion_is_cleared_without_current_finding(tmp_path):
    """没有本轮问题时清空所有旧意见；tmp_path 为测试目录。"""
    from risk_audit.writer import write_outputs
    path = tmp_path / '09信通公司三清单.xlsx'; book = Workbook(); ws = book.active; ws.title = '岗位内控责任清单'
    ws.append(['控制措施编号', '部门', '岗位名称', '人员姓名', '岗位职责', '角色', '审核意见'])
    ws.append(['M1', '财务部', '核算专责', '张三', '对资料真实性负主体责任', '经办'])
    for text in ('请核实', '请核实并补充材料'):
        ws['G2'] = CellRichText([TextBlock(InlineFont(color='FF0000'), text)]); book.save(path)
        file, _ = file_record(path); write_outputs([file], [], tmp_path / 'output', baselines={}, metadata_dir=tmp_path / 'metadata')
    value = load_workbook(tmp_path / 'output' / path.name, rich_text=True).active['G2'].value
    assert value is None


def test_old_opinion_quoting_program_text_is_replaced_on_repeat(tmp_path):
    """旧意见即使引用程序文字也由本轮结果替换；tmp_path 为测试目录。"""
    from risk_audit.writer import write_outputs
    path = tmp_path / '09信通公司三清单.xlsx'; book = Workbook(); ws = book.active; ws.title = '岗位内控责任清单'
    ws.append(['控制措施编号', '部门', '岗位名称', '人员姓名', '岗位职责', '角色', '审核意见'])
    ws.append(['M1', '财务部', '核算专责', '张三', '对资料真实性负主体责任', '经办'])
    human = '人工说明：【第5条】程序意见仅作参考'
    ws['G2'] = CellRichText([TextBlock(InlineFont(color='FF0000'), human)]); book.save(path)
    file, _ = file_record(path)
    finding = Finding('r', 's', 'x', 'position_specific', 'R05', 1, 'violation', '205H', '信通公司', '09', 'default', path.name, ws.title, 2, '【第5条】程序意见', {}, 'row')
    for _ in range(2):
        write_outputs([file], [finding], tmp_path / 'output', baselines={}, metadata_dir=tmp_path / 'metadata')
        value = load_workbook(tmp_path / 'output' / path.name, rich_text=True).active['G2'].value
        assert str(value) == '【第5条】程序意见'
