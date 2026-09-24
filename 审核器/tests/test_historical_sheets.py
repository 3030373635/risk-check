from pathlib import Path
from openpyxl import Workbook
from risk_audit.models import FileRecord
from risk_audit.configuration.loader import load_pack
from risk_audit.readers.excel import parse_workbook

ROOT=Path(__file__).resolve().parents[2]
ALIASES=load_pack(ROOT/'审核器/rulepacks/releases/1.2.5')['field_aliases']

def sheet(w,title,*,complete=False,old=False,hidden=False):
    s=w.create_sheet(title);s.sheet_state='hidden' if hidden else 'visible'
    s.append(['岗位内控责任清单'])
    headers=['部门','岗位名称','岗位职责编号','岗位职责','控制措施编号']
    values=['财务部','财务专责',1,'对报表准确性负主体责任','M1']
    if complete:headers+=['人员姓名','角色'];values+=['甲','经办']
    if old:headers+=['旧的'];values+=['旧M1']
    s.append(headers);s.append(values)
    return s

def parse(tmp_path,specs,*,include_hidden=False):
    w=Workbook();w.remove(w.active)
    for title,options in specs:sheet(w,title,**options)
    if all(s.sheet_state=='hidden' for s in w):w.create_sheet('填报说明')
    p=tmp_path/'book.xlsx';w.save(p)
    f=FileRecord(p,Path(p.name),'h','xlsx','205H',[],False,'02','default','three_lists')
    result=parse_workbook(f,p,ALIASES,include_hidden=include_hidden)
    return [s.title for s in result],f

def test_crosswalk_and_transition_are_not_current_submission(tmp_path):
    titles,f=parse(tmp_path,[('岗位职责清单（含有红色补充的）',{'old':True,'hidden':True}),('岗位职责清单-过渡',{'hidden':True}),('岗位内控责任清单9.8',{'complete':True})])
    assert titles==['岗位内控责任清单9.8']
    assert len(f.preservation['hidden_sheets'])==2
    assert all(s['action']=='skipped_hidden' for s in f.preservation['hidden_sheets'])

def test_hidden_current_sheet_is_skipped_by_default(tmp_path):
    titles,f=parse(tmp_path,[('岗位内控责任清单',{'complete':True,'hidden':True})])
    assert titles==[]
    assert not f.preservation.get('excluded_historical_sheets')

def test_hidden_current_sheet_requires_explicit_inclusion(tmp_path):
    titles,f=parse(tmp_path,[('岗位内控责任清单',{'complete':True,'hidden':True})],include_hidden=True)
    assert titles==['岗位内控责任清单']
    assert f.preservation['hidden_sheets'][0]['action']=='audited'

def test_only_transition_sheet_is_not_silently_discarded(tmp_path):
    titles,_=parse(tmp_path,[('岗位职责清单-过渡',{})])
    assert titles==['岗位职责清单-过渡']

def test_incomplete_ordinary_sheet_is_not_excused_as_history(tmp_path):
    titles,_=parse(tmp_path,[('岗位职责清单',{}),('岗位内控责任清单9.8',{'complete':True})])
    assert len(titles)==2

def test_multiple_current_versions_require_retaining_evidence(tmp_path):
    titles,f=parse(tmp_path,[('岗位职责清单-过渡',{}),('岗位内控责任清单9.8',{'complete':True}),('岗位内控责任清单9.9',{'complete':True})])
    assert len(titles)==3
    assert not f.preservation.get('excluded_historical_sheets')

def test_complete_supplement_sheet_is_retained_when_explicitly_included(tmp_path):
    titles,_=parse(tmp_path,[('岗位职责清单（含有红色补充的）',{'complete':True,'hidden':True}),('岗位内控责任清单9.8',{'complete':True})],include_hidden=True)
    assert len(titles)==2


def test_empty_current_template_cannot_supersede_history(tmp_path):
    w=Workbook();w.remove(w.active)
    sheet(w,'岗位职责清单-过渡')
    current=sheet(w,'岗位内控责任清单9.8',complete=True)
    current.delete_rows(3)
    p=tmp_path/'empty-current.xlsx';w.save(p)
    f=FileRecord(p,Path(p.name),'h','xlsx','205H',[],False,'02','default','three_lists')
    parsed=parse_workbook(f,p,ALIASES)
    assert any(s.title=='岗位职责清单-过渡' and s.records for s in parsed)
    assert not f.preservation.get('excluded_historical_sheets')
