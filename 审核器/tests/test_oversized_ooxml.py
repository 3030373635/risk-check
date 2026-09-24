"""异常膨胀 OOXML 工作簿的安全净化回归测试。"""

from __future__ import annotations

from copy import copy
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from risk_audit.models import FileRecord
from risk_audit.output_0916 import sort_duties
from risk_audit.readers.excel import parse_files
from risk_audit.readers import ooxml
from risk_audit.util import sha256_file


def prepare_oversized_workbook(*args, **kwargs):
    """调用待实现接口；args/kwargs 为工作簿净化参数，并确保缺失能力表现为断言失败。"""
    implementation = getattr(ooxml, "prepare_oversized_workbook", None)
    assert callable(implementation), "缺少异常膨胀工作簿净化接口"
    return implementation(*args, **kwargs)


def create_style_expanded_duty_workbook(path: Path, *, far_value: str | None = None) -> None:
    """创建格式扩散的岗位清单；path 为目标路径，far_value 为极远列可选业务值。"""
    book = Workbook()
    sheet = book.active
    sheet.title = "岗位职责清单"
    sheet.append(["部门", "岗位名称", "人员姓名", "岗位角色", "岗位职责", "控制措施编号"])
    sheet.append(["财务部", "会计", "张三", "经办", "负责设备核算", "M1"])
    # 模拟 WPS 将空白格式扩散到 Excel 最大列；业务区内样式必须保留。
    sheet["F2"].font = copy(sheet["A2"].font)
    sheet["XFD2"].font = copy(sheet["A2"].font)
    if far_value is not None:
        sheet["XFD2"] = far_value
    book.save(path)


def create_compact_extreme_dimension_workbook(path: Path) -> None:
    """创建 XML 体积很小但声明范围接近 Excel 上限的工作簿。

    Args:
        path: 目标工作簿路径。
    """

    book = Workbook()
    sheet = book.active
    sheet.title = "岗位职责清单"
    sheet.append(["部门", "岗位名称", "人员姓名", "岗位角色", "岗位职责", "控制措施编号"])
    sheet.append(["财务部", "会计", "张三", "经办", "负责设备核算", "M1"])
    # 只写入一个远端样式单元格，模拟真实文件的 XFB1048572 虚假边界。
    sheet.cell(1_048_572, 16_382).font = copy(sheet["A2"].font)
    book.save(path)
    book.close()


def create_far_merged_range_workbook(path: Path) -> None:
    """创建保留真实合并且含远端空白合并的工作簿。

    Args:
        path: 目标工作簿路径。
    """

    book = Workbook()
    sheet = book.active
    sheet.title = "岗位职责清单"
    sheet.append(["部门", None, "人员姓名", "岗位角色", "岗位职责", "控制措施编号"])
    sheet.append(["财务部", "会计", "张三", "经办", "负责设备核算", "M1"])
    sheet.merge_cells("A1:B1")
    # 模拟 WPS 在真实业务区之外写入的纯样式合并区域。
    sheet.merge_cells("XEZ100000:XFA100001")
    book.save(path)
    book.close()


def append_trailing_style_rows(path: Path, final_row: int) -> None:
    """在工作表 XML 末尾追加纯样式空行。

    path 为待修改的测试工作簿，final_row 为模拟的虚假末行号。
    """
    with ZipFile(path) as archive:
        parts = {info.filename: (info, archive.read(info.filename)) for info in archive.infolist()}
    info, worksheet = parts["xl/worksheets/sheet1.xml"]
    rows = b"".join(
        f'<row r="{row}"><c r="A{row}" s="1"/></row>'.encode()
        for row in range(3, final_row + 1)
    )
    worksheet = worksheet.replace(b"</sheetData>", rows + b"</sheetData>")
    worksheet = worksheet.replace(b'A1:XFD2', f'A1:XFD{final_row}'.encode())
    parts["xl/worksheets/sheet1.xml"] = (info, worksheet)
    with ZipFile(path, "w") as archive:
        for part_info, content in parts.values():
            archive.writestr(part_info, content)


def test_prepare_oversized_workbook_removes_only_far_style_cells(tmp_path):
    """临时净化须缩回虚假范围并保持源文件及真实业务内容；tmp_path 为测试目录。"""
    source = tmp_path / "source.xlsx"
    destination = tmp_path / "work/sanitized.xlsx"
    create_style_expanded_duty_workbook(source)
    original = source.read_bytes()

    prepared, reports = prepare_oversized_workbook(
        source,
        destination,
        worksheet_size_limit=1,
        safe_content_column_limit=256,
    )

    assert prepared == destination
    assert source.read_bytes() == original
    assert reports == [{
        "sheet": "岗位职责清单",
        "original_dimension": "A1:XFD2",
        "sanitized_dimension": "A1:F2",
        "content_max_column": 6,
        "content_max_row": 2,
    }]
    loaded = load_workbook(prepared)
    original_loaded = load_workbook(source)
    assert loaded["岗位职责清单"].max_column == 6
    assert loaded["岗位职责清单"]["F2"].value == "M1"
    assert loaded["岗位职责清单"]["F2"].style_id == original_loaded["岗位职责清单"]["F2"].style_id
    loaded.close()
    original_loaded.close()


def test_prepare_oversized_workbook_removes_trailing_style_rows(tmp_path):
    """临时净化须删除真实内容之后的纯样式空行；tmp_path 为测试目录。"""
    source = tmp_path / "trailing-rows.xlsx"
    destination = tmp_path / "work/trailing-rows.xlsx"
    create_style_expanded_duty_workbook(source)
    append_trailing_style_rows(source, 5000)
    original = source.read_bytes()

    prepared, reports = prepare_oversized_workbook(
        source,
        destination,
        worksheet_size_limit=1,
        safe_content_column_limit=256,
    )

    assert source.read_bytes() == original
    assert reports == [{
        "sheet": "岗位职责清单",
        "original_dimension": "A1:XFD5000",
        "sanitized_dimension": "A1:F2",
        "content_max_column": 6,
        "content_max_row": 2,
    }]
    with ZipFile(prepared) as archive:
        worksheet = archive.read("xl/worksheets/sheet1.xml")
    assert worksheet.count(b"<row ") == 2
    loaded = load_workbook(prepared)
    assert loaded["岗位职责清单"].max_row == 2
    assert loaded["岗位职责清单"]["F2"].value == "M1"
    loaded.close()


def test_prepare_oversized_workbook_detects_compact_extreme_dimension(tmp_path):
    """XML 体积未超限时，巨大声明维度仍必须触发临时净化。

    Args:
        tmp_path: pytest 提供的隔离目录。
    """

    source = tmp_path / "compact-extreme.xlsx"
    destination = tmp_path / "work/compact-extreme.xlsx"
    create_compact_extreme_dimension_workbook(source)
    original = source.read_bytes()

    prepared, reports = prepare_oversized_workbook(source, destination)

    assert prepared == destination
    assert source.read_bytes() == original
    assert reports == [{
        "sheet": "岗位职责清单",
        "original_dimension": "A1:XFB1048572",
        "sanitized_dimension": "A1:F2",
        "content_max_column": 6,
        "content_max_row": 2,
    }]
    loaded = load_workbook(prepared)
    assert loaded["岗位职责清单"].max_row == 2
    assert loaded["岗位职责清单"].max_column == 6
    loaded.close()


def test_prepare_oversized_workbook_removes_only_out_of_bounds_merges(tmp_path):
    """净化必须删除真实内容外的合并区域并保留业务区合并。

    Args:
        tmp_path: pytest 提供的隔离目录。
    """

    source = tmp_path / "far-merges.xlsx"
    destination = tmp_path / "work/far-merges.xlsx"
    create_far_merged_range_workbook(source)

    prepared, reports = prepare_oversized_workbook(source, destination)

    assert reports[0]["sanitized_dimension"] == "A1:F2"
    loaded = load_workbook(prepared)
    assert {str(cell_range) for cell_range in loaded.active.merged_cells.ranges} == {"A1:B1"}
    loaded.close()


def test_prepare_oversized_workbook_rejects_real_content_in_far_column(tmp_path):
    """极远列存在真实内容时不得猜测删除；tmp_path 为隔离目录。"""
    source = tmp_path / "unsafe.xlsx"
    create_style_expanded_duty_workbook(source, far_value="必须保留")

    error_type = getattr(ooxml, "OversizedWorksheetError", RuntimeError)
    with pytest.raises(error_type, match="XFD.*真实内容"):
        prepare_oversized_workbook(
            source,
            tmp_path / "work/unsafe.xlsx",
            worksheet_size_limit=1,
            safe_content_column_limit=256,
        )


def test_prepare_oversized_workbook_logs_streaming_progress(tmp_path, caplog):
    """耗时净化须及时记录开始和完成；tmp_path/caplog 为文件及日志夹具。"""
    source = tmp_path / "logged.xlsx"
    create_style_expanded_duty_workbook(source)

    prepare_oversized_workbook(
        source,
        tmp_path / "work/logged.xlsx",
        worksheet_size_limit=1,
        safe_content_column_limit=256,
    )

    messages = [record.getMessage() for record in caplog.records]
    assert any("流式净化开始" in message and "岗位职责清单" in message for message in messages)
    assert any("流式净化完成" in message and "岗位职责清单" in message for message in messages)


def test_parse_files_uses_sanitized_copy_for_expanded_workbook(tmp_path, pack, monkeypatch):
    """正式读取链路须自动采用净化副本；tmp_path/pack/monkeypatch 为测试依赖。"""
    source = tmp_path / "06三清单.xlsx"
    create_style_expanded_duty_workbook(source)
    file = FileRecord(
        source,
        Path("06三清单.xlsx"),
        sha256_file(source),
        "xlsx",
        "205H",
        [],
        False,
        "06",
        "default",
        "three_lists",
    )
    file._rules_version = "1.9.16"
    monkeypatch.setattr(ooxml, "WORKSHEET_XML_SIZE_LIMIT", 1, raising=False)

    parse_files(
        [file],
        pack["field_aliases"],
        tmp_path / "run",
        parser_policy=pack.get("parser_policy"),
        semantic_lexicon=pack.get("semantic_lexicon"),
    )

    assert not file.parse_errors
    assert [sheet.sheet_type for sheet in file.sheets] == ["position_duty"]
    assert file.preservation["oversized_worksheet_cleanup"][0]["sanitized_dimension"] == "A1:F2"
    assert file.preservation["oversized_worksheet_cleanup"][0]["content_max_row"] == 2
    assert Path(file._preprocessed_path).is_file()
    prepared = load_workbook(file._preprocessed_path)
    assert prepared["岗位职责清单"].max_column == 6
    prepared.close()


def test_duty_sorting_continues_from_sanitized_copy(tmp_path, pack, monkeypatch):
    """岗位排序不得重新打开膨胀原件；tmp_path/pack/monkeypatch 为测试依赖。"""
    source = tmp_path / "06三清单.xlsx"
    create_style_expanded_duty_workbook(source)
    source_book = load_workbook(source)
    source_sheet = source_book["岗位职责清单"]
    source_sheet.append(["财务部", "出纳", "李四", "经办", "负责资金支付", "M0"])
    source_sheet["XFD3"].font = copy(source_sheet["A3"].font)
    source_book.save(source)
    source_book.close()
    file = FileRecord(
        source,
        Path("06三清单.xlsx"),
        sha256_file(source),
        "xlsx",
        "205H",
        [],
        False,
        "06",
        "default",
        "three_lists",
    )
    file._rules_version = "1.9.16"
    monkeypatch.setattr(ooxml, "WORKSHEET_XML_SIZE_LIMIT", 1, raising=False)
    work_dir = tmp_path / "run"
    parse_files(
        [file],
        pack["field_aliases"],
        work_dir,
        parser_policy=pack.get("parser_policy"),
        semantic_lexicon=pack.get("semantic_lexicon"),
    )
    # 净化完成后让原件不可读，真实排序仍应只依赖已冻结的临时副本。
    source.write_bytes(b"original workbook must not be reopened")

    sort_duties(file, {}, work_dir, pack["field_aliases"])

    prepared = load_workbook(file._preprocessed_path)
    sheet = prepared["岗位职责清单"]
    assert sheet.max_column == 6
    assert [sheet[f"F{row}"].value for row in (2, 3)] == ["M0", "M1"]
    prepared.close()
