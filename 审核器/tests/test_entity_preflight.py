"""验证主体冲突终止，未匹配名册时仍可通过命令行完成审核。"""
import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from risk_audit import runner


ROOT = Path(__file__).resolve().parents[2]


@pytest.mark.parametrize('write', [True, False])
def test_conflicting_entity_stops_before_parsing(tmp_path, monkeypatch, capsys, write):
    """tmp_path 为测试目录，monkeypatch/capsys 为替换及捕获工具，write 控制审核或试跑。"""
    input_root = tmp_path / 'input'
    input_root.mkdir()
    entities_path = tmp_path / 'entities.xlsx'
    entities = Workbook()
    entities.active.append(['单位名称', '单位代码', '上级单位', '对应主业单位', '是否存续'])
    entities.active.append(['测试主体甲有限公司', 'A001', '', '', '是'])
    entities.active.append(['测试主体乙有限公司', 'B001', '', '', '是'])
    entities.save(entities_path)
    name = '测试主体甲有限公司测试主体乙有限公司'
    source = input_root / f'06风控矩阵-{name}.xlsx'
    book = Workbook()
    book.active.append(['控制措施编号', '控制措施', '责任主体'])
    book.active.append(['M1', '核对报表', '财务部'])
    book.save(source)

    def unexpected_step(*args, **kwargs):
        """args/kwargs 为后续步骤参数；主体未确认时不允许解析、读取范围或加载基准。"""
        pytest.fail('主体未匹配时仍继续执行后续步骤')

    for step in ('parse_files', 'resolve_submission_scopes', 'load_baselines'):
        monkeypatch.setattr(runner, step, unexpected_step)
    reason = '主体证据冲突'
    with pytest.raises(ValueError, match=reason):
        runner.audit(input_root, tmp_path / 'output', ROOT / '审核器/rulepacks/releases/1.8.0',
                     entities_path, ROOT, tmp_path / 'runs', run_id='entity-preflight', write=write)
    log = (tmp_path / 'runs/entity-preflight/audit.log').read_text()
    assert 'ERROR' in log and reason in log and source.name in log
    assert '开始解析整批材料' not in log
    assert not (tmp_path / 'output').exists()
    captured = capsys.readouterr()
    assert captured.out == ''
    assert captured.err == ''
    run_directory = tmp_path / 'runs/entity-preflight'
    report_directory = run_directory / '_risk_audit' if write else run_directory
    alerts = json.loads((report_directory / 'entity_alerts.json').read_text())
    assert alerts[0]['files'][0]['file'] == source.name
    assert len(alerts[0]['files'][0]['evidence']) == 2


def test_cli_unmatched_entity_continues_with_notice(tmp_path, monkeypatch, capsys):
    """tmp_path 为测试目录，monkeypatch/capsys 为基准替换及终端捕获器；验证命令行成功和提示。"""
    from risk_audit.cli import main
    from test_business_incremental import create_business_batch, use_business_baseline
    input_root, roster_path, sources = create_business_batch(tmp_path, [('06', 'default')])
    use_business_baseline(input_root, sources, monkeypatch)
    book = Workbook()
    book.active.append(['单位名称', '单位代码', '上级单位', '对应主业单位', '是否存续'])
    book.save(roster_path)
    code = main(['audit', '--input', str(input_root), '--output', str(tmp_path / 'output'),
                 '--pack', str(ROOT / '审核器/rulepacks/releases/1.8.0'),
                 '--entities', str(roster_path),
                 '--project-root', str(ROOT), '--runs-root', str(tmp_path / 'runs')])
    captured = capsys.readouterr()
    assert code == 0
    result = json.loads(captured.out)
    assert result['write_completed'] and result['entity_code_alerts']['files'] == 2
    assert captured.err == ''
    log_text = Path(result['log_file']).read_text(encoding='utf-8')
    assert 'WARNING' in log_text and '主体不在会计主体清单中' in log_text
    assert (tmp_path / 'output').exists()
