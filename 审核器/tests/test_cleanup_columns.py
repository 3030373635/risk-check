"""审核结果中程序输出列的清理验收测试。"""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path
from zipfile import BadZipFile

import pytest
from openpyxl import Workbook, load_workbook

from risk_audit import cleanup
from risk_audit.cleanup import clean_audit_directory


TARGET_HEADERS = (
    "审核意见",
    "省公司版本责任主体（核对后删除）",
    "岗位清单已有的控制措施编号",
)


def _write_audited_workbook(path: Path) -> None:
    """生成包含三个程序输出列的工作簿。

    Args:
        path: 测试工作簿的保存路径。
    """

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    matrix = workbook.active
    matrix.title = "风控矩阵"
    matrix.append([
        "控制措施编号",
        TARGET_HEADERS[0],
        "控制措施",
        TARGET_HEADERS[1],
        TARGET_HEADERS[2],
        "责任主体",
    ])
    matrix.append(["M01", "程序意见", "核对资料", "省公司", "M01", "财务部"])

    duty = workbook.create_sheet("岗位清单")
    duty.append(["岗位名称", "历史审核意见", "备注"])
    duty.append(["会计", "人工意见", "保留"])
    workbook.save(path)


def test_cleanup_removes_exact_program_columns_and_preserves_source(tmp_path: Path) -> None:
    """清理必须删除三个精确列名且不修改源文件。

    Args:
        tmp_path: pytest 提供的隔离临时目录。
    """

    source_root = tmp_path / "审核结果"
    output_root = tmp_path / "清理结果"
    source = source_root / "业务目录" / "矩阵.xlsx"
    _write_audited_workbook(source)
    source_bytes = source.read_bytes()

    summary = clean_audit_directory(source_root, output_root)

    cleaned = load_workbook(output_root / "业务目录" / "矩阵.xlsx")
    matrix = cleaned["风控矩阵"]
    assert [matrix.cell(1, column).value for column in range(1, matrix.max_column + 1)] == [
        "控制措施编号",
        "控制措施",
        "责任主体",
    ]
    assert [matrix.cell(2, column).value for column in range(1, matrix.max_column + 1)] == [
        "M01",
        "核对资料",
        "财务部",
    ]
    assert cleaned["岗位清单"]["B1"].value == "历史审核意见"
    assert source.read_bytes() == source_bytes
    assert summary == {
        "processed_workbooks": 1,
        "unchanged_files": 0,
        "removed_columns": 3,
    }


def test_cleanup_copies_unmatched_and_non_excel_files_without_rewriting(tmp_path: Path) -> None:
    """无目标列的工作簿和非 Excel 文件必须按原字节复制。

    Args:
        tmp_path: pytest 提供的隔离临时目录。
    """

    source_root = tmp_path / "审核结果"
    output_root = tmp_path / "清理结果"
    workbook_path = source_root / "无意见列.xlsx"
    workbook_path.parent.mkdir(parents=True)
    workbook = Workbook()
    workbook.active.append(["审核意见说明", "备注"])
    workbook.save(workbook_path)
    note_path = source_root / "附件" / "说明.txt"
    note_path.parent.mkdir(parents=True)
    note_path.write_bytes(b"keep-original-bytes")
    workbook_bytes = workbook_path.read_bytes()

    summary = clean_audit_directory(source_root, output_root)

    assert (output_root / "无意见列.xlsx").read_bytes() == workbook_bytes
    assert (output_root / "附件" / "说明.txt").read_bytes() == b"keep-original-bytes"
    assert summary == {
        "processed_workbooks": 0,
        "unchanged_files": 2,
        "removed_columns": 0,
    }


def test_cleanup_skips_office_temporary_files(tmp_path: Path) -> None:
    """Office/WPS 临时文件即使使用 .xlsx 后缀也不得被读取或复制。

    Args:
        tmp_path: pytest 提供的隔离临时目录。
    """

    source_root = tmp_path / "审核结果"
    output_root = tmp_path / "清理结果"
    _write_audited_workbook(source_root / "矩阵.xlsx")
    for name in (".~WPS锁文件.xlsx", "~$Office锁文件.xlsx", ".DS_Store"):
        (source_root / name).write_bytes(b"not-an-excel-workbook")

    summary = clean_audit_directory(source_root, output_root)

    assert (output_root / "矩阵.xlsx").is_file()
    assert not (output_root / ".~WPS锁文件.xlsx").exists()
    assert not (output_root / "~$Office锁文件.xlsx").exists()
    assert not (output_root / ".DS_Store").exists()
    assert summary == {
        "processed_workbooks": 1,
        "unchanged_files": 0,
        "removed_columns": 3,
    }


def test_cleanup_discards_partial_output_after_workbook_failure(tmp_path: Path) -> None:
    """任一工作簿失败时必须清掉临时产物，不得留下半成品输出目录。

    Args:
        tmp_path: pytest 提供的隔离临时目录。
    """

    source_root = tmp_path / "审核结果"
    output_root = tmp_path / "清理结果"
    _write_audited_workbook(source_root / "01-正常.xlsx")
    (source_root / "02-损坏.xlsx").write_bytes(b"not-an-excel-workbook")

    with pytest.raises(BadZipFile):
        clean_audit_directory(source_root, output_root)

    assert not output_root.exists()
    assert not list(tmp_path.glob(".清理结果-*"))


def test_cleanup_sanitizes_expanded_workbook_before_deleting_columns(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """清理必须先净化远端空白合并，再删除程序输出列。

    Args:
        tmp_path: pytest 提供的隔离临时目录。
        monkeypatch: 用于降低测试净化触发阈值的 pytest 工具。
    """

    source_root = tmp_path / "审核结果"
    output_root = tmp_path / "清理结果"
    source = source_root / "异常范围.xlsx"
    _write_audited_workbook(source)
    workbook = load_workbook(source)
    workbook["风控矩阵"].merge_cells("Z20:AA21")
    workbook.save(source)
    workbook.close()
    monkeypatch.setattr(cleanup, "CLEANUP_DIMENSION_CELL_LIMIT", 100, raising=False)

    clean_audit_directory(source_root, output_root)

    cleaned = load_workbook(output_root / "异常范围.xlsx")
    matrix = cleaned["风控矩阵"]
    assert "Z20:AA21" not in {str(cell_range) for cell_range in matrix.merged_cells.ranges}
    assert [matrix.cell(1, column).value for column in range(1, matrix.max_column + 1)] == [
        "控制措施编号",
        "控制措施",
        "责任主体",
    ]
    cleaned.close()


@pytest.mark.parametrize("relationship", ["same", "output_inside_input", "input_inside_output"])
def test_cleanup_rejects_overlapping_directories(tmp_path: Path, relationship: str) -> None:
    """输入输出目录相同或相互包含时必须在写入前拒绝。

    Args:
        tmp_path: pytest 提供的隔离临时目录。
        relationship: 待验证的目录位置关系。
    """

    if relationship == "same":
        source_root = output_root = tmp_path / "材料"
    elif relationship == "output_inside_input":
        source_root = tmp_path / "材料"
        output_root = source_root / "清理后"
    else:
        output_root = tmp_path / "材料"
        source_root = output_root / "审核后"
    source_root.mkdir(parents=True)

    with pytest.raises(ValueError, match="不能相同或相互包含"):
        clean_audit_directory(source_root, output_root)


def test_cleanup_rejects_existing_output_directory(tmp_path: Path) -> None:
    """已存在的输出目录必须被拒绝，防止混入旧结果。

    Args:
        tmp_path: pytest 提供的隔离临时目录。
    """

    source_root = tmp_path / "审核结果"
    output_root = tmp_path / "清理结果"
    source_root.mkdir()
    output_root.mkdir()

    with pytest.raises(FileExistsError, match="输出目录已存在"):
        clean_audit_directory(source_root, output_root)


def test_cleanup_rejects_missing_input_directory(tmp_path: Path) -> None:
    """输入目录不存在时必须报错，不得生成伪成功的空结果。

    Args:
        tmp_path: pytest 提供的隔离临时目录。
    """

    source_root = tmp_path / "不存在"
    output_root = tmp_path / "清理结果"

    with pytest.raises(FileNotFoundError, match="输入目录不存在"):
        clean_audit_directory(source_root, output_root)
    assert not output_root.exists()


def test_cleanup_rejects_input_file(tmp_path: Path) -> None:
    """--input 指向单个文件时必须报错，避免生成空目录。

    Args:
        tmp_path: pytest 提供的隔离临时目录。
    """

    source_file = tmp_path / "审核结果.xlsx"
    source_file.write_bytes(b"not-used")
    output_root = tmp_path / "清理结果"

    with pytest.raises(ValueError, match="输入路径必须是目录"):
        clean_audit_directory(source_file, output_root)
    assert not output_root.exists()


def test_cleanup_script_creates_clean_directory_and_prints_summary(
    tmp_path: Path,
    project_root: Path,
) -> None:
    """根目录脚本必须能直接执行并输出清理汇总。

    Args:
        tmp_path: pytest 提供的隔离临时目录。
        project_root: 项目根目录。
    """

    source_root = tmp_path / "审核结果"
    output_root = tmp_path / "清理结果"
    _write_audited_workbook(source_root / "矩阵.xlsx")

    completed = subprocess.run(
        [
            sys.executable,
            str(project_root / "clean_audit_columns.py"),
            "--input",
            str(source_root),
            "--output",
            str(output_root),
        ],
        cwd=project_root,
        text=True,
        capture_output=True,
        check=False,
    )

    assert completed.returncode == 0, completed.stderr
    assert json.loads(completed.stdout) == {
        "processed_workbooks": 1,
        "removed_columns": 3,
        "unchanged_files": 0,
    }
    assert (output_root / "矩阵.xlsx").is_file()
