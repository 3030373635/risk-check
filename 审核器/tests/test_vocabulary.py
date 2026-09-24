from pathlib import Path
from types import SimpleNamespace
import copy
import json
import pytest
from openpyxl import Workbook
from risk_audit.inventory import identify_business
from risk_audit.vocabulary import carrier_candidates,admissible_term,select_mining_sheets
from risk_audit.configuration.validator import validate_pack,ConfigError

def test_numbered_company_folder_is_not_business_number():
    assert identify_business(Path('10湘西自治州德能电力建设有限公司泸溪县分公司/06 设备（资产）管理/06风控矩阵.xlsx'))==('06','default')
    assert identify_business(Path('01 营销售电/三清单-02交易与购电.xlsx'))[0] is None

def test_carrier_candidate_parenthesis_and_audit_isolation():
    assert carrier_candidates('1.采购验收记录（采购记录；验收单）；2.验收与移交记录。')==['采购验收记录（采购记录；验收单）','验收与移交记录']
    assert not admissible_term('9.11初审：请明确到具体岗位。','position')
    assert not admissible_term('对记录准确性負责任。','carrier')
    assert admissible_term('项目负责人','position')  # Observation does not mean approved or invalid.
    assert carrier_candidates('1．《中标通知书》、合同；2．验收单（含移交、清点）') == ['中标通知书','合同','验收单（含移交、清点）']

def test_hidden_is_excluded_from_mining_even_without_a_visible_counterpart(tmp_path):
    p=tmp_path/'a.xlsx';w=Workbook();w.active.title='岗位职责清单';w.active.sheet_state='hidden';w.create_sheet('说明');w.save(p)
    f=SimpleNamespace(source=p,sheets=[SimpleNamespace(title='岗位职责清单',sheet_type='position_duty')])
    assert select_mining_sheets(f)[0]==set()
    assert '隐藏' in select_mining_sheets(f)[1]['岗位职责清单']
    w.create_sheet('岗位职责清单9.8');w.save(p)
    f.sheets.append(SimpleNamespace(title='岗位职责清单9.8',sheet_type='position_duty'))
    assert select_mining_sheets(f)[0]=={'岗位职责清单9.8'}

def test_vocabulary_configuration_is_optional_but_strict(pack,registry):
    validate_pack(pack,registry)
    pack['terminology']={'schema_version':'1.0','terms':{'carrier':['验收单']},'metadata':{}}
    validate_pack(pack,registry)
    pack['terminology']['terms']['carrier'].append('审核意见')
    with pytest.raises(ConfigError):validate_pack(pack,registry)
