"""第17条删除修复的发布与交付验收。"""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

from risk_audit import __version__
from risk_audit.configuration.loader import load_pack
from risk_audit.util import sha256_file


ROOT = Path(__file__).resolve().parents[1]


def test_active_rulepack_removes_rule_17() -> None:
    """当前程序与激活规则包必须继承1.9.10的第17条删除结果。"""
    assert __version__ == "1.9.19"
    active = json.loads((ROOT / "rulepacks/active.json").read_text(encoding="utf-8"))
    assert active["version"] == "1.9.19"
    pack = load_pack(ROOT / "rulepacks/releases" / active["version"])

    assert all(rule["rule_id"] != "matrices.responsibility_department" for rule in pack["rules"])
    assert all(rule["display_code"] != "R17" for rule in pack["rules"])
    assert pack["manifest"]["source_document_sha256"] == sha256_file(
        ROOT / "规则来源/0923-2.docx"
    )


def test_remove_rule_17_publish_script_rebuilds_release() -> None:
    """1.9.10规则包必须可从1.9.9基线稳定重建，且只读复验不改变激活版本。"""
    result = subprocess.run(
        [sys.executable, str(ROOT / "tools/publish_remove_rule_17_1910.py"), "--verify-only"],
        cwd=ROOT.parent,
        capture_output=True,
        text=True,
        check=False,
    )

    assert result.returncode == 0, result.stderr
    payload = json.loads(result.stdout)
    assert payload["verified"] is True
    assert payload["activated"] is False
    assert json.loads((ROOT / "rulepacks/active.json").read_text(encoding="utf-8"))["version"] == "1.9.19"


def test_audited_excel_does_not_contain_rule_17(tmp_path: Path) -> None:
    """tmp_path 为临时交付目录；审核意见 Excel 不得再写入 R17 或第17条。"""
    from run_audit import configure_soffice
    from risk_audit.runner import audit

    configure_soffice()
    input_root = tmp_path / "input/信通公司"
    input_root.mkdir(parents=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "风控矩阵"
    worksheet.append(["控制措施编号", "控制措施", "责任主体", "是否适用及原因"])
    worksheet.append(["设备业务-1.设备验收-控制措施01", "核对设备", "信通公司-财务部", "适用"])
    workbook.save(input_root / "06信通公司风控矩阵.xlsx")

    result = audit(
        input_root.parent,
        tmp_path / "output",
        ROOT / "rulepacks/releases/1.9.10",
        ROOT.parent / "审核/会计主体清单20260907.xlsx",
        ROOT.parent,
        tmp_path / "runs",
    )

    output_file = next((tmp_path / "output").rglob("06信通公司风控矩阵.xlsx"))
    audited = load_workbook(output_file, read_only=True, data_only=True)
    values = [str(cell) for row in audited["风控矩阵"].iter_rows(values_only=True) for cell in row if cell]
    audited.close()
    assert not any("R17" in value or "第17条" in value for value in values)
    assert "feedback_report" not in result
