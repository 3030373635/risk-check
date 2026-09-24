from pathlib import Path
import pytest
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from risk_audit.models import FileRecord
from risk_audit.readers.excel import parse_workbook
from risk_audit.util import sha256_file


def parse(path, wb, pack, material='three_lists'):
    wb.save(path)
    f=FileRecord(path,Path(path.name),sha256_file(path),'xlsx','205H',['确认简称:信通公司'],False,'06','default',material)
    f.sheets=parse_workbook(f,path,pack['field_aliases'])
    assert sha256_file(path)==f.sha256
    return f,f.sheets[0]


def duty_book(row=2):
    w=Workbook();s=w.active;s.title='清单2-岗位责任清单'
    s.cell(1,1,'模板使用说明')
    for col,text in enumerate(['部门','岗位名称','部门','岗位名称','姓名','角色','岗位职责','控制措施编号'],1):s.cell(row,col,text)
    for col,text in enumerate(['旧部门','旧岗位','新部门','新岗位','张三','经办','对记录准确性负主体责任','M1'],1):s.cell(row+1,col,text)
    return w,s


@pytest.mark.parametrize('value',[None,'=1'])
def test_visible_duplicate_wins_without_row_fallback(tmp_path,pack,value):
    w,s=duty_book(7);s.column_dimensions.group('A','B',hidden=True)
    s['C8']=value
    f,p=parse(tmp_path/'a.xlsx',w,pack)
    assert p.columns['department']==3 and p.columns['position']==4
    assert p.records[0].value('department')==''
    assert p.records[0].value('position')=='新岗位'
    assert p.records[0].fields['department'].state==('empty' if value is None else 'formula_no_cache')
    assert any(x['reason']=='visible_equivalent_preferred' for x in f.preservation['column_selections'])


def test_only_hidden_field_is_still_required_input(tmp_path,pack):
    w,s=duty_book();s.delete_cols(3,2)
    s.column_dimensions.group('A','B',hidden=True)
    f,p=parse(tmp_path/'b.xlsx',w,pack)
    assert p.records[0].value('department')=='旧部门'
    assert p.records[0].value('position')=='旧岗位'
    assert {x['reason'] for x in f.preservation['column_selections']}=={'hidden_field_required'}


def test_visible_current_unit_department_selects_its_position_pair(tmp_path,pack):
    w,s=duty_book();s['A3']='别的公司-财务部';s['C3']='信通公司-项目管理中心'
    _,p=parse(tmp_path/'c.xlsx',w,pack)
    assert p.columns['department']==3 and p.columns['position']==4


def test_unresolved_visible_duplicates_are_not_guessed(tmp_path,pack):
    w,s=duty_book()
    f,p=parse(tmp_path/'d.xlsx',w,pack)
    assert 'department' not in p.columns and 'position' not in p.columns
    assert {x['reason'] for x in f.preservation['column_selections']}=={'column_conflict'}


def test_hidden_strike_source_is_kept_only_when_no_visible_equivalent(tmp_path,pack):
    w=Workbook();s=w.active;s.title='风控矩阵';s['A1']='风控矩阵'
    for c,h in enumerate(['控制措施编号','控制措施','责任主体','控制载体','控制载体'],1):s.cell(2,c,h)
    for c,v in enumerate(['M1','现行措施','信通公司-财务部-主任','旧参考载体','现行载体'],1):s.cell(3,c,v)
    s['D3']=CellRichText([TextBlock(InlineFont(strike=True),'参考删除原文')])
    s.column_dimensions['D'].hidden=True
    _,p=parse(tmp_path/'e.xlsx',w,pack,'matrix')
    assert p.records[0].value('carrier')=='现行载体' and not p.records[0].fields['carrier'].deleted_spans
    s.delete_cols(5)
    _,p=parse(tmp_path/'f.xlsx',w,pack,'matrix')
    assert p.records[0].fields['carrier'].deleted_spans[0]['text']=='参考删除原文'


def test_wrong_unit_applicability_is_not_relabelled(tmp_path,pack):
    w=Workbook();s=w.active;s.title='风控矩阵';s['A1']='风控矩阵'
    for c,h in enumerate(['控制措施编号','控制措施','责任主体','其他公司适用情况'],1):s.cell(2,c,h)
    for c,v in enumerate(['M1','措施','信通公司-财务部-主任','是'],1):s.cell(3,c,v)
    f,p=parse(tmp_path/'g.xlsx',w,pack,'matrix')
    assert 'applicability' not in p.columns
    assert f.preservation['column_selections'][0]['reason']=='owner_unconfirmed'


def test_generic_named_sheet_is_recognized_by_complete_business_headers(tmp_path,pack):
    w,s=duty_book(12);s.delete_cols(1,2);s.title='Sheet1';s['A12']='部门（综合管理部）'
    _,p=parse(tmp_path/'h.xlsx',w,pack)
    assert p.sheet_type=='position_duty' and p.records[0].row==13
