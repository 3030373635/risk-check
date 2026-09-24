from __future__ import annotations

import zipfile
from pathlib import Path

from lxml import etree
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Font

from risk_audit.inventory import true_format
from risk_audit.models import FileRecord
from risk_audit.readers.excel import current_text, parse_workbook
from risk_audit.readers.xls import convert_xls
from risk_audit.util import sha256_file


def fixture_book(path: Path, manual="人工A", sparse=False):
    wb = Workbook(); ws = wb.active; ws.title = "岗位职责清单 "
    ws["A1"] = "岗位内控责任清单"
    headers = ["部门", "岗位名称", "姓名", "岗位职责编号", "角色", "岗位职责", "控制措施编号", "9.11初审"]
    for c, v in enumerate(headers, 1): ws.cell(2, c, v)
    values = ["部门", "岗位", "张三", "", "经办", "对事项负有主体责任", "", manual]
    for c, v in enumerate(values, 1): ws.cell(3, c, v)
    if sparse: ws.cell(1, 16384).font = Font(bold=True)
    wb.save(path)


def aliases(pack): return pack["field_aliases"]


def parse_fixture(path, pack):
    f = FileRecord(path, Path(path.name), sha256_file(path), "xlsx", "205H", [], False, "06", "default", "three_lists")
    return parse_workbook(f, path, aliases(pack))[0]


def test_missing_measure_row_is_not_filtered_and_sparse_width_ignored(tmp_path, pack):
    path = tmp_path / "06三清单.xlsx"; fixture_book(path, sparse=True)
    sheet = parse_fixture(path, pack)
    assert len(sheet.records) == 1 and sheet.records[0].value("measure_id") == ""
    assert sheet.output_column == 9


def test_manual_answer_isolated_from_business_records(tmp_path, pack):
    a = tmp_path / "a.xlsx"; b = tmp_path / "b.xlsx"; fixture_book(a, "人工答案A"); fixture_book(b, "完全改写")
    ra, rb = parse_fixture(a, pack).records[0], parse_fixture(b, pack).records[0]
    assert {k: v.current for k, v in ra.fields.items()} == {k: v.current for k, v in rb.fields.items()}
    assert "manual_review" not in ra.fields


def test_rich_text_strike_removes_only_explicit_struck_run():
    wb = Workbook(); c = wb.active["A1"]
    c.value = CellRichText([TextBlock(InlineFont(strike=True), "旧段"), TextBlock(InlineFont(color="FFFF0000"), "新红段")])
    c.font = Font(strike=True)
    assert current_text(c) == "新红段"


def test_real_xls_signature_and_cache_preserving_conversion(tmp_path, project_root):
    source = project_root / "审核/国网湖南信通公司第一批风控矩阵应用落地资料-9.11初审/国网湖南信通公司第一批风控矩阵应用落地资料-9.11初审/09 职工福利保障与薪酬管理-9.11初审/国网湖南信通公司 09风控矩阵-职工福利保障与薪酬管理.xls"
    assert true_format(source) == "xls"
    converted, report = convert_xls(source, tmp_path)
    assert converted.suffix == ".xlsx" and report["passed"] and report["formula_caches_restored"] > 0

