from __future__ import annotations

import copy
import json
import math
import os
import unicodedata
import re
import shutil
import tempfile
import zipfile
from bisect import bisect_left
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable

from lxml import etree
from openpyxl.utils import get_column_letter, column_index_from_string

from risk_audit.models import Entity, FileRecord, Finding
from risk_audit.readers.ooxml import load_compatible_workbook
from risk_audit.util import sha256_file, write_json

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKGREL = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"m": MAIN, "r": REL}


@dataclass
class OutputState:
    """增量状态；ownership 为本轮归属，pending/unparsed 为累计诊断，previous_ownership 用于变更检测。"""
    ownership: dict[str, list[dict[str, Any]]] = field(default_factory=dict)
    pending: list[dict[str, Any]] = field(default_factory=list)
    unparsed: list[dict[str, Any]] = field(default_factory=list)
    previous_ownership: dict[str, list[dict[str, Any]]] | None = None


class WorkbookOutputCompatibilityError(ValueError):
    """表示单个工作簿的原始布局无法安全写入审核结果。"""

    def __init__(self, message: str, *, sheet: str, cell_range: str) -> None:
        """初始化兼容性异常；message 为原因，sheet 为工作表，cell_range 为冲突区域。"""
        super().__init__(message)
        self.sheet = sheet
        self.cell_range = cell_range


def _sheet_paths(zf: zipfile.ZipFile) -> dict[str, str]:
    wb = etree.fromstring(zf.read("xl/workbook.xml"))
    rels = etree.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    targets = {x.get("Id"): x.get("Target") for x in rels.findall(f"{{{PKGREL}}}Relationship")}
    out = {}
    for s in wb.xpath("//m:sheets/m:sheet", namespaces=NS):
        target = targets[s.get(f"{{{REL}}}id")]
        out[s.get("name")] = "xl/" + target.lstrip("/").removeprefix("xl/")
    return out


def _coord_col(ref: str) -> int:
    return column_index_from_string(re.match(r"[A-Z]+", ref).group())


def _inline_cell(row: etree._Element, coord: str, text: str, style: str | None) -> etree._Element:
    for old in row.findall(f"{{{MAIN}}}c"):
        if old.get("r") == coord:
            row.remove(old); break
    c = etree.Element(f"{{{MAIN}}}c", r=coord, t="inlineStr")
    if style is not None: c.set("s", style)
    isel = etree.SubElement(c, f"{{{MAIN}}}is")
    t = etree.SubElement(isel, f"{{{MAIN}}}t")
    if text.startswith(" ") or text.endswith(" ") or "\n" in text: t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    cells = list(row.findall(f"{{{MAIN}}}c"))
    before = next((x for x in cells if _coord_col(x.get("r")) > _coord_col(coord)), None)
    if before is None: row.append(c)
    else: row.insert(row.index(before), c)
    return c


def _index_rows(sheet_data: etree._Element) -> tuple[dict[int, etree._Element], list[int]]:
    """建立工作表行索引。

    sheet_data 为 OOXML sheetData 节点；返回按行号定位的映射和有序行号列表。
    """
    row_index = {
        int(row.get("r")): row
        for row in sheet_data.findall(f"{{{MAIN}}}row")
    }
    return row_index, sorted(row_index)


def _ensure_row(
        sheet_data: etree._Element,
        number: int,
        row_index: dict[int, etree._Element] | None = None,
        row_numbers: list[int] | None = None,
) -> etree._Element:
    """返回或创建指定行。

    sheet_data 为 OOXML sheetData 节点，number 为目标行号，row_index 和
    row_numbers 为可复用的行索引；未传入索引时为单次调用自动建立。
    """
    if row_index is None or row_numbers is None:
        row_index, row_numbers = _index_rows(sheet_data)
    if number in row_index:
        return row_index[number]
    position = bisect_left(row_numbers, number)
    row = etree.Element(f"{{{MAIN}}}row", r=str(number))
    # sheetData 的子节点即为按行号排列的 row，插入后同步索引。
    sheet_data.insert(position, row)
    row_numbers.insert(position, number)
    row_index[number] = row
    return row


def _select_output_header_row(
        header_rows: list[int], is_hidden: Callable[[int], bool], default_row: int = 1) -> int:
    """选择程序新增列的可见表头行。

    header_rows 为读取器识别的表头行，is_hidden 用于判断指定行是否隐藏，
    default_row 为未识别到表头行时的回退行号。多层表头优先使用最后一个可见行，
    全部隐藏时保留原有末行回退行为。
    """
    rows = sorted(set(header_rows)) or [default_row]
    return next((row for row in reversed(rows) if not is_hidden(row)), rows[-1])


def _hidden_rows_by_sheet(workbook_path: Path) -> dict[str, set[int]]:
    """读取工作簿中每张表的隐藏行。

    workbook_path 为 OOXML 工作簿路径，返回值按工作表名保存隐藏行号集合。
    """
    with zipfile.ZipFile(workbook_path) as archive:
        paths = _sheet_paths(archive)
        hidden = {}
        for title, path in paths.items():
            root = etree.fromstring(archive.read(path))
            hidden[title] = {
                int(row.get('r')) for row in root.xpath('//m:sheetData/m:row', namespaces=NS)
                if row.get('hidden') in {'1', 'true', 'True'}
            }
        return hidden


def _nearest_style(row: etree._Element, target_col: int) -> str | None:
    candidates = [c for c in row.findall(f"{{{MAIN}}}c") if _coord_col(c.get("r")) < target_col]
    return candidates[-1].get("s") if candidates else None


def patch_ooxml(source: Path, destination: Path, operations: dict[str, dict[str, Any]]) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    same_path = source.resolve() == destination.resolve()
    fd, temporary = tempfile.mkstemp(prefix=".risk-audit-", suffix=".xlsx", dir=destination.parent)
    os.close(fd)
    actual_destination = Path(temporary)
    with zipfile.ZipFile(source, "r") as zin:
        paths = _sheet_paths(zin)
        replacements: dict[str, bytes] = {}
        styles = etree.fromstring(zin.read("xl/styles.xml"))
        cell_xfs = styles.find(f"{{{MAIN}}}cellXfs")
        style_cache = {}
        def wrapped_style(base_style: str | None) -> str:
            base = int(base_style or 0)
            if base in style_cache: return style_cache[base]
            xf = copy.deepcopy(cell_xfs[base] if base < len(cell_xfs) else cell_xfs[0])
            for old in xf.findall(f"{{{MAIN}}}alignment"): xf.remove(old)
            etree.SubElement(xf, f"{{{MAIN}}}alignment", wrapText="1", vertical="top")
            xf.set("applyAlignment", "1"); cell_xfs.append(xf); cell_xfs.set("count", str(len(cell_xfs)))
            style_cache[base] = str(len(cell_xfs) - 1); return style_cache[base]
        for sheet_name, spec in operations.items():
            path = paths[sheet_name]
            root = etree.fromstring(zin.read(path))
            sheet_data = root.find(f"{{{MAIN}}}sheetData")
            row_index, row_numbers = _index_rows(sheet_data)
            col = spec["column"]; letter = get_column_letter(col)
            header_row = spec["header_row"]
            # 同一工作表有多块清单区域时，各区域表头在同一意见列标明用途。
            for opinion_header_row in spec.get('header_rows', [header_row]):
                hrow = _ensure_row(sheet_data, opinion_header_row, row_index, row_numbers)
                _inline_cell(hrow, f"{letter}{opinion_header_row}", "审核意见", wrapped_style(_nearest_style(hrow, col)))
            for row_num, text in sorted(spec["values"].items()):
                row = _ensure_row(sheet_data, row_num, row_index, row_numbers)
                if text == "":
                    for existing in list(row.findall(f"{{{MAIN}}}c")):
                        if existing.get("r") == f"{letter}{row_num}": row.remove(existing)
                    continue
                original = next((cell for cell in row.findall(f'{{{MAIN}}}c') if cell.get('r') == f'{letter}{row_num}'), None)
                style = original.get('s') if original is not None else _nearest_style(row, col)
                cell = _inline_cell(row, f"{letter}{row_num}", text, wrapped_style(style))
                if row_num in spec.get('rich_values', {}):
                    # 人工意见的富文本直接复制，程序意见另加明确的正常字体段落。
                    cell.remove(cell.find(f'{{{MAIN}}}is'))
                    content = etree.fromstring(spec['rich_values'][row_num].encode())
                    for child in list(content):
                        if child.tag == f'{{{MAIN}}}t':
                            position = content.index(child); content.remove(child)
                            run = etree.Element(f'{{{MAIN}}}r'); run.append(child); content.insert(position, run)
                    program = spec.get('program_values', {}).get(row_num, '')
                    if program:
                        run = etree.SubElement(content, f'{{{MAIN}}}r'); properties = etree.SubElement(run, f'{{{MAIN}}}rPr')
                        etree.SubElement(properties, f'{{{MAIN}}}strike', val='0'); etree.SubElement(properties, f'{{{MAIN}}}color', rgb='FF000000')
                        value = etree.SubElement(run, f'{{{MAIN}}}t'); value.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                        value.text = '\n' + program
                    cell.append(content)
                # Account for visual wrapping, including CJK full-width glyphs.
                base_style = int(_nearest_style(row, col) or 0)
                xf = cell_xfs[base_style] if base_style < len(cell_xfs) else cell_xfs[0]
                font = styles.find(f"{{{MAIN}}}fonts")[int(xf.get("fontId", 0))]
                size = font.find(f"{{{MAIN}}}sz")
                points = float(size.get("val", 11)) if size is not None else 11
                available_points = (45 * 7 - 12) * 0.75
                line_count = sum(max(1, math.ceil(sum(points if unicodedata.east_asian_width(ch) in "WF" else points * 0.58 for ch in line) / available_points)) for line in text.split("\n"))
                current_height = float(row.get("ht", 0) or 0)
                row.set("ht", str(max(current_height, min(409, points * 1.5 * line_count + 5)))); row.set("customHeight", "1")
            cols = root.find(f"{{{MAIN}}}cols")
            if cols is None:
                cols = etree.Element(f"{{{MAIN}}}cols"); sheet_data.addprevious(cols)
            for old_col in list(cols.findall(f"{{{MAIN}}}col")):
                if old_col.get("min") == str(col) and old_col.get("max") == str(col): cols.remove(old_col)
            new_col = etree.SubElement(cols, f"{{{MAIN}}}col", min=str(col), max=str(col), width="45", customWidth="1")
            dim = root.find(f"{{{MAIN}}}dimension")
            if dim is not None:
                ref = dim.get("ref", "A1")
                first, _, last = ref.partition(":")
                last_row = int(re.search(r"\d+", last or first).group())
                # 复用意见列可能不在末尾，不能截断右侧原业务列；空表首条意见也须在范围内。
                last_column = get_column_letter(max(col, _coord_col(last or first)))
                last_row = max(last_row, *spec.get('header_rows', [header_row]), *spec['values'].keys())
                dim.set("ref", f"{first}:{last_column}{last_row}")
            replacements[path] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        replacements["xl/styles.xml"] = etree.tostring(styles, xml_declaration=True, encoding="UTF-8", standalone=True)
        with zipfile.ZipFile(actual_destination, "w") as zout:
            for info in zin.infolist():
                data = replacements[info.filename] if info.filename in replacements else zin.read(info.filename)
                zout.writestr(info, data)
    actual_destination.replace(destination)


def compare_ooxml_preservation(source: Path, destination: Path, operations: dict[str, dict[str, Any]]) -> dict[str, Any]:
    differences = []
    with zipfile.ZipFile(source) as za, zipfile.ZipFile(destination) as zb:
        paths_a = _sheet_paths(za); paths_b = _sheet_paths(zb)
        touched = {paths_a[x] for x in operations} | {"xl/styles.xml"}
        unchanged_parts = 0
        for name in set(za.namelist()) & set(zb.namelist()):
            if name in touched: continue
            if za.read(name) != zb.read(name): differences.append({"part": name, "reason": "unexpected package-part change"})
            else: unchanged_parts += 1
        for sheet, spec in operations.items():
            ra = etree.fromstring(za.read(paths_a[sheet])); rb = etree.fromstring(zb.read(paths_b[sheet]))
            old_cells = {c.get("r"): etree.tostring(c, method="c14n") for c in ra.xpath("//m:sheetData/m:row/m:c", namespaces=NS)}
            new_cells = {c.get("r"): etree.tostring(c, method="c14n") for c in rb.xpath("//m:sheetData/m:row/m:c", namespaces=NS)}
            audit_letter = get_column_letter(spec["column"])
            for coord, xml in old_cells.items():
                if _coord_col(coord) == spec["column"]: continue
                if new_cells.get(coord) != xml: differences.append({"sheet": sheet, "cell": coord, "reason": "existing cell XML changed"})
    return {"passed": not differences, "differences": differences[:100], "unchanged_package_parts": unchanged_parts, "allowed_changes": "仅正式业务表新增/重建单一审核意见列、列宽及dimension"}


def _detect_owned_edits(destination: Path, ownership: dict[str, Any], relative: str) -> list[dict[str, Any]]:
    """检测旧副本中的程序意见是否被修改；参数为副本、旧归属和相对路径。"""
    if not destination.exists():
        return []
    wb = load_compatible_workbook(
        destination, read_only=False, data_only=False, rich_text=True,
    )
    edits = []
    for item in ownership.get(relative, []):
        if item["sheet"] not in wb.sheetnames:
            continue
        current = str(wb[item["sheet"]][item["cell"]].value or "")
        expected = item.get("rendered_text", item.get("text", ""))
        if current != expected:
            edits.append({**item, "current": current})
    wb.close()
    return edits


def _join_opinions(messages: list[str], order: dict[str, tuple[int, int]]) -> str:
    """按规则及检查顺序合并意见；messages 为文本，order 为对应的规则编号和检查次序。"""
    return '\n'.join(dict.fromkeys(sorted(messages, key=lambda message: order.get(message, (0, 0)))))


def _output_path(file: FileRecord) -> Path:
    """确定副本相对路径；file 为已解析或原样保留的输入文件，只有解析成功的旧格式转为 XLSX。"""
    if file.sheets and (file.true_format == 'xls' or file.relative_path.suffix.lower() == '.et'):
        return file.relative_path.with_suffix('.xlsx')
    return file.relative_path


def validate_output_paths(files: list[FileRecord], *, anticipate_conversion: bool = False) -> None:
    """检查整批目标碰撞；files 为输入，anticipate_conversion 控制解析前预检旧格式转为 XLSX 的目标。

    名称统一按大小写不敏感及 NFC 规范比较。
    """
    targets = {}
    for file in files:
        relative = _output_path(file)
        # 逐主体解析时，首份副本写出前也必须检查尚未解析的旧格式潜在目标。
        if anticipate_conversion and (file.true_format == 'xls' or file.relative_path.suffix.lower() == '.et'):
            relative = file.relative_path.with_suffix('.xlsx')
        # macOS/Windows可能将不同大小写或Unicode组合形式视作同一个文件，须防止覆盖。
        key = unicodedata.normalize('NFC', str(relative)).casefold()
        if key in targets:
            raise ValueError(f'审核副本输出路径重复：{targets[key]} 与 {file.relative_path} 均对应 {relative}，请明确正式文件或调整文件名。')
        targets[key] = file.relative_path


def write_outputs(files: list[FileRecord], findings: list[Finding], output_root: Path, *, metadata_dir: Path, entities: dict[str, Entity] | None = None, output_state: OutputState | None = None, baselines: dict | None = None) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """写审核副本。

    files 为输入，findings 为意见，output_root 为副本目标，entities 为名册，
    output_state 为同批累计状态，baselines 为新版核对列基准，metadata_dir 为审核辅助资料目录。
    """
    if any(f.evidence.get('publication_channel') == 'internal' for f in findings):
        raise ValueError('内部未完成检查不得写入单位审核意见；请修正结果分流后重试。')
    validate_output_paths(files)
    output_state = output_state if output_state is not None else OutputState()
    ownership_path = metadata_dir / "ownership.json"
    if output_state.previous_ownership is None:
        # 首次写出前冻结旧归属，逐主体写出时据此检测旧副本中的人工改写。
        output_state.previous_ownership = json.loads(ownership_path.read_text(encoding="utf-8")) if ownership_path.exists() else {}
        # 尚未处理主体的旧归属暂存到本批状态，避免增量写出提前删除其变更检测依据。
        output_state.ownership = dict(output_state.previous_ownership)
    old_ownership = output_state.previous_ownership
    by_location: dict[tuple[str, str, int], list[str]] = defaultdict(list)
    message_order = {}
    for index, finding in enumerate(findings):
        number = re.search(r'\d+', finding.display_code)
        # 不同主体或行可产生相同提示；按规则编号排序并保留首次检查次序，禁止后行改写权重。
        message_order.setdefault(finding.message, (int(number[0]) if number else 999, index))
    material = [x for x in findings if x.location_policy == "material" or not x.file_path]
    writable = [(f, s) for f in files for s in f.sheets if s.first_data_row]
    fallback = writable[0] if writable else None
    pending_material = []
    for finding in findings:
        if finding.location_policy != "material" and finding.file_path and finding.row:
            by_location[(finding.file_path, finding.sheet, finding.row)].append(finding.message)
    for finding in material:
        if fallback:
            candidates = [(f, s) for f, s in writable
                          if (not finding.entity_code or f.entity_code == finding.entity_code)
                          and (not finding.business_id or (f.business_id or f.business_code) == finding.business_id)
                          and (not finding.business_code or f.business_code == finding.business_code)
                          and (not finding.variant_id or f.variant_id == finding.variant_id)]
            if not candidates: candidates = [(f, s) for f, s in writable if not finding.entity_code or f.entity_code == finding.entity_code]
            if not candidates and entities is not None:
                entity = entities.get(finding.entity_code)
                # 所属主体全部缺件时只允许写入同一报送单位的材料，禁止跨市误记意见。
                group = entity.parent if entity else ''
                candidates = [(f, s) for f, s in writable if group and f.entity_code in entities and entities[f.entity_code].parent == group]
            if not candidates and entities is None:
                candidates = [fallback]
            if not candidates:
                pending_material.append(finding.to_dict())
                continue
            ff, ss = next(((f, s) for f, s in candidates if s.sheet_type == "matrix"), candidates[0])
            # 规则归属仅由意见原文中的“第X条”表达，整套材料只是写回位置。
            text = finding.message
            message_order[text] = message_order[finding.message]
            by_location[(str(ff.relative_path), ss.title, ss.first_data_row)].append(text)
        else:
            pending_material.append(finding.to_dict())
    warnings = [] ; ownership: dict[str, list[dict[str, Any]]] = {}
    # 整个包没有可承载意见的表格时保留待回填事项，runner 据此禁止标记完成。
    output_state.pending.extend(pending_material)
    write_json(metadata_dir / '待回填资料级意见.json', output_state.pending)
    if pending_material: warnings.append({"type": "no_writable_sheet", "message": "缺件主体及所属报送单位无可承载意见的正式业务表，仍有待回填事项，流程未完成", "report": str(metadata_dir / '待回填资料级意见.json')})
    for file in files:
        if not file.sheets:
            # 未识别文件也必须保留原样副本；不能因无业务表而从交付目录消失。
            destination = output_root / _output_path(file)
            destination.parent.mkdir(parents=True, exist_ok=True)
            if destination.exists() and sha256_file(destination) != file.sha256:
                raise FileExistsError(f'未审核文件副本已存在且内容不同，请使用新输出目录：{destination}')
            shutil.copy2(file.source, destination)
            if sha256_file(destination) != file.sha256:
                raise RuntimeError(f'未审核文件副本与源文件不一致：{destination}')
            if file.material_type != 'explanation':
                warnings.append({'type': 'unparsed_business_file', 'file': str(file.relative_path),
                                 'message': '未识别到可审核的业务表，仅原样保留文件，未添加审核意见，不能视为已审核。',
                                 'parse_errors': file.parse_errors})
            continue
        rel_out = _output_path(file)
        destination = output_root / rel_out
        source = Path(file.preservation.get("converted_path", "")) if file.true_format == "xls" else file.source
        if file.true_format == "xls":
            # parse_files stores converted workbook under work dir; resolve from converted_from sibling metadata set by runner.
            source = Path(getattr(file, "_converted_path", ""))
        source = Path(getattr(file, '_preprocessed_path', source))
        edits = _detect_owned_edits(destination, old_ownership, str(rel_out))
        # 每次使用本次源材料业务数据；旧输出只用于变更告警，不能回滚整改或继承意见。
        base = source
        if edits: warnings.extend({"type": "managed_cell_modified", "file": str(rel_out), **e} for e in edits)
        # 按实际 OOXML 内容读取，原冻结规则允许的 ET 副本不依赖扩展名。
        source_book = load_compatible_workbook(base, read_only=True, data_only=False)
        hidden_rows = _hidden_rows_by_sheet(base)
        ops = {}
        written_sheets = set()
        for sheet in file.sheets:
            if sheet.title in written_sheets:
                continue
            written_sheets.add(sheet.title)
            related_sheets = [part for part in file.sheets if part.title == sheet.title]
            source_sheet = source_book[sheet.title]
            opinion_header_rows = [
                _select_output_header_row(
                    part.header_rows,
                    hidden_rows.get(sheet.title, set()).__contains__,
                    default_row=2,
                )
                for part in related_sheets
            ]
            values: dict[int, str] = {}
            for (fp, st, row), messages in by_location.items():
                if fp == str(file.relative_path) and st == sheet.title: values[row] = _join_opinions(messages, message_order)
            # 重新审核必须重建整列结果：旧审核意见无论来源和归属都不得进入本轮输出。
            first_header_row = min(opinion_header_rows)
            for row_number in range(first_header_row + 1, source_sheet.max_row + 1):
                if row_number not in opinion_header_rows:
                    values.setdefault(row_number, '')
            ops[sheet.title] = {"column": sheet.output_column, "header_row": opinion_header_rows[0], "header_rows": opinion_header_rows, "values": values}
            if baselines is not None:
                ops[sheet.title]['program_values'] = {row: _join_opinions(by_location.get((str(file.relative_path), sheet.title, row), []), message_order) for row in values}
            letter = get_column_letter(sheet.output_column)
            for row, text in values.items():
                program = _join_opinions(by_location.get((str(file.relative_path), sheet.title, row), []), message_order)
                if program:
                    ownership.setdefault(str(rel_out), []).append({
                        "sheet": sheet.title,
                        "cell": f"{letter}{row}",
                        "rendered_text": text,
                        "program_text": program,
                    })
        source_book.close()
        patch_ooxml(base, destination, ops)
        report = compare_ooxml_preservation(base, destination, ops)
        file.preservation["writeback"] = report
        if not report["passed"]: raise RuntimeError(f"OOXML增量写回保真失败: {destination}: {report['differences'][:3]}")
        if baselines is not None:
            from risk_audit.output_0916 import enrich_matrices
            try:
                moves = enrich_matrices(file, destination, baselines, findings, files)
            except WorkbookOutputCompatibilityError as error:
                # 丢弃已写入意见的半成品，以未改动的读取基线作为交付副本。
                shutil.copy2(base, destination)
                if sha256_file(destination) != sha256_file(base):
                    raise RuntimeError(f'未审核文件副本与读取基线不一致：{destination}') from error
                ownership.pop(str(rel_out), None)
                warnings.append({
                    'type': 'output_incompatible_file',
                    'file': str(file.relative_path),
                    'sheet': error.sheet,
                    'range': error.cell_range,
                    'message': '工作簿结构与审核列写回规则冲突，已跳过自动写回并保留原始副本，不能视为已审核。',
                    'error': str(error),
                })
                if file.true_format != 'xls' and sha256_file(file.source) != file.sha256:
                    raise RuntimeError(f'input changed during run: {file.source}')
                continue
            # 核对列插入后同步意见归属，复核仍能分清人工与程序意见。
            for state in ownership.get(str(rel_out), []):
                if state['sheet'] in moves:
                    old_column, new_column = moves[state['sheet']]
                    if _coord_col(state['cell']) == old_column:
                        row_number = re.search(r'\d+', state['cell'])[0]
                        state['cell'] = f'{get_column_letter(new_column)}{row_number}'
        if file.true_format != "xls" and sha256_file(file.source) != file.sha256: raise RuntimeError(f"input changed during run: {file.source}")
    # 同批各主体共享状态，后续主体不能覆盖已经输出的归属与诊断清单。
    output_state.unparsed.extend(item for item in warnings if item['type'] in {'unparsed_business_file', 'output_incompatible_file'})
    write_json(metadata_dir / '未审核文件.json', output_state.unparsed)
    for file in files:
        output_state.ownership.pop(str(_output_path(file)), None)
    output_state.ownership.update(ownership)
    write_json(ownership_path, output_state.ownership)
    return warnings, ownership
