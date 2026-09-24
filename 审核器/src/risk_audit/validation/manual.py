from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from risk_audit.util import norm_text, write_json

AUDIT_RE = re.compile(r"(?:9[.月-]11.*初审|9\.11初审)")


def split_atomic(text: str) -> list[str]:
    cleaned = str(text).strip()
    if not cleaned: return []
    parts = re.split(r"(?:\n\s*\n+|(?=\s*\d+[.、]\s*)|；(?=\s*\d*[.、]?))", cleaned)
    return [re.sub(r"^\s*\d+[.、]\s*", "", x).strip() for x in parts if re.sub(r"^\s*\d+[.、]\s*", "", x).strip()]


def classify_label(text: str, sheet_type: str) -> dict[str, Any]:
    n = norm_text(text)
    if sheet_type == "system_rule" and ("只保留" in n or "本单位" in n or "本层级" in n):
        return {"classification": "scope_excluded", "rule_id": "systems.content", "reason": "R07本期关闭"}
    if "部门" in n and ("具体" in n or "明确" in n): return {"classification": "in_scope", "rule_id": "duties.completeness", "check_id": "department_specific", "problem_type": "department"}
    if "岗位" in n and ("具体" in n or "明确" in n): return {"classification": "in_scope", "rule_id": "duties.completeness", "check_id": "position_specific", "problem_type": "position"}
    if "姓名" in n or "人员" in n and "明确" in n: return {"classification": "in_scope", "rule_id": "duties.completeness", "check_id": "person_required", "problem_type": "person_names"}
    if "经办" in n and "审核" in n: return {"classification": "in_scope", "rule_id": "duties.separation", "check_id": "handler_reviewer_overlap", "problem_type": "overlap"}
    if "句式" in n or "负有" in n or "负主体责任" in n: return {"classification": "legacy_scope_difference", "rule_id": "duties.responsibility_phrase", "reason": "人工旧严格句式与本期宽松口径不同"}
    if "角色" in n: return {"classification": "legacy_scope_difference", "rule_id": "duties.role_mapping", "reason": "人工角色必填口径与本期角色责任对应口径不同"}
    return {"classification": "unmapped_review", "reason": "需人工映射，未用于算法"}


def extract_manual_labels(input_root: str | Path, output: str | Path | None = None) -> dict[str, Any]:
    root = Path(input_root).resolve(); cells = []; atoms = []
    for path in sorted(root.rglob("*.xlsx")):
        if path.name.startswith("~$") or "三清单" not in path.name: continue
        wb = load_workbook(path, data_only=False, read_only=False, rich_text=True)
        for ws in wb.worksheets:
            sheet_norm = norm_text(ws.title + str(ws.cell(1, 1).value or ""))
            sheet_type = "system_rule" if "系统" in sheet_norm and "规则" in sheet_norm else "position_duty" if "岗位" in sheet_norm and "职责" in sheet_norm else "other"
            if sheet_type == "other": continue
            audit_cols = []
            for row in range(1, min(ws.max_row, 8) + 1):
                for cell in ws[row]:
                    if cell.value and AUDIT_RE.search(str(cell.value)): audit_cols.append(cell.column)
            for col in sorted(set(audit_cols)):
                for row in range(1, ws.max_row + 1):
                    value = ws.cell(row, col).value
                    if value is None or not str(value).strip() or AUDIT_RE.search(str(value)): continue
                    item = {"file": str(path.relative_to(root)), "sheet": ws.title, "row": row, "cell": ws.cell(row, col).coordinate, "sheet_type": sheet_type, "original": str(value)}
                    cells.append(item)
                    for atom in split_atomic(str(value)):
                        atoms.append({**item, "atom": atom, **classify_label(atom, sheet_type)})
    counts = {"manual_nonempty_cells": len(cells), "atomic_labels": len(atoms)}
    for key in sorted({x["classification"] for x in atoms}): counts[key] = sum(x["classification"] == key for x in atoms)
    report = {"counts": counts, "cells": cells, "atoms": atoms, "warning": "未标注行不视为负例；本报告不计算总体准确率。人工内容仅在算法结束后读取。"}
    if output: write_json(Path(output), report)
    return report


def compare_findings(labels: dict[str, Any], findings: list[dict[str, Any]]) -> dict[str, Any]:
    keys = [(x["rule_id"], x["check_id"], x["file_path"], x["sheet"], x["row"], x.get("evidence", {}).get("field"), x.get("severity")) for x in findings]
    compared = []; consumed: set[int] = set()
    for atom in labels["atoms"]:
        if atom["classification"] != "in_scope": continue
        problem_type = atom.get("problem_type")
        match_index = next((i for i, k in enumerate(keys) if i not in consumed and k[0] == atom.get("rule_id") and k[1] == atom.get("check_id") and k[2] == atom["file"] and k[3] == atom["sheet"] and k[4] == atom["row"] and (problem_type not in {"department", "position", "person_names"} or k[5] == problem_type)), None)
        if match_index is not None: consumed.add(match_index)
        compared.append({**atom, "matched": match_index is not None, "finding_severity": keys[match_index][6] if match_index is not None else None})
    problem_types = sorted({x.get("problem_type", "other") for x in compared})
    breakdown = {p: {"known_positive": sum(x.get("problem_type", "other") == p for x in compared), "matched": sum(x.get("problem_type", "other") == p and x["matched"] for x in compared)} for p in problem_types}
    return {"known_positive_atoms": len(compared), "matched": sum(x["matched"] for x in compared), "unmatched": sum(not x["matched"] for x in compared), "matched_confirmed": sum(x["matched"] and x["finding_severity"] == "violation" for x in compared), "matched_review": sum(x["matched"] and x["finding_severity"] == "review" for x in compared), "scope_excluded_atoms": sum(x.get("classification") == "scope_excluded" for x in labels.get("atoms", [])), "legacy_scope_difference_atoms": sum(x.get("classification") == "legacy_scope_difference" for x in labels.get("atoms", [])), "by_problem_type": breakdown, "items": compared, "warning": "仅报告已知正例覆盖，不把未标注行当负例；算法新增问题须复核，不宣称总体准确率。"}
