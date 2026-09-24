"""0916 规则的数据预处理和矩阵核对列。"""
from __future__ import annotations

from collections import defaultdict
from copy import deepcopy
from pathlib import Path
from typing import Any
import os
import posixpath
import re
import tempfile
import zipfile

from lxml import etree
from openpyxl.formula.translate import Translator, TranslatorError
from openpyxl.utils import get_column_letter, range_boundaries

from risk_audit.checks.confirmed_v180 import measure_numbers
from risk_audit.models import FileRecord, Finding
from risk_audit.util import measure_id_key, natural_key, norm_text
from risk_audit.writer import (MAIN, NS, PKGREL, _coord_col, _ensure_row, _index_rows, _inline_cell,
                               _select_output_header_row, _sheet_paths)

REFERENCE_HEADER = '省公司版本责任主体（核对后删除）'
EXISTING_HEADER = '岗位清单已有的控制措施编号'
MATRIX_VISIBLE_KEYWORDS = ('关键控制点', '控制措施编号', '控制措施', '不相容岗位', '控制方式', '控制系统',
                           '控制载体', '责任主体', '是否适用', '不适用原因')


def _save_zip(source: Path, destination: Path, replacements: dict[str, bytes]) -> None:
    """原子替换指定 OOXML 部件；source/destination 为源和目标，replacements 为允许修改的部件。"""
    destination.parent.mkdir(parents=True, exist_ok=True)
    descriptor, name = tempfile.mkstemp(suffix='.xlsx', dir=destination.parent)
    os.close(descriptor)
    try:
        with zipfile.ZipFile(source) as old, zipfile.ZipFile(name, 'w') as new:
            for info in old.infolist():
                new.writestr(info, replacements.get(info.filename, old.read(info.filename)))
        Path(name).replace(destination)
    finally:
        Path(name).unlink(missing_ok=True)


def _xml(root: Any) -> bytes:
    """序列化 OOXML；root 为 XML 根元素。"""
    return etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)


def _update_matrix_auto_filter(root: Any, first_column: int, last_column: int,
                               header_row: int, last_row: int) -> None:
    """为矩阵核对列设置筛选范围。

    root 为工作表 XML 根元素，first_column/last_column 为新增列边界，
    header_row 为表头行，last_row 为最后一条矩阵明细行。
    """
    auto_filter = root.find(f'{{{MAIN}}}autoFilter')
    if auto_filter is None:
        sheet_data = root.find(f'{{{MAIN}}}sheetData')
        if sheet_data is None:
            return
        auto_filter = etree.Element(f'{{{MAIN}}}autoFilter')
        predecessors = {
            'sheetPr', 'dimension', 'sheetViews', 'sheetFormatPr', 'cols', 'sheetData',
            'sheetCalcPr', 'sheetProtection', 'protectedRanges', 'scenarios',
        }
        # 放在规范允许的最后一个前置节点后，避免破坏工作表保护等可选节点顺序。
        insert_at = max(index for index, child in enumerate(root)
                        if etree.QName(child).localname in predecessors) + 1
        root.insert(insert_at, auto_filter)
        auto_filter.set(
            'ref',
            f'{get_column_letter(first_column)}{header_row}:{get_column_letter(last_column)}{last_row}',
        )
        return

    current_ref = auto_filter.get('ref')
    if not current_ref:
        auto_filter.set(
            'ref',
            f'{get_column_letter(first_column)}{header_row}:{get_column_letter(last_column)}{last_row}',
        )
        return

    left, _, right, bottom = range_boundaries(current_ref)
    # 扩展现有筛选的右边界，并让起始行与程序选定的可见表头一致；已选条件保持不变。
    auto_filter.set(
        'ref',
        f'{get_column_letter(left)}{header_row}:{get_column_letter(max(right, last_column))}{bottom}',
    )


def _sort_key(row: Any) -> tuple:
    """确定稳定的数字措施顺序；row 为岗位清单明细行。"""
    numbers = measure_id_key(row.value('measure_id'))
    return (0, numbers) if all(isinstance(value, int) for value in numbers) else (1, natural_key(row.value('measure_id')))


def _cell_has_content(cell: Any) -> bool:
    """判断 OOXML 单元格是否包含值或公式；cell 为工作表 XML 单元格节点。"""
    if cell.find(f'{{{MAIN}}}f') is not None:
        return True
    value = cell.find(f'{{{MAIN}}}v')
    if value is not None and value.text not in (None, ''):
        return True
    return bool(''.join(cell.xpath('.//m:is//m:t/text()', namespaces=NS)))


def _hidden_columns(root: Any) -> set[int]:
    """返回工作表中所有隐藏列号；root 为工作表 XML 根节点。"""
    columns = root.find(f'{{{MAIN}}}cols')
    if columns is None:
        return set()
    hidden = set()
    for column in columns.findall(f'{{{MAIN}}}col'):
        if column.get('hidden') not in {'1', 'true', 'True'}:
            continue
        hidden.update(range(int(column.get('min')), int(column.get('max')) + 1))
    return hidden


def _expand_shared_formulas(cells: dict[str, Any], title: str) -> None:
    """展开整表共享公式；cells 为原坐标到 XML 单元格的映射，title 为错误定位用表名。"""
    shared = []
    masters = {}
    for ref, cell in cells.items():
        formula = cell.find(f'{{{MAIN}}}f')
        if formula is None or formula.get('t') != 'shared':
            continue
        index = formula.get('si')
        if index is None:
            raise ValueError(f'岗位排序共享公式缺少组编号：{title}!{ref}')
        shared.append((ref, formula, index))
        if formula.text:
            if index in masters:
                raise ValueError(f'岗位排序共享公式存在多个主公式：{title}!{ref}')
            masters[index] = (ref, '=' + formula.text)
    # 先收集全部主公式，支持主公式在明细之外或位于从属单元格之后。
    for ref, formula, index in shared:
        if index not in masters:
            raise ValueError(f'岗位排序共享公式缺少主公式：{title}!{ref}')
        origin, text = masters[index]
        formula.text = Translator(text, origin=origin).translate_formula(ref)[1:]
        # 整组解除共享，未移动成员也不再依赖主公式的位置及共享范围。
        for attribute in ('t', 'si', 'ref'):
            formula.attrib.pop(attribute, None)


def sort_duties(file: FileRecord, baselines: dict, work_dir: Path, aliases: dict) -> None:
    """按数字编号排序岗位有效明细；file 为已解析材料，baselines 为保留接口参数，work_dir 为副本目录，aliases 为表头别名。

    仅在临时副本移动业务单元格；按实际合并范围补齐后，意见与批注随原记录移动。
    """
    source = Path(getattr(file, '_preprocessed_path', getattr(file, '_converted_path', file.source)))
    replacements = {}; reports = []
    with zipfile.ZipFile(source) as archive:
        paths = _sheet_paths(archive)
        for title in dict.fromkeys(sheet.title for sheet in file.sheets if sheet.sheet_type == 'position_duty'):
            root = etree.fromstring(archive.read(paths[title])); data = root.find(f'{{{MAIN}}}sheetData')
            hidden_columns = _hidden_columns(root)
            rows = {int(row.get('r')): row for row in data}
            # XLS 转换可能保留大量带样式的空行，只有值或公式的行才是业务边界。
            last_content_row = max(
                (number for number, row in rows.items()
                 if any(_cell_has_content(cell) for cell in row.findall(f'{{{MAIN}}}c'))),
                default=0,
            )
            cells = {cell.get('r'): cell for row in data for cell in row.findall(f'{{{MAIN}}}c')}
            cell_map = {}
            for sheet in (part for part in file.sheets if part.title == title and part.sheet_type == 'position_duty'):
                details = [row for row in sheet.records if not re.match(r'^(?:填表说明|填报说明|说明[:：]|注[:：]|合计)', row.value('measure_id'))]
                if len(details) < 2:
                    continue
                # 0916-4 明确按编号数字排序，旧基准清单的行次不再覆盖数字顺序。
                ordered = sorted(details, key=_sort_key)
                mapping = {row.row: slot for row, slot in zip(ordered, sorted(row.row for row in details))}
                if all(old == new for old, new in mapping.items()):
                    continue
                # 使用业务区域的实际表头边界，字段识别冲突时也必须移动部门等原始业务列。
                header_columns = {path['column'] for path in sheet.business_header_paths}
                recognized_columns = set(sheet.columns.values())
                first_col = min(header_columns or recognized_columns)
                last_col = max(header_columns or recognized_columns)
                other_columns = {
                    column
                    for part in file.sheets
                    if part.title == title and part is not sheet
                    for column in (set(part.columns.values()) | {path['column'] for path in part.business_header_paths})
                }
                # 隐藏且未选中的重复字段是核对辅助区起点，不得随左侧业务记录移动。
                detached_columns = {
                    candidate['column']
                    for selection in file.preservation.get('column_selections', [])
                    if selection.get('sheet') == title and selection.get('selected_column') is not None
                    for candidate in selection.get('candidates', [])
                    if not candidate.get('selected') and candidate.get('hidden')
                    and candidate.get('column', 0) > max(recognized_columns)
                }
                # 未识别的隐藏列位于业务列右侧时视为核对辅助区；已识别的隐藏业务列仍随记录移动。
                detached_columns.update(
                    column for column in hidden_columns
                    if column > max(recognized_columns) and column not in recognized_columns
                )
                detached_start = min(detached_columns, default=None)
                # 空表头列不会进入字段映射，但明细数据仍属于同一条业务记录。
                # 仅在相邻业务区域和隐藏辅助区之间扩展，避免排序时夹带其他表格。
                detail_columns = {
                    _coord_col(cell.get('r'))
                    for number in mapping
                    for cell in rows[number].findall(f'{{{MAIN}}}c')
                    if _cell_has_content(cell)
                }
                if sheet.region_bounds is not None:
                    _, _, left_limit, right_limit = sheet.region_bounds
                else:
                    # 旧解析结果没有区域边界时，才按同表其他字段列限定横向范围。
                    left_limit = max((column for column in other_columns if column < first_col), default=0) + 1
                    right_barriers = [column for column in other_columns if column > last_col]
                    right_limit = min(right_barriers) - 1 if right_barriers else max(detail_columns | {last_col})
                if detached_start is not None:
                    right_limit = min(right_limit, detached_start - 1)
                first_col = max(first_col, left_limit)
                last_col = min(last_col, right_limit)
                record_columns = {
                    column for column in detail_columns
                    if left_limit <= column <= right_limit
                }
                if record_columns:
                    first_col = min(first_col, min(record_columns))
                    last_col = max(last_col, max(record_columns))
                # 单区表仍包括核对区之前的原始附加列，包括只在少数明细中填写的内容。
                if sheet.region_bounds is None and not other_columns:
                    last_col = (detached_start - 1 if detached_start is not None
                                else max((_coord_col(ref) for ref in cells), default=last_col))
                if sheet.region_bounds is not None:
                    # 空映射表示当前区域确实没有审核列，不能回退到同表其他区域的统一映射。
                    region_audit_columns = set(sheet.region_audit_columns.values())
                    _, _, region_first_col, region_last_col = sheet.region_bounds
                    region_audit_columns = {
                        column for column in region_audit_columns
                        if region_first_col <= column <= region_last_col
                    }
                else:
                    region_audit_columns = set((sheet.region_audit_columns or sheet.audit_columns).values())
                move_columns = set(range(first_col, last_col + 1)) | region_audit_columns
                # 在复制及移动前还原各单元格公式，随后按普通公式调整相对引用。
                _expand_shared_formulas(cells, title)
                materialized = {ref: deepcopy(cell) for ref, cell in cells.items()}
                merges = root.find(f'{{{MAIN}}}mergeCells')
                if merges is not None:
                    for merge in list(merges):
                        left, top, right, bottom = range_boundaries(merge.get('ref'))
                        # Excel/WPS 可能把末组合并到最大行；空白样式尾部不属于其他业务区域。
                        effective_bottom = min(bottom, last_content_row)
                        merge_rows = range(top, effective_bottom + 1)
                        merge_columns = range(left, right + 1)
                        if not any(number in mapping for number in merge_rows) or not any(column in move_columns for column in merge_columns):
                            continue
                        if not all(number in mapping for number in merge_rows) or not all(column in move_columns for column in merge_columns):
                            raise ValueError(f'岗位排序合并范围跨越有效明细或其他区域：{file.relative_path}#{title}!{merge.get("ref")}')
                        anchor = cells.get(f'{get_column_letter(left)}{top}')
                        if anchor is not None:
                            for number in merge_rows:
                                for column in merge_columns:
                                    ref = f'{get_column_letter(column)}{number}'
                                    materialized[ref] = deepcopy(anchor); materialized[ref].set('r', ref)
                        # 展开明细合并，确保同措施的各部门和人员在移动后仍完整。
                        merges.remove(merge)
                    merges.set('count', str(len(merges)))
                    if not len(merges):
                        root.remove(merges)
                moved = []
                for old, new in mapping.items():
                    for column in move_columns:
                        old_ref = f'{get_column_letter(column)}{old}'; new_ref = f'{get_column_letter(column)}{new}'
                        cell_map[old_ref] = new_ref
                        cell = materialized.get(old_ref)
                        if cell is not None:
                            cell.set('r', new_ref)
                            formula = cell.find(f'{{{MAIN}}}f')
                            if formula is not None and old != new:
                                if formula.get('t') in {'shared', 'array'}:
                                    raise ValueError(f'岗位排序暂不能安全移动共享或数组公式：{title}!{old_ref}')
                                source_formula = '=' + (formula.text or '')
                                try:
                                    formula.text = Translator(source_formula, origin=old_ref).translate_formula(new_ref)[1:]
                                except TranslatorError:
                                    # ROW(A1) 等序号公式平移后可能越过首行；原样保留才能让公式随所属记录移动。
                                    formula.text = source_formula[1:]
                            moved.append((new, cell))
                for number in mapping:
                    for cell in list(rows[number].findall(f'{{{MAIN}}}c')):
                        if _coord_col(cell.get('r')) in move_columns:
                            rows[number].remove(cell)
                for number, cell in moved:
                    rows[number].append(cell)
                for number in mapping:
                    children = sorted(rows[number].findall(f'{{{MAIN}}}c'), key=lambda cell: _coord_col(cell.get('r')))
                    for cell in children:
                        rows[number].remove(cell); rows[number].append(cell)
                if not other_columns:
                    attributes = {number: dict(rows[number].attrib) for number in mapping}
                    for old, new in mapping.items():
                        rows[new].attrib.clear(); rows[new].attrib.update(attributes[old]); rows[new].set('r', str(new))
                reports.append({'sheet': title, 'row_mapping': mapping, 'basis': 'numeric'})
            if cell_map:
                replacements[paths[title]] = _xml(root)
                _move_annotations(archive, paths[title], cell_map, replacements)
    if not replacements:
        return
    destination = work_dir / 'preprocessed' / file.relative_path.with_suffix('.xlsx')
    _save_zip(source, destination, replacements)
    file._preprocessed_path = str(destination)
    # 再读取排序副本，使审核证据、输出意见和业务记录使用相同的新坐标。
    from risk_audit.readers.confirmed_v180 import parse_workbook_v180
    file.sheets = parse_workbook_v180(file, destination, aliases, include_hidden=getattr(file, '_include_hidden', False))
    file.preservation['sorting'] = reports


def _move_annotations(archive: zipfile.ZipFile, sheet_path: str, mapping: dict, replacements: dict) -> None:
    """移动人工批注和超链接；archive 为原文件，sheet_path 为业务表部件，mapping 为单元格移动表。"""
    root = etree.fromstring(replacements[sheet_path])
    for link in root.xpath('//m:hyperlink', namespaces=NS):
        link.set('ref', mapping.get(link.get('ref'), link.get('ref')))
    replacements[sheet_path] = _xml(root)
    relationship = posixpath.join(posixpath.dirname(sheet_path), '_rels', posixpath.basename(sheet_path) + '.rels')
    if relationship not in archive.namelist():
        return
    for rel in etree.fromstring(archive.read(relationship)):
        target = rel.get('Target', '')
        path = target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(sheet_path), target))
        if path not in archive.namelist():
            continue
        if rel.get('Type', '').endswith('/comments'):
            comments = etree.fromstring(archive.read(path))
            for comment in comments.xpath('//m:comment', namespaces=NS):
                comment.set('ref', mapping.get(comment.get('ref'), comment.get('ref')))
            replacements[path] = _xml(comments)
        elif rel.get('Type', '').endswith('/vmlDrawing'):
            drawing = etree.fromstring(archive.read(path))
            office = {'x': 'urn:schemas-microsoft-com:office:excel'}
            for client in drawing.xpath('//x:ClientData', namespaces=office):
                row = client.find('x:Row', office); column = client.find('x:Column', office)
                if row is None or column is None:
                    continue
                old = f'{get_column_letter(int(column.text) + 1)}{int(row.text) + 1}'
                new = mapping.get(old, old); delta = int(re.search(r'\d+', new)[0]) - int(row.text) - 1
                column_delta = _coord_col(new) - int(column.text) - 1
                row.text = str(int(row.text) + delta)
                column.text = str(int(column.text) + column_delta)
                anchor = client.find('x:Anchor', office)
                if anchor is not None:
                    values = [value.strip() for value in anchor.text.split(',')]
                    for index in (2, 6):
                        values[index] = str(int(values[index]) + delta)
                    for index in (0, 4):
                        values[index] = str(int(values[index]) + column_delta)
                    anchor.text = ', '.join(values)
            replacements[path] = _xml(drawing)


def _baseline_cell(item: dict, target_styles: Any, imported: dict) -> Any:
    """复制基准原责任单元格及样式；item 为来源定位，target_styles 为目标样式，imported 为本次样式缓存。"""
    source = item['source_path']
    cache_key = (source, 'source_data')
    if cache_key not in imported:
        with zipfile.ZipFile(source) as archive:
            source_data = {title: {cell.get('r'): cell for cell in etree.fromstring(archive.read(path)).xpath('//m:sheetData/m:row/m:c', namespaces=NS)} for title, path in _sheet_paths(archive).items()}
            strings = etree.fromstring(archive.read('xl/sharedStrings.xml')) if 'xl/sharedStrings.xml' in archive.namelist() else []
            imported[cache_key] = (source_data, strings, etree.fromstring(archive.read('xl/styles.xml')))
    source_data, strings, source_styles = imported[cache_key]
    cell = deepcopy(source_data[item['sheet']].get(item['cell']))
    if cell is not None:
        if cell is None:
            return None
        if cell.get('t') == 's':
            index = int(cell.find(f'{{{MAIN}}}v').text)
            for child in list(cell):
                cell.remove(child)
            cell.set('t', 'inlineStr'); inline = etree.SubElement(cell, f'{{{MAIN}}}is')
            for child in strings[index]:
                inline.append(deepcopy(child))
        style_id = int(cell.get('s', 0)); key = (source, style_id)
        if key not in imported:
            xf = deepcopy(source_styles.find(f'{{{MAIN}}}cellXfs')[style_id])
            for section, attribute in [('fonts', 'fontId'), ('fills', 'fillId'), ('borders', 'borderId')]:
                old_id = int(xf.get(attribute, 0)); source_part = source_styles.find(f'{{{MAIN}}}{section}')[old_id]
                target = target_styles.find(f'{{{MAIN}}}{section}')
                new_id = next((index for index, value in enumerate(target) if etree.tostring(value) == etree.tostring(source_part)), None)
                if new_id is None:
                    new_id = len(target); target.append(deepcopy(source_part)); target.set('count', str(len(target)))
                xf.set(attribute, str(new_id))
            # 责任主体是文字字段；不依赖基准自定义数值格式及命名样式索引。
            xf.set('numFmtId', '0'); xf.set('xfId', '0')
            alignment = xf.find(f'{{{MAIN}}}alignment')
            if alignment is None:
                alignment = etree.SubElement(xf, f'{{{MAIN}}}alignment')
            alignment.set('wrapText', '1'); xf.set('applyAlignment', '1')
            xfs = target_styles.find(f'{{{MAIN}}}cellXfs'); imported[key] = str(len(xfs))
            xfs.append(xf); xfs.set('count', str(len(xfs)))
        cell.set('s', imported[key])
        return cell
    return None


def _set_matrix_column_visibility(columns: Any, maximum: int, visible: set[int]) -> None:
    """重建矩阵列显示属性；columns 为 OOXML 列定义，maximum 为末列，visible 为允许显示的列号。"""
    definitions = list(columns)
    for definition in definitions:
        columns.remove(definition)
    for number in range(1, maximum + 1):
        source = next((definition for definition in reversed(definitions)
                       if int(definition.get('min')) <= number <= int(definition.get('max'))), None)
        attributes = dict(source.attrib) if source is not None else {}
        attributes.update({'min': str(number), 'max': str(number), 'hidden': '0' if number in visible else '1'})
        columns.append(etree.Element(f'{{{MAIN}}}col', **attributes))


def _select_matrix_output_start(root: Any, cells: dict[str, Any], old_opinion: int,
                                occupied: int) -> int:
    """选择三列矩阵输出的起始列。

    root 为工作表 XML，cells 为坐标到单元格的映射，old_opinion 为当前审核意见列，
    occupied 为当前最右单元格列。空白合并区自动取消，有内容或公式的合并区保留，
    三列输出整体向右移动到连续安全位置。
    """
    candidate = old_opinion if old_opinion == occupied else occupied + 1
    merges = root.find(f'{{{MAIN}}}mergeCells')
    while True:
        moved = False
        block_left, block_right = candidate, candidate + 2
        if merges is not None:
            for merge in list(merges):
                left, top, right, bottom = range_boundaries(merge.get('ref'))
                if left == right:
                    continue
                contains_source = left <= old_opinion <= right
                overlaps_target = left <= block_right and right >= block_left
                if not overlaps_target and not contains_source:
                    continue
                preserved_content = any(
                    _coord_col(reference) != old_opinion and _cell_has_content(cell)
                    and top <= int(re.search(r'\d+', reference)[0]) <= bottom
                    and left <= _coord_col(reference) <= right
                    for reference, cell in cells.items()
                )
                if preserved_content:
                    # 原表内容不拆分也不覆盖，整个输出块跳到合并区右侧。
                    if overlaps_target or (contains_source and candidate <= right):
                        candidate = right + 1
                        moved = True
                        break
                    continue
                # 空白合并只是模板格式，取消后可安全使用这些列。
                merges.remove(merge)
        if moved:
            continue
        conflicting_columns = {
            _coord_col(reference)
            for reference, cell in cells.items()
            if block_left <= _coord_col(reference) <= block_right
            and _coord_col(reference) != old_opinion
            and _cell_has_content(cell)
        }
        if conflicting_columns:
            candidate = max(conflicting_columns) + 1
            continue
        break
    if merges is not None:
        merges.set('count', str(len(merges)))
        if not len(merges):
            root.remove(merges)
    return candidate


def enrich_matrices(file: FileRecord, destination: Path, baselines: dict, findings: list[Finding],
                    all_files: list[FileRecord] | None = None) -> dict[str, tuple[int, int]]:
    """新增矩阵核对列；file 为材料记录，destination 为审核副本，baselines 为基准，findings 为当前意见，all_files 为同批材料。

    返回每张矩阵表的旧、新审核意见列，供程序意见归属同步更新。
    """
    sheets = [sheet for sheet in file.sheets if sheet.sheet_type == 'matrix']
    if not sheets:
        return {}
    _ = findings  # 意见写入由 writer 负责；本阶段只使用同批清单事实生成核对列。
    present = {(*_scope_key(record), measure_numbers(record.value('measure_id'))) for item in (all_files or [file])
               for sheet in item.sheets if sheet.sheet_type == 'position_duty' for record in sheet.records
               if norm_text(record.value('measure_id'))}
    replacements = {}; moves = {}; imported = {}; auxiliary_columns = {}
    with zipfile.ZipFile(destination) as archive:
        paths = _sheet_paths(archive); styles = etree.fromstring(archive.read('xl/styles.xml'))
        strings = etree.fromstring(archive.read('xl/sharedStrings.xml')) if 'xl/sharedStrings.xml' in archive.namelist() else []
        for title in dict.fromkeys(sheet.title for sheet in sheets):
            parts = [sheet for sheet in sheets if sheet.title == title]
            root = etree.fromstring(archive.read(paths[title])); data = root.find(f'{{{MAIN}}}sheetData')
            cells = {cell.get('r'): cell for cell in root.xpath('//m:sheetData/m:row/m:c', namespaces=NS)}
            headers = {}
            for cell in cells.values():
                text = ''.join(cell.xpath('.//m:t/text()', namespaces=NS))
                if cell.get('t') == 's':
                    text = ''.join(strings[int(cell.find(f'{{{MAIN}}}v').text)].itertext())
                if text in {REFERENCE_HEADER, EXISTING_HEADER, '审核意见'} and int(re.search(r'\d+', cell.get('r'))[0]) in {number for part in parts for number in part.header_rows}:
                    headers[text] = _coord_col(cell.get('r'))
            old_opinion = headers.get('审核意见', parts[0].output_column)
            auxiliary_columns[title] = {headers[name] for name in (REFERENCE_HEADER, EXISTING_HEADER) if name in headers}
            occupied = max((_coord_col(ref) for ref in cells), default=1)
            if REFERENCE_HEADER in headers and EXISTING_HEADER in headers:
                reference_col = headers[REFERENCE_HEADER]; existing_col = headers[EXISTING_HEADER]
                opinion_col = old_opinion
                matrix_header_rows = {number for part in parts for number in part.header_rows}
                first_matrix_header_row = min(matrix_header_rows)
                # 修正版保留三列表头时，先删除两列核对数据区的全部旧值，再按本轮事实重建。
                for row in data:
                    row_number = int(row.get('r'))
                    if row_number <= first_matrix_header_row or row_number in matrix_header_rows:
                        continue
                    for cell in list(row.findall(f'{{{MAIN}}}c')):
                        if _coord_col(cell.get('r')) in {reference_col, existing_col}:
                            row.remove(cell)
            else:
                # 原意见为最末列时挪到两列核对内容之后，其余业务数据完全保留。
                if REFERENCE_HEADER not in headers and EXISTING_HEADER not in headers:
                    next_column = _select_matrix_output_start(root, cells, old_opinion, occupied)
                else:
                    next_column = old_opinion if old_opinion == occupied else occupied + 1
                reference_col = headers.get(REFERENCE_HEADER, next_column)
                if REFERENCE_HEADER not in headers:
                    next_column += 1
                existing_col = headers.get(EXISTING_HEADER, next_column)
                if EXISTING_HEADER not in headers:
                    next_column += 1
                opinion_col = next_column
                for row in data:
                    for cell in row.findall(f'{{{MAIN}}}c'):
                        if _coord_col(cell.get('r')) == old_opinion:
                            cell.set('r', f'{get_column_letter(opinion_col)}{row.get("r")}')
            moves[title] = (old_opinion, opinion_col)
            if old_opinion != opinion_col:
                mapping = {f'{get_column_letter(old_opinion)}{row.get("r")}': f'{get_column_letter(opinion_col)}{row.get("r")}' for row in data}
                replacements[paths[title]] = _xml(root)
                _move_annotations(archive, paths[title], mapping, replacements)
                root = etree.fromstring(replacements[paths[title]]); data = root.find(f'{{{MAIN}}}sheetData')
                merges = root.find(f'{{{MAIN}}}mergeCells')
                if merges is not None:
                    for merge in merges:
                        left, top, right, bottom = range_boundaries(merge.get('ref'))
                        if left == right == old_opinion:
                            merge.set('ref', f'{get_column_letter(opinion_col)}{top}:{get_column_letter(opinion_col)}{bottom}')
            sources = baselines.get((file.business_id or file.business_code, file.variant_id), {}).get('responsibilities', [])
            indexed = defaultdict(list)
            for item in sources:
                indexed[measure_numbers(item['measure_id'])].append(item)
            hidden_rows = {
                int(row.get('r')) for row in data
                if row.get('hidden') in {'1', 'true', 'True'}
            }
            row_index, row_numbers = _index_rows(data)
            for part in parts:
                header = _select_output_header_row(part.header_rows, hidden_rows.__contains__)
                header_row = _ensure_row(data, header, row_index, row_numbers)
                _inline_cell(header_row, f'{get_column_letter(reference_col)}{header}', REFERENCE_HEADER, None)
                _inline_cell(header_row, f'{get_column_letter(existing_col)}{header}', EXISTING_HEADER, None)
                opinion_ref = f'{get_column_letter(opinion_col)}{header}'
                if not any(cell.get('r') == opinion_ref for cell in header_row.findall(f'{{{MAIN}}}c')):
                    _inline_cell(header_row, opinion_ref, '审核意见', None)
                for record in part.records:
                    row = _ensure_row(data, record.row, row_index, row_numbers); ref = f'{get_column_letter(reference_col)}{record.row}'
                    matches = indexed[measure_numbers(record.value('measure_id'))]
                    # 同编号多条基准责任主体不可任取一条；完整基准应提供唯一措施记录。
                    for cell in list(row.findall(f'{{{MAIN}}}c')):
                        if _coord_col(cell.get('r')) in {reference_col, existing_col}:
                            row.remove(cell)
                    if len(matches) == 1:
                        cell = _baseline_cell(matches[0], styles, imported)
                        if cell is not None:
                            cell.set('r', ref); row.append(cell)
                    if (*_scope_key(record), measure_numbers(record.value('measure_id'))) in present:
                        _inline_cell(row, f'{get_column_letter(existing_col)}{record.row}', record.value('measure_id'), None)
                part.output_column = opinion_col; part.audit_columns['审核意见'] = opinion_col
            for row in data:
                for cell in sorted(row.findall(f'{{{MAIN}}}c'), key=lambda value: _coord_col(value.get('r'))):
                    row.remove(cell); row.append(cell)
            dim = root.find(f'{{{MAIN}}}dimension')
            if dim is not None:
                first = dim.get('ref', 'A1').split(':')[0]; last_row = max(int(row.get('r')) for row in data)
                dim.set('ref', f'{first}:{get_column_letter(max(occupied, opinion_col))}{last_row}')
            columns = root.find(f'{{{MAIN}}}cols')
            if columns is None:
                columns = etree.Element(f'{{{MAIN}}}cols'); data.addprevious(columns)
            for column in (reference_col, existing_col, opinion_col):
                for old in list(columns):
                    if old.get('min') == old.get('max') == str(column):
                        columns.remove(old)
                etree.SubElement(columns, f'{{{MAIN}}}col', min=str(column), max=str(column), width='45', customWidth='1')
            visible = {reference_col, existing_col, opinion_col}
            for part in parts:
                visible.update(column for field, column in part.columns.items() if field in {
                    'key_control_point', 'measure_id', 'control_measure', 'incompatible_control', 'control_mode',
                    'control_system', 'carrier', 'responsibility', 'applicability', 'applicability_reason',
                })
                visible.update(path['column'] for path in part.business_header_paths
                               if any(keyword in node['name'] for node in path['path'] for keyword in MATRIX_VISIBLE_KEYWORDS))
            # 仅改变审核副本的可见性，隐藏列仍保留原值、样式和公式。
            _set_matrix_column_visibility(columns, max(occupied, opinion_col), visible)
            filter_header = min(
                (_select_output_header_row(part.header_rows, hidden_rows.__contains__) for part in parts),
                default=1,
            )
            filter_last_row = max((record.row for part in parts for record in part.records), default=filter_header)
            _update_matrix_auto_filter(
                root,
                min(reference_col, existing_col, opinion_col),
                max(reference_col, existing_col, opinion_col),
                filter_header,
                filter_last_row,
            )
            replacements[paths[title]] = _xml(root)
        replacements['xl/styles.xml'] = _xml(styles)
        _verify_enrichment(archive, paths, replacements, moves, auxiliary_columns)
    _save_zip(destination, destination, replacements)
    return moves


def _verify_enrichment(archive: zipfile.ZipFile, paths: dict, replacements: dict, moves: dict, auxiliary_columns: dict) -> None:
    """核验核对列阶段业务单元格；archive/paths 为原副本，replacements 为新部件，moves 为意见列移动表，auxiliary_columns 为可更新的原核对列。"""
    for title, (old_column, new_column) in moves.items():
        before = etree.fromstring(archive.read(paths[title])); after = etree.fromstring(replacements[paths[title]])
        new_cells = {cell.get('r'): cell for cell in after.xpath('//m:sheetData/m:row/m:c', namespaces=NS)}
        auxiliary = auxiliary_columns[title]
        for cell in before.xpath('//m:sheetData/m:row/m:c', namespaces=NS):
            coordinate = cell.get('r'); column = _coord_col(coordinate)
            if column in auxiliary:
                continue
            row_number = re.search(r'\d+', coordinate)[0]
            target = f'{get_column_letter(new_column)}{row_number}' if column == old_column else coordinate
            expected = deepcopy(cell); expected.set('r', target)
            if target not in new_cells or etree.tostring(expected, method='c14n', exclusive=True) != etree.tostring(new_cells[target], method='c14n', exclusive=True):
                raise RuntimeError(f'核对列输出改变原业务数据或人工意见：{title}!{coordinate}')


def _scope_key(record: Any) -> tuple[str, str, str]:
    """返回记录隔离范围；record 为记录或文件，依次使用主体、稳定业务 ID 和矩阵变体。"""
    return record.entity_code or '', record.business_id or record.business_code or '', record.variant_id
