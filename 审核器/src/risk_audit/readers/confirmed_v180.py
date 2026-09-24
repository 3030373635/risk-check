"""确认稿业务区域识别，字段读取仍复用成熟 Excel 解析器。"""

from __future__ import annotations

from copy import copy
import hashlib
import re
from pathlib import Path
from typing import Any

from openpyxl.cell.cell import Cell
from openpyxl.utils.cell import coordinate_from_string
from openpyxl.worksheet.cell_range import CellRange, MultiCellRange

from risk_audit.models import FileRecord, ParsedSheet
from risk_audit.readers import excel
from risk_audit.readers.header_semantics import HEADER_REQUIREMENTS, header_row_matches
from risk_audit.readers.ooxml import load_compatible_workbook
from risk_audit.util import norm_text


TITLE_NAMES = {
    "matrix": ("风控矩阵", "风险控制矩阵"),
    "position_duty": ("岗位内控责任清单", "岗位职责清单"),
    "incompatible_position": ("不相容岗位清单",),
    "system_rule": ("系统控制规则清单", "系统规则清单"),
}
AUXILIARY_NAMES = ("核对过程", "核对版本", "意见建议", "业务流程", "流程图", "填报说明", "WpsReserved")


def _title_kind(text: str, *, marker: bool = False) -> str | None:
    """识别明确业务标题；text 为标题文本，marker 指定是否为表内区域标题。"""
    text = norm_text(text)
    kinds = []
    for kind, names in TITLE_NAMES.items():
        for name in names:
            if marker:
                # 仅完整标题可划分区域，措施正文提及清单不构成区域边界。
                matched = re.fullmatch(r"(?:[一二三四五六七八九十\d]+[、.．)）:：]?)*" + re.escape(name) + r"(?:填报区域|填报区|填报表)?(?:\([^()]*\))?", text)
            else:
                matched = name in text
            if matched:
                kinds.append(kind)
                break
    return kinds[0] if len(set(kinds)) == 1 else None


def _heading_kind(text: str) -> str | None:
    """识别表内完整标题；text 为首格有效文字，业务后缀必须以分隔符或括号连接。"""
    text = norm_text(text)
    for kind, names in TITLE_NAMES.items():
        if any(re.fullmatch(re.escape(name) + r'(?:[-—:：].+|\([^()]+\))?', text) for name in names):
            return kind
    return None


def _header_fields(ws: Any, kind: str, aliases: dict, bounds: tuple[int, int, int, int]) -> dict[int, set[str]]:
    """收集区域内整行字段命中；参数为工作表、类型、别名和行列边界。"""
    first_row, last_row, first_col, last_col = bounds
    reverse = {norm_text(value): field for field, values in aliases.get(kind, {}).items() for value in values}
    rows: dict[int, set[str]] = {}
    cells_by_row: dict[int, list[Any]] = {}
    for cell in ws._cells.values():
        if first_row <= cell.row <= last_row and first_col <= cell.column <= last_col and cell.value not in (None, ""):
            cells_by_row.setdefault(cell.row, []).append(cell)
    for row, cells in cells_by_row.items():
        values = [excel.current_text(cell) for cell in sorted(cells, key=lambda cell: cell.column)]
        business_values = ["" if excel.AUDIT_HEADER.search(value) else value for value in values]
        matches = header_row_matches(business_values, kind, reverse, allow_fuzzy=row < first_row + 5)
        fields = {field for field, _ in matches if field}
        if fields:
            rows[row] = fields
    return rows


def _infer_kind(ws: Any, aliases: dict, bounds: tuple[int, int, int, int]) -> str | None:
    """以同一行的业务字段识别单区表；参数为工作表、别名和行列边界。"""
    candidates = [kind for kind, required in HEADER_REQUIREMENTS.items()
                  if any(required <= fields for fields in _header_fields(ws, kind, aliases, bounds).values())]
    return candidates[0] if len(candidates) == 1 else None


def _regions(ws: Any, aliases: dict) -> list[tuple[str, tuple[int, int, int, int]]]:
    """划分明确标题的纵向或并排区域；ws 为物理表，aliases 为字段别名。"""
    if any(name in ws.title for name in AUXILIARY_NAMES):
        return []
    last_row = max((cell.row for cell in ws._cells.values()), default=1)
    last_col = max((cell.column for cell in ws._cells.values() if cell.value not in (None, "")), default=1)
    # 标题可并排或带明确填报元信息，明细里的清单/载体名称不是边界。
    title_rows = {cell.row for cell in ws._cells.values()
                  if cell.value not in (None, '') and not _title_kind(excel.current_text(cell), marker=True)
                  and not re.fullmatch(r'(?:填报单位|编制单位|会计主体|单位名称|填报日期|编制日期|填报人|填报部门|日期)[:：].+', norm_text(excel.current_text(cell)))}
    markers = sorted((cell.row, cell.column, kind) for cell in ws._cells.values()
                     if cell.row not in title_rows and (kind := _title_kind(excel.current_text(cell), marker=True)))
    regions = []
    if markers:
        row_starts = sorted({row for row, _, _ in markers})
        for row, col, kind in markers:
            next_rows = [start for start in row_starts if start > row]
            next_cols = [column for start, column, _ in markers if start == row and column > col]
            same_row = [column for start, column, _ in markers if start == row]
            first_col = 1 if len(same_row) == 1 else col
            bounds = (row, min(next_rows) - 1 if next_rows else last_row,
                      first_col, min(next_cols) - 1 if next_cols else last_col)
            regions.append((kind, bounds))
    else:
        bounds = (1, last_row, 1, last_col)
        # 正式矩阵常用业务名作为表名；表内标题也可确认类型，缺列留给后续检查。
        kind = _title_kind(ws.title) or _heading_kind(excel.current_text(ws.cell(1, 1))) or _infer_kind(ws, aliases, bounds)
        if kind:
            fields = _header_fields(ws, kind, aliases, bounds)
            # 独立不相容清单即使没有表头也表示存在；岗位名册仍需业务字段。
            if kind == "incompatible_position" or any(len(row_fields) >= 2 for row_fields in fields.values()):
                regions.append((kind, bounds))
    result = []
    for kind, bounds in regions:
        fields = _header_fields(ws, kind, aliases, bounds)
        primary = max(fields, key=lambda row: len(fields[row]), default=bounds[0])
        # 将晚出现的表头移入成熟解析器30行窗口，保留紧邻的分组表头。
        start = max(bounds[0], primary - 2) if len(fields.get(primary, set())) >= 2 else bounds[0]
        result.append((kind, (start, bounds[1], bounds[2], bounds[3])))
    return result


def classify_physical_sheets(workbook: Any, aliases: dict) -> dict[str, set[str]]:
    """识别工作簿中每个物理工作表的业务类型。

    参数 workbook 为 openpyxl 工作簿，aliases 为字段别名配置；
    返回以工作表名为键、识别到的业务类型集合为值的字典。
    """
    return {
        worksheet.title: {sheet_type for sheet_type, _ in _regions(worksheet, aliases)}
        for worksheet in workbook.worksheets
    }


class _WorksheetView:
    """保留原始单元格内容、样式和列号，仅将区域行号平移到读取窗口。"""

    def __init__(self, ws: Any, title: str, kind: str, bounds: tuple[int, int, int, int]):
        """参数：ws 为原表，title 为内部唯一名称，kind 为类型，bounds 为区域边界。"""
        self.source = ws
        self.title = title
        self.kind = kind
        self.bounds = bounds
        self.row_offset = bounds[0] - 1
        self.sheet_state = ws.sheet_state
        self.column_dimensions = ws.column_dimensions
        self._cells = {}
        for cell in ws._cells.values():
            if bounds[0] <= cell.row <= bounds[1] and bounds[2] <= cell.column <= bounds[3]:
                value = copy(cell)
                value.row -= self.row_offset
                self._cells[(value.row, value.column)] = value
        self.merged_cells = MultiCellRange()
        for merged in ws.merged_cells.ranges:
            # 跨边界合并不传入视图，禁止从前一个业务区域下填锚点。
            if (bounds[0] <= merged.min_row <= merged.max_row <= bounds[1]
                    and bounds[2] <= merged.min_col <= merged.max_col <= bounds[3]):
                self.merged_cells.add(CellRange(min_col=merged.min_col, max_col=merged.max_col,
                                               min_row=merged.min_row - self.row_offset,
                                               max_row=merged.max_row - self.row_offset))

    def cell(self, row: int, column: int) -> Any:
        """读取视图单元格；row 为平移行号，column 为原始物理列号。"""
        key = (row, column)
        if key not in self._cells:
            self._cells[key] = Cell(self.source, row=row, column=column)
        return self._cells[key]


class _WorkbookView:
    """向成熟读取器提供多个独立区域，内部唯一标题对应缓存视图。"""

    def __init__(self, worksheets: list[_WorksheetView]):
        """参数 worksheets 为按物理工作簿顺序构造的区域视图。"""
        self.worksheets = worksheets
        self.by_title = {ws.title: ws for ws in worksheets}

    def __getitem__(self, title: str) -> _WorksheetView:
        """按内部名称定位区域；title 为内部唯一工作表名称。"""
        return self.by_title[title]

    def close(self) -> None:
        """视图没有文件句柄，原始工作簿由外层关闭；无参数。"""


def _resolve_view_type(ws: _WorksheetView, material_type: str, aliases: dict) -> str:
    """返回已确认的区域类型；参数为视图、材料类型及字段别名。"""
    return ws.kind


def _physical_coordinate(coordinate: str, row_offset: int) -> str:
    """恢复物理坐标；coordinate 为视图坐标，row_offset 为行平移量。"""
    column, row = coordinate_from_string(coordinate)
    return f"{column}{row + row_offset}"


def _restore_evidence(value: Any, row_offset: int) -> None:
    """恢复列选择证据内的坐标和行号；value 为嵌套证据，row_offset 为平移量。"""
    if isinstance(value, dict):
        for key, item in value.items():
            if key == "coordinate" and isinstance(item, str):
                value[key] = _physical_coordinate(item, row_offset)
            elif key == "row" and isinstance(item, int):
                value[key] = item + row_offset
            elif key.endswith("_rows") and isinstance(item, list) and all(isinstance(row, int) for row in item):
                # business_rows等整数计数不能平移，只有明确的行号数组需要转换。
                value[key] = [row + row_offset for row in item]
            else:
                _restore_evidence(item, row_offset)
    elif isinstance(value, list):
        for item in value:
            _restore_evidence(item, row_offset)


def parse_workbook_v180(file: FileRecord, workbook_path: Path, aliases: dict, *, include_hidden: bool = False, input_overrides: list | None = None) -> list[ParsedSheet]:
    """按确认稿识别矩阵及三清单，复用成熟解析器读取每个区域。

    参数：file 为已确认归属的文件记录，workbook_path 为真实Excel路径，aliases 为字段别名；
    include_hidden 控制是否审核隐藏表，input_overrides 为已经确认的业务列对应。
    """
    formula_wb = load_compatible_workbook(
        workbook_path, data_only=False, rich_text=True, keep_links=True,
    )
    cache_wb = load_compatible_workbook(
        workbook_path, data_only=True, rich_text=True, keep_links=True,
    )
    try:
        formula_views = []
        cache_views = []
        physical = {}
        overrides = []
        file.preservation["audit_column_conflicts"] = []
        for index, ws in enumerate(formula_wb.worksheets):
            for region_index, (kind, bounds) in enumerate(_regions(ws, aliases)):
                title = f"{ws.title}__区域{index}_{region_index}"
                view = _WorksheetView(ws, title, kind, bounds)
                formula_views.append(view)
                cache_views.append(_WorksheetView(cache_wb[ws.title], title, kind, bounds))
                physical[title] = view
                for item in input_overrides or []:
                    if item["sheet"] == ws.title and item["field"] in aliases.get(kind, {}):
                        # 并排区域只接收自身列号的确认，避免跨区覆盖业务字段。
                        if bounds[2] <= item["column"] <= bounds[3]:
                            overrides.append({**item, "sheet": title})
        sheets = excel.parse_workbook(file, workbook_path, aliases, include_hidden=include_hidden,
                                      input_overrides=overrides,
                                      loaded_workbooks=(_WorkbookView(formula_views), _WorkbookView(cache_views)),
                                      sheet_type_resolver=_resolve_view_type)
        for sheet in sheets:
            view = physical[sheet.title]
            # 保留物理区域的行列边界，供排序等写回流程隔离共表区域。
            sheet.region_bounds = view.bounds
            # 后续会统一同一物理表的输出列，排序前须保留当前区域原有的人工审核列。
            sheet.region_audit_columns = dict(sheet.audit_columns)
            sheet.title = view.source.title
            sheet.header_rows = [row + view.row_offset for row in sheet.header_rows]
            if sheet.first_data_row is not None:
                sheet.first_data_row += view.row_offset
            for record in sheet.records:
                record.sheet = sheet.title
                record.row += view.row_offset
                seed = f"{file.entity_code}|{file.business_id}|{file.variant_id}|{file.relative_path}|{sheet.title}|{record.row}|{sheet.sheet_type}"
                record.record_id = hashlib.sha256(seed.encode()).hexdigest()[:24]
                for field in record.fields.values():
                    field.coordinate = _physical_coordinate(field.coordinate, view.row_offset)
            for path in sheet.business_header_paths:
                path["coordinate"] = _physical_coordinate(path["coordinate"], view.row_offset)
                for segment in path["path"]:
                    segment["coordinate"] = _physical_coordinate(segment["coordinate"], view.row_offset)
        for item in file.preservation.get("column_selections", []):
            view = physical[item["sheet"]]
            item["sheet"] = view.source.title
            _restore_evidence(item, view.row_offset)
        for notice in file.preservation.get("sheet_owner_notices", []):
            notice["sheet"] = physical[notice["sheet"]].source.title
        for item in file.preservation.get("excluded_historical_sheets", []):
            item["title"] = physical[item["title"]].source.title
            item["replacement_sheet"] = physical[item["replacement_sheet"]].source.title
        for ws in formula_wb.worksheets:
            same_sheet = [sheet for sheet in sheets if sheet.title == ws.title]
            if not same_sheet:
                continue
            audits = {name: column for sheet in same_sheet for name, column in sheet.audit_columns.items()}
            opinion_cols = {sheet.audit_columns["审核意见"] for sheet in same_sheet if "审核意见" in sheet.audit_columns}
            business_cols = {column for sheet in same_sheet for column in sheet.columns.values()}
            occupied = max((cell.column for cell in ws._cells.values() if cell.value not in (None, "")), default=1)
            # 同一列可能在另一区域承担业务字段，统一意见列不得覆盖这些字段。
            safe_opinion_cols = opinion_cols - business_cols
            output = next(iter(safe_opinion_cols)) if len(safe_opinion_cols) == 1 and len(opinion_cols) == 1 else occupied + 1
            output = excel._next_safe_output_column(ws, output)
            if opinion_cols and output not in opinion_cols:
                conflicts = [{"sheet_type": sheet.sheet_type, "header_rows": sheet.header_rows,
                              "column": column, "fields": sorted(field for field, value in sheet.columns.items() if value == column)}
                             for sheet in same_sheet for column in sorted(opinion_cols & set(sheet.columns.values()))]
                file.preservation["audit_column_conflicts"].append({
                    "sheet": ws.title, "original_columns": sorted(opinion_cols),
                    "conflicting_regions": conflicts, "output_column": output,
                    "reason": "原人工审核意见列在其他填报区域同时属于业务列，统一采用表末尾新列，保留原意见和业务数据，列布局需人工确认。" if conflicts
                    else "同一物理工作表的人工审核意见分布于多个列，统一采用表末尾新列并保留原意见，列布局需人工确认。",
                })
            for sheet in same_sheet:
                sheet.output_column = output
                sheet.audit_columns = dict(audits)
        hidden = []
        historical = {item["title"] for item in file.preservation.get("excluded_historical_sheets", [])}
        audited = {sheet.title for sheet in sheets}
        for ws in formula_wb.worksheets:
            if ws.sheet_state != "visible":
                action = "skipped_hidden" if not include_hidden else "audited" if ws.title in audited else "excluded_historical" if ws.title in historical else "not_business"
                hidden.append({"title": ws.title, "state": ws.sheet_state, "action": action})
        file.preservation["hidden_sheets"] = hidden
        return sheets
    finally:
        formula_wb.close()
        cache_wb.close()
