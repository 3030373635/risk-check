from __future__ import annotations

import hashlib
import json
import logging
import re
from pathlib import Path
from typing import Any

from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.utils import column_index_from_string, get_column_letter

from risk_audit.models import FieldValue, FileRecord, ParsedSheet, Record
from risk_audit.readers.ooxml import load_compatible_workbook, prepare_oversized_workbook
from risk_audit.readers.xls import convert_xls
from risk_audit.readers.header_semantics import HEADER_REQUIREMENTS, header_row_matches, header_spec, unit_names, header_owner_names
from risk_audit.util import norm_text


AUDIT_HEADER = re.compile(r"(?:\d{1,2}[.月-]\d{1,2}.*初审|初审|审核意见|复审意见|省公司版本责任主体[（(]核对后删除|岗位清单已有的控制措施编号)")
COLUMN_VISIBILITY_POLICY = "prefer_visible_equivalent_else_hidden"


def _next_safe_output_column(ws: Any, candidate: int) -> int:
    """查找可安全写入的列；ws 为工作表，candidate 为首个候选列。

    候选列落入包含文字、数值或公式的横向合并区时，移到该合并区右侧并继续检查。
    空白合并区留给后续输出阶段取消合并。
    """
    while True:
        conflict = next((cell_range for cell_range in ws.merged_cells.ranges
                         if cell_range.min_col < cell_range.max_col
                         and cell_range.min_col <= candidate <= cell_range.max_col
                         and any(cell.value not in (None, '')
                                 for cell in ws._cells.values()
                                 if cell_range.min_row <= cell.row <= cell_range.max_row
                                 and cell_range.min_col <= cell.column <= cell_range.max_col)), None)
        if conflict is None:
            return candidate
        # 整个输出区从合并区右侧重新选址，不覆盖原表内容。
        candidate = conflict.max_col + 1


def deleted_spans(cell: Any) -> list[dict[str, Any]]:
    """Keep contiguous struck text, joining font runs but not retained text."""
    value = cell.value
    if value is None: return []
    if isinstance(value, CellRichText):
        parts = [(part.text, part.font.strike is True) if isinstance(part, TextBlock) else (str(part), False) for part in value]
    else:
        parts = [(str(value), getattr(cell.font, "strike", False) is True)]
    spans = []; offset = 0
    for text, struck in parts:
        if struck and text:
            if spans and spans[-1]["end"] == offset:
                spans[-1]["text"] += text; spans[-1]["end"] += len(text)
            else: spans.append({"text": text, "start": offset, "end": offset + len(text)})
        offset += len(text)
    return spans


def _is_red_font(font: Any) -> bool:
    """判断字体是否为明确红色；font 为 openpyxl 字体或富文本行内字体。"""
    color = getattr(font, 'color', None)
    rgb = str(getattr(color, 'rgb', '') or '').upper()
    return rgb[-6:] == 'FF0000'


def red_spans(cell: Any) -> list[dict[str, Any]]:
    """提取连续红字范围；cell 为工作表单元格，返回原文偏移及文本。"""
    value = cell.value
    if value is None:
        return []
    if isinstance(value, CellRichText):
        parts = [(part.text, _is_red_font(part.font)) if isinstance(part, TextBlock)
                 else (str(part), _is_red_font(cell.font)) for part in value]
    else:
        parts = [(str(value), _is_red_font(cell.font))]
    spans = []; offset = 0
    for text, red in parts:
        if red and text:
            if spans and spans[-1]['end'] == offset:
                spans[-1]['text'] += text; spans[-1]['end'] += len(text)
            else:
                spans.append({'text': text, 'start': offset, 'end': offset + len(text)})
        offset += len(text)
    return spans


def current_text(cell: Any) -> str:
    value = cell.value
    if value is None: return ""
    if isinstance(value, CellRichText):
        chunks = []
        for part in value:
            if isinstance(part, TextBlock):
                # A run with strike=None remains effective even if the cell-level font is struck.
                if getattr(part.font, "strike", None) is True: continue
                chunks.append(part.text)
            else: chunks.append(str(part))
        return "".join(chunks).strip()
    if getattr(cell.font, "strike", False) is True: return ""
    return str(value).strip()


def _merged_value(ws: Any, row: int, col: int) -> Any:
    cell = ws.cell(row, col)
    if cell.__class__.__name__ != "MergedCell": return cell
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return ws.cell(rng.min_row, rng.min_col)
    return cell


def _sheet_type(ws: Any, material_type: str, aliases: dict | None = None) -> str | None:
    title = norm_text(ws.title)
    a1 = norm_text(ws.cell(1, 1).value)
    if any(x in title for x in ("核对过程", "核对版本", "意见建议", "业务流程", "流程图", "填报说明", "WpsReserved")): return None
    if aliases:
        # Require several business fields on the same row. A title alone or a
        # staff directory without measure/duty columns is not a formal list.
        max_col = _effective_header_max(ws, range(1, 31))
        for kind,required in HEADER_REQUIREMENTS.items():
            if kind=='matrix' and material_type!='matrix':continue
            reverse={norm_text(v):k for k,vs in aliases.get(kind,{}).items() for v in vs}
            for row in range(1, 31):
                values = [current_text(ws.cell(row, col)) for col in range(1, max_col + 1)]
                fields = {field for field, _ in header_row_matches(values, kind, reverse, allow_fuzzy=row <= 5) if field}
                if required <= fields:return kind
    if re.fullmatch(r"Sheet\d+",ws.title,re.I):return None
    header_probe = {norm_text(c.value) for c in ws._cells.values() if c.row <= 8 and c.value not in (None, "")}
    if material_type == "matrix":
        if "风控矩阵" in a1 or ("矩阵" in title and "意见" not in title) or ({"控制措施", "责任主体"} <= header_probe and ("控制措施编号" in header_probe or "控制措施/编号" in header_probe)):
            return "matrix"
    combined = title + a1
    if ("岗位职责清单" in combined or "岗位内控责任清单" in combined) and {"部门", "岗位名称", "岗位职责"} <= header_probe: return "position_duty"
    if ("系统规则清单" in combined or "系统控制规则清单" in combined) and ("规则内容" in header_probe or "系统规则" in header_probe): return "system_rule"
    if "不相容岗位清单" in combined and ("不相容业务角色" in header_probe or "岗位A" in header_probe): return "incompatible_position"
    return None


def _effective_header_max(ws: Any, rows: range) -> int:
    cols = [c.column for c in ws._cells.values() if c.row in rows and c.value not in (None, "")]
    return max(cols, default=1)


def _column_hidden(ws: Any, column: int) -> bool:
    # A grouped range can have a single dimension entry for several columns.
    # Do not access dimensions by column key: that would create new entries.
    return any(
        dim.hidden and (dim.min or column_index_from_string(key)) <= column
        <= (dim.max or dim.min or column_index_from_string(key))
        for key, dim in ws.column_dimensions.items()
    )


def _map_columns(ws: Any, sheet_type: str, aliases: dict[str, dict[str, list[str]]], *, selection_log: list[dict[str, Any]] | None = None, file: FileRecord | None = None, cache_ws=None) -> tuple[dict[str, int], dict[str, int], list[int]]:
    policy_version = (getattr(file,'_parser_policy',None) or {}).get('version', 0)
    extended = policy_version in {2, 3}
    reverse = {}
    for fid, values in aliases.get(sheet_type, {}).items():
        for value in values: reverse[norm_text(value)] = fid
    columns: dict[str, int] = {}; audit: dict[str, int] = {}; header_rows = set(); candidates: dict[str, list[int]] = {}
    headers: dict[tuple[str, int], dict[str, Any]] = {}
    max_probe = min(30, max((c.row for c in ws._cells.values()), default=1))
    max_col = _effective_header_max(ws, range(1, max_probe + 1))
    row_matches = {}
    for row in range(1, max_probe + 1):
        values = [current_text(ws.cell(row, col)) for col in range(1, max_col + 1)]
        business_values = ["" if AUDIT_HEADER.search(value) else value for value in values]
        row_matches[row] = header_row_matches(business_values, sheet_type, reverse, extended=extended, allow_fuzzy=row <= 5)
    hit_counts = {row: sum(bool(field) for field, _ in matches) for row, matches in row_matches.items()}
    primary = max(hit_counts,key=lambda row:hit_counts[row]) if max(hit_counts.values(),default=0)>=2 else 2
    scan_rows = [primary]
    if sheet_type in {"matrix", "system_rule"} and primary + 1 <= max_probe and hit_counts[primary + 1] >= 2: scan_rows.append(primary + 1)
    for row in scan_rows:
        for col in range(1, max_col + 1):
            value = current_text(ws.cell(row, col))
            nv = norm_text(value)
            if not nv: continue
            if AUDIT_HEADER.search(value): audit[value] = col; header_rows.add(row); continue
            fid,owner=row_matches[row][col - 1]
            if fid:
                header_rows.add(row)
                candidates.setdefault(fid, []).append(col)
                headers.setdefault((fid, col), {"header": value, "coordinate": ws.cell(row, col).coordinate, "owner":owner})
    eligible = {}
    own_names=header_owner_names(file) if file else set()
    body_names=unit_names(file) if file else set()
    reasons={}
    data_rows=sorted({c.row for c in ws._cells.values() if c.row>max(scan_rows) and c.value not in (None,'')})
    def signature(col):
        return tuple((current_text(_merged_value(ws,r,col)),json.dumps(deleted_spans(_merged_value(ws,r,col)),ensure_ascii=False)) for r in data_rows)
    for fid in sorted(candidates,key=lambda x:({'department':0,'responsibility':0,'position':2,'carrier':2}.get(x,1),x)):
        cols=list(dict.fromkeys(candidates[fid]))
        valid=[c for c in cols if not headers[(fid,c)]['owner'] or file is None or file.entity_code=='BASELINE' or headers[(fid,c)]['owner'] in own_names
               # 表头明确包含“是否适用”时直接保留为候选，单位前缀仅用于多列冲突时的优选。
               or (fid == 'applicability' and '是否适用' in norm_text(headers[(fid,c)]['header']))
               or (not own_names and norm_text(headers[(fid,c)]['header']) in reverse)]
        if (policy_version == 3 and fid == 'responsibility' and not valid and len(cols) == 1
                and file.entity_code and file.entity_code != 'BASELINE' and not file.entity_conflict):
            from risk_audit.readers.header_semantics import body_owner_evidence
            ids = list(dict.fromkeys(candidates.get('measure_id', [])))
            body_rows = [r for r in data_rows if len(ids) == 1 and current_text(_merged_value(ws, r, ids[0]))]
            proof = body_owner_evidence(file, [(r, current_text(_merged_value(ws, r, cols[0]))) for r in body_rows])
            headers[(fid, cols[0])]['owner_resolution'] = {'status': 'template_owner_corrected' if proof['confirmed'] else 'unconfirmed', **proof}
            if proof['confirmed']: valid = cols
        visible=[c for c in valid if not _column_hidden(ws,c)]
        options=visible or valid
        reason='visible_equivalent_preferred' if visible and any(_column_hidden(ws,c) for c in valid) else 'hidden_field_required' if options and not visible else 'single_field'
        if options and headers[(fid, options[0])].get('owner_resolution', {}).get('status') == 'template_owner_corrected':
            reason = 'template_owner_corrected'
        if len(options)>1 and file and file.entity_code!='BASELINE':
            explicitly_owned=[c for c in options if headers[(fid,c)]['owner'] and headers[(fid,c)]['owner'] in own_names]
            owned=[c for c in options if any(any(name in current_text(_merged_value(ws,r,c)) for name in body_names) for r in data_rows)]
            responsibility_header=headers.get(('responsibility',columns.get('responsibility')),{}).get('header','')
            template_kinds=[kind for kind in ('原集体企业','市公司') if f'({kind})' in norm_text(ws.title) or f'({kind})' in norm_text(responsibility_header)]
            typed=[c for c in options if len(template_kinds)==1 and f'({template_kinds[0]})' in norm_text(headers[(fid,c)]['header'])]
            if (fid=='applicability' or extended and fid in {'control_measure','carrier'}) and len(explicitly_owned)==1:
                options=explicitly_owned;reason='current_unit_header'
            elif len(owned)==1: options=owned;reason='current_unit_values'
            elif fid=='applicability_reason' and columns.get('applicability',-2)+1 in options:
                options=[columns['applicability']+1];reason='paired_with_applicability'
            elif fid=='carrier' and len(typed)==1:
                options=typed;reason='same_template_as_responsibility'
            elif fid=='position' and columns.get('department',-2)+1 in options:
                options=[columns['department']+1];reason='paired_with_department'
            elif len({signature(c) for c in options})==1:
                reason='equivalent_columns'
            else:
                options=[];reason='column_conflict'
        if not options and reason!='column_conflict':reason='owner_unconfirmed'
        eligible[fid]=options;reasons[fid]=reason
        if options:columns[fid]=max(options) if fid=='applicability' else options[0]
    if sheet_type == "matrix" and "applicability" in columns:
        # A prior generic modification note must not mask the unit's paired reason column.
        app_col = max(eligible["applicability"])
        columns["applicability"] = app_col
        if app_col + 1 in eligible.get("applicability_reason", []):
            columns["applicability_reason"] = app_col + 1
    if sheet_type == 'matrix' and policy_version != 3 and 'applicability' not in candidates and {'measure_id','control_measure'} <= columns.keys():
        # A note column can carry applicability values, but only with an
        # unambiguous whole-column profile. Never fill blanks or mix columns.
        proposed=[]
        last=max((c.column for c in ws._cells.values() if c.value not in (None,'')),default=0)
        for col in range(max(columns['measure_id'],columns['control_measure'])+1,last+1):
            if col in columns.values() or col in audit.values():continue
            header=norm_text(current_text(_merged_value(ws,primary,col)))
            if extended:
                header=next((norm_text(current_text(_merged_value(ws,r,col))) for r in reversed(scan_rows)
                             if current_text(_merged_value(ws,r,col))), '')
            if header not in {'','备注'}:continue
            cells=[_merged_value(ws,row,col) for row in data_rows]
            values=[norm_text(current_text(c)) for c in cells if current_text(c)]
            if len(values)<2 or any(c.data_type=='f' for c in cells) or not set(values)<={'适用','不适用','是','否'}:continue
            proposed.append(col)
            headers[('applicability',col)]={'header':header,'coordinate':ws.cell(primary,col).coordinate,'owner':'',
                                          'profile':{'nonempty':len(values),'values':sorted(set(values))}}
        options=[c for c in proposed if not _column_hidden(ws,c)] or proposed
        if proposed:
            candidates['applicability']=proposed
            reasons['applicability']='content_profile_conflict'
            if len(options)==1 or len({signature(c) for c in options})==1:
                columns['applicability']=options[0];header_rows.add(primary)
                reasons['applicability']='explicit_applicability_content_profile'
    if sheet_type=='matrix' and policy_version in {2, 3}:
        # Nonstandard reporting headers require actual answer values. Names or
        # test results alone are not evidence of applicability.
        from risk_audit.readers.applicability_columns import discover_answer_columns
        discovery_rows = data_rows
        if policy_version == 3:
            provisional = {fid: cols[0] for fid, cols in candidates.items() if len(set(cols)) == 1}
            discovery_rows = _record_rows(ws, {**provisional, **columns}, sheet_type, max(header_rows, default=primary) + 1)
        discover_answer_columns(ws, cache_ws, file, columns, candidates, headers, reasons,
                                scan_rows, discovery_rows, audit, _merged_value, current_text)
    if sheet_type=='matrix' and getattr(file,'_parser_policy',None):
        from risk_audit.readers.applicability_columns import resolve
        provisional={fid:cols[0] for fid,cols in candidates.items() if len(set(cols))==1}
        actual_rows=_record_rows(ws,{**provisional,**columns},sheet_type,max(header_rows,default=primary)+1)
        def read_cell(sheet,row,col,cached=None):
            cell=_merged_value(sheet,row,col)
            if cell.data_type=='f':
                value=_merged_value(cached,row,col).value if cached is not None else None
                return (str(value) if value is not None else '',value is None)
            return current_text(cell),False
        resolve(ws,cache_ws,file,columns,candidates,headers,reasons,actual_rows,
                file._parser_policy,getattr(file,'_semantic_lexicon',{}),read_cell,_column_hidden)
    if selection_log is not None:
        for fid, cols in candidates.items():
            cols = list(dict.fromkeys(cols))
            hidden = [col for col in cols if _column_hidden(ws, col)]
            if not hidden and len(cols)==1 and reasons[fid]=='single_field':
                continue
            selection_log.append({
                "sheet": ws.title, "field": fid, "selected_column": columns.get(fid),
                "reason": reasons[fid],
                "candidates": [{"column": col, **headers[(fid, col)], "hidden": col in hidden,
                                "selected": col == columns.get(fid)} for col in cols],
            })
    return columns, audit, sorted(header_rows)


def _business_header_paths(ws: Any, header_rows: list[int]) -> list[dict[str, Any]]:
    """Return every business header column, retaining merged group ancestry and coordinates."""
    if not header_rows:
        return []
    first, last = min(header_rows), max(header_rows)
    max_col = _effective_header_max(ws, range(first, last + 1))
    paths = []
    for col in range(1, max_col + 1):
        segments = []
        seen_coordinates = set()
        audit_column = False
        for row in range(first, last + 1):
            cell = _merged_value(ws, row, col)
            label = current_text(cell)
            if label and AUDIT_HEADER.search(label):
                audit_column = True
                break
            if label and cell.coordinate not in seen_coordinates:
                segments.append({"name": label, "coordinate": cell.coordinate})
                seen_coordinates.add(cell.coordinate)
        if audit_column or not segments:
            continue
        paths.append({
            "coordinate": f"{get_column_letter(col)}{last}",
            "column": col,
            "path": segments,
        })
    return paths


def _persons(text: str, identifiers: str = "") -> list[dict[str, Any]]:
    names = [x.strip() for x in re.split(r"[、,，;；\n/]+", text) if x.strip()]
    ids = [x.strip() for x in re.split(r"[、,，;；\n/]+", identifiers) if x.strip()]
    # Multiple names and IDs in a cell do not prove their one-to-one ordering.
    single_id = ids[0] if len(names) == len(ids) == 1 and ids[0] not in {'-', '无', '待定', '…', '……'} else None
    return [{"name": x, "person_id": single_id, "identity_reliable": bool(single_id)} for x in names]


def _record_rows(ws: Any, columns: dict[str, int], sheet_type: str, start: int) -> list[int]:
    signal = {
        "matrix": ["measure_id", "control_measure", "key_control_point", "business_field"],
        "position_duty": ["measure_id", "department", "position", "person_names", "duty", "role", "duty_id"],
        "system_rule": ["measure_id", "system_name", "system_rule_name", "system_rule_content"],
        "incompatible_position": ["measure_id", "department", "incompatible_roles", "position_a", "position_b", "position", "duty"],
    }[sheet_type]
    signal_cols = [columns[x] for x in signal if x in columns]
    if not signal_cols: return []
    actual_rows = sorted({c.row for c in ws._cells.values() if c.row >= start and c.column in signal_cols and c.value not in (None, "")})
    return [r for r in actual_rows if any(current_text(_merged_value(ws, r, c)) for c in signal_cols)]


def _superseded_duty_sheets(workbook: Any, material_type: str, aliases: dict, *, include_hidden: bool = False) -> list[dict[str, str]]:
    """Require both historical evidence and a complete current counterpart.

    Only sheets within the selected visibility scope may replace another sheet.
    Older crosswalks have an explicit transition title / old-value header, lack
    named people or roles, and coexist with a complete current counterpart.
    """
    candidates = []
    for ws in workbook.worksheets:
        if ws.sheet_state != 'visible' and not include_hidden:
            continue
        if _sheet_type(ws, material_type, aliases) != 'position_duty':
            continue
        columns, _, header_rows = _map_columns(ws, 'position_duty', aliases)
        historical = bool(re.search(r'(?:过渡|旧版|历史版本)', norm_text(ws.title))) or any(
            norm_text(c.value) in {'旧的', '旧编号', '原控制措施编号'}
            for c in ws._cells.values() if c.row <= 8 and c.value is not None)
        complete = {'person_names', 'role', 'duty', 'measure_id'} <= set(columns)
        # An empty current-looking template cannot supersede submitted records.
        populated = complete and bool(_record_rows(ws, columns, 'position_duty', max(header_rows, default=2) + 1))
        candidates.append((ws.title, historical, complete, populated))
    current = [title for title, historical, complete, populated in candidates if complete and populated and not historical]
    if len(current) != 1:
        return []  # Do not choose between multiple possible current versions.
    return [{'title': title, 'sheet_type': 'position_duty', 'replacement_sheet': current[0],
             'reason': '具有过渡/旧版标识或旧编号对照列，且缺少人员/角色字段；同一工作簿存在唯一包含人员、角色、职责和措施编号的当前岗位清单'}
            for title, historical, complete, populated in candidates if historical and not complete]


def parse_workbook(file: FileRecord, workbook_path: Path, aliases: dict, *, include_hidden: bool = False, input_overrides: list | None = None, loaded_workbooks: tuple | None = None, sheet_type_resolver: Any = None) -> list[ParsedSheet]:
    """读取工作簿，复用字段、缓存和删除线处理。

    参数：file 为文件归属记录，workbook_path 为源路径，aliases 为字段别名；
    include_hidden 控制隐藏表范围，input_overrides 为已确认的列对应；
    loaded_workbooks 可注入公式和缓存工作簿视图，sheet_type_resolver 可注入分类函数。
    """
    # Format detection precedes parsing; binary handles avoid rejecting OOXML
    # content merely because WPS saved it with an .et extension.
    if loaded_workbooks is None:
        formula_wb = load_compatible_workbook(
            workbook_path, data_only=False, read_only=False, rich_text=True, keep_links=True,
        )
        cache_wb = load_compatible_workbook(
            workbook_path, data_only=True, read_only=False, rich_text=True, keep_links=True,
        )
    else:
        # 区域读取器注入内存视图，避免另存源文件导致公式缓存消失。
        formula_wb, cache_wb = loaded_workbooks
    classify_sheet = sheet_type_resolver or _sheet_type
    hidden = [{'title': ws.title, 'state': ws.sheet_state} for ws in formula_wb.worksheets if ws.sheet_state != 'visible']
    file.preservation['hidden_sheets'] = hidden
    file.preservation['sheet_visibility_policy'] = 'include_hidden' if include_hidden else 'visible_only'
    file.preservation['column_visibility_policy'] = COLUMN_VISIBILITY_POLICY
    file.preservation['column_selections'] = []
    file.preservation['sheet_owner_notices'] = []
    superseded = _superseded_duty_sheets(formula_wb, file.material_type or '', aliases, include_hidden=include_hidden)
    if superseded:
        file.preservation['excluded_historical_sheets'] = superseded
    else:
        file.preservation.pop('excluded_historical_sheets', None)
    superseded_titles = {x['title'] for x in superseded}
    sheets = []
    for ws in formula_wb.worksheets:
        if ws.sheet_state != 'visible' and not include_hidden:
            continue
        if ws.title in superseded_titles:
            continue
        st = classify_sheet(ws, file.material_type or "", aliases)
        if not st: continue
        for owner in re.findall(r'\(([^()]+(?:公司|分公司|本部))\)',norm_text(ws.title)):
            names=unit_names(file)
            if names and owner not in names and file.entity_code!='BASELINE':
                file.preservation['sheet_owner_notices'].append({'sheet':ws.title,'title_owner':owner,
                    'message':f'工作表名称标注“{owner}”，与本文件主体不一致，请确认是否为模板遗留表名，并核实表内资料归属。'})
        columns, audit_cols, header_rows = _map_columns(ws, st, aliases, selection_log=file.preservation['column_selections'], file=file, cache_ws=cache_wb[ws.title])
        for item in input_overrides or []:
            if item['file_path']!=str(file.relative_path) or item['sheet']!=ws.title:continue
            fid=item['field']
            # 当前表头已明确定位字段时采用动态识别结果，旧人工确认只用于补足无法识别的字段。
            if fid in columns:continue
            if item['sha256']!=file.sha256:raise ValueError(f'已确认列对应的源文件发生变化，请重新核实：{file.relative_path}#{ws.title}')
            col=item['column']
            if fid not in aliases.get(st,{}) or col in audit_cols.values():raise ValueError('列确认只能对应本表业务字段，不能读取人工审核意见')
            if any(AUDIT_HEADER.search(current_text(ws.cell(row,col))) for row in header_rows):raise ValueError('人工审核意见列不能作为业务字段')
            columns[fid]=col
            file.preservation['column_selections']=[x for x in file.preservation['column_selections'] if (x['sheet'],x['field'])!=(ws.title,fid)]
            file.preservation['column_selections'].append({'sheet':ws.title,'field':fid,'selected_column':col,'reason':'confirmed_input_override',
                                                         'confirmation_reason':item['reason'],'source_sha256':item['sha256'],
                                                         'candidates':[{'column':col,'hidden':_column_hidden(ws,col),'selected':True}]})
        header_end = max(header_rows, default=2)
        rows = _record_rows(ws, columns, st, header_end + 1)
        active_rows = set(rows)
        if st == "matrix":
            deleted_columns = {columns[f] for f in ("control_measure", "carrier") if f in columns}
            extra = {c.row for c in ws._cells.values() if c.row > header_end and c.column in deleted_columns and deleted_spans(c)}
            rows = sorted(active_rows | extra)
        header_max = _effective_header_max(ws, range(1, header_end + 1))
        occupied_max = max((c.column for c in ws._cells.values() if c.value not in (None, "")), default=1)
        output_col = max([header_max, occupied_max, *columns.values(), *audit_cols.values()], default=1) + 1
        # 确认稿要求复用明确的同用途列，历史初审列仍单独保留。
        if '审核意见' in audit_cols:
            output_col = audit_cols['审核意见']
        output_col = _next_safe_output_column(ws, output_col)
        records = []
        cache_ws = cache_wb[ws.title]
        for row in rows:
            fields = {}
            for fid, col in columns.items():
                cell = _merged_value(ws, row, col); cached_cell = _merged_value(cache_ws, row, col)
                raw = cell.value; formula = str(raw) if isinstance(raw, str) and raw.startswith("=") else None
                cached = cached_cell.value if formula else None
                current = str(cached).strip() if formula and cached not in (None, "") else "" if formula else current_text(cell)
                state = "formula_cached" if formula and cached not in (None, "") else "formula_no_cache" if formula else "empty" if not current else "value"
                removed = deleted_spans(cell); highlighted = red_spans(cell)
                unavailable = False
                if formula and removed:
                    unavailable = cached in (None, "")
                    removed = [] if unavailable else [{"text": str(cached), "start": 0, "end": len(str(cached))}]
                if formula and highlighted:
                    highlighted = [] if cached in (None, "") else [{"text": str(cached), "start": 0, "end": len(str(cached))}]
                fields[fid] = FieldValue(raw, current, cell.coordinate, formula, cached, state, removed, unavailable, highlighted)
            seed = f"{file.entity_code}|{file.business_id}|{file.variant_id}|{file.relative_path}|{ws.title}|{row}"
            record_type = "deleted_matrix" if st == "matrix" and row not in active_rows else st
            rec = Record(record_type, file.entity_code or "", file.business_code or "", file.variant_id,
                         str(file.relative_path), ws.title, row, fields,
                         hashlib.sha256(seed.encode()).hexdigest()[:24], business_id=file.business_id or "")
            if st == "position_duty": rec.person_keys = _persons(rec.value("person_names"), rec.value("person_ids"))
            records.append(rec)
        header_paths = _business_header_paths(ws, header_rows)
        sheets.append(ParsedSheet(ws.title, st, header_rows, columns, audit_cols, output_col, rows[0] if rows else header_end + 1, records, header_paths))
    parsed_titles = {s.title for s in sheets}
    for entry in hidden:
        entry['action'] = ('skipped_hidden' if not include_hidden else 'audited' if entry['title'] in parsed_titles
                           else 'excluded_historical' if entry['title'] in superseded_titles else 'not_business')
    formula_wb.close(); cache_wb.close()
    return sheets


def parse_files(files: list[FileRecord], aliases: dict, work_dir: Path, *, include_hidden: bool = False, input_overrides: list | None = None, entity_aliases: dict | None = None, parser_policy: dict | None = None, semantic_lexicon: dict | None = None) -> list[FileRecord]:
    """逐文件解析。

    files 为输入记录，aliases 为字段别名，work_dir 为转换目录，include_hidden 控制隐藏表，
    input_overrides 为文件覆盖配置，entity_aliases 为主体别名，parser_policy 为解析策略，
    semantic_lexicon 为语义词典；日志按 BASELINE 主体标记区分公共基准与报送材料，
    解析失败写入文件记录并记录日志。
    """
    logger = logging.getLogger(__name__)
    conversion_dir = work_dir / "converted"
    for i, file in enumerate(files):
        # 公共基准与报送材料共用读取器，但日志必须明确当前处理的材料来源。
        phase = '公共基准加载' if file.entity_code == 'BASELINE' else '报送材料解析'
        file._parser_policy=parser_policy;file._semantic_lexicon=semantic_lexicon or {}
        file._confirmed_unit_names=[name for name,info in (entity_aliases or {}).items() if file.entity_code and (info.get('entity_code') if isinstance(info,dict) else info)==file.entity_code]
        if file.material_type == "explanation":
            logger.debug('保留说明材料：%s', file.relative_path)
            continue
        if file.true_format not in {"xlsx", "xls"} or file.entity_conflict:
            logger.warning('%s跳过：%s，格式=%s，主体冲突=%s', phase, file.relative_path, file.true_format, file.entity_conflict)
            continue
        logger.debug('%s开始 [%d/%d]：%s', phase, i + 1, len(files), file.relative_path)
        try:
            path = file.source
            if file.true_format == "xls":
                path, report = convert_xls(file.source, conversion_dir / f"{i:03d}")
                file.converted_from = str(file.source)
                # 转换诊断追加到现有元数据，保留扫描阶段的主体匹配状态。
                file.preservation.update(report)
                file._converted_path = str(path)
            if file.true_format in {"xlsx", "xls"}:
                # 仅对临时副本移除真实内容右侧的空白样式单元格，防止 openpyxl 构造数千万对象。
                prepared, cleanup = prepare_oversized_workbook(
                    path,
                    work_dir / 'sanitized' / file.relative_path.with_suffix('.xlsx'),
                )
                if cleanup:
                    path = prepared
                    file._preprocessed_path = str(prepared)
                    file.preservation['oversized_worksheet_cleanup'] = cleanup
                    logger.warning('%s使用临时净化副本：%s，工作表=%s', phase, file.relative_path,
                                   [item['sheet'] for item in cleanup])
            if getattr(file, '_rules_version', '') in {'1.8.0', '1.9.0', '1.9.1', '1.9.2', '1.9.3', '1.9.4', '1.9.5', '1.9.6', '1.9.7', '1.9.8', '1.9.9', '1.9.10', '1.9.11', '1.9.12', '1.9.13', '1.9.14', '1.9.15', '1.9.16', '1.9.17', '1.9.18', '1.9.19'}:
                from risk_audit.readers.confirmed_v180 import parse_workbook_v180
                file.sheets = parse_workbook_v180(file, path, aliases, include_hidden=include_hidden, input_overrides=input_overrides)
                if file.material_type == 'unclassified' and file.sheets:
                    # 以实际独立业务表补充类别，使内容识别的材料仍参与主体命名核查。
                    file.material_type = 'matrix' if any(sheet.sheet_type == 'matrix' for sheet in file.sheets) else 'three_lists'
            else:
                file.sheets = parse_workbook(file, path, aliases, include_hidden=include_hidden, input_overrides=input_overrides)
            if (not file.sheets
                    and not any(s['action'] == 'skipped_hidden' for s in file.preservation.get('hidden_sheets', []))):
                file.parse_errors.append("未识别到正式业务工作表")
            logger.log(logging.WARNING if file.parse_errors or not file.sheets else logging.DEBUG,
                       '%s完成：%s，业务表=%d，记录=%d，解析提示=%s', phase, file.relative_path,
                       len(file.sheets), sum(len(sheet.records) for sheet in file.sheets), file.parse_errors)
        except Exception as exc:
            file.parse_errors.append(f"解析失败: {exc}")
            logger.exception('%s失败：%s', phase, file.relative_path)
    return files
