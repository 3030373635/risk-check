"""验证重复展示编号的业务在审核范围和执行单元中保持隔离。"""
from pathlib import Path

from openpyxl import Workbook, load_workbook

from risk_audit.models import Entity, FileRecord
from risk_audit.runner import _business_keys
from risk_audit.submission_scope import resolve_submission_scopes
from risk_audit.util import read_json


def _file(root: Path, business_id: str, business_name: str) -> FileRecord:
    """构造业务材料；root 为输入根目录，其余参数为稳定身份和目录名。"""
    relative = Path("测试主体有限公司") / business_name / f"01风控矩阵-{business_name}.xlsx"
    return FileRecord(
        root / relative,
        relative,
        "hash",
        "xlsx",
        "A001",
        ["测试主体有限公司"],
        False,
        "01",
        "default",
        "matrix",
        business_id=business_id,
    )


def test_duplicate_business_codes_create_two_scope_entries_and_execution_units(tmp_path):
    """tmp_path 为输入根目录；两个 01 业务不得在范围或 runner 中合并。"""
    files = [
        _file(tmp_path, "sales_electricity", "营销售电"),
        _file(tmp_path, "equity_management", "股权产权管理"),
    ]
    entities = {"A001": Entity("A001", "测试主体有限公司")}

    scopes, _ = resolve_submission_scopes(tmp_path, files, entities, {}, "第二批")

    businesses = scopes["A001"]["businesses"]
    assert {(item["business_id"], item["business_code"], item["variant_id"]) for item in businesses} == {
        ("sales_electricity", "01", "default"),
        ("equity_management", "01", "default"),
    }
    assert set(_business_keys(scopes["A001"], files)) == {
        ("sales_electricity", "01", "default"),
        ("equity_management", "01", "default"),
    }


def test_second_batch_business_directory_is_not_mistaken_for_entity_directory(tmp_path):
    """tmp_path 为输入根目录；28 业务目录必须跳过并匹配其上层真实单位。"""
    from risk_audit.inventory_v180 import scan_package_v180

    root = Path(__file__).resolve().parents[2]
    registry = read_json(root / "risk-audit/rulepacks/releases/1.9.19/baseline_registry.json")
    material = tmp_path / "测试主体有限公司" / "28 教育培训-省公司版" / "28风控矩阵-教育培训.xlsx"
    material.parent.mkdir(parents=True)
    workbook = Workbook()
    workbook.active.append(["控制措施编号", "控制措施", "责任主体"])
    workbook.save(material)

    files = scan_package_v180(
        tmp_path,
        {"A001": Entity("A001", "测试主体有限公司")},
        {},
        registry,
    )

    assert len(files) == 1
    assert files[0].entity_code == "A001"
    assert files[0].business_id == "education_training"


def test_entity_name_containing_business_alias_is_not_skipped(tmp_path):
    """单位全称包含“供电服务”时仍必须优先作为正式主体目录。"""
    from risk_audit.inventory_v180 import scan_package_v180

    root = Path(__file__).resolve().parents[2]
    registry = read_json(root / "risk-audit/rulepacks/releases/1.9.19/baseline_registry.json")
    entity_name = "湖南供电服务有限公司"
    material = tmp_path / entity_name / "31 供电服务-省公司版" / "31风控矩阵-供电服务.xlsx"
    material.parent.mkdir(parents=True)
    workbook = Workbook()
    workbook.active.append(["控制措施编号", "控制措施", "责任主体"])
    workbook.save(material)

    files = scan_package_v180(
        tmp_path,
        {"A031": Entity("A031", entity_name)},
        {},
        registry,
    )

    assert len(files) == 1
    assert files[0].entity_code == "A031"
    assert files[0].business_id == "power_supply_service"


def test_full_audit_uses_the_matching_01_template_for_each_business(tmp_path):
    """tmp_path 为审核目录；两个 01 输出必须写入各自省公司模板责任主体。"""
    from run_audit import configure_soffice
    from risk_audit.runner import audit

    configure_soffice()
    root = Path(__file__).resolve().parents[2]
    input_root = tmp_path / "input"
    unit = input_root / "测试主体有限公司"
    measure_ids = {
        "01 营销售电": "营销售电业务-1.系统客户信息创建-控制措施01",
        "01 股权（产权）管理": "股权（产权）管理－1.股权投资项目提出和遴选－控制措施01",
    }
    sources = []
    for directory_name, measure_id in measure_ids.items():
        directory = unit / directory_name
        directory.mkdir(parents=True)
        path = directory / f"01风控矩阵-{directory_name}-测试主体有限公司.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "风控矩阵"
        sheet.append(["控制措施编号", "控制措施", "控制载体", "责任主体", "是否适用及原因"])
        sheet.append([measure_id, "核对资料", "审批单", "测试主体有限公司-测试部-测试专责", "适用"])
        workbook.save(path)
        sources.append(path)
    roster_path = tmp_path / "entities.xlsx"
    roster = Workbook()
    roster.active.append(["单位名称", "单位代码", "上级单位", "对应主业单位", "是否存续"])
    roster.active.append(["测试主体有限公司", "A001", "", "", "是"])
    roster.save(roster_path)

    result = audit(
        input_root,
        tmp_path / "output",
        root / "risk-audit/rulepacks/releases/1.9.19",
        roster_path,
        root,
        tmp_path / "runs",
        run_id="duplicate-01",
    )

    assert {(item["business_id"], item["business_code"]) for item in result["business_results"]} == {
        ("sales_electricity", "01"),
        ("equity_management", "01"),
    }
    reference_values = {}
    for source in sources:
        output = load_workbook(tmp_path / "output" / source.relative_to(input_root), data_only=False)
        sheet = output["风控矩阵"]
        headers = {cell.value: cell.column for cell in sheet[1]}
        reference_values[source.parent.name] = sheet.cell(2, headers["省公司版本责任主体（核对后删除）"]).value
    assert reference_values["01 营销售电"] != reference_values["01 股权（产权）管理"]
    assert "营销" in str(reference_values["01 营销售电"])
    assert "投资" in str(reference_values["01 股权（产权）管理"])
