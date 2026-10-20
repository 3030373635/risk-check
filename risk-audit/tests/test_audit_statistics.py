"""审核统计表按单位、业务、规则组合和意见内容汇总的行为验证。"""
from __future__ import annotations

import json
from dataclasses import FrozenInstanceError
from pathlib import Path

from lxml import html
from openpyxl import load_workbook
import pytest

from risk_audit import audit_statistics
from risk_audit.audit_statistics import (
    collect_audit_statistics,
    write_audit_statistics as _write_audit_statistics,
    write_audit_statistics_html as _write_audit_statistics_html,
)
from risk_audit.models import FileRecord, Finding


def write_audit_statistics(
    path: Path,
    files: list[FileRecord],
    ownership: dict[str, list[dict]],
    package_name: str,
    **options,
) -> None:
    """为既有行为用例归集一次并写出 Excel。

    Args:
        path: Excel 输出路径。
        files: 审核文件列表。
        ownership: 写回意见归属。
        package_name: 审核包名称。
        options: 归集函数的可选参数。
    """
    statistics = collect_audit_statistics(
        files,
        ownership,
        package_name,
        **options,
    )
    _write_audit_statistics(path, statistics)


def write_audit_statistics_html(
    path: Path,
    files: list[FileRecord],
    ownership: dict[str, list[dict]],
    package_name: str,
    **options,
) -> None:
    """为既有行为用例归集一次并写出 HTML。

    Args:
        path: HTML 输出路径。
        files: 审核文件列表。
        ownership: 写回意见归属。
        package_name: 审核包名称。
        options: 归集函数的可选参数。
    """
    statistics = collect_audit_statistics(
        files,
        ownership,
        package_name,
        **options,
    )
    _write_audit_statistics_html(path, statistics)


def _file(tmp_path: Path, relative_path: str, material_type: str, business_code: str | None) -> FileRecord:
    """构建审核统计样例文件；tmp_path 为测试目录，relative_path 为包内路径。"""
    return FileRecord(
        tmp_path / Path(relative_path).name,
        Path(relative_path),
        f'hash-{relative_path}',
        'xlsx',
        '205H',
        [],
        False,
        business_code,
        'default',
        material_type,
    )


def _finding(
    message: str,
    file_path: str,
    sheet: str,
    row: int | None,
    *,
    entity_code: str = "205H",
    business_id: str = "sales_electricity",
    location_policy: str = "row",
) -> Finding:
    """构建统计索引测试使用的审核结果。

    Args:
        message: 审核意见原文。
        file_path: 意见所属文件路径。
        sheet: 意见所属工作表。
        row: 意见所属行号。
        entity_code: 会计主体代码。
        business_id: 标准业务标识。
        location_policy: 意见定位策略。
    """
    return Finding(
        finding_key=f"{file_path}:{sheet}:{row}:{business_id}",
        finding_id="finding-id",
        rule_id="test.rule",
        check_id="test_check",
        display_code="R1",
        revision=1,
        severity="violation",
        entity_code=entity_code,
        entity_name="测试单位",
        business_code="01",
        variant_id="default",
        file_path=file_path,
        sheet=sheet,
        row=row,
        message=message,
        evidence={},
        location_policy=location_policy,
        business_id=business_id,
    )


def test_occurrence_index_selects_exact_file_sheet_and_row(tmp_path: Path) -> None:
    """相同文本必须按文件、工作表和行号精确命中。

    Args:
        tmp_path: pytest 提供的隔离目录。
    """
    file = _file(tmp_path, "单位/01业务/目标.xlsx", "matrix", "01")
    message = "【第1条】请核实。"
    wrong = _finding(message, str(file.relative_path), "风控矩阵", 8)
    expected = _finding(message, str(file.relative_path), "风控矩阵", 9)

    index = audit_statistics._build_finding_occurrence_index([wrong, expected])
    actual = audit_statistics._find_occurrence(
        index,
        message,
        {"sheet": "风控矩阵", "cell": "H9"},
        file,
    )

    assert actual is expected


def test_occurrence_index_does_not_scan_same_message_candidates(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """高频相同文本归集不得回退到旧候选列表扫描。

    Args:
        tmp_path: pytest 提供的隔离目录。
        monkeypatch: pytest 提供的补丁工具。
    """
    file = _file(tmp_path, "单位/01业务/目标.xlsx", "matrix", "01")
    file.business_id = "sales_electricity"
    message = "【第1条】高频意见。"
    findings = [
        _finding(message, f"单位/01业务/文件-{index}.xlsx", "风控矩阵", index + 2)
        for index in range(10_000)
    ]
    findings.append(_finding(message, str(file.relative_path), "风控矩阵", 5))
    ownership = {
        str(file.relative_path): [{
            "sheet": "风控矩阵",
            "cell": "H5",
            "program_text": message,
        }],
    }

    def fail_old_scan(*_args, **_kwargs):
        """旧扫描入口一旦被调用即失败；参数为历史函数调用值。"""
        raise AssertionError("不得扫描相同文本候选列表")

    monkeypatch.setattr(
        audit_statistics,
        "_finding_for_occurrence",
        fail_old_scan,
        raising=False,
    )
    statistics = audit_statistics.collect_audit_statistics(
        [file],
        ownership,
        "报送包",
        findings=findings,
    )

    assert statistics.units[0].businesses[0].records == ("【第1条】高频意见。 1条。",)


def test_occurrence_index_rejects_conflicting_routes_at_same_location() -> None:
    """同一精确位置的不同业务路由必须显式报错；无参数。"""
    message = "【第1条】请核实。"
    first = _finding(message, "单位/01业务/目标.xlsx", "风控矩阵", 9)
    second = _finding(
        message,
        "单位/01业务/目标.xlsx",
        "风控矩阵",
        9,
        business_id="trade_purchase",
    )

    with pytest.raises(ValueError, match=r"目标\.xlsx.*风控矩阵.*9"):
        audit_statistics._build_finding_occurrence_index([first, second])


@pytest.mark.parametrize("finding_path", ["", "缺失业务/来源.xlsx"])
def test_occurrence_index_routes_material_finding_by_entity_fallback(
    tmp_path: Path,
    finding_path: str,
) -> None:
    """资料级意见必须按主体回退到真实 Finding。

    Args:
        tmp_path: pytest 提供的隔离目录。
        finding_path: 资料级 Finding 的可选来源路径。
    """
    file = _file(tmp_path, "单位/01业务/承载.xlsx", "matrix", "01")
    message = "【第17条】缺少目标业务材料。"
    expected = _finding(
        message,
        finding_path,
        "缺失表" if finding_path else "",
        None,
        business_id="missing_business",
        location_policy="material",
    )

    index = audit_statistics._build_finding_occurrence_index([expected])
    actual = audit_statistics._find_occurrence(
        index,
        message,
        {"sheet": "风控矩阵", "cell": "J4"},
        file,
    )

    assert actual is expected


def test_shared_statistics_data_drives_excel_and_html(tmp_path: Path) -> None:
    """同一不可变统计对象必须同时驱动 Excel 和 HTML。

    Args:
        tmp_path: pytest 提供的隔离目录。
    """
    file = _file(tmp_path, "单位/01业务/目标.xlsx", "matrix", "01")
    message = "【第1条】请核实。"
    ownership = {
        str(file.relative_path): [{
            "sheet": "风控矩阵",
            "cell": "H5",
            "program_text": message,
        }],
    }
    statistics = audit_statistics.collect_audit_statistics(
        [file],
        ownership,
        "报送包",
        findings=[_finding(message, str(file.relative_path), "风控矩阵", 5)],
    )
    excel_target = tmp_path / "审核统计表.xlsx"
    html_target = tmp_path / "审核统计表.html"

    audit_statistics.write_audit_statistics(excel_target, statistics)
    audit_statistics.write_audit_statistics_html(html_target, statistics)

    excel_rows = list(load_workbook(excel_target)["审核统计表"].values)
    document = html.fromstring(html_target.read_text(encoding="utf-8"))
    html_source = json.loads(document.get_element_by_id("audit-source-data").text)
    assert excel_rows == [
        ("单位", "备注"),
        ("报送包/单位", "01业务：\n【第1条】请核实。 1条。"),
    ]
    assert html_source == {
        "units": [{
            "path": "报送包/单位",
            "businesses": [{
                "name": "01业务",
                "records": ["【第1条】请核实。 1条。"],
            }],
        }],
    }
    with pytest.raises(FrozenInstanceError):
        statistics.units = ()


def test_finalization_benchmark_reports_stable_metrics() -> None:
    """基准入口必须返回稳定字段和非负计时；无参数。"""
    from tools.benchmark_task_finalization import run_benchmark

    metrics = run_benchmark(findings_count=20, ownership_count=10)

    assert metrics.keys() == {
        "findings",
        "ownership",
        "collect_seconds",
        "excel_seconds",
        "html_seconds",
        "total_seconds",
    }
    assert metrics["findings"] == 20
    assert metrics["ownership"] == 10
    assert all(
        isinstance(metrics[key], float) and metrics[key] >= 0
        for key in (
            "collect_seconds",
            "excel_seconds",
            "html_seconds",
            "total_seconds",
        )
    )


def test_statistics_groups_same_rule_combination_by_opinion_text(tmp_path):
    """tmp_path 为隔离目录；同一规则组合的不同意见必须分别统计。"""
    matrix = _file(tmp_path, 'a/b/c/d/风控矩阵.xlsx', 'matrix', '09')
    lists = _file(tmp_path, 'a/b/c/d/三清单.xlsx', 'three_lists', '09')
    system = _file(tmp_path, 'a/b/c/e/风控矩阵.xlsx', 'matrix', '07')
    ownership = {
        str(matrix.relative_path): [
            {'sheet': '岗位清单', 'cell': 'H2', 'program_text': '【第10条】岗位职责请按照国网标准句式编制。'},
            {'sheet': '风控矩阵', 'cell': 'H4', 'program_text': '【第5条】【第3条】待核实：未能读取必需字段。'},
        ],
        str(lists.relative_path): [
            {'sheet': '岗位清单', 'cell': 'I3', 'program_text': '【第10条】岗位职责请按照国网标准句式编制。'},
            {'sheet': '岗位清单', 'cell': 'I5', 'program_text': '【第3条】【第5条】字段对应待核实。'},
        ],
        str(system.relative_path): [
            {'sheet': '系统规则', 'cell': 'G2', 'program_text': '【第7条】请核实与矩阵对应的系统规则是否应该做出修改。'},
        ],
    }
    target = tmp_path / '审核统计表.xlsx'

    write_audit_statistics(target, [matrix, lists, system], ownership, '包名')

    sheet = load_workbook(target)['审核统计表']
    assert list(sheet.values) == [
        ('单位', '备注'),
        ('包名/a/b/c',
         'd：\n【第3条】【第5条】字段对应待核实。 1条。\n'
         '【第3条】【第5条】待核实：未能读取必需字段。 1条。\n'
         '【第10条】岗位职责请按照国网标准句式编制。 2条。'),
        ('包名/a/b/c',
         'e：\n【第7条】请核实与矩阵对应的系统规则是否应该做出修改。 1条。'),
    ]


def test_statistics_routes_material_finding_to_its_actual_business(tmp_path: Path) -> None:
    """资料级缺报意见必须显示在实际缺失业务下，而不是承载意见的首个业务下。

    Args:
        tmp_path: pytest 提供的隔离目录。
    """

    sales = _file(
        tmp_path,
        '大足供电分公司/01风控矩阵-营销售电/01风控矩阵-营销售电.xlsx',
        'matrix',
        '01',
    )
    sales.business_id = 'sales_electricity'
    message = (
        '【第17条】检修运维成本项目管理缺少供电分公司适用清单要求的关键控制点：1、2。'
        '请补充；如确实不适用，请说明原因。'
    )
    ownership = {
        str(sales.relative_path): [{
            'sheet': '风控矩阵',
            'cell': 'J4',
            'program_text': message,
        }],
    }
    finding = Finding(
        finding_key='finding-key',
        finding_id='finding-id',
        rule_id='chongqing.power_supply_branch_key_points',
        check_id='power_supply_branch_key_points',
        display_code='R17',
        revision=3,
        severity='violation',
        entity_code='205H',
        entity_name='大足供电分公司',
        business_code='',
        variant_id='default',
        file_path='',
        sheet='',
        row=None,
        message=message,
        evidence={
            'issue_type': 'missing_key_points',
            'business_name': '检修运维成本项目管理',
            'missing_key_points': [[1], [2]],
        },
        location_policy='material',
        business_id='maintenance_cost_project_management',
    )
    target = tmp_path / '审核统计表.xlsx'

    write_audit_statistics(
        target,
        files=[sales],
        ownership=ownership,
        package_name='报送包',
        findings=[finding],
    )

    assert list(load_workbook(target)['审核统计表'].values) == [
        ('单位', '备注'),
        ('报送包/大足供电分公司', '01风控矩阵-营销售电：\n无'),
        (
            '报送包/大足供电分公司',
            '检修运维成本项目管理：\n'
            '【第17条】缺少供电分公司适用清单要求的关键控制点 2项：\n'
            '1\n'
            '2',
        ),
    ]


def test_statistics_reuses_alias_matched_business_for_material_finding(tmp_path: Path) -> None:
    """资料级缺失意见应复用通过目录别名归一的实际业务分组。

    Args:
        tmp_path: pytest 提供的隔离目录。
    """

    finance = _file(
        tmp_path,
        '大足供电分公司/11.财务管理矩阵及三清单/财务管理风控矩阵.xlsx',
        'matrix',
        None,
    )
    finance.business_id = 'unmatched_1234'
    message = (
        '【第17条】财务管理缺少供电分公司适用清单要求的关键控制点：45。'
        '请补充；如确实不适用，请说明原因。'
    )
    ownership = {
        str(finance.relative_path): [{
            'sheet': '风控矩阵',
            'cell': 'J4',
            'program_text': message,
        }],
    }
    finding = Finding(
        finding_key='finding-key-alias',
        finding_id='finding-id-alias',
        rule_id='chongqing.power_supply_branch_key_points',
        check_id='power_supply_branch_key_points',
        display_code='R17',
        revision=3,
        severity='violation',
        entity_code='205H',
        entity_name='大足供电分公司',
        business_code='',
        variant_id='default',
        file_path='',
        sheet='',
        row=None,
        message=message,
        evidence={
            'issue_type': 'missing_key_points',
            'business_name': '财务管理',
            'business_file_path': str(finance.relative_path),
            'missing_key_points': [[45]],
        },
        location_policy='material',
        business_id='financial_management',
    )
    target = tmp_path / '审核统计表.xlsx'

    write_audit_statistics(
        target,
        files=[finance],
        ownership=ownership,
        package_name='报送包',
        findings=[finding],
    )

    assert list(load_workbook(target)['审核统计表'].values) == [
        ('单位', '备注'),
        (
            '报送包/大足供电分公司',
            '11.财务管理矩阵及三清单：\n'
            '【第17条】缺少供电分公司适用清单要求的关键控制点 1项：\n'
            '45',
        ),
    ]


def test_statistics_groups_extra_key_points_with_names_and_rows(tmp_path: Path) -> None:
    """适用清单外控制点应按类型汇总，并展示名称和矩阵定位行。

    Args:
        tmp_path: pytest 提供的隔离目录。
    """

    sales = _file(
        tmp_path,
        '大足供电分公司/01风控矩阵-营销售电/01风控矩阵-营销售电.xlsx',
        'matrix',
        '01',
    )
    sales.business_id = 'sales_electricity'
    messages = [
        (
            '【第17条】存在供电分公司适用清单外的关键控制点：4.供用电合同签订。'
            '请核实是否适用；适用则保留并说明依据，不适用则删除。'
        ),
        (
            '【第17条】存在供电分公司适用清单外的关键控制点：6.电价码制定与执行。'
            '请核实是否适用；适用则保留并说明依据，不适用则删除。'
        ),
    ]
    ownership = {
        str(sales.relative_path): [
            {'sheet': '营销售电', 'cell': 'J13', 'program_text': messages[0]},
            {'sheet': '营销售电', 'cell': 'J20', 'program_text': messages[1]},
        ],
    }
    findings = [
        Finding(
            finding_key=f'finding-key-{index}',
            finding_id=f'finding-id-{index}',
            rule_id='chongqing.power_supply_branch_key_points',
            check_id='power_supply_branch_key_points',
            display_code='R17',
            revision=3,
            severity='violation',
            entity_code='205H',
            entity_name='大足供电分公司',
            business_code='01',
            variant_id='default',
            file_path=str(sales.relative_path),
            sheet='营销售电',
            row=row,
            message=message,
            evidence={
                'issue_type': 'extra_key_point',
                'business_name': '营销售电',
                'key_point': [point],
                'key_point_name': point_name,
            },
            location_policy='row',
            business_id='sales_electricity',
        )
        for index, (row, point, point_name, message) in enumerate([
            (13, 4, '4.供用电合同签订', messages[0]),
            (20, 6, '6.电价码制定与执行', messages[1]),
        ])
    ]
    target = tmp_path / '审核统计表.xlsx'

    write_audit_statistics(
        target,
        files=[sales],
        ownership=ownership,
        package_name='报送包',
        findings=findings,
    )

    assert load_workbook(target)['审核统计表']['B2'].value == (
        '01风控矩阵-营销售电：\n'
        '【第17条】存在供电分公司适用清单外的关键控制点 2项：\n'
        '4.供用电合同签订（营销售电第13行）\n'
        '6.电价码制定与执行（营销售电第20行）'
    )

    html_target = tmp_path / '审核统计表.html'
    write_audit_statistics_html(
        html_target,
        files=[sales],
        ownership=ownership,
        package_name='报送包',
        findings=findings,
    )

    html_text = html_target.read_text(encoding='utf-8')
    document = html.fromstring(html_text)
    source = document.get_element_by_id('audit-source-data').text
    assert json.loads(source)['units'][0]['businesses'][0]['records'] == [
        '【第17条】存在供电分公司适用清单外的关键控制点 2项：\n'
        '4.供用电合同签订（营销售电第13行）\n'
        '6.电价码制定与执行（营销售电第20行）',
    ]
    assert 'white-space: pre-wrap' in html_text


def test_html_statistics_uses_same_grouped_content_as_excel(tmp_path: Path) -> None:
    """HTML 必须以与 Excel 相同的主体、业务和意见数据生成；tmp_path 为隔离目录。"""
    matrix = _file(tmp_path, '一区/07物资采购/风控矩阵.xlsx', 'matrix', '07')
    ownership = {
        str(matrix.relative_path): [{
            'sheet': '风控矩阵',
            'cell': 'H2',
            'program_text': '【第10条】请补充岗位职责。</script><script>unsafe()</script>',
        }],
    }
    target = tmp_path / '审核统计表.html'

    write_audit_statistics_html(target, [matrix], ownership, '报送包')

    document = html.fromstring(target.read_text(encoding='utf-8'))
    source = document.get_element_by_id('audit-source-data').text
    assert json.loads(source) == {
        'units': [{
            'path': '报送包/一区',
            'businesses': [{
                'name': '07物资采购',
                'records': ['【第10条】请补充岗位职责。</script><script>unsafe()</script> 1条。'],
            }],
        }],
    }
    assert '<script>unsafe()</script>' not in target.read_text(encoding='utf-8')


def test_statistics_keeps_different_business_identities_in_one_directory_separate(tmp_path: Path) -> None:
    """同一物理目录中的不同标准业务必须分开统计。

    Args:
        tmp_path: pytest 提供的隔离目录。
    """

    matrix = _file(
        tmp_path,
        '8.数字化与研发矩阵及“三清单”/08风控矩阵-数字化与研发投入.xlsx',
        'matrix',
        '08',
    )
    matrix.business_id = 'digitalization_rd_investment'
    matrix.preservation['business_name'] = '数字化与研发投入'
    duties = _file(
        tmp_path,
        '8.数字化与研发矩阵及“三清单”/岗位内控责任清单.xlsx',
        'three_lists',
        '30',
    )
    duties.business_id = 'property_logistics'
    ownership = {
        str(matrix.relative_path): [{
            'sheet': '风控矩阵',
            'cell': 'H2',
            'program_text': '【第3条】请核实矩阵字段。',
        }],
        str(duties.relative_path): [{
            'sheet': '岗位职责清单',
            'cell': 'I2',
            'program_text': '【第5条】请补充人员姓名。',
        }],
    }
    target = tmp_path / '审核统计表.html'

    write_audit_statistics_html(target, [matrix, duties], ownership, '新建文件')

    document = html.fromstring(target.read_text(encoding='utf-8'))
    source = document.get_element_by_id('audit-source-data').text
    assert json.loads(source) == {
        'units': [{
            'path': '新建文件',
                'businesses': [
                    {
                        'name': '8.数字化与研发矩阵及“三清单”',
                        'records': ['【第3条】请核实矩阵字段。 1条。'],
                    },
                    {
                        'name': '8.数字化与研发矩阵及“三清单”',
                        'records': ['【第5条】请补充人员姓名。 1条。'],
                    },
                ],
        }],
    }


def test_statistics_merges_nested_materials_by_canonical_business_directory(
    tmp_path: Path,
) -> None:
    """统计表必须使用规范业务目录，不得把附件子目录当成业务或主体。

    Args:
        tmp_path: pytest 提供的隔离目录。
    """

    business_directory = Path('1.营销售电矩阵及“三清单”')
    matrix = _file(
        tmp_path,
        f'{business_directory}/01风控矩阵-营销售电.xlsx',
        'matrix',
        '01',
    )
    duties = _file(
        tmp_path,
        f'{business_directory}/三清单/岗位内控责任清单.xlsx',
        'three_lists',
        '01',
    )
    for file in (matrix, duties):
        file.business_id = 'sales_electricity'
        file.business_directory = business_directory
    ownership = {
        str(matrix.relative_path): [{
            'sheet': '风控矩阵',
            'cell': 'H2',
            'program_text': '【第3条】请核实矩阵字段。',
        }],
        str(duties.relative_path): [{
            'sheet': '岗位内控责任清单',
            'cell': 'I2',
            'program_text': '【第5条】请补充人员姓名。',
        }],
    }
    target = tmp_path / '审核统计表.html'

    write_audit_statistics_html(target, [matrix, duties], ownership, '永川公司矩阵')

    document = html.fromstring(target.read_text(encoding='utf-8'))
    source = json.loads(document.get_element_by_id('audit-source-data').text)
    assert source == {
        'units': [{
            'path': '永川公司矩阵',
            'businesses': [{
                'name': '1.营销售电矩阵及“三清单”',
                'records': [
                    '【第3条】请核实矩阵字段。 1条。',
                    '【第5条】请补充人员姓名。 1条。',
                ],
            }],
        }],
    }


def test_statistics_writes_each_business_on_separate_row_without_truncation(tmp_path: Path) -> None:
    """每个业务必须独占一行，避免主体汇总超过 Excel 单元格字符上限。

    Args:
        tmp_path: pytest 提供的隔离目录。
    """

    first = _file(tmp_path, '某公司/01营销售电/风控矩阵.xlsx', 'matrix', '01')
    second = _file(tmp_path, '某公司/02交易与购电/风控矩阵.xlsx', 'matrix', '02')
    first_message = f'【第1条】第一项问题：{"甲" * 18000}'
    second_message = f'【第2条】第二项问题：{"乙" * 18000}'
    ownership = {
        str(first.relative_path): [{
            'sheet': '风控矩阵',
            'cell': 'H2',
            'program_text': first_message,
        }],
        str(second.relative_path): [{
            'sheet': '风控矩阵',
            'cell': 'H2',
            'program_text': second_message,
        }],
    }
    target = tmp_path / '审核统计表.xlsx'

    write_audit_statistics(target, [first, second], ownership, '报送包')

    rows = list(load_workbook(target)['审核统计表'].values)
    assert len(rows) == 3
    assert rows[1][0] == rows[2][0] == '报送包/某公司'
    assert rows[1][1] == f'01营销售电：\n{first_message} 1条。'
    assert rows[2][1] == f'02交易与购电：\n{second_message} 1条。'


def test_statistics_writes_none_for_recognized_business_without_opinions(tmp_path):
    """tmp_path 为隔离目录；已识别但无程序意见的业务仍须输出“无”。"""
    sales = _file(tmp_path, '1.某供电公司/01营销售电/风控矩阵.xlsx', 'matrix', '01')
    trade = _file(tmp_path, '1.某供电公司/02交易与购电/风控矩阵.xlsx', 'matrix', '02')
    ownership = {
        str(sales.relative_path): [
            {'sheet': '风控矩阵', 'cell': 'H2', 'program_text': '【第10条】请修改岗位职责。'},
        ],
    }
    target = tmp_path / '审核统计表.xlsx'

    write_audit_statistics(target, [sales, trade], ownership, '包名')

    assert list(load_workbook(target)['审核统计表'].values) == [
        ('单位', '备注'),
        ('包名/1.某供电公司',
         '01营销售电：\n【第10条】请修改岗位职责。 1条。'),
        ('包名/1.某供电公司', '02交易与购电：\n无'),
    ]


def test_statistics_writes_none_for_unmapped_business_directory(tmp_path):
    """未映射规则身份的业务目录也必须出现在统计中，不能因无意见而消失。"""
    file = _file(tmp_path, '测试主体/17.全新专项矩阵及“三清单”/风控矩阵.xlsx', 'matrix', None)
    file.business_id = 'unmatched_directory_identity'
    target = tmp_path / '审核统计表.xlsx'

    write_audit_statistics(target, [file], {}, '包名')

    assert list(load_workbook(target)['审核统计表'].values) == [
        ('单位', '备注'),
        ('包名/测试主体', '17.全新专项矩阵及“三清单”：\n无'),
    ]


def test_statistics_uses_selected_business_root_and_its_parent_subject(tmp_path):
    """直接选择业务目录时，统计必须分别显示外层主体和业务根目录。"""
    sales = _file(tmp_path, '01风控矩阵-营销售电.xlsx', 'matrix', '01')
    sales.business_id = 'sales_electricity'
    sales.preservation['business_name'] = '营销售电'
    equity = _file(tmp_path, '01风控矩阵-股权管理.xlsx', 'matrix', '01')
    equity.business_id = 'equity_management'
    equity.preservation['business_name'] = '股权（产权）管理'
    target = tmp_path / '审核统计表.xlsx'

    write_audit_statistics(
        target,
        [sales, equity],
        {},
        '11.财务管理矩阵及“三清单”',
        root_subject_name='重庆报送主题',
    )

    sheet = load_workbook(target)['审核统计表']
    assert list(sheet.values) == [
        ('单位', '备注'),
        ('重庆报送主题', '股权（产权）管理：\n无'),
        ('重庆报送主题', '营销售电：\n无'),
    ]


def test_statistics_merges_same_business_across_directories(tmp_path: Path) -> None:
    """同一主体、业务和变体分散在多个目录时必须合并统计。

    Args:
        tmp_path: pytest 提供的隔离输出目录。
    """

    matrix = _file(tmp_path, '测试主体/首次报送/02 财务管理/矩阵.xlsx', 'matrix', '02')
    duties = _file(tmp_path, '测试主体/补充报送/财务管理补充/岗位清单.xlsx', 'three_lists', '02')
    for file in (matrix, duties):
        file.entity_code = 'CQ001'
        file.business_id = 'financial_management'
    ownership = {
        str(matrix.relative_path): [{
            'sheet': '风控矩阵',
            'cell': 'H2',
            'program_text': '【第7条】请核实岗位清单。',
        }],
        str(duties.relative_path): [{
            'sheet': '岗位内控责任清单',
            'cell': 'I2',
            'program_text': '【第14条】请核实是否为一人多岗。',
        }],
    }
    target = tmp_path / '审核统计表.html'

    write_audit_statistics_html(target, [matrix, duties], ownership, '报送包')

    document = html.fromstring(target.read_text(encoding='utf-8'))
    source = json.loads(document.get_element_by_id('audit-source-data').text)
    assert source == {
        'units': [{
            'path': '报送包/测试主体/补充报送、报送包/测试主体/首次报送',
            'businesses': [{
                'name': '02 财务管理、财务管理补充',
                'records': [
                    '【第7条】请核实岗位清单。 1条。',
                    '【第14条】请核实是否为一人多岗。 1条。',
                ],
            }],
        }],
    }


def test_statistics_formats_header_and_protects_excel_text(tmp_path):
    """tmp_path 为隔离目录；表头样式明显，外部目录文本不能触发 Excel 公式。"""
    file = _file(tmp_path, 'a/=1+1/风控矩阵.xlsx', 'matrix', '09')
    ownership = {str(file.relative_path): [
        {'sheet': '风控矩阵', 'cell': 'H2', 'program_text': '【第10条】请修改。'},
    ]}
    target = tmp_path / '审核统计表.xlsx'

    write_audit_statistics(target, [file], ownership, '=PACKAGE')

    sheet = load_workbook(target, data_only=False)['审核统计表']
    assert sheet['A1'].font.bold is True
    assert sheet['A1'].fill.fill_type == 'solid'
    assert sheet['A1'].alignment.horizontal == 'center'
    assert sheet['B2'].alignment.wrap_text is True
    assert sheet.freeze_panes == 'A2'
    assert sheet.auto_filter.ref == 'A1:B2'
    assert sheet['A2'].data_type == 's' and sheet['A2'].value.startswith("'=PACKAGE/a")
    assert sheet['B2'].data_type == 's' and sheet['B2'].value.startswith("'=1+1：")


def test_statistics_uses_subject_business_file_directory_order(tmp_path):
    """tmp_path 为隔离目录；统计必须按主体、业务、文件的目录顺序归组。"""
    sales = _file(tmp_path, '1.国网某供电分公司/1.营销售电-ok/01风控矩阵.xlsx', 'matrix', '01')
    trade = _file(tmp_path, '1.国网某供电分公司/2.交易与购电-ok/02风控矩阵.xlsx', 'matrix', '02')
    ownership = {
        str(sales.relative_path): [
            {'sheet': '风控矩阵', 'cell': 'H2', 'program_text': '【第10条】请修改岗位职责。'},
        ],
        str(trade.relative_path): [
            {'sheet': '风控矩阵', 'cell': 'H2', 'program_text': '【第7条】请核实系统规则。'},
        ],
    }
    target = tmp_path / '审核统计表.xlsx'

    write_audit_statistics(target, [sales, trade], ownership, '包名')

    assert list(load_workbook(target)['审核统计表'].values) == [
        ('单位', '备注'),
        ('包名/1.国网某供电分公司',
         '1.营销售电-ok：\n【第10条】请修改岗位职责。 1条。'),
        ('包名/1.国网某供电分公司',
         '2.交易与购电-ok：\n【第7条】请核实系统规则。 1条。'),
    ]


def test_statistics_does_not_infer_inverted_directory_order(tmp_path):
    """tmp_path 为隔离目录；统计不得根据名称猜测倒置的主体与业务目录。"""
    file = _file(tmp_path, '01营销售电/1.某技术技能培训中心/风控矩阵.xlsx', 'matrix', '01')
    ownership = {str(file.relative_path): [
        {'sheet': '风控矩阵', 'cell': 'H2', 'program_text': '【第10条】请修改岗位职责。'},
    ]}
    target = tmp_path / '审核统计表.xlsx'

    write_audit_statistics(target, [file], ownership, '包名')

    assert list(load_workbook(target)['审核统计表'].values) == [
        ('单位', '备注'),
        ('包名/01营销售电', '1.某技术技能培训中心：\n【第10条】请修改岗位职责。 1条。'),
    ]


def test_statistics_warns_and_keeps_header_when_written_file_is_unknown(tmp_path, caplog):
    """tmp_path/caplog 为测试工具；无法归属的写回记录告警而不中断输出。"""
    ownership = {'未知文件.xlsx': [
        {'sheet': '风控矩阵', 'cell': 'H2', 'program_text': '【第1条】缺少业务材料。'},
    ]}
    target = tmp_path / '审核统计表.xlsx'

    write_audit_statistics(target, [], ownership, '包名')

    assert '审核统计表未能定位意见所属业务目录' in caplog.text
    assert list(load_workbook(target)['审核统计表'].values) == [('单位', '备注')]


def test_statistics_splits_actual_program_opinions_from_one_written_cell(tmp_path):
    """tmp_path 为隔离目录；同一审核意见单元格内的独立意见分别按规则组合统计。"""
    file = _file(tmp_path, 'a/b/c/d/风控矩阵.xlsx', 'matrix', '09')
    ownership = {
        str(file.relative_path): [{
            'sheet': '风控矩阵',
            'cell': 'H2',
            'program_text': '【第10条】请修改岗位职责。\n【第3条】【第5条】请核实必需字段。',
        }],
    }
    target = tmp_path / '审核统计表.xlsx'

    write_audit_statistics(target, files=[file], ownership=ownership, package_name='包名')

    assert load_workbook(target)['审核统计表']['B2'].value == (
        'd：\n【第3条】【第5条】请核实必需字段。 1条。\n'
        '【第10条】请修改岗位职责。 1条。'
    )


def test_statistics_includes_preprocessing_opinions(tmp_path):
    """系统类型预处理提示必须按单位和业务写入审核统计表。

    Args:
        tmp_path: pytest 提供的隔离目录。
    """

    file = _file(tmp_path, 'a/b/c/d/三清单.xlsx', 'three_lists', '09')
    ownership = {
        str(file.relative_path): [{
            'sheet': '系统控制规则清单',
            'cell': 'G2',
            'program_text': (
                '【预处理】请补充系统类型，该列填报枚举值：一级部署系统、'
                '二级部署系统、三级部署系统，请根据系统的实际情况填报。'
            ),
        }],
    }
    target = tmp_path / '审核统计表.xlsx'

    write_audit_statistics(target, files=[file], ownership=ownership, package_name='包名')

    assert load_workbook(target)['审核统计表']['B2'].value == (
        'd：\n【预处理】请补充系统类型，该列填报枚举值：一级部署系统、'
        '二级部署系统、三级部署系统，请根据系统的实际情况填报。 1条。'
    )


def test_statistics_includes_skipped_business_failure_details(tmp_path):
    """跳过业务必须在统计表中展示阶段、文件、异常、原因和实际日志引用。

    Args:
        tmp_path: pytest 提供的隔离目录。
    """

    matrix = _file(tmp_path, '测试主体/07物资采购/07风控矩阵.xlsx', 'matrix', '07')
    lists = _file(tmp_path, '测试主体/07物资采购/07三清单.xlsx', 'three_lists', '07')
    failure = {
        'entity_key': '205H',
        'business_id': 'material_procurement',
        'business_code': '07',
        'variant_id': 'default',
        'stage': '岗位排序预处理',
        'files': [str(lists.relative_path)],
        'error_type': 'ValueError',
        'reason': '岗位排序合并范围跨越有效明细或其他区域：岗位职责清单!A32:A82',
        'log_reference': '.task/report/audit.log',
    }
    target = tmp_path / '审核统计表.xlsx'

    write_audit_statistics(
        target,
        files=[matrix, lists],
        ownership={},
        package_name='报送包',
        business_failures=[failure],
    )

    assert load_workbook(target)['审核统计表']['B2'].value == (
        '07物资采购：\n'
        '【审核状态】已跳过，未完成审核\n'
        '【失败阶段】岗位排序预处理\n'
        '【涉及文件】测试主体/07物资采购/07三清单.xlsx\n'
        '【异常类型】ValueError\n'
        '【跳过原因】岗位排序合并范围跨越有效明细或其他区域：岗位职责清单!A32:A82\n'
        '【详细日志】.task/report/audit.log'
    )
