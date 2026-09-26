from __future__ import annotations

import json
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from risk_audit.models import FileRecord, Finding
from risk_audit.readers.excel import parse_workbook
from risk_audit.readers.ooxml import load_compatible_workbook
from risk_audit.util import sha256_file
from risk_audit.writer import write_outputs


def test_writer_requires_explicit_metadata_directory(tmp_path):
    """tmp_path 为隔离目录；写入器不得默认在审核副本目录创建 _risk_audit。"""
    with pytest.raises(TypeError):
        write_outputs([], [], tmp_path / 'output')
    assert not (tmp_path / 'output/_risk_audit').exists()


def matrix(path):
    wb = Workbook(); ws = wb.active; ws.title = "设备管理"; ws["A1"] = "风控矩阵-设备管理"
    headers = ["业务条线", "控制措施编号", "控制措施", "控制系统", "控制载体", "责任主体", "信通公司是否适用"]
    for c, v in enumerate(headers, 1): ws.cell(2, c, v)
    for c, v in enumerate(["设备", "M01", "措施", "系统", "载体", "部门-岗位", "是"], 1): ws.cell(3, c, v)
    ws["C3"].comment = __import__("openpyxl").comments.Comment("原批注", "a")
    wb.save(path)


def add_row_only_sort_state(path):
    """写入 openpyxl 不接受的整行排序范围；参数 path 为测试工作簿。"""
    book = load_workbook(path)
    sheet = book['设备管理']
    sheet.auto_filter.ref = 'A2:G3'
    sheet.auto_filter.add_sort_condition('A3:A3')
    book.save(path)
    with ZipFile(path) as archive:
        parts = {info.filename: (info, archive.read(info.filename)) for info in archive.infolist()}
    info, worksheet_xml = parts['xl/worksheets/sheet1.xml']
    invalid_xml = worksheet_xml.replace(b'<sortState ref="A2:G3">', b'<sortState ref="2:3">')
    assert invalid_xml != worksheet_xml
    parts['xl/worksheets/sheet1.xml'] = (info, invalid_xml)
    with ZipFile(path, 'w', ZIP_DEFLATED) as archive:
        for part_info, content in parts.values():
            archive.writestr(part_info, content)


def test_minimal_writer_preserves_source_comment_and_is_idempotent(tmp_path, pack):
    source = tmp_path / "source.xlsx"; matrix(source); digest = sha256_file(source)
    f = FileRecord(source, Path("06/source.xlsx"), digest, "xlsx", "205H", [], False, "06", "default", "matrix")
    f.sheets = parse_workbook(f, source, pack["field_aliases"])
    finding = Finding("k", "i", "x", "c", "RX", 1, "violation", "205H", "信通", "06", "default", str(f.relative_path), "设备管理", 3, "测试意见", {}, "row")
    out = tmp_path / "out"
    warnings, _ = write_outputs([f], [finding], out, metadata_dir=tmp_path / 'metadata')
    assert warnings == [] and sha256_file(source) == digest
    wb = load_workbook(out / "06/source.xlsx")
    assert wb["设备管理"]["H2"].value == "审核意见" and wb["设备管理"]["H3"].value == "测试意见"
    assert wb["设备管理"]["C3"].comment.text == "原批注"
    warnings, _ = write_outputs([f], [finding], out, metadata_dir=tmp_path / 'metadata')
    assert warnings == []
    wb = load_workbook(out / "06/source.xlsx")
    assert wb["设备管理"]["H3"].value == "测试意见"


def test_writer_accepts_row_only_sort_state_without_modifying_source(tmp_path, pack):
    """写回应兼容整行排序范围且不修改源件；参数为目录和规则包。"""
    source = tmp_path / 'source.xlsx'
    matrix(source)
    add_row_only_sort_state(source)
    source_bytes = source.read_bytes()
    file = FileRecord(source, Path('06/source.xlsx'), sha256_file(source), 'xlsx', '205H', [], False,
                      '06', 'default', 'matrix')
    file.sheets = parse_workbook(file, source, pack['field_aliases'])
    finding = Finding('k', 'i', 'x', 'c', 'RX', 1, 'violation', '205H', '信通', '06', 'default',
                      str(file.relative_path), '设备管理', 3, '测试意见', {}, 'row')

    warnings, _ = write_outputs([file], [finding], tmp_path / 'out', metadata_dir=tmp_path / 'metadata')
    rerun_warnings, _ = write_outputs([file], [finding], tmp_path / 'out', metadata_dir=tmp_path / 'metadata')

    output = load_compatible_workbook(tmp_path / 'out/06/source.xlsx')
    assert warnings == []
    assert rerun_warnings == []
    assert output['设备管理']['H3'].value == '测试意见'
    assert source.read_bytes() == source_bytes


def test_material_opinion_keeps_rule_number_without_scope_prefix(tmp_path, pack):
    """tmp_path/pack 为测试目录和规则包；整套材料问题不能用范围标签覆盖规则归属。"""
    source = tmp_path / 'source.xlsx'; matrix(source)
    file = FileRecord(source, Path('06/source.xlsx'), sha256_file(source), 'xlsx', '205H', [], False,
                      '06', 'default', 'matrix')
    file.sheets = parse_workbook(file, source, pack['field_aliases'])
    finding = Finding('k', 'i', 'documents.completeness', 'required_matrix_and_lists', 'R01', 1,
                      'violation', '205H', '信通公司', '06', 'default', '', '', None,
                      '【第1条】缺少三清单。', {}, 'material')

    write_outputs([file], [finding], tmp_path / 'output', metadata_dir=tmp_path / 'metadata')

    value = load_workbook(tmp_path / 'output/06/source.xlsx')['设备管理']['H3'].value
    assert value == '【第1条】缺少三清单。'


def test_manual_edit_of_owned_opinion_is_replaced_and_warned(tmp_path, pack):
    """重跑替换旧副本中的人工改写并报告变更；参数为隔离目录和规则包。"""
    source = tmp_path / "source.xlsx"; matrix(source)
    f = FileRecord(source, Path("06/source.xlsx"), sha256_file(source), "xlsx", "205H", [], False, "06", "default", "matrix")
    f.sheets = parse_workbook(f, source, pack["field_aliases"])
    finding = Finding("k", "i", "x", "c", "RX", 1, "violation", "205H", "信通", "06", "default", str(f.relative_path), "设备管理", 3, "程序意见", {}, "row")
    out = tmp_path / "out"; write_outputs([f], [finding], out, metadata_dir=tmp_path / 'metadata')
    wb = load_workbook(out / "06/source.xlsx"); wb["设备管理"]["H3"] = "人工改写"; wb.save(out / "06/source.xlsx")
    warnings, _ = write_outputs([f], [finding], out, metadata_dir=tmp_path / 'metadata')
    assert any(x["type"] == "managed_cell_modified" for x in warnings)
    assert load_workbook(out / "06/source.xlsx")["设备管理"]["H3"].value == "程序意见"


def test_unparsed_workbook_is_preserved_without_claiming_audit(tmp_path):
    """未识别的工作簿须保留并提示未审核；tmp_path 为隔离输入和输出目录。"""
    source = tmp_path / 'source.xlsx'
    matrix(source)
    file = FileRecord(source, Path('01/source.xlsx'), sha256_file(source), 'xlsx', '20JQ', [], False, '01', 'default', 'matrix')
    warnings, ownership = write_outputs([file], [], tmp_path / 'output', metadata_dir=tmp_path / 'metadata')
    destination = tmp_path / 'output/01/source.xlsx'
    assert destination.is_file()
    assert sha256_file(destination) == sha256_file(source)
    assert not ownership
    assert any(item['type'] == 'unparsed_business_file' and item['file'] == '01/source.xlsx' for item in warnings)


def test_converted_xls_cannot_overwrite_case_equivalent_unparsed_file(tmp_path, pack):
    """转换后目标与未审核副本碰撞时须在写入前拒绝；参数为目录和规则包。"""
    source = tmp_path / 'book.XLSX'
    matrix(source)
    unparsed = FileRecord(source, Path('06/book.XLSX'), sha256_file(source), 'xlsx', '205H', [], False, '06', 'default', 'matrix')
    xls = next((Path(__file__).resolve().parents[2] / 'templates').rglob('*.xls'))
    parsed = FileRecord(xls, Path('06/book.xls'), sha256_file(xls), 'xls', '205H', [], False, '06', 'default', 'matrix')
    # 写回边界接收实际业务表和转换路径；碰撞检查必须早于读取转换文件及任何覆盖。
    parsed.sheets = parse_workbook(parsed, source, pack['field_aliases'])
    parsed._converted_path = str(source)
    output = tmp_path / 'output'
    with pytest.raises(ValueError, match='输出路径重复'):
        write_outputs([unparsed, parsed], [], output, metadata_dir=tmp_path / 'metadata')
    assert not output.exists()
    assert sha256_file(source) == unparsed.sha256
