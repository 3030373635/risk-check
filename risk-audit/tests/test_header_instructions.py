"""所有业务字段的填写说明识别与原始列读取回归测试。"""

import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from risk_audit.models import FileRecord
from risk_audit.readers.confirmed_v180 import parse_workbook_v180
from risk_audit.readers.header_semantics import header_spec
from risk_audit.util import norm_text


ALIASES = json.loads((Path(__file__).parents[1] / "rulepacks/releases/1.9.0/field_aliases.json").read_text())
FIELDS = [(kind, field, values[0]) for kind, fields in ALIASES.items() for field, values in fields.items()]


@pytest.mark.parametrize("kind,field,label", FIELDS)
@pytest.mark.parametrize("suffix", ["\n未预先收录的说明文字", "（任意说明文本XYZ）", "：2027版自由备注", "【外部模板自定义内容】"])
def test_instructions_supported_for_every_business_field(kind, field, label, suffix):
    """每个字段都支持任意说明内容；参数为表类型、字段、别名和说明后缀。"""
    reverse = {norm_text(value): key for key, values in ALIASES[kind].items() for value in values}
    assert header_spec(label + suffix, kind, reverse, extended=True)[0] == field


@pytest.mark.parametrize("text", ["岗位名称\n角色", "原岗位名称", "对岗位职责的准确性负责", "岗位名称角色"])
def test_unconfirmed_labels_are_not_guessed(text):
    """不截断多字段或任意正文；text 为不能确认的表头文本。"""
    reverse = {norm_text(value): key for key, values in ALIASES["position_duty"].items() for value in values}
    assert header_spec(text, "position_duty", reverse) == (None, "")


@pytest.mark.parametrize("text", ["负责人姓名", "实际经办人姓名", "姓名及联系方式", "责任人姓名/工号"])
def test_position_duty_header_containing_person_name_is_recognized(text):
    """岗位职责清单表头包含“姓名”即识别为人员姓名；text 为待识别表头。"""
    reverse = {norm_text(value): key for key, values in ALIASES["position_duty"].items() for value in values}

    assert header_spec(text, "position_duty", reverse) == ("person_names", "")


def test_unit_qualifier_survives_instruction_removal():
    """说明剥离后仍保留单位限定；无参数。"""
    assert header_spec("责任主体（其他公司）\n未定义的新说明", "matrix", {}, extended=True) == ("responsibility", "其他公司")
    assert header_spec("其他公司控制措施（任意新文字）", "matrix", {}, extended=True) == ("control_measure", "其他公司")


def test_applicability_header_accepts_comma_separated_instruction():
    """适用情况表头逗号后的填写说明不得阻断字段识别；无参数。"""
    reverse = {norm_text(value): key for key, values in ALIASES["matrix"].items() for value in values}
    header = "县公司适用情况，每一条措施选择：适用/不适用/修改。修改的可写清楚修改什么"

    assert header_spec(header, "matrix", reverse, extended=True) == ("applicability", "")


def test_applicability_keyword_anywhere_in_header_is_a_candidate():
    """“是否适用”位于表头中间时也必须识别为候选列；无参数。"""
    reverse = {norm_text(value): key for key, values in ALIASES["matrix"].items() for value in values}

    assert header_spec("电网工程公司是否适用填写结果", "matrix", reverse, extended=True) == (
        "applicability",
        "电网工程公司",
    )


def test_county_applicability_instruction_is_read_from_matrix(tmp_path):
    """县公司适用情况属于层级字段，不得因“县公司”不是单位简称而丢弃；tmp_path 为临时目录。"""
    book = Workbook()
    sheet = book.active
    sheet.title = "营销售电"
    sheet.append(["风控矩阵"])
    sheet.append([
        "控制措施编号",
        "控制措施",
        "责任主体",
        "县公司适用情况，每一条措施选择：适用/不适用/修改。修改的可写清楚修改什么",
    ])
    sheet.append(["业务-1.节点-控制措施01", "核对资料。", "城步公司-营销部-业务人员", "适用"])
    path = tmp_path / "县公司矩阵.xlsx"
    book.save(path)
    file = FileRecord(
        path,
        Path(path.name),
        "digest",
        "xlsx",
        "@名称:城步风控矩阵0916",
        ["报送名称:城步风控矩阵0916"],
        False,
        "01",
        "default",
        "matrix",
    )
    file._parser_policy = {"version": 3}
    file._semantic_lexicon = {}

    parsed = parse_workbook_v180(file, path, ALIASES)

    assert parsed[0].columns["applicability"] == 4
    assert parsed[0].records[0].value("applicability") == "适用"


def test_header_containing_applicability_is_always_a_candidate(tmp_path):
    """包含“是否适用”的表头必须读为适用性候选列；tmp_path 为临时目录。"""
    book = Workbook()
    sheet = book.active
    sheet.title = "物资（服务）采购与实施"
    sheet.append(["风控矩阵"])
    sheet.append(["控制措施编号", "控制措施", "责任主体", "电网工程公司是否适用"])
    sheet.append([
        "物资管理-仓储与配送-27.销售回款及物资出库准确性控制-控制措施02",
        "出库时需过磅或清点。",
        "物资部门-废旧物资管理人员",
        "不适用",
    ])
    path = tmp_path / "电网工程公司矩阵.xlsx"
    book.save(path)
    file = FileRecord(
        path,
        Path(path.name),
        "digest",
        "xlsx",
        "205Q",
        ["正式单位目录全称:湖南省电网工程有限公司本部"],
        False,
        "07",
        "default",
        "matrix",
    )
    file._parser_policy = {"version": 3}
    file._semantic_lexicon = {}

    parsed = parse_workbook_v180(file, path, ALIASES)

    assert parsed[0].columns["applicability"] == 4
    assert parsed[0].records[0].value("applicability") == "不适用"


def test_generic_sheet_reads_all_position_fields_with_instructions(tmp_path):
    """通过实际工作簿验证分类、映射和原始内容；tmp_path 为临时目录。"""
    book = Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    sheet.append(["部门（任意新说明A）", "岗位名称\n任意新说明B", "人员姓名：任意新说明C", "岗位职责编号【任意新说明D】", "角色\n（任意新说明E）", "岗位职责\n任意新说明F", "控制措施编号\n任意新说明G", "审核意见"])
    values = ["业务监控部", "主任", "张三", "1", "审批", "对异常事项闭环销号负责。", "M1", "人工意见"]
    sheet.append(values)
    path = tmp_path / "资料.xlsx"
    book.save(path)
    original = path.read_bytes()
    file = FileRecord(path, Path(path.name), "digest", "xlsx", "E1", [], False, "01", "default", "three_lists")
    parsed = parse_workbook_v180(file, path, ALIASES)
    assert len(parsed) == 1 and parsed[0].sheet_type == "position_duty"
    record = parsed[0].records[0]
    for column, field in enumerate(["department", "position", "person_names", "duty_id", "role", "duty", "measure_id"], 1):
        assert record.value(field) == values[column - 1]
        assert record.fields[field].coordinate == f"{chr(64 + column)}2"
    assert len(record.fields) == 7
    assert path.read_bytes() == original


def fuzzy_position_book(header_row: int = 3, *, complete: bool = True) -> Workbook:
    """构造模糊岗位表头工作簿；header_row 为表头行，complete 控制核心字段是否齐全。"""
    book = Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    headers = ["本单位部门任意文字", "实际岗位名称任意文字", "人员姓名任意文字", "岗位职责编号任意文字", "角色任意文字", "岗位职责任意文字", "控制措施编号任意文字"]
    if not complete:
        headers.pop()
    for column, value in enumerate(headers, 1):
        sheet.cell(header_row, column, value)
    values = ["业务监控部", "主任", "张三", "1", "审批", "对异常事项负责。", "M1"]
    for column, value in enumerate(values[:len(headers)], 1):
        sheet.cell(header_row + 1, column, value)
    return book


def test_first_five_rows_use_coherent_fuzzy_header_fallback(tmp_path):
    """前五行同一行命中全部核心字段后启用模糊列映射；tmp_path 为临时目录。"""
    path = tmp_path / "模糊表头.xlsx"
    fuzzy_position_book().save(path)
    file = FileRecord(path, Path(path.name), "digest", "xlsx", "E1", [], False, "01", "default", "three_lists")
    parsed = parse_workbook_v180(file, path, ALIASES)
    record = parsed[0].records[0]
    assert parsed[0].header_rows == [3]
    assert {field: value.coordinate for field, value in record.fields.items()} == {
        "department": "A4", "position": "B4", "person_names": "C4", "duty_id": "D4",
        "role": "E4", "duty": "F4", "measure_id": "G4",
    }


@pytest.mark.parametrize("header_row,complete", [(3, False), (6, True)])
def test_fuzzy_header_requires_complete_core_fields_within_first_five_rows(tmp_path, header_row, complete):
    """模糊回退受完整字段集和前五行限制；参数指定表头行与字段完整性。"""
    path = tmp_path / f"不应识别-{header_row}-{complete}.xlsx"
    fuzzy_position_book(header_row, complete=complete).save(path)
    file = FileRecord(path, Path(path.name), "digest", "xlsx", "E1", [], False, "01", "default", "three_lists")
    assert parse_workbook_v180(file, path, ALIASES) == []
