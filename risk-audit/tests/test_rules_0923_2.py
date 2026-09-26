"""0923-2 系统类型预处理与未决提示规则验收。"""

from __future__ import annotations

from copy import copy
import json
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table

from risk_audit.checks.registry import CheckContext, build_registry
from risk_audit import __version__
from risk_audit.configuration.loader import load_pack
from risk_audit.models import FieldValue, FileRecord, Record
from risk_audit.readers.confirmed_v180 import parse_workbook_v180
from risk_audit.util import sha256_file
from test_rules_0917_6 import active_pack_for, execute, make_file, make_record


ROOT = Path(__file__).resolve().parents[1]
SYSTEM_TYPE_OPINION = (
    "请补充系统类型，该列填报枚举值：一级部署系统、二级部署系统、"
    "三级部署系统，请根据系统的实际情况填报。"
)


def make_system_rule_file(tmp_path: Path, rows: list[list[str]], headers: list[str]) -> tuple[FileRecord, dict]:
    """构造并解析系统规则工作簿。

    Args:
        tmp_path: pytest 提供的隔离目录。
        rows: 待写入的系统规则明细。
        headers: 系统规则表头。
    """

    path = tmp_path / "系统控制规则清单.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "系统控制规则清单"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    workbook.save(path)

    aliases = json.loads(
        (ROOT / "rulepacks/releases/1.9.18/field_aliases.json").read_text(encoding="utf-8")
    )
    source_file = FileRecord(
        path,
        Path(path.name),
        sha256_file(path),
        "xlsx",
        "205H",
        [],
        False,
        "06",
        "default",
        "three_lists",
    )
    source_file.sheets = parse_workbook_v180(source_file, path, aliases)
    return source_file, aliases


def test_preprocessing_inserts_system_type_and_fills_known_owners(tmp_path: Path) -> None:
    """缺少系统类型列时应插入到系统名称与管理主体之间并自动填值。

    Args:
        tmp_path: pytest 提供的隔离目录。
    """

    source_file, aliases = make_system_rule_file(
        tmp_path,
        [
            ["M1", "财务系统", "总部", "规则1", "内容1"],
            ["M2", "营销系统", "国网湖南省电力有限公司", "规则2", "内容2"],
            ["M3", "物资系统", "长沙供电分公司", "规则3", "内容3"],
        ],
        ["控制措施编号", "系统名称", "系统规则管理主体", "规则名称", "规则内容"],
    )

    from risk_audit.system_type_preprocessing import preprocess_system_types

    preprocess_system_types(source_file, tmp_path / "work", aliases)

    result = load_workbook(Path(source_file._preprocessed_path))
    sheet = result.active
    assert [sheet.cell(1, column).value for column in range(1, 7)] == [
        "控制措施编号",
        "系统名称",
        "系统类型",
        "系统规则管理主体",
        "规则名称",
        "规则内容",
    ]
    assert [sheet.cell(row, 3).value for row in range(2, 5)] == [
        "一级部署系统",
        "二级部署系统",
        None,
    ]


def test_adjust_workbook_references_keeps_sparse_sheets_sparse() -> None:
    """公式引用调整只得访问已存在单元格，不得展开稀疏工作表。"""

    from risk_audit.system_type_preprocessing import _adjust_workbook_references

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "系统控制规则清单"
    worksheet["A1"] = "=D2"
    # 远端样式单元格仅用于放大声明维度，不应导致中间空白单元格被实例化。
    worksheet.cell(300, 300).font = copy(worksheet["A1"].font)
    existing_coordinates = set(worksheet._cells)

    _adjust_workbook_references(workbook, worksheet.title, [3])

    assert worksheet["A1"].value == "=E2"
    assert set(worksheet._cells) == existing_coordinates
    workbook.close()


def test_preprocessing_reuses_enum_column_and_preserves_third_level(tmp_path: Path) -> None:
    """无表头但已有枚举值的列应改名复用，三级部署系统不得覆盖。

    Args:
        tmp_path: pytest 提供的隔离目录。
    """

    source_file, aliases = make_system_rule_file(
        tmp_path,
        [["M1", "财务系统", "三级部署系统", "总部", "规则1", "内容1"]],
        ["控制措施编号", "系统名称", "", "系统规则管理主体", "规则名称", "规则内容"],
    )

    from risk_audit.system_type_preprocessing import preprocess_system_types

    preprocess_system_types(source_file, tmp_path / "work", aliases)

    result = load_workbook(Path(source_file._preprocessed_path))
    sheet = result.active
    assert sheet.max_column == 6
    assert sheet["C1"].value == "系统类型"
    assert sheet["C2"].value == "三级部署系统"


def test_preprocessing_recomputes_first_or_second_level_from_owner(tmp_path: Path) -> None:
    """旧的一、二级分类必须按管理主体重算，无法判定时清空。

    Args:
        tmp_path: pytest 提供的隔离目录。
    """

    source_file, aliases = make_system_rule_file(
        tmp_path,
        [
            ["M1", "财务系统", "一级部署系统", "长沙供电分公司", "规则1", "内容1"],
            ["M2", "营销系统", "一级部署系统", "国网湖南省电力有限公司", "规则2", "内容2"],
            ["M3", "物资系统", "二级部署系统", "总部", "规则3", "内容3"],
        ],
        ["控制措施编号", "系统名称", "系统类型", "系统规则管理主体", "规则名称", "规则内容"],
    )

    from risk_audit.system_type_preprocessing import preprocess_system_types

    preprocess_system_types(source_file, tmp_path / "work", aliases)

    result = load_workbook(Path(source_file._preprocessed_path))
    assert [result.active.cell(row, 3).value for row in range(2, 5)] == [
        None,
        "二级部署系统",
        "一级部署系统",
    ]


def test_preprocessing_expands_merged_title_when_inserting_column(tmp_path: Path) -> None:
    """新增系统类型列时，跨越插入点的合并标题必须同步扩展。

    Args:
        tmp_path: pytest 提供的隔离目录。
    """

    path = tmp_path / "合并标题三清单.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "汇总"
    sheet["A1"] = "系统控制规则清单"
    sheet.merge_cells("A1:E1")
    sheet.append(["控制措施编号", "系统名称", "系统规则管理主体", "规则名称", "规则内容"])
    sheet.append(["M1", "财务系统", "总部", "规则1", "内容1"])
    workbook.save(path)

    aliases = json.loads(
        (ROOT / "rulepacks/releases/1.9.18/field_aliases.json").read_text(encoding="utf-8")
    )
    source_file = FileRecord(
        path,
        Path(path.name),
        sha256_file(path),
        "xlsx",
        "205H",
        [],
        False,
        "06",
        "default",
        "three_lists",
    )
    source_file.sheets = parse_workbook_v180(source_file, path, aliases)

    from risk_audit.system_type_preprocessing import preprocess_system_types

    preprocess_system_types(source_file, tmp_path / "work", aliases)

    result = load_workbook(Path(source_file._preprocessed_path))
    assert [str(cell_range) for cell_range in result.active.merged_cells.ranges] == ["A1:F1"]


def test_preprocessing_reuses_vertically_merged_system_type_header(tmp_path: Path) -> None:
    """精确系统类型表头纵向合并时应复用合并锚点。

    Args:
        tmp_path: pytest 提供的隔离目录。
    """

    path = tmp_path / "纵向合并表头.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "系统控制规则清单"
    headers = ["控制措施编号", "系统名称", "系统类型", "系统规则管理主体"]
    for column, header in enumerate(headers, 1):
        sheet.cell(2, column).value = header
        sheet.merge_cells(start_row=2, start_column=column, end_row=3, end_column=column)
    sheet["E2"] = "规则"
    sheet["F2"] = "规则"
    sheet["E3"] = "规则名称"
    sheet["F3"] = "规则内容"
    sheet.append(["M1", "财务系统", "", "总部", "规则1", "内容1"])
    workbook.save(path)

    aliases = json.loads(
        (ROOT / "rulepacks/releases/1.9.18/field_aliases.json").read_text(encoding="utf-8")
    )
    source_file = FileRecord(
        path, Path(path.name), sha256_file(path), "xlsx", "205H", [], False, "06", "default", "three_lists"
    )
    source_file.sheets = parse_workbook_v180(source_file, path, aliases)

    from risk_audit.system_type_preprocessing import preprocess_system_types

    preprocess_system_types(source_file, tmp_path / "work", aliases)

    result = load_workbook(Path(source_file._preprocessed_path))
    assert "C2:C3" in {str(cell_range) for cell_range in result.active.merged_cells.ranges}
    assert result.active["C2"].value == "系统类型"
    assert result.active["C4"].value == "一级部署系统"


def test_preprocessing_renames_enum_column_with_vertically_merged_blank_header(tmp_path: Path) -> None:
    """枚举列的空表头纵向合并时，应在合并锚点写入系统类型。

    Args:
        tmp_path: pytest 提供的隔离目录。
    """

    path = tmp_path / "纵向合并空表头.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "系统控制规则清单"
    headers = ["控制措施编号", "系统名称", None, "系统规则管理主体"]
    for column, header in enumerate(headers, 1):
        sheet.cell(2, column).value = header
        sheet.merge_cells(start_row=2, start_column=column, end_row=3, end_column=column)
    sheet["E2"] = "规则"
    sheet["F2"] = "规则"
    sheet["E3"] = "规则名称"
    sheet["F3"] = "规则内容"
    sheet.append(["M1", "财务系统", "三级部署系统", "长沙供电分公司", "规则1", "内容1"])
    workbook.save(path)

    aliases = json.loads(
        (ROOT / "rulepacks/releases/1.9.18/field_aliases.json").read_text(encoding="utf-8")
    )
    source_file = FileRecord(
        path, Path(path.name), sha256_file(path), "xlsx", "205H", [], False, "06", "default", "three_lists"
    )
    source_file.sheets = parse_workbook_v180(source_file, path, aliases)

    from risk_audit.system_type_preprocessing import preprocess_system_types

    preprocess_system_types(source_file, tmp_path / "work", aliases)

    result = load_workbook(Path(source_file._preprocessed_path))
    assert "C2:C3" in {str(cell_range) for cell_range in result.active.merged_cells.ranges}
    assert result.active["C2"].value == "系统类型"
    assert result.active["C4"].value == "三级部署系统"


def test_preprocessing_splits_horizontal_blank_header_for_enum_column(tmp_path: Path) -> None:
    """枚举列位于横向合并空表头时，应拆分合并并准确命名枚举列。

    Args:
        tmp_path: pytest 提供的隔离目录。
    """

    path = tmp_path / "横向合并空表头.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "系统控制规则清单"
    for column, header in enumerate(["控制措施编号", "系统名称"], 1):
        sheet.cell(2, column).value = header
        sheet.merge_cells(start_row=2, start_column=column, end_row=3, end_column=column)
    sheet.merge_cells("C2:D3")
    sheet["E2"] = "系统规则管理主体"
    sheet.merge_cells("E2:E3")
    sheet["F2"] = "规则"
    sheet["G2"] = "规则"
    sheet["F3"] = "规则名称"
    sheet["G3"] = "规则内容"
    sheet.append(["M1", "财务系统", "", "三级部署系统", "长沙供电分公司", "规则1", "内容1"])
    workbook.save(path)

    aliases = json.loads(
        (ROOT / "rulepacks/releases/1.9.18/field_aliases.json").read_text(encoding="utf-8")
    )
    source_file = FileRecord(
        path, Path(path.name), sha256_file(path), "xlsx", "205H", [], False, "06", "default", "three_lists"
    )
    source_file.sheets = parse_workbook_v180(source_file, path, aliases)

    from risk_audit.system_type_preprocessing import preprocess_system_types

    preprocess_system_types(source_file, tmp_path / "work", aliases)

    result = load_workbook(Path(source_file._preprocessed_path))
    assert "C2:D3" not in {str(cell_range) for cell_range in result.active.merged_cells.ranges}
    assert result.active["D2"].value == "系统类型"
    assert result.active["D4"].value == "三级部署系统"


def test_preprocessing_updates_formula_reference_after_column_insertion(tmp_path: Path) -> None:
    """新增系统类型列后，跨越插入点的公式引用必须同步右移。

    Args:
        tmp_path: pytest 提供的隔离目录。
    """

    source_file, aliases = make_system_rule_file(
        tmp_path,
        [["M1", "财务系统", "总部", "规则1", "=C2"]],
        ["控制措施编号", "系统名称", "系统规则管理主体", "规则名称", "规则内容"],
    )

    from risk_audit.system_type_preprocessing import preprocess_system_types

    preprocess_system_types(source_file, tmp_path / "work", aliases)

    result = load_workbook(Path(source_file._preprocessed_path), data_only=False)
    assert result.active["F2"].value == "=D2"


def test_preprocessing_updates_range_structures_after_column_insertion(tmp_path: Path) -> None:
    """插列后表格、筛选、数据验证和条件格式范围必须同步平移。

    Args:
        tmp_path: pytest 提供的隔离目录。
    """

    path = tmp_path / "带范围结构的系统规则.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "系统控制规则清单"
    sheet.append(["控制措施编号", "系统名称", "系统规则管理主体", "规则名称", "规则内容"])
    sheet.append(["M1", "财务系统", "总部", "规则1", "内容1"])
    sheet.auto_filter.ref = "A1:E2"
    validation = DataValidation(type="list", formula1="=$C$2:$C$10")
    validation.add("C2:C10")
    sheet.add_data_validation(validation)
    sheet.conditional_formatting.add("C2:C10", FormulaRule(formula=['C2="总部"']))
    sheet.add_table(Table(displayName="SystemRules", ref="A1:E2"))
    workbook.defined_names.add(
        DefinedName("OwnerValues", attr_text="'系统控制规则清单'!$C$2:$C$10")
    )
    workbook.save(path)

    aliases = json.loads(
        (ROOT / "rulepacks/releases/1.9.18/field_aliases.json").read_text(encoding="utf-8")
    )
    source_file = FileRecord(
        path, Path(path.name), sha256_file(path), "xlsx", "205H", [], False, "06", "default", "three_lists"
    )
    source_file.sheets = parse_workbook_v180(source_file, path, aliases)

    from risk_audit.system_type_preprocessing import preprocess_system_types

    preprocess_system_types(source_file, tmp_path / "work", aliases)

    result = load_workbook(Path(source_file._preprocessed_path), data_only=False)
    result_sheet = result.active
    assert result_sheet.auto_filter.ref == "A1:F2"
    assert str(result_sheet.data_validations.dataValidation[0].sqref) == "D2:D10"
    assert result_sheet.data_validations.dataValidation[0].formula1 == "=$D$2:$D$10"
    conditional_format = next(iter(result_sheet.conditional_formatting._cf_rules))
    assert str(conditional_format.sqref) == "D2:D10"
    assert result_sheet.conditional_formatting._cf_rules[conditional_format][0].formula == ['D2="总部"']
    assert result_sheet.tables["SystemRules"].ref == "A1:F2"
    assert [column.name for column in result_sheet.tables["SystemRules"].tableColumns] == [
        "控制措施编号",
        "系统名称",
        "系统类型",
        "系统规则管理主体",
        "规则名称",
        "规则内容",
    ]
    assert result.defined_names["OwnerValues"].attr_text == "'系统控制规则清单'!$D$2:$D$10"


def test_table_at_insertion_boundary_moves_without_extra_column_definition(tmp_path: Path) -> None:
    """表格左边界恰为插入点时，表格整体右移且列定义数保持一致。

    Args:
        tmp_path: pytest 提供的隔离目录。
    """

    path = tmp_path / "表格位于插入边界.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "系统控制规则清单"
    sheet.append(["控制措施编号", "系统名称", "系统规则管理主体", "规则名称", "规则内容"])
    sheet.append(["M1", "财务系统", "总部", "规则1", "内容1"])
    sheet.add_table(Table(displayName="OwnerRules", ref="C1:E2"))
    workbook.save(path)

    aliases = json.loads(
        (ROOT / "rulepacks/releases/1.9.18/field_aliases.json").read_text(encoding="utf-8")
    )
    source_file = FileRecord(
        path, Path(path.name), sha256_file(path), "xlsx", "205H", [], False, "06", "default", "three_lists"
    )
    source_file.sheets = parse_workbook_v180(source_file, path, aliases)

    from risk_audit.system_type_preprocessing import preprocess_system_types

    preprocess_system_types(source_file, tmp_path / "work", aliases)

    result = load_workbook(Path(source_file._preprocessed_path))
    table = result.active.tables["OwnerRules"]
    assert table.ref == "D1:F2"
    assert [column.name for column in table.tableColumns] == ["系统规则管理主体", "规则名称", "规则内容"]


def test_merged_range_adjustment_uses_original_column_coordinates() -> None:
    """并排逻辑区多次插列时，局部合并范围不得重复累加左边界偏移。"""

    from risk_audit.system_type_preprocessing import _adjusted_range

    assert _adjusted_range("G1:I1", [3, 8]) == "H1:K1"


def test_owner_classification_handles_suffixes_and_negative_headquarters() -> None:
    """省级主体后缀应命中二级，否定的总部语义不得误判为一级。"""

    from risk_audit.system_type_preprocessing import classify_system_owner

    assert classify_system_owner("湖南省电力公司本部") == "二级部署系统"
    assert classify_system_owner("非总部湖南省公司") == "二级部署系统"
    assert classify_system_owner("长沙供电公司总部") is None
    assert classify_system_owner("国网总部财务部") == "一级部署系统"
    assert classify_system_owner("国家电网有限公司总部业务部门") == "一级部署系统"


def field(value: str, coordinate: str) -> FieldValue:
    """构造测试字段。

    Args:
        value: 字段当前值。
        coordinate: 字段单元格坐标。
    """

    return FieldValue(value, value, coordinate)


def system_rule_record(row: int, owner: str, system_type: str = "") -> Record:
    """构造系统规则记录。

    Args:
        row: 明细行号。
        owner: 系统规则管理主体。
        system_type: 当前系统类型。
    """

    return Record(
        "system_rule",
        "205H",
        "06",
        "default",
        "系统控制规则清单.xlsx",
        "系统控制规则清单",
        row,
        {
            "system_owner": field(owner, f"D{row}"),
            "system_type": field(system_type, f"C{row}"),
        },
        f"system-rule-{row}",
    )


def test_unknown_owner_reports_once_but_preserved_third_level_does_not() -> None:
    """无法判断管理主体时逐行提示，已填三级部署系统的记录优先通过。"""

    from risk_audit.checks.confirmed_v180 import system_type_completion_v1

    rows = [
        system_rule_record(2, "长沙供电分公司"),
        system_rule_record(3, "长沙供电分公司", "三级部署系统"),
        system_rule_record(4, "总部", "一级部署系统"),
    ]
    context = CheckContext(rows, [], rows, {}, {}, {}, ["205H"], [])

    issues = system_type_completion_v1(context, {})

    assert len(issues) == 1
    assert issues[0]["record"].row == 2
    assert issues[0]["evidence"]["issue_type"] == "system_type_unresolved"
    assert issues[0]["evidence"]["opinion"] == SYSTEM_TYPE_OPINION


def test_registry_exposes_system_type_completion_operator() -> None:
    """规则注册表必须提供系统类型补全后的未决检查能力。"""

    capability = build_registry().get("system_type_completion", 1)

    assert capability.record_types == frozenset({"system_rule"})
    assert "opinion" in capability.evidence_variables


def test_runner_preprocessing_applies_0923_system_type_rules(tmp_path: Path) -> None:
    """1.9.18 业务预处理入口必须实际执行系统类型规范化。

    Args:
        tmp_path: pytest 提供的隔离目录。
    """

    source_file, aliases = make_system_rule_file(
        tmp_path,
        [["M1", "财务系统", "总部", "规则1", "内容1"]],
        ["控制措施编号", "系统名称", "系统规则管理主体", "规则名称", "规则内容"],
    )

    from risk_audit.runner import preprocess_business_file

    preprocess_business_file(
        source_file,
        {"manifest": {"version": "1.9.18"}, "field_aliases": aliases},
        {},
        tmp_path / "run",
    )

    result = load_workbook(Path(source_file._preprocessed_path))
    assert result.active["C1"].value == "系统类型"
    assert result.active["C2"].value == "一级部署系统"


def test_active_rule_reports_unresolved_system_type_with_unified_opinion() -> None:
    """无法识别管理主体时，激活规则必须输出 0923-2 统一提示语。"""

    row = make_record(
        "system_rule",
        system_owner="长沙供电分公司",
        system_type="",
    )

    findings = execute(
        active_pack_for("systems.type_completion"),
        [make_file("system_rule", [row])],
    )

    assert [finding.message for finding in findings] == [SYSTEM_TYPE_OPINION]


def test_0923_2_release_is_preserved_and_rebuildable() -> None:
    """0923-2 来源和 1.9.18 规则包必须保留，当前程序使用后续发布版本。"""

    assert __version__ == "1.9.19"
    active = json.loads((ROOT / "rulepacks/active.json").read_text(encoding="utf-8"))
    assert active["version"] == "1.9.19"
    pack = load_pack(ROOT / "rulepacks/releases/1.9.18")
    assert pack["manifest"]["source_document_sha256"] == sha256_file(
        ROOT / "规则来源/0923-2.docx"
    )

    result = subprocess.run(
        [sys.executable, str(ROOT / "tools/publish_rules_0923_2.py"), "--verify-only"],
        cwd=ROOT.parent,
        capture_output=True,
        text=True,
        check=False,
    )

    assert result.returncode == 0, result.stderr
    assert json.loads(result.stdout) == {
        "release": str(ROOT / "rulepacks/releases/1.9.18"),
        "activated": False,
        "verified": True,
    }
