"""验证重庆六条规则、标准库比对及独立报告。"""

from pathlib import Path
import json
import shutil

from openpyxl import Workbook, load_workbook
import pytest

from risk_audit.checks.registry import build_registry
from risk_audit.configuration.loader import load_pack
from risk_audit.engine import run_engine
from risk_audit.models import FileRecord, ParsedSheet
from test_rules_1015_chongqing import make_context, make_record


ROOT = Path(__file__).resolve().parents[1]
RELEASE = ROOT / "rulepacks/releases/1.9.36"


def standard_library() -> dict:
    """返回手工确定的两个业务、三条措施和两种单位性质样本。"""
    return {
        "companies": [
            {"company_name": "国网重庆市电力公司北碚供电分公司", "unit_nature": "供电分公司"},
            {"company_name": "重庆市送变电工程有限公司", "unit_nature": "其他"},
        ],
        "measures": [
            {"matrix_name": "财务管理", "matrix_type": "通用业务", "standard_code": "TY-CW-01-01",
             "measure_id": "财务管理-1.资金管理-控制措施01", "business_id": "financial_management",
             "business_code": "02", "applicable_cells": {"供电分公司": "Sheet1!G6", "其他": "Sheet1!R6"}},
            {"matrix_name": "财务管理", "matrix_type": "通用业务", "standard_code": "TY-CW-01-02",
             "measure_id": "财务管理-1.资金管理-控制措施02", "business_id": "financial_management",
             "business_code": "02", "applicable_cells": {"供电分公司": "Sheet1!G7"}},
            {"matrix_name": "合同管理", "matrix_type": "通用业务", "standard_code": "TY-HT-01-01",
             "measure_id": "合同管理-1.合同-控制措施01", "business_id": "contract_management",
             "business_code": "03", "applicable_cells": {"供电分公司": "Sheet1!G8"}},
        ],
    }


def capability_issues(name: str, records: list, resources: dict | None = None) -> list[dict]:
    """运行真实注册能力；name 为能力名称，records 为记录，resources 为冻结附件。"""
    capability = build_registry().get(name, 1)
    assert capability is not None, f"尚未实现新版能力：{name}"
    context = make_context(records)
    context.resources.update(resources or {})
    return capability.runner(context, {})


def test_incompatible_positions_allow_cross_department_and_plain_job_names() -> None:
    """新版不因跨部门或未使用部门前缀报错，但仍检查同名岗位。"""
    record = make_record("incompatible_position", 2, measure_id="措施1-01", position_a="财务部-出纳", position_b="审计部-审计岗")
    assert capability_issues("chongqing_incompatible_alignment_1024", [record]) == []
    record.fields["position_a"].current = "出纳"
    record.fields["position_b"].current = "出纳"
    issues = capability_issues("chongqing_incompatible_alignment_1024", [record])
    assert [issue["evidence"]["opinion"] for issue in issues] == ["岗位名称相同，请确认是否为不相容岗位"]


def test_incompatible_jobs_must_exist_in_same_measure_and_business() -> None:
    """另一个措施或业务中的同名岗位不得满足当前岗位清单关联。"""
    incompatible = make_record("incompatible_position", 2, measure_id="措施01-02", position_a="财务部-出纳", position_b="会计岗")
    matching = make_record("position_duty", 3, measure_id="另一标题1-2", department="财务部", position="出纳")
    other_measure = make_record("position_duty", 4, measure_id="措施1-3", position="会计岗")
    other_business = make_record("position_duty", 5, measure_id="措施1-2", position="会计岗")
    other_business.business_id = "contract_management"
    issues = capability_issues("chongqing_incompatible_duty_alignment", [incompatible, matching, other_measure, other_business])
    assert len(issues) == 1
    assert issues[0]["record"] is incompatible
    assert issues[0]["evidence"]["missing_positions"] == ["会计岗"]
    assert issues[0]["evidence"]["opinion"] == "请核实不相容岗位清单中填列的岗位名称在岗位内控责任清单中是否填列"


def test_incompatible_empty_list_is_tolerated() -> None:
    """明确填写无、不涉及的记录不触发岗位关联。"""
    record = make_record("incompatible_position", 2, position_a="不涉及", position_b="无")
    assert capability_issues("chongqing_incompatible_duty_alignment", [record]) == []


def test_matrix_alone_still_requires_three_lists(tmp_path: Path) -> None:
    """tmp_path 为材料目录；有矩阵不得掩盖缺少三清单。"""
    context = make_context([])
    context.files = [FileRecord(tmp_path / "矩阵.xlsx", Path("矩阵.xlsx"), "", "xlsx", "CQ001", [], False, "02", "default", "matrix", business_id="financial_management")]
    capability = build_registry().get("chongqing_required_documents_1024", 1)
    assert capability is not None
    issues = capability.runner(context, {"material_types": ["matrix", "three_lists"]})
    assert [issue["evidence"].get("missing_material") for issue in issues] == ["three_lists"]


def test_standard_library_does_not_mix_businesses_or_subjects() -> None:
    """措施数字相同也必须分别按业务和主体比对。"""
    from risk_audit.standard_library import compare_standard_library

    record = make_record("matrix", 2, measure_id="任意标题01-001")
    another_subject = make_record("matrix", 3, measure_id="措施1-2")
    another_subject.entity_code = "OTHER"
    result = compare_standard_library(standard_library(), "北碚供电分公司", [record, another_subject], entity_code="CQ001")
    assert [row["opinion"] for row in result["rows"]] == [
        "无异常", "需核实标准库中适用的控制措施是否完整", "需核实标准库中适用的控制措施是否完整",
    ]
    assert [row["applicable_cell"] for row in result["rows"]] == ["Sheet1!G6", "Sheet1!G7", "Sheet1!G8"]


@pytest.mark.parametrize("company_name, status, unit_nature", [
    ("7.送变电", "matched", "其他"),
    ("国网重庆市电力公司北碚供电分公司", "matched", "供电分公司"),
    ("未知公司", "unmatched", ""),
])
def test_company_match_uses_unique_names(company_name: str, status: str, unit_nature: str) -> None:
    """company_name 为报送目录名，status/unit_nature 为手工确定的匹配结果。"""
    from risk_audit.standard_library import compare_standard_library

    result = compare_standard_library(standard_library(), company_name, [], entity_code="CQ001")
    assert result["match_status"] == status
    assert result["unit_nature"] == unit_nature
    if status != "matched":
        assert result["rows"] == []
        assert "核实公司名称与单位性质" in result["opinion"]


def test_company_short_name_ambiguity_does_not_guess_nature() -> None:
    """两个公司产生同一简称时应提示确认，不能任选一个单位性质。"""
    from risk_audit.standard_library import compare_standard_library

    library = standard_library()
    library["companies"].append({"company_name": "重庆市北碚供电分公司", "unit_nature": "其他"})
    result = compare_standard_library(library, "北碚供电分公司", [], entity_code="CQ001")
    assert result["match_status"] == "ambiguous"
    assert result["unit_nature"] == ""


def test_failed_business_is_neither_missing_nor_normal() -> None:
    """读取失败的业务没有足够证据认定缺失或正常。"""
    from risk_audit.standard_library import compare_standard_library

    result = compare_standard_library(standard_library(), "北碚供电分公司", [], entity_code="CQ001", blocked_business_ids={"financial_management"})
    assert [row["status"] for row in result["rows"]] == ["unavailable", "unavailable", "missing"]
    assert "无法完成比对" in result["rows"][0]["opinion"]


def test_report_contains_normal_and_missing_results_without_formulas(tmp_path: Path) -> None:
    """tmp_path 为输出目录；独立表必须保留全部适用措施及七个指定字段。"""
    from risk_audit.standard_library import compare_standard_library, write_standard_library_report

    record = make_record("matrix", 2, measure_id="措施1-1")
    result = compare_standard_library(standard_library(), "北碚供电分公司", [record], entity_code="CQ001")
    path = write_standard_library_report(tmp_path, result)
    workbook = load_workbook(path)
    sheet = workbook.active
    assert path.name == "标准库比对结果-北碚供电分公司.xlsx"
    assert list(next(sheet.values)) == ["矩阵名称", "矩阵类型", "控制措施编码", "重庆控制措施编号", "对应的单位性质", "标记为“是”的单元格", "提示内容"]
    assert [sheet.cell(row, 7).value for row in range(2, 5)] == ["无异常", "需核实标准库中适用的控制措施是否完整", "需核实标准库中适用的控制措施是否完整"]
    # 重跑必须覆盖同一报告，不能累加旧结果。
    write_standard_library_report(tmp_path, result)
    assert len(list(tmp_path.glob("标准库比对结果-*.xlsx"))) == 1
    workbook.close()


def test_release_emits_only_six_rules_and_no_normal_findings() -> None:
    """新版运行不可继续产生已删除人员检查，正常比对项不进入问题统计。"""
    pack = load_pack(RELEASE)
    assert {rule["display_code"] for rule in pack["rules"] if rule["display_code"].startswith("R")} == {"R01", "R02", "R03", "R04", "R05", "R06"}
    pack["standard_library"] = standard_library()
    pack["_entity_names"] = {"CQ001": "送变电"}
    pack["rules"] = [rule for rule in pack["rules"] if rule["display_code"] == "R06"]
    pack["submission_scope"] = {"entity_codes": ["CQ001"], "businesses": [{"business_id": "financial_management", "business_code": "02", "variant_id": "default", "required": True}]}
    record = make_record("matrix", 2, measure_id="措施1-1", responsibility="各级单位", person_names="多人")
    file = FileRecord(Path("矩阵.xlsx"), Path("矩阵.xlsx"), "", "xlsx", "CQ001", [], False, "02", "default", "matrix", sheets=[ParsedSheet("矩阵", "matrix", [1], {}, {}, 2, 2, [record])], business_id="financial_management")
    findings, statuses = run_engine(pack, build_registry(), [file], {}, {})
    assert findings == []
    assert all(status.status == "executed" for status in statuses)


@pytest.mark.parametrize("write", [True, False])
def test_runner_generates_one_company_report_across_businesses(tmp_path: Path, write: bool) -> None:
    """tmp_path 为隔离目录，write 控制试跑；第6条只在主体汇总后运行一次。"""
    from risk_audit.runner import audit

    pack_path = tmp_path / "pack"
    shutil.copytree(RELEASE, pack_path)
    manifest = json.loads((pack_path / "manifest.json").read_text())
    manifest["status"] = "draft"
    (pack_path / "manifest.json").write_text(json.dumps(manifest))
    (pack_path / "standard_library.json").write_text(json.dumps(standard_library(), ensure_ascii=False))
    inputs = tmp_path / "input"
    for business in ("02 财务管理", "03 合同管理"):
        path = inputs / "北碚供电分公司" / business / f"风控矩阵-{business}.xlsx"
        path.parent.mkdir(parents=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "风控矩阵"
        sheet.append(["控制措施编号", "控制措施", "责任主体", "不相容岗位", "控制方式", "审核意见"])
        sheet.append(["措施1-1", "管理资料", "各级单位-部门-岗位", "", "人工控制", "人工意见\n【第8条】旧人员检查意见"])
        workbook.save(path)
        workbook.close()
    result = audit(inputs, tmp_path / "output", pack_path, None, tmp_path / "optional-templates", tmp_path / "runs", run_id="1024-integration", entity_mode="parent_directory", write=write)
    statuses = json.loads((tmp_path / "runs/1024-integration/statuses.json").read_text())
    assert sum(status["check_id"] == "standard_library_completeness" for status in statuses) == 1
    findings = json.loads((tmp_path / "runs/1024-integration/findings.json").read_text())
    standard_findings = [finding for finding in findings if finding["display_code"] == "R06"]
    assert len(standard_findings) == 1
    assert "1条" in standard_findings[0]["message"]
    assert "无异常" not in "".join(finding["message"] for finding in findings)
    assert result["standard_library"]["missing"] == 1
    if write:
        reports = list((tmp_path / "output").glob("标准库比对结果-*.xlsx"))
        assert len(reports) == 1
        workbook = load_workbook(reports[0])
        assert [workbook.active.cell(row, 7).value for row in range(2, 5)] == ["无异常", "需核实标准库中适用的控制措施是否完整", "无异常"]
        workbook.close()
        matrix = load_workbook(tmp_path / "output/北碚供电分公司/02 财务管理/风控矩阵-02 财务管理.xlsx")
        opinion = matrix.active.cell(2, 6).value
        assert "人工意见" in opinion
        assert "第8条" not in opinion
        assert "【第6条】" in opinion
        matrix.close()
    else:
        assert not (tmp_path / "output").exists()


def test_rulepack_rejects_missing_standard_library_and_invalid_number() -> None:
    """新版规则不能在附件缺失或标准编号无法解析时静默运行。"""
    from risk_audit.configuration.validator import ConfigError, validate_pack

    pack = load_pack(RELEASE)
    pack["manifest"]["status"] = "draft"
    pack.pop("standard_library")
    with pytest.raises(ConfigError, match="standard_library"):
        validate_pack(pack, build_registry())
    pack["standard_library"] = standard_library()
    pack["standard_library"]["measures"][0]["measure_id"] = "没有数字的编号"
    with pytest.raises(ConfigError, match="measure_id"):
        validate_pack(pack, build_registry())


def test_same_confirmed_company_merges_submissions_across_directories(tmp_path: Path) -> None:
    """tmp_path 为输入根；附件全称和唯一简称指向同一公司时合并跨目录材料。"""
    from risk_audit.standard_library import bind_company_groups

    files = []
    groups = {}
    for index, name in enumerate(("北碚供电分公司", "国网重庆市电力公司北碚供电分公司")):
        key = f"@目录主体:批次{index}/{name}"
        identity = {"entity_name": name, "entity_registry_status": "directory_matched", "group_key": key}
        groups[key] = identity
        files.append(FileRecord(tmp_path / f"{index}.xlsx", Path(f"{index}.xlsx"), "", "xlsx", key, [], False, "02", "default", "matrix", preservation={"entity_identity": identity}))
    bound = bind_company_groups(files, groups, standard_library())
    assert len(bound) == 1
    assert files[0].entity_code == files[1].entity_code


def test_standard_missing_opinions_keep_each_business_in_engine() -> None:
    """两项业务均缺适用措施时，引擎不得把同类缺项意见去重成一条。"""
    pack = load_pack(RELEASE)
    pack["standard_library"] = standard_library()
    pack["_entity_names"] = {"CQ001": "北碚供电分公司"}
    pack["rules"] = [rule for rule in pack["rules"] if rule["display_code"] == "R06"]
    pack["submission_scope"] = {"entity_codes": ["CQ001"], "businesses": [{"business_id": "financial_management", "business_code": "02", "variant_id": "default", "required": True}]}
    findings, _ = run_engine(pack, build_registry(), [], {}, {})
    assert {finding.business_id for finding in findings} == {"financial_management", "contract_management"}


def test_failed_capability_does_not_become_missing_standard_measures(tmp_path: Path, monkeypatch) -> None:
    """tmp_path/monkeypatch 为测试依赖；引擎返回失败状态时标准库应显示无法比对。"""
    from risk_audit import runner
    from risk_audit.models import CheckStatus

    pack_path = tmp_path / "pack"
    shutil.copytree(RELEASE, pack_path)
    manifest = json.loads((pack_path / "manifest.json").read_text())
    manifest["status"] = "draft"
    (pack_path / "manifest.json").write_text(json.dumps(manifest))
    library = standard_library()
    library["measures"] = library["measures"][:1]
    (pack_path / "standard_library.json").write_text(json.dumps(library))
    source = tmp_path / "input/北碚供电分公司/02 财务管理/风控矩阵.xlsx"
    source.parent.mkdir(parents=True)
    workbook = Workbook()
    workbook.active.title = "风控矩阵"
    workbook.active.append(["控制措施编号", "控制措施", "责任主体", "不相容岗位", "控制方式"])
    workbook.active.append(["措施1-1", "办理资料", "公司-部门-岗位", "", "人工控制"])
    workbook.save(source)
    workbook.close()
    real_engine = runner.run_engine

    def failing_business_engine(pack, *args, **kwargs):
        """模拟能力失败；pack 为真实规则包，args/kwargs 传给原引擎，其余行为不替换。"""
        findings, statuses = real_engine(pack, *args, **kwargs)
        if not any(rule["display_code"] == "R06" for rule in pack["rules"]):
            statuses.append(CheckStatus("test", "failed_capability", "failed", "能力执行失败"))
        return findings, statuses

    monkeypatch.setattr(runner, "run_engine", failing_business_engine)
    result = runner.audit(tmp_path / "input", tmp_path / "output", pack_path, None, tmp_path / "templates", tmp_path / "runs", entity_mode="parent_directory")
    assert result["parsed_files"] == 1
    assert result["standard_library"]["missing"] == 0
    assert result["standard_library"]["unavailable"] == 1


def test_cli_trial_uses_chongqing_directory_identity_without_legacy_registry(tmp_path: Path, monkeypatch, capsys) -> None:
    """tmp_path/monkeypatch/capsys 为测试依赖；重庆命令行试跑不应要求湖南名册。"""
    from risk_audit.cli import main

    source = tmp_path / "input/北碚供电分公司/02 财务管理/风控矩阵.xlsx"
    source.parent.mkdir(parents=True)
    workbook = Workbook()
    workbook.active.title = "风控矩阵"
    workbook.active.append(["控制措施编号", "控制措施", "责任主体", "不相容岗位", "控制方式"])
    workbook.active.append(["措施1-1", "办理资料", "公司-部门-岗位", "", "人工控制"])
    workbook.save(source)
    workbook.close()
    monkeypatch.chdir(tmp_path)
    code = main(["--rulepacks-root", str(ROOT / "rulepacks"), "trial", "--pack", str(RELEASE), "--input", str(tmp_path / "input"), "--runs-root", str(tmp_path / "runs"), "--project-root", str(tmp_path / "templates")])
    assert code == 0
    assert json.loads(capsys.readouterr().out)["standard_library"]["companies"][0]["unit_nature"] == "供电分公司"
    assert not (tmp_path / "runs/_trial_no_output").exists()
