from pathlib import Path
from dataclasses import replace
import json
import pytest
from openpyxl import Workbook
from risk_audit.applicability import interpret,interpret_record
from risk_audit.configuration.loader import load_pack
from risk_audit.models import FileRecord,Record,FieldValue
from risk_audit.readers.excel import parse_workbook
from risk_audit.checks.registry import CheckContext
from risk_audit.checks.applicability import applicability_values_v2
from risk_audit.checks.quality_checks import responsibility_coverage_v5
from risk_audit.util import sha256_file

ROOT=Path(__file__).resolve().parents[1]

@pytest.fixture
def semantic_pack():return load_pack(ROOT/'rulepacks/drafts/semantic-v1.4.0')


@pytest.mark.parametrize('text,state,decision',[
 ('是','explicit_positive','applicable'),('否','explicit_negative','not_applicable'),
 ('适用，本单位暂无此业务','explicit_positive','applicable'),
 ('是，本单位尚未开展该项业务','explicit_positive','applicable'),
 ('不适用，无此业务','explicit_negative','not_applicable'),
 ('暂无此业务，参照省公司','business_absent',None),
 ('县公司暂无此业务，参照省公司执行','business_absent',None),
 ('本单位尚未开展该项业务','business_absent',None),
 ('直接引用','template_adopted',None),('引用，修改控制措施','template_modified',None),
 ('修改控制措施\n、责任主体、控制载体','template_modified',None),
 ('县公司无业务核算岗位，参照省公司','scope_explanation',None),
 ('无对应系统，空着（8-1也是空着的）','scope_explanation',None),
 ('本条适用于本单位','explicit_positive','applicable'),
 ('不适用于本公司，业务由总部处理','explicit_negative','not_applicable'),
 ('是否适用','unknown',None),('不一定适用','unknown',None),('并非不适用','unknown',None),
 ('若开展业务则适用','unknown',None),('是，但不适用','conflict',None),
 ('暂无此业务，但已经开始开展','unknown',None),('适用性待确认','unknown',None),
 ('','empty',None),
])
def test_meaning_is_separate_from_applicability(semantic_pack,text,state,decision):
 result=interpret(text,semantic_pack['semantic_lexicon'])
 assert (result.state,result.decision)==(state,decision)


def matrix(tmp_path,pack,headers,rows,unit='湘西自治州德能电力建设有限公司凤凰分公司',code='20CI',body=None,hidden=(),footer=None):
 w=Workbook();s=w.active;s.title='风控矩阵';s.append(['风控矩阵'])
 s.append(['控制措施编号','控制措施','责任主体','控制载体',*headers])
 for i,tail in enumerate(rows,3):s.append([f'M{i}','措施',body if body is not None else unit+'-综合管理室-主任','审批单',*tail])
 if footer:s.cell(50,footer[0],footer[1])
 for col in hidden:s.column_dimensions[col].hidden=True
 path=tmp_path/'matrix.xlsx';w.save(path)
 f=FileRecord(path,Path('matrix.xlsx'),sha256_file(path),'xlsx',code,['标准全称:'+unit],False,'06','default','matrix')
 f._parser_policy=pack['parser_policy'];f._semantic_lexicon=pack['semantic_lexicon']
 f.sheets=parse_workbook(f,path,pack['field_aliases'])
 return f,f.sheets[0]


def test_change_notes_cannot_mask_real_answers(tmp_path,semantic_pack):
 f,s=matrix(tmp_path,semantic_pack,['是否适用','凤凰分公司适用情况'],[['是','修改责任主体'],['否','修改控制措施']])
 assert s.columns['applicability']==5
 assert [r.value('applicability') for r in s.records]==['是','否']
 assert any(x['reason']=='answer_values_preferred_to_change_notes' for x in f.preservation['column_selections'])


def test_identical_business_rows_ignore_nonbusiness_footer(tmp_path,semantic_pack):
 f,s=matrix(tmp_path,semantic_pack,['是否适用','是否适用'],[['是','是'],['否','否']],footer=(6,'模板填报说明'))
 assert s.columns['applicability']==6
 assert any(x['reason']=='equivalent_business_rows' for x in f.preservation['column_selections'])


def test_real_conflicts_remain_unselected(tmp_path,semantic_pack):
 _,s=matrix(tmp_path,semantic_pack,['凤凰公司适用情况','凤凰分公司适用情况'],[['是','否'],['否','是']])
 assert 'applicability' not in s.columns


def test_literal_yes_no_and_applicability_are_equivalent(tmp_path,semantic_pack):
 f,s=matrix(tmp_path,semantic_pack,['是否适用','凤凰分公司适用情况'],[['是','适用'],['否','不适用']])
 assert s.columns['applicability']==6
 assert any(x['reason']=='equivalent_business_rows' for x in f.preservation['column_selections'])


def test_current_unit_answers_override_generic_template_answers(tmp_path,semantic_pack):
 f,s=matrix(tmp_path,semantic_pack,['是否适用','凤凰分公司适用情况'],[['是','不适用'],['否','适用']])
 assert s.columns['applicability']==6
 assert any(x['reason']=='current_unit_answer_values' for x in f.preservation['column_selections'])


def test_foreign_header_can_be_corrected_only_with_body_proof(tmp_path,semantic_pack):
 f,s=matrix(tmp_path,semantic_pack,['湘西自治州德能电力建设有限公司古丈分公司适用情况'],[['是'],['否'],[None]])
 assert s.columns['applicability']==5
 log=next(x for x in f.preservation['column_selections'] if x['field']=='applicability')
 assert log['reason']=='template_owner_corrected'
 assert '古丈' in log['candidates'][0]['header']
 assert log['candidates'][0]['owner_resolution']['current_unit_body_rows']==[3,4,5]
 assert s.records[-1].value('applicability')==''
 _,s=matrix(tmp_path,semantic_pack,['古丈分公司适用情况'],[['是'],['否']],body='古丈分公司-综合管理室-主任')
 assert 'applicability' not in s.columns
 _,s=matrix(tmp_path,semantic_pack,['古丈分公司适用情况'],[['是'],['否']],body='各级单位-综合管理室-主任')
 assert 'applicability' not in s.columns


def test_local_abbreviation_and_hidden_only_source(tmp_path,semantic_pack):
 _,s=matrix(tmp_path,semantic_pack,['龙山公司适用情况'],[['是'],['否']],unit='湘西自治州德能电力建设有限公司龙山分公司',code='20CG',hidden=['E'])
 assert s.columns['applicability']==5
 _,s=matrix(tmp_path,semantic_pack,['是否适用','是否适用'],[['是','否'],['是','否']],hidden=['F'])
 assert s.columns['applicability']==5


def test_hidden_not_applicable_reason_header_supplies_reason(tmp_path, semantic_pack):
 """tmp_path 提供临时工作簿目录，semantic_pack 提供真实解析配置；隐藏的理由列必须参与第3条判断。"""
 _, sheet = matrix(
  tmp_path,
  semantic_pack,
  ['是否适用', '不适用理由'],
  [['不适用', '该项业务由县公司负责'], ['适用', None]],
  hidden=['F'],
 )
 assert sheet.columns['applicability_reason'] == 6
 assert sheet.records[0].value('applicability_reason') == '该项业务由县公司负责'
 context = CheckContext(sheet.records, [], sheet.records, {}, {}, semantic_pack, [], [])
 assert applicability_values_v2(context, {'placeholders': ['-']}) == []


def test_shared_interpretation_does_not_skip_no_business(semantic_pack):
 f=lambda text:FieldValue(text,text,'A3')
 r=Record('matrix','20CI','06','default','a.xlsx','矩阵',3,{'applicability':f('暂无此业务，参照省公司'),'measure_id':f('M1')},'r')
 ctx=CheckContext([r],[],[r],{}, {},semantic_pack,[],[])
 assert applicability_values_v2(ctx,{'placeholders':['-']})[0]['evidence']['interpretation']['state']=='business_absent'
 issues=responsibility_coverage_v5(ctx,{'generic_responsibilities':[],'confirmed_mappings':{},'confirmed_aliases':{},'generic_unit_names':['各级单位']})
 assert len(issues)==1 and issues[0]['evidence']['issue_type']=='coverage_applicability_unknown'
 assert issues[0]['evidence']['applicability_interpretation']['state']=='business_absent'


def test_explicit_negative_reason_and_formula_sources(semantic_pack):
 r=Record('matrix','20CI','06','default','a.xlsx','矩阵',3,{'applicability':FieldValue('不适用于本公司，业务由总部处理','不适用于本公司，业务由总部处理','A3')},'r')
 ctx=CheckContext([r],[],[r],{}, {},semantic_pack,[],[])
 assert not applicability_values_v2(ctx,{'placeholders':['-']})
 r.fields['applicability'].state='formula_no_cache'
 assert interpret_record(r,semantic_pack).state=='source_unavailable'


@pytest.mark.parametrize('difference,expected',[
 ('process','duty_record_missing'),
 ('measure','duty_record_missing'),('entity','duty_record_missing'),
 ('business','duty_record_missing'),('variant','duty_record_missing'),
])
def test_measure_numbers_or_scope_difference_remains_missing(semantic_pack,difference,expected):
 """difference 为数字或范围差异；真实不同措施仍须报岗位职责缺失。"""
 semantic_pack['semantic_config']['enabled']=False
 field=lambda text:FieldValue(text,text,'A3')
 matrix_id='职工福利保障与薪酬管理业务-12.职工疗休养组织-控制措施05'
 duty_id='职工福利保障与薪酬管理业务-12.职工疗养组织-控制措施05'
 if difference=='process':duty_id=duty_id.replace('-12.','-13.')
 if difference=='measure':duty_id=duty_id.replace('措施05','措施06')
 if difference=='prefix':duty_id=duty_id.replace('薪酬管理业务','其他业务')
 m=Record('matrix','20CI','09','default','matrix.xlsx','矩阵',3,
          {'measure_id':field(matrix_id),'applicability':field('适用')},'matrix')
 d=Record('position_duty','other' if difference=='entity' else '20CI',
          '06' if difference=='business' else '09','other' if difference=='variant' else 'default',
          'duty.xlsx','岗位责任清单',3,{'measure_id':field(duty_id)},'duty')
 ctx=CheckContext([m],[],[m,d],{}, {},semantic_pack,[],[])
 params={'generic_responsibilities':[],'confirmed_mappings':{},'confirmed_aliases':{},'generic_unit_names':[]}
 issues=responsibility_coverage_v5(ctx,params)
 assert len(issues)==1 and issues[0]['evidence']['issue_type']==expected
 assert m.value('measure_id')==matrix_id and d.value('measure_id')==duty_id


@pytest.mark.parametrize('difference',['name','prefix'])
def test_measure_wording_difference_joins_by_numbers(semantic_pack,difference):
 """difference 为编号文字差异；两组数字相同必须当作同一措施。"""
 semantic_pack['semantic_config']['enabled']=False
 field=lambda text:FieldValue(text,text,'A3')
 matrix_id='职工福利保障与薪酬管理业务-12.职工疗休养组织-控制措施05'
 duty_id='职工福利保障与薪酬管理业务-12.职工疗养组织-控制措施05'
 if difference=='prefix':duty_id='xxx福利保障与xxx管理业务-12.职工疗养组织-控制措施05'
 m=Record('matrix','20CI','09','default','matrix.xlsx','矩阵',3,
          {'measure_id':field(matrix_id),'applicability':field('适用'),
           'responsibility':field('财务部-薪酬专责')},'matrix')
 d=Record('position_duty','20CI','09','default','duty.xlsx','岗位责任清单',3,
          {'measure_id':field(duty_id),'department':field('财务部'),'position':field('薪酬专责')},'duty')
 ctx=CheckContext([m],[],[m,d],{}, {},semantic_pack,[],[])
 params={'generic_responsibilities':[],'confirmed_mappings':{},'confirmed_aliases':{},'generic_unit_names':[]}

 assert responsibility_coverage_v5(ctx,params)==[]


def test_grammar_repairs_do_not_allow_unlisted_documents():
 from risk_audit.checks.carriers_v5 import parse_carrier_references_v5 as parse
 assert parse('对使用合同约定账户支付的准确性负审核责任。',allowed={'合同'})['status']=='matched'
 assert parse('对使用合同约定账户支付的准确性负审核责任。',allowed={'合同或协议'})['status']=='matched'
 p=parse('对使用合同约定账户支付的准确性负审核责任。',allowed={'付款单'})
 assert p['unresolved']==['合同']
 assert parse('对出险报案与后续配合及时性负审核责任。',allowed={'出险通知书'})['status']=='matched'
 p=parse('对出险报案报告与后续配合及时性负审核责任。',allowed={'合同'})
 assert p['status']!='matched'


def test_unverified_semantic_relation_cannot_be_activated(semantic_pack,registry):
 from risk_audit.configuration.validator import validate_pack,ConfigError
 semantic_pack['semantic_config']['candidate_only']=False
 with pytest.raises(ConfigError):validate_pack(semantic_pack,registry)


def test_local_model_missing_does_not_attempt_download(tmp_path,monkeypatch):
 import socket
 from risk_audit.semantic import LocalEncoder,SemanticUnavailable
 def denied(*args,**kwargs):raise AssertionError('network attempted')
 monkeypatch.setattr(socket.socket,'connect',denied)
 with pytest.raises(SemanticUnavailable):LocalEncoder(tmp_path/'missing',tmp_path/'cache.sqlite3')


def test_model_integrity_checked_before_loading(tmp_path):
 from risk_audit.semantic import LocalEncoder,SemanticUnavailable
 (tmp_path/'manifest.json').write_text('{}')
 with pytest.raises(SemanticUnavailable):LocalEncoder(tmp_path,tmp_path/'cache.sqlite3',expected_manifest='0'*64)


def test_embedding_never_overrides_a_real_difference(semantic_pack):
 from risk_audit.checks.carriers_v5 import parse_carrier_references_v5
 class VeryConfident:
  def suggest(self,text,entries,**kwargs):
   return {'candidates':[{'text':entries[0]['text'],'score':1.0}],'decision':'candidate_only'}
 p=parse_carrier_references_v5('对收款申请单的准确性负有审核责任。',allowed={'付款申请单'},assistant=VeryConfident())
 assert p['status']!='matched' and '收款申请单' in p['unresolved']


def test_encoder_offline_cache_and_no_truncation(tmp_path,monkeypatch):
 pytest.importorskip('onnxruntime');pytest.importorskip('tokenizers')
 import socket
 from risk_audit.semantic import LocalEncoder,SemanticUnavailable
 model=ROOT/'models/bge-small-zh-v1.5'
 if not model.exists():pytest.skip('integration model is not installed')
 def denied(*args,**kwargs):raise AssertionError('network attempted')
 monkeypatch.setattr(socket.socket,'connect',denied)
 encoder=LocalEncoder(model,tmp_path/'cache.sqlite3')
 first=encoder.encode(['暂无此业务','暂无此业务','直接引用'])
 assert first.shape==(3,512) and encoder.stats['encoded_texts']==2
 assert all(v.flags.owndata for v in encoder.memory.values())
 assert (first[0]==first[1]).all()
 again=encoder.encode(['暂无此业务']);assert encoder.stats['encoded_texts']==2
 assert (again[0]==first[0]).all()
 encoder.cache.execute("UPDATE embeddings SET vector=?,digest=? WHERE text=?",(b'bad','bad','暂无此业务'));encoder.cache.commit()
 # The validated in-memory copy remains usable; clear it to exercise disk repair.
 assert (encoder.encode(['暂无此业务'])[0]==first[0]).all()
 encoder.memory.clear()
 encoder.encode(['暂无此业务']);assert encoder.stats['encoded_texts']==3
 with pytest.raises(SemanticUnavailable):encoder.encode(['控制措施'*600])
 encoder.close()


def test_invalid_ownership_never_gets_repaired(tmp_path,semantic_pack):
 f,s=matrix(tmp_path,semantic_pack,['古丈分公司适用情况'],[['是'],['否']])
 f.entity_conflict=True
 f.sheets=parse_workbook(f,f.source,semantic_pack['field_aliases'])
 assert 'applicability' not in f.sheets[0].columns
