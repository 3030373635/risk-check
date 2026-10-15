"""验证 1015 重庆确认稿的编号、规则映射和专项审核能力。"""

from __future__ import annotations

import copy
import json
import re
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

from risk_audit.checks.registry import CheckContext, build_registry
from risk_audit.configuration.loader import load_pack
from risk_audit.engine import SafeFormat, _render_opinion, run_engine
from risk_audit.models import Entity, FieldValue, FileRecord, Finding, ParsedSheet, Record
from risk_audit.pipeline_versions import (
    CONFIRMED_RULE_VERSIONS,
    PREPROCESSING_RULE_VERSIONS,
    SYSTEM_TYPE_PREPROCESSING_RULE_VERSIONS,
)
from risk_audit.readers.excel import parse_workbook
from risk_audit.runner import audit
from risk_audit.util import sha256_file
from risk_audit.writer import OutputState, write_outputs


ROOT = Path(__file__).resolve().parents[1]
BASE_RELEASE = ROOT / "rulepacks/releases/1.9.32"
RELEASE = ROOT / "rulepacks/releases/1.9.33"
SOURCE = ROOT / "规则来源/风控矩阵1015-重庆.docx"
ENTITY_CODE = "CQ001"
BUSINESS_CODE = "02"
BUSINESS_ID = "financial_management"


def make_record(record_type: str, row: int, **values: str) -> Record:
    """构造重庆专项能力测试记录。

    Args:
        record_type: 矩阵、岗位责任清单或系统规则清单记录类型。
        row: 原始工作表行号。
        values: 按逻辑字段名给出的单元格文字。

    Returns:
        带固定主体、业务及可追溯单元格坐标的记录。
    """

    fields = {
        name: FieldValue(value, value, f"A{row}")
        for name, value in values.items()
    }
    record = Record(
        record_type,
        ENTITY_CODE,
        BUSINESS_CODE,
        "default",
        "02财务管理.xlsx",
        record_type,
        row,
        fields,
        f"{record_type}-{row}",
        business_id=BUSINESS_ID,
    )
    names = values.get("person_names", "")
    record.person_keys = [
        {"name": name, "person_id": name, "identity_reliable": True}
        for name in names.replace("，", "、").split("、")
        if name
    ]
    return record


def make_context(records: list[Record], selected: list[Record] | None = None) -> CheckContext:
    """构造重庆专项能力运行上下文。

    Args:
        records: 跨表核对使用的全部记录。
        selected: 当前规则直接检查的记录；为空时使用全部记录。

    Returns:
        仅包含一个重庆主体和一个财务管理业务的审核上下文。
    """

    return CheckContext(
        selected if selected is not None else records,
        [],
        records,
        {ENTITY_CODE: Entity(ENTITY_CODE, "国网重庆市电力公司")},
        {},
        {},
        [ENTITY_CODE],
        [{
            "business_id": BUSINESS_ID,
            "business_code": BUSINESS_CODE,
            "variant_id": "default",
            "required": True,
        }],
    )


def run_capability(name: str, records: list[Record]) -> list[dict]:
    """执行一个无额外参数的重庆 1015 专项能力。

    Args:
        name: 能力注册名称。
        records: 参与跨表核对的全部业务记录。

    Returns:
        能力产生的原始结构化问题。
    """

    capability = build_registry().get(name, 1)
    assert capability is not None
    return capability.runner(make_context(records), {})


def test_display_code_controls_word_rule_number() -> None:
    """规则重新排序后，意见前缀必须使用 display_code，而不是模板里的旧序号。"""

    pack = copy.deepcopy(load_pack(BASE_RELEASE))
    rule = next(rule for rule in pack["rules"] if rule["rule_id"] == "chongqing.person_name_single")
    rule["display_code"] = "R14"
    pack["rules"] = [rule]
    pack["submission_scope"] = {
        "batch": "chongqing",
        "entity_codes": [ENTITY_CODE],
        "businesses": [{
            "business_id": BUSINESS_ID,
            "business_code": BUSINESS_CODE,
            "variant_id": "default",
            "required": True,
        }],
    }
    record = make_record("position_duty", 2, person_names="张三 李四")
    sheet = ParsedSheet("岗位内控责任清单", "position_duty", [1], {}, {}, 7, 2, [record])
    file = FileRecord(
        Path("02财务管理.xlsx"),
        Path("02财务管理.xlsx"),
        "hash",
        "xlsx",
        ENTITY_CODE,
        [],
        False,
        BUSINESS_CODE,
        "default",
        "three_lists",
        [sheet],
        business_id=BUSINESS_ID,
    )

    findings, _ = run_engine(
        pack,
        build_registry(),
        [file],
        {ENTITY_CODE: Entity(ENTITY_CODE, "国网重庆市电力公司")},
        {},
        run_id="1015-numbering",
    )

    assert len(findings) == 1
    assert findings[0].message.startswith("【第14条】")
    assert "第20条" not in findings[0].message


def test_material_opinion_uses_material_and_word_rule_labels() -> None:
    """资料级意见必须同时显示资料级标识和 Word 规则编号。"""

    message = _render_opinion(
        "本次应报材料中缺少风控矩阵",
        SafeFormat(),
        "R01",
        "material",
    )

    assert message == "【资料级】【第1条】本次应报材料中缺少风控矩阵"


def test_release_matches_word_rule_order_and_removes_deleted_rules() -> None:
    """1.9.33 只保留 Word 接受修订后的 14 条规则及两项预处理。"""

    pack = load_pack(RELEASE)
    numbered = [
        (rule["display_code"], rule["title"])
        for rule in pack["rules"]
        if rule["display_code"].startswith("R")
    ]

    assert [code for code, _ in numbered] == [f"R{number:02d}" for number in range(1, 15)]
    assert [title for _, title in numbered] == [
        "资料提交完整性",
        "会计主体与提交资料对应情况",
        "基准字段及适用性填写情况",
        "岗位清单中部门、岗位名称和人员姓名",
        "不相容岗位清单存在性",
        "经办与审核、监督人员分离",
        "岗位内控责任清单与矩阵一致性审核",
        "岗位职责基本责任表述",
        "角色与责任类型对应",
        "岗位职责清单中重复项筛查",
        "系统控制规则清单控制方式匹配性判断",
        "矩阵中责任主体规范性审核",
        "矩阵中责任主体与岗位内控责任清单一致性审核",
        "岗位责任清单中一人多岗审核",
    ]
    assert {rule["display_code"] for rule in pack["rules"] if rule["display_code"].startswith("P")} == {"P01", "P02"}
    assert not {
        "systems.forbidden_platform",
        "matrices.forbidden_platform",
        "deleted_matrix_text.in_duties",
        "deleted_matrix_text.in_incompatible_duties",
        "chongqing.power_supply_branch_key_points",
        "chongqing.control_measure_completeness",
        "chongqing.person_name_single",
    } & {rule["rule_id"] for rule in pack["rules"]}
    assert all(
        not re.match(r"^【第\d+条】", template)
        for rule in pack["rules"]
        if rule["display_code"].startswith("R")
        for check in rule["checks"]
        for template in check["message"].values()
    )
    assert pack["manifest"]["source_document_sha256"] == sha256_file(SOURCE)


def test_rule_6_checks_handler_against_reviewer_and_supervisor() -> None:
    """第6条应分别识别经办与审核、经办与监督由同一人承担。"""

    records = [
        make_record(
            "position_duty",
            2,
            measure_id="财务管理-12-控制措施05",
            department="财务部",
            position="核算专责",
            person_names="张三",
            role="经办",
            duty="对资料完整性负主体责任",
        ),
        make_record(
            "position_duty",
            3,
            measure_id="其他名称-12-控制措施05",
            department="财务部",
            position="复核专责",
            person_names="张三",
            role="审核",
            duty="对资料完整性负审核责任",
        ),
        make_record(
            "position_duty",
            4,
            measure_id="其他名称-12-控制措施05",
            department="财务部",
            position="监督专责",
            person_names="张三",
            role="监督",
            duty="对资料完整性负监督责任",
        ),
    ]

    issues = run_capability("chongqing_duty_separation", records)

    conflict_types = {
        item["evidence"]["conflict_type"]
        for item in issues
        if item["evidence"]["issue_type"] == "same_duty_person_overlap"
    }
    assert conflict_types == {"经办审核", "经办监督"}


def test_rule_7_compares_measure_ids_in_both_directions() -> None:
    """第7条应把矩阵缺岗位和岗位缺矩阵分别落在对应原表行。"""

    matrix_only = make_record("matrix", 2, measure_id="财务管理-12-控制措施05")
    duty_only = make_record("position_duty", 3, measure_id="财务管理-18-控制措施02")

    issues = run_capability("chongqing_matrix_duty_alignment", [matrix_only, duty_only])

    assert {(item["record"].record_id, item["evidence"]["opinion"]) for item in issues} == {
        (matrix_only.record_id, "请核实岗位内控责任清单编制是否完整"),
        (duty_only.record_id, "矩阵中没有该条控制措施请核实是否需要保留"),
    }


def test_number_alignment_and_separation_ignore_missing_measure_numbers() -> None:
    """没有任何数字的措施编号不得被当作同一措施或双向缺失编号。"""

    matrix = make_record("matrix", 2, measure_id="")
    handler = make_record(
        "position_duty",
        3,
        measure_id="待补充",
        department="财务部",
        person_names="张三",
        role="经办",
        duty="对资料完整性负主体责任",
    )
    reviewer = make_record(
        "position_duty",
        4,
        measure_id="待补充",
        department="财务部",
        person_names="张三",
        role="审核",
        duty="对资料完整性负审核责任",
    )

    alignment = run_capability("chongqing_matrix_duty_alignment", [matrix, handler, reviewer])
    separation = run_capability("chongqing_duty_separation", [handler, reviewer])

    assert alignment == []
    assert all(item["evidence"]["issue_type"] != "same_duty_person_overlap" for item in separation)


def test_rule_6_falls_back_to_control_measure_text_when_number_is_unavailable() -> None:
    """第6条无法按编号关联时，应按相同控制措施正文继续核对人员分离。"""

    handler = make_record(
        "position_duty",
        3,
        measure_id="待补充",
        control_measure="复核付款资料",
        department="财务部",
        person_names="张三",
        role="经办",
        duty="对资料完整性负主体责任",
    )
    reviewer = make_record(
        "position_duty",
        4,
        measure_id="待补充",
        control_measure="复核付款资料",
        department="财务部",
        person_names="张三",
        role="审核",
        duty="对资料完整性负审核责任",
    )

    issues = run_capability("chongqing_duty_separation", [handler, reviewer])

    assert {
        item["evidence"]["conflict_type"]
        for item in issues
        if item["evidence"]["issue_type"] == "same_duty_person_overlap"
    } == {"经办审核"}


def test_rule_11_checks_status_and_system_name() -> None:
    """第11条应同时核对建设状态、矩阵控制方式和系统名称。"""

    matrix = make_record(
        "matrix",
        2,
        measure_id="财务管理-12-控制措施05",
        control_mode="系统已建",
        control_system="财务管控系统",
    )
    system = make_record(
        "system_rule",
        3,
        measure_id="其他名称-12-控制措施05",
        build_status="系统待建",
        system_name="资金管理系统",
    )

    issues = run_capability("chongqing_system_control_consistency", [matrix, system])

    assert {item["evidence"]["opinion"] for item in issues} == {
        "需核实建设状态是否正确",
        "需核实系统规则清单中系统名称与矩阵中控制系统的一致性",
    }


def test_rules_12_and_13_check_format_and_bidirectional_names() -> None:
    """第12、13条应区分责任主体格式问题和部门岗位双向缺失问题。"""

    matrix = make_record(
        "matrix",
        2,
        measure_id="财务管理-12-控制措施05",
        responsibility="各级单位-重庆公司财务部-核算专责、国网重庆市电力公司-审计部-审计专责",
    )
    duty = make_record(
        "position_duty",
        3,
        measure_id="财务管理-12-控制措施05",
        department="资金部",
        position="资金专责",
        person_names="李四",
    )

    format_issues = run_capability("chongqing_responsibility_format", [matrix, duty])
    alignment_issues = run_capability("chongqing_responsibility_name_alignment", [matrix, duty])

    assert {item["evidence"]["opinion"] for item in format_issues} == {
        "请将责任主体修改为‘公司名称-部门名称-岗位名称’的形式",
        "请核实部门名称是否准确",
    }
    assert {item["evidence"]["opinion"] for item in alignment_issues} == {
        "请核实岗位内控责任清单中的部门名称是否完整",
        "请核实岗位内控责任清单中的岗位名称是否完整",
        "请核实矩阵中的部门名称是否完整",
        "请核实矩阵中的岗位名称是否完整",
    }


def test_rule_14_reports_every_row_of_one_person_with_multiple_positions() -> None:
    """第14条应在同一人员对应多个不同岗位时标记所有关联行。"""

    first = make_record("position_duty", 2, person_names="张三", position="核算专责")
    second = make_record("position_duty", 3, person_names="张三", position="资金专责")
    same_position = make_record("position_duty", 4, person_names="李四", position="核算专责")

    issues = run_capability("chongqing_person_multiple_positions", [first, second, same_position])

    assert {item["record"].record_id for item in issues} == {first.record_id, second.record_id}
    assert {item["evidence"]["opinion"] for item in issues} == {"请核实是否为一人多岗。"}


def test_rule_14_compares_one_person_across_businesses() -> None:
    """第14条按主体核对人员，不得把不同业务中的不同岗位拆开。"""

    finance = make_record("position_duty", 2, person_names="张三", position="核算专责")
    contract = make_record("position_duty", 3, person_names="张三", position="合同专责")
    contract.business_id = "contract_management"
    contract.business_code = "03"
    contract.file_path = "03合同管理.xlsx"

    issues = run_capability("chongqing_person_multiple_positions", [finance, contract])

    assert {item["record"].record_id for item in issues} == {finance.record_id, contract.record_id}


def test_runner_executes_rule_14_once_with_all_entity_businesses(tmp_path: Path) -> None:
    """运行器必须在主体级汇总不同业务的岗位清单后执行第14条。

    Args:
        tmp_path: pytest 提供的隔离输入、输出和运行目录。
    """

    input_root = tmp_path / "input"
    source_paths: list[Path] = []
    for directory, filename, position in (
        ("02 财务管理", "02财务管理三清单.xlsx", "核算专责"),
        ("03 合同管理", "03合同管理三清单.xlsx", "合同专责"),
    ):
        source = input_root / "测试公司" / directory / filename
        source.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "岗位内控责任清单"
        worksheet.append(["控制措施编号", "控制措施", "部门", "岗位名称", "人员姓名", "人员编号", "岗位职责", "角色"])
        worksheet.append(["M1", "复核资料", "财务部", position, "张三", "E1", "对资料完整性负主体责任", "经办"])
        workbook.save(source)
        source_paths.append(source)

    result = audit(
        input_root,
        tmp_path / "output",
        RELEASE,
        None,
        ROOT / "baselines/chongqing/templates",
        tmp_path / "runs",
        run_id="entity-rule-14",
        enabled_rules=["R14"],
        entity_mode="parent_directory",
    )

    findings = json.loads(
        (tmp_path / "runs/entity-rule-14/findings.json").read_text(encoding="utf-8")
    )
    rule_14_findings = [item for item in findings if item["display_code"] == "R14"]
    assert len(rule_14_findings) == 2
    for source in source_paths:
        output = load_workbook(tmp_path / "output" / source.relative_to(input_root))
        worksheet = output["岗位内控责任清单"]
        opinion_column = next(
            cell.column for cell in worksheet[1] if cell.value == "审核意见"
        )
        assert "【第14条】请核实是否为一人多岗。" in worksheet.cell(2, opinion_column).value


def test_rule_14_does_not_reinclude_failed_business(
    tmp_path: Path,
    monkeypatch,
) -> None:
    """业务级处理失败后，第14条不得再审核或写出该业务。

    Args:
        tmp_path: pytest 提供的隔离输入、输出和运行目录。
        monkeypatch: pytest 提供的运行时替换工具。
    """

    from risk_audit import runner

    input_root = tmp_path / "input"
    source_paths: dict[str, Path] = {}
    for directory, filename, position in (
        ("02 财务管理", "02财务管理三清单.xlsx", "核算专责"),
        ("03 合同管理", "03合同管理三清单.xlsx", "合同专责"),
    ):
        source = input_root / "测试公司" / directory / filename
        source.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "岗位内控责任清单"
        worksheet.append(["控制措施编号", "控制措施", "部门", "岗位名称", "人员姓名", "人员编号", "岗位职责", "角色"])
        worksheet.append(["M1", "复核资料", "财务部", position, "张三", "E1", "对资料完整性负主体责任", "经办"])
        workbook.save(source)
        source_paths[directory] = source

    original_preprocess = runner.preprocess_business_file

    def fail_contract_preprocess(file, pack, baselines, work_dir) -> None:
        """模拟合同管理业务预处理失败。

        Args:
            file: 当前业务文件。
            pack: 当前规则包。
            baselines: 业务基准。
            work_dir: 当前业务工作目录。
        """

        if file.business_id == "contract_management":
            raise RuntimeError("模拟合同业务预处理失败")
        original_preprocess(file, pack, baselines, work_dir)

    monkeypatch.setattr(runner, "preprocess_business_file", fail_contract_preprocess)

    result = audit(
        input_root,
        tmp_path / "output",
        RELEASE,
        None,
        ROOT / "baselines/chongqing/templates",
        tmp_path / "runs",
        run_id="failed-entity-rule-14",
        entity_mode="parent_directory",
    )

    findings = json.loads(
        (tmp_path / "runs/failed-entity-rule-14/findings.json").read_text(encoding="utf-8")
    )
    assert not [item for item in findings if item["display_code"] == "R14"]
    assert result["skipped_businesses"] == 1
    assert (tmp_path / "output" / source_paths["02 财务管理"].relative_to(input_root)).is_file()
    assert not (tmp_path / "output" / source_paths["03 合同管理"].relative_to(input_root)).exists()


def test_rule_13_applies_exception_lists_by_containment() -> None:
    """第13条的部门和岗位例外名单必须按“包含”口径执行。"""

    matrix = make_record(
        "matrix",
        2,
        responsibility="国网重庆市电力公司-项目管理部门（基建）-项目管理人员（建设）",
    )
    duty = make_record(
        "position_duty",
        3,
        department="项目管理部门-建设部",
        position="项目管理人员（建设）",
        person_names="张三",
    )
    pack = load_pack(RELEASE)
    rule = next(item for item in pack["rules"] if item["display_code"] == "R13")
    check = rule["checks"][0]
    capability = build_registry().get(check["operator"], check["operator_version"])
    assert capability is not None

    issues = capability.runner(make_context([matrix, duty]), check["params"])

    assert {item["evidence"]["opinion"] for item in issues} == {
        "请明确到具体部门",
        "请明确到具体岗位",
    }


def test_publish_script_rebuilds_release_deterministically() -> None:
    """1015 重庆发布脚本必须可以从 1.9.32 稳定重建 1.9.33。"""

    result = subprocess.run(
        [sys.executable, str(ROOT / "tools/publish_rules_1015_chongqing_1933.py"), "--verify-only"],
        cwd=ROOT.parent,
        capture_output=True,
        text=True,
        check=False,
    )

    assert result.returncode == 0, result.stderr
    assert json.loads(result.stdout)["verified"] is True


def test_1933_pipeline_remains_available_after_region_upgrade() -> None:
    """地区升级后，历史 1.9.33 管线仍应可复验。"""

    index = json.loads((ROOT / "regions/index.json").read_text(encoding="utf-8"))
    profile = next(item for item in index["regions"] if item["region_code"] == "chongqing")
    assert profile["rulepack_version"] == "1.9.34"
    assert "1.9.33" in CONFIRMED_RULE_VERSIONS
    assert "1.9.33" in PREPROCESSING_RULE_VERSIONS
    assert "1.9.33" in SYSTEM_TYPE_PREPROCESSING_RULE_VERSIONS


def test_reaudit_clears_opinion_from_deleted_rule(tmp_path: Path) -> None:
    """重新审核无本轮问题时，应清除源表里已删除规则留下的旧编号意见。

    Args:
        tmp_path: pytest 提供的隔离输入和输出目录。
    """

    source = tmp_path / "02财务管理三清单.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "岗位内控责任清单"
    worksheet.append(["控制措施编号", "部门", "岗位名称", "人员姓名", "岗位职责", "角色", "审核意见"])
    worksheet.append([
        "财务管理-12-控制措施05",
        "财务部",
        "核算专责",
        "张三",
        "对资料完整性负主体责任",
        "经办",
        "【第17条】旧版已删除规则意见",
    ])
    workbook.save(source)
    pack = load_pack(RELEASE)
    file = FileRecord(
        source,
        Path(source.name),
        sha256_file(source),
        "xlsx",
        ENTITY_CODE,
        [],
        False,
        BUSINESS_CODE,
        "default",
        "three_lists",
        business_id=BUSINESS_ID,
    )
    file.sheets = parse_workbook(file, source, pack["field_aliases"])

    write_outputs(
        [file],
        [],
        tmp_path / "output",
        metadata_dir=tmp_path / "metadata",
    )

    output = load_workbook(tmp_path / "output" / source.name)
    assert output["岗位内控责任清单"]["G2"].value is None


def test_reaudit_removes_deleted_rule_but_keeps_manual_opinion(tmp_path: Path) -> None:
    """重新审核只清理可识别的旧程序意见，人工复核文字必须保留。

    Args:
        tmp_path: pytest 提供的隔离输入和输出目录。
    """

    source = tmp_path / "02财务管理三清单.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "岗位内控责任清单"
    worksheet.append(["控制措施编号", "部门", "岗位名称", "人员姓名", "岗位职责", "角色", "审核意见"])
    worksheet.append([
        "财务管理-12-控制措施05",
        "财务部",
        "核算专责",
        "张三",
        "对资料完整性负主体责任",
        "经办",
        "人工复核：请保留\n【第17条】旧版已删除规则意见",
    ])
    workbook.save(source)
    pack = load_pack(RELEASE)
    file = FileRecord(
        source,
        Path(source.name),
        sha256_file(source),
        "xlsx",
        ENTITY_CODE,
        [],
        False,
        BUSINESS_CODE,
        "default",
        "three_lists",
        business_id=BUSINESS_ID,
    )
    file.sheets = parse_workbook(file, source, pack["field_aliases"])

    write_outputs(
        [file],
        [],
        tmp_path / "output",
        metadata_dir=tmp_path / "metadata",
    )

    output = load_workbook(tmp_path / "output" / source.name)
    assert output["岗位内控责任清单"]["G2"].value == "人工复核：请保留"


def test_reaudit_streams_opinions_without_random_read_only_access(
    tmp_path: Path,
    monkeypatch,
) -> None:
    """复审应顺序读取意见列，避免只读工作簿逐行随机访问。

    Args:
        tmp_path: pytest 提供的隔离输入和输出目录。
        monkeypatch: pytest 提供的运行时替换工具。
    """

    source = tmp_path / "02财务管理三清单.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "岗位内控责任清单"
    worksheet.append(["控制措施编号", "部门", "岗位名称", "人员姓名", "岗位职责", "角色", "审核意见"])
    worksheet.append([
        "财务管理-12-控制措施05",
        "财务部",
        "核算专责",
        "张三",
        "对资料完整性负主体责任",
        "经办",
        "人工复核：请保留\n【第17条】旧版已删除规则意见",
    ])
    workbook.save(source)
    pack = load_pack(RELEASE)
    file = FileRecord(
        source,
        Path(source.name),
        sha256_file(source),
        "xlsx",
        ENTITY_CODE,
        [],
        False,
        BUSINESS_CODE,
        "default",
        "three_lists",
        business_id=BUSINESS_ID,
    )
    file.sheets = parse_workbook(file, source, pack["field_aliases"])

    from openpyxl.worksheet._read_only import ReadOnlyWorksheet

    def reject_random_cell_access(self, row: int, column: int) -> None:
        """拒绝只读工作表的随机单元格访问。

        Args:
            self: openpyxl 只读工作表实例。
            row: 待读取的行号。
            column: 待读取的列号。
        """

        raise AssertionError(f"不得随机读取只读工作表：{row}, {column}")

    monkeypatch.setattr(ReadOnlyWorksheet, "cell", reject_random_cell_access)

    write_outputs(
        [file],
        [],
        tmp_path / "output",
        metadata_dir=tmp_path / "metadata",
    )

    output = load_workbook(tmp_path / "output" / source.name)
    assert output["岗位内控责任清单"]["G2"].value == "人工复核：请保留"


def test_reaudit_keeps_untouched_manual_rich_text(tmp_path: Path) -> None:
    """没有旧程序意见和本轮问题时，人工富文本单元格必须原样保留。

    Args:
        tmp_path: pytest 提供的隔离输入和输出目录。
    """

    source = tmp_path / "02财务管理三清单.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "岗位内控责任清单"
    worksheet.append(["控制措施编号", "部门", "岗位名称", "人员姓名", "岗位职责", "角色", "审核意见"])
    worksheet.append(["M1", "财务部", "核算专责", "张三", "对资料完整性负主体责任", "经办"])
    worksheet["G2"] = CellRichText([
        TextBlock(InlineFont(color="FFFF0000", b=True), "人工复核："),
        "请保留",
    ])
    workbook.save(source)
    pack = load_pack(RELEASE)
    file = FileRecord(
        source,
        Path(source.name),
        sha256_file(source),
        "xlsx",
        ENTITY_CODE,
        [],
        False,
        BUSINESS_CODE,
        "default",
        "three_lists",
        business_id=BUSINESS_ID,
    )
    file.sheets = parse_workbook(file, source, pack["field_aliases"])

    write_outputs([file], [], tmp_path / "output", metadata_dir=tmp_path / "metadata")

    output = load_workbook(tmp_path / "output" / source.name, rich_text=True)
    output_value = output["岗位内控责任清单"]["G2"].value
    assert isinstance(output_value, CellRichText)
    assert str(output_value) == "人工复核：请保留"


def test_later_writeback_removes_resolved_material_pending_state(tmp_path: Path) -> None:
    """资料级意见在后续找到承载表后，必须从待回填状态中移除。

    Args:
        tmp_path: pytest 提供的隔离输入、输出和元数据目录。
    """

    source = tmp_path / "02财务管理.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "风控矩阵"
    worksheet.append(["控制措施编号", "控制措施", "审核意见"])
    worksheet.append(["M1", "复核资料", None])
    workbook.save(source)
    pack = load_pack(RELEASE)
    file = FileRecord(
        source,
        Path(source.name),
        sha256_file(source),
        "xlsx",
        ENTITY_CODE,
        [],
        False,
        BUSINESS_CODE,
        "default",
        "matrix",
        business_id=BUSINESS_ID,
    )
    file.sheets = parse_workbook(file, source, pack["field_aliases"])
    finding = Finding(
        "material-key",
        "material-id",
        "documents.completeness",
        "required_matrix_and_lists",
        "R01",
        1,
        "violation",
        ENTITY_CODE,
        "国网重庆市电力公司",
        BUSINESS_CODE,
        "default",
        "",
        "",
        None,
        "【资料级】【第1条】缺少三清单。",
        {},
        "material",
        business_id=BUSINESS_ID,
    )
    state = OutputState()
    metadata = tmp_path / "metadata"

    warnings, _ = write_outputs(
        [], [finding], tmp_path / "output", metadata_dir=metadata, output_state=state
    )
    assert state.pending and warnings[0]["finding_keys"] == [finding.finding_key]

    write_outputs(
        [file], [finding], tmp_path / "output", metadata_dir=metadata, output_state=state
    )

    assert state.pending == []
    assert json.loads((metadata / "待回填资料级意见.json").read_text(encoding="utf-8")) == []
