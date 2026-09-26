from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

import pytest
import xlrd
from lxml import etree
from openpyxl import Workbook, load_workbook

from risk_audit.readers import xls
from risk_audit.readers.xls import _same_value


class SourceSheetWithoutMerges:
    """提供恢复逻辑所需字段的无合并 BIFF 工作表。"""

    name = 'Sheet'
    visibility = 0
    nrows = 0
    ncols = 0
    merged_cells = []
    rich_text_runlist_map = {}


class SourceWorkbookWithoutMerges:
    """提供恢复逻辑所需字段的单工作表 BIFF 工作簿。"""

    nsheets = 1
    colour_map = {}

    def __init__(self):
        """初始化唯一的源工作表；无参数。"""
        self.sheet = SourceSheetWithoutMerges()

    def sheet_by_index(self, index: int) -> SourceSheetWithoutMerges:
        """按索引返回源工作表；index 为从零开始的工作表索引。"""
        assert index == 0
        return self.sheet


def open_source_workbook_without_merges(path: Path, formatting_info: bool = False):
    """返回无合并源工作簿；参数为占位文件路径及格式读取开关。"""
    assert isinstance(path, Path)
    assert formatting_info is True
    return SourceWorkbookWithoutMerges()


def create_error_formula_xls(temp_dir: Path, soffice: Path) -> Path:
    """创建带 #N/A 缓存的 XLS。

    参数 temp_dir 为测试临时目录，soffice 为 LibreOffice 可执行文件路径；
    返回生成的 XLS 路径。
    """
    source_xlsx = temp_dir / 'error-formula.xlsx'
    workbook = Workbook()
    workbook.active['A1'] = '=NA()'
    workbook.save(source_xlsx)

    xls_dir = temp_dir / 'xls'
    xls_dir.mkdir()
    # 先由真实转换器生成 BIFF 错误缓存，避免用模拟值验证转换链路。
    result = subprocess.run(
        [
            str(soffice), '--headless', '--convert-to', 'xls',
            '--outdir', str(xls_dir), str(source_xlsx),
        ],
        check=False,
        capture_output=True,
        text=True,
    )
    assert result.returncode == 0, result.stderr
    return xls_dir / 'error-formula.xls'


def test_biff_error_code_matches_equivalent_excel_error_text():
    """BIFF 错误码与转换后的同义 Excel 错误文本应视为相同值。"""
    assert _same_value(42, '#N/A', source_type=xlrd.XL_CELL_ERROR)


def test_conversion_accepts_equivalent_biff_error_cache(tmp_path, monkeypatch):
    """保真校验应接受等价错误缓存；参数为临时目录和属性替换器。"""
    soffice = shutil.which('soffice')
    if not soffice:
        pytest.skip('LibreOffice 不可用')
    soffice_path = Path(soffice).resolve()
    source = create_error_formula_xls(tmp_path, soffice_path)
    monkeypatch.setattr(xls, 'SOFFICE', soffice_path)

    converted, report = xls.convert_xls(source, tmp_path / 'converted')

    workbook = load_workbook(converted, data_only=True)
    try:
        assert report['passed'] is True
        assert workbook.active['A1'].value == '#N/A'
    finally:
        workbook.close()


@pytest.mark.parametrize(
    ('source_ranges', 'expected_refs'),
    [
        ([], []),
        ([(1, 3, 2, 5)], ['C2:E3']),
    ],
)
def test_restore_merge_ranges_uses_biff_source_exactly(source_ranges, expected_refs):
    """转换结果应精确采用 BIFF 合并范围；参数为源范围及预期 OOXML 引用。"""
    root = etree.fromstring(
        f'<worksheet xmlns="{xls.MAIN}"><sheetData/>'
        '<mergeCells count="1"><mergeCell ref="A1:K1"/></mergeCells>'
        '</worksheet>'
    )

    xls._restore_merge_ranges(root, source_ranges)

    merge_cells = root.find(f'{{{xls.MAIN}}}mergeCells')
    actual_refs = [] if merge_cells is None else [
        node.get('ref') for node in merge_cells.findall(f'{{{xls.MAIN}}}mergeCell')
    ]
    assert actual_refs == expected_refs


def test_cache_restore_removes_merge_added_by_converter(tmp_path, monkeypatch):
    """缓存恢复应移除转换器新增合并；参数为临时目录和属性替换器。"""
    converted = tmp_path / 'converted.xlsx'
    workbook = Workbook()
    workbook.active.merge_cells('A1:K1')
    workbook.save(converted)
    monkeypatch.setattr(xls.xlrd, 'open_workbook', open_source_workbook_without_merges)

    xls.restore_biff_formula_caches(tmp_path / 'source.xls', converted)

    restored = load_workbook(converted, data_only=False)
    try:
        assert list(restored.active.merged_cells.ranges) == []
    finally:
        restored.close()
