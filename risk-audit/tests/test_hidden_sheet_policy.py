from pathlib import Path
import zipfile
import pytest
from openpyxl import Workbook, load_workbook
from risk_audit.models import FileRecord
from risk_audit.readers.excel import parse_files, parse_workbook
from risk_audit.hidden_sheets import hidden_sheet_alerts, write_hidden_sheet_alerts
from risk_audit.snapshot import build_snapshot
from risk_audit.util import sha256_file, sha256_json
from risk_audit.writer import write_outputs
from test_historical_sheets import sheet, ALIASES


@pytest.mark.parametrize('state',['hidden','veryHidden'])
def test_hidden_only_business_sheet_has_notice_and_no_business_records(tmp_path,state):
    w=Workbook();w.active.title='填报说明'
    h=sheet(w,'岗位内控责任清单',complete=True);h.sheet_state=state
    p=tmp_path/'三清单.xlsx';w.save(p)
    f=FileRecord(p,Path(p.name),sha256_file(p),'xlsx','205H',[],False,'06','default','three_lists')
    parse_files([f],ALIASES,tmp_path/'work')
    assert f.sheets==[] and not f.parse_errors
    alerts=hidden_sheet_alerts([f]);summary=write_hidden_sheet_alerts(tmp_path/'report',alerts)
    assert summary['skipped']==1 and summary['audited']==0
    assert state in Path(summary['report']).read_text()
    assert '是否' not in str(f.parse_errors)
    assert sha256_file(p)==f.sha256
    included=parse_workbook(f,p,ALIASES,include_hidden=True)
    assert len(included)==1 and included[0].records
    assert f.preservation['hidden_sheets'][0]['action']=='audited'


def test_hidden_current_sheet_cannot_supersede_visible_transition_sheet(tmp_path):
    w=Workbook();w.remove(w.active)
    sheet(w,'岗位职责清单-过渡')
    sheet(w,'岗位内控责任清单9.8',complete=True,hidden=True)
    p=tmp_path/'三清单.xlsx';w.save(p)
    f=FileRecord(p,Path(p.name),sha256_file(p),'xlsx','205H',[],False,'06','default','three_lists')
    parsed=parse_workbook(f,p,ALIASES)
    assert [s.title for s in parsed]==['岗位职责清单-过渡']
    assert not f.preservation.get('excluded_historical_sheets')


def test_writer_does_not_touch_hidden_content_or_visibility(tmp_path):
    w=Workbook();w.remove(w.active)
    sheet(w,'岗位内控责任清单',complete=True)
    h=sheet(w,'岗位内控责任清单隐藏',complete=True);h.sheet_state='veryHidden';h['F3']='隐藏原文'
    p=tmp_path/'三清单.xlsx';w.save(p)
    f=FileRecord(p,Path(p.name),sha256_file(p),'xlsx','205H',[],False,'06','default','three_lists')
    f.sheets=parse_workbook(f,p,ALIASES)
    out=tmp_path/'out';write_outputs([f],[],out,metadata_dir=tmp_path/'metadata')
    with zipfile.ZipFile(p) as before,zipfile.ZipFile(out/p.name) as after:
        assert before.read('xl/worksheets/sheet2.xml')==after.read('xl/worksheets/sheet2.xml')
        assert before.read('xl/workbook.xml')==after.read('xl/workbook.xml')
    assert sha256_file(p)==f.sha256


def test_visibility_choice_is_inside_snapshot_hash(pack,registry):
    a=build_snapshot(pack,[],registry.versions(),'r')
    b=build_snapshot(pack,[],registry.versions(),'r',include_hidden=True)
    assert a['sheet_visibility_policy']=='visible_only' and b['sheet_visibility_policy']=='include_hidden'
    for s in (a,b):
        digest=s.pop('snapshot_hash');assert digest==sha256_json(s)
