"""验证 LibreOffice 转换修复和运行时隔离。"""

import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook


def test_normalize_converted_auto_filter_ranges_repairs_row_only_reference(
    tmp_path: Path,
) -> None:
    """LibreOffice 产生的纯行号筛选范围必须转为合法单元格范围。

    Args:
        tmp_path: pytest 提供的隔离目录。
    """
    from risk_audit.readers import xls

    converted = tmp_path / "converted.xlsx"
    workbook = Workbook()
    historical_sheet = workbook.active
    historical_sheet.title = "历史版本"
    historical_sheet["A1"] = "起点"
    historical_sheet["IU423"] = "终点"
    historical_sheet.auto_filter.ref = "A16:IU423"
    current_sheet = workbook.create_sheet("正式版本")
    current_sheet["A1"] = "标题"
    current_sheet["G10"] = "内容"
    current_sheet.auto_filter.ref = "A1:G10"
    workbook.save(converted)

    invalid_workbook = converted.with_name("invalid-filter.xlsx")
    with zipfile.ZipFile(converted, "r") as source_zip:
        with zipfile.ZipFile(invalid_workbook, "w") as target_zip:
            for item in source_zip.infolist():
                content = source_zip.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    # 复现 LibreOffice 从旧版 XLS 产生的非法 OOXML 筛选范围。
                    content = content.replace(
                        b'<autoFilter ref="A16:IU423"/>',
                        b'<autoFilter ref="16:423"/>',
                    )
                target_zip.writestr(item, content)

    with pytest.raises(ValueError, match="Unable to read workbook"):
        load_workbook(invalid_workbook)

    normalize_ranges = getattr(xls, "normalize_converted_auto_filter_ranges", None)
    assert normalize_ranges is not None, "缺少转换后筛选范围规范化函数"
    changes = normalize_ranges(invalid_workbook)

    assert changes == [
        {
            "sheet_path": "xl/worksheets/sheet1.xml",
            "original_ref": "16:423",
            "normalized_ref": "A16:IU423",
        }
    ]
    repaired_workbook = load_workbook(invalid_workbook)
    assert repaired_workbook["历史版本"].auto_filter.ref == "A16:IU423"
    assert repaired_workbook["正式版本"].auto_filter.ref == "A1:G10"
    repaired_workbook.close()


def test_convert_xls_uses_short_explicit_destination(monkeypatch, tmp_path: Path) -> None:
    """转换命令必须使用短目标名、配置的可执行文件和独立 profile。

    Args:
        monkeypatch: pytest 提供的补丁工具。
        tmp_path: pytest 提供的隔离目录。
    """
    from risk_audit.readers import xls

    soffice = tmp_path / "runtime/libreoffice/program/soffice.com"
    soffice.parent.mkdir(parents=True)
    soffice.write_bytes(b"exe")
    profile = tmp_path / "output/.task/libreoffice-profile"
    source = tmp_path / ("超长业务路径" * 20) / ("超长原文件名" * 12 + ".xls")
    source.parent.mkdir(parents=True)
    source.write_bytes(b"xls")
    destination = tmp_path / "work/converted/0123456789abcdef0123456789abcdef.xlsx"
    captured = {}
    stage_calls: list[str] = []
    normalized_filters = [
        {
            "sheet_path": "xl/worksheets/sheet9.xml",
            "original_ref": "16:423",
            "normalized_ref": "A16:IU423",
        }
    ]

    def fake_run(command, **kwargs):
        """记录真实命令边界并生成转换目标。

        Args:
            command: LibreOffice 命令参数。
            kwargs: 子进程运行选项。
        """
        captured["command"] = command
        captured["kwargs"] = kwargs
        converted_source = Path(command[-1])
        converted_source.with_suffix(".xlsx").write_bytes(b"xlsx")
        return SimpleNamespace(returncode=0, stdout="", stderr="")

    def fake_restore(_source: Path, _destination: Path) -> int:
        """记录公式缓存恢复阶段。

        Args:
            _source: 原始 XLS 路径。
            _destination: 转换后的 XLSX 路径。
        """
        stage_calls.append("restore")
        return 0

    def fake_normalize(_converted: Path) -> list[dict[str, str]]:
        """记录筛选范围规范化阶段。

        Args:
            _converted: 转换后的 XLSX 路径。
        """
        stage_calls.append("normalize")
        return normalized_filters

    def fake_compare(_source: Path, _converted: Path) -> dict[str, bool]:
        """记录转换保真校验阶段。

        Args:
            _source: 原始 XLS 路径。
            _converted: 转换后的 XLSX 路径。
        """
        stage_calls.append("compare")
        return {"passed": True}

    monkeypatch.setattr(xls.subprocess, "run", fake_run)
    monkeypatch.setattr(xls, "restore_biff_formula_caches", fake_restore)
    monkeypatch.setattr(
        xls,
        "normalize_converted_auto_filter_ranges",
        fake_normalize,
    )
    monkeypatch.setattr(xls, "compare_conversion", fake_compare)
    monkeypatch.setattr(xls, "sys", SimpleNamespace(platform="win32"), raising=False)

    xls.configure_conversion_runtime(soffice, profile)
    converted, report = xls.convert_xls(source, destination)

    assert converted == destination
    assert stage_calls == ["restore", "normalize", "compare"]
    assert report["normalized_auto_filters"] == normalized_filters
    assert captured["command"][0] == str(soffice.resolve())
    assert Path(captured["command"][-1]).name == "0123456789abcdef0123456789abcdef.xls"
    assert Path(captured["command"][-1]).parent == destination.parent
    assert captured["command"][-3:-1] == ["--outdir", str(destination.parent)]
    assert f"-env:UserInstallation={profile.resolve().as_uri()}" in captured["command"]
    assert captured["kwargs"].get("shell") is None
    assert captured["kwargs"]["creationflags"] == 0x08000000
    settings = profile / "user/registrymodifications.xcu"
    assert "DisableMacrosExecution" in settings.read_text(encoding="utf-8")
    assert captured["kwargs"]["env"]["SAL_DISABLE_MACROS"] == "1"
    assert not destination.with_suffix(".xls").exists()


def test_convert_xls_reports_paths_when_soffice_creates_no_output(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """LibreOffice 无输出时必须返回可定位的路径诊断。

    Args:
        monkeypatch: pytest 提供的补丁工具。
        tmp_path: pytest 提供的隔离目录。
    """
    from risk_audit.readers import xls

    soffice = tmp_path / "runtime/libreoffice/program/soffice.com"
    soffice.parent.mkdir(parents=True)
    soffice.write_bytes(b"exe")
    source = tmp_path / "source.xls"
    source.write_bytes(b"xls")
    destination = tmp_path / "converted/result.xlsx"
    profile = tmp_path / "profile"

    def fake_run(_command, **_kwargs):
        """模拟退出成功但未生成目标文件的 LibreOffice。

        Args:
            _command: LibreOffice 命令参数。
            _kwargs: 子进程运行选项。
        """
        return SimpleNamespace(returncode=0, stdout="", stderr="")

    monkeypatch.setattr(xls.subprocess, "run", fake_run)
    xls.configure_conversion_runtime(soffice, profile)

    with pytest.raises(RuntimeError) as captured:
        xls.convert_xls(source, destination)

    message = str(captured.value)
    assert "exit_code=0" in message
    assert f"source_path={destination.with_suffix('.xls')}" in message
    assert f"source_path_length={len(str(destination.with_suffix('.xls')))}" in message
    assert f"destination_path={destination}" in message
    assert f"destination_path_length={len(str(destination))}" in message
    assert "stdout=<empty>" in message
    assert "stderr=<empty>" in message


def test_convert_xls_rejects_unsafe_windows_runtime_path(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """Windows 转换路径过长时必须在启动 LibreOffice 前失败。

    Args:
        monkeypatch: pytest 提供的补丁工具。
        tmp_path: pytest 提供的隔离目录。
    """
    from risk_audit.readers import xls

    soffice = tmp_path / "runtime/libreoffice/program/soffice.com"
    soffice.parent.mkdir(parents=True)
    soffice.write_bytes(b"exe")
    source = tmp_path / "source.xls"
    source.write_bytes(b"xls")
    destination = (
        tmp_path
        / ("a" * 100)
        / ("b" * 100)
        / ("c" * 100)
        / "result.xlsx"
    )

    def unexpected_run(_command, **_kwargs):
        """LibreOffice 不应在路径校验失败后启动。

        Args:
            _command: LibreOffice 命令参数。
            _kwargs: 子进程运行选项。
        """
        raise AssertionError("LibreOffice was started with an unsafe path")

    def unexpected_copy(_source: Path, _destination: Path) -> None:
        """Windows 长路径必须在复制文件前拒绝。

        Args:
            _source: 原始 XLS 路径。
            _destination: 临时 XLS 路径。
        """
        raise AssertionError("XLS was copied to an unsafe path")

    monkeypatch.setattr(xls.subprocess, "run", unexpected_run)
    monkeypatch.setattr(xls.shutil, "copy2", unexpected_copy)
    monkeypatch.setattr(xls, "sys", SimpleNamespace(platform="win32"), raising=False)
    xls.configure_conversion_runtime(soffice, tmp_path / "profile")

    with pytest.raises(RuntimeError, match="Windows conversion path is too long"):
        xls.convert_xls(source, destination)


def test_configure_runtime_reports_long_windows_soffice_before_file_access(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """Windows 安装路径过长时必须优先返回路径诊断。

    Args:
        monkeypatch: pytest 提供的补丁工具。
        tmp_path: pytest 提供的隔离目录。
    """
    from risk_audit.readers import xls

    soffice = (
        tmp_path
        / ("a" * 100)
        / ("b" * 100)
        / ("c" * 100)
        / "soffice.com"
    )
    original_resolve = Path.resolve

    def reject_long_path_resolution(path: Path, *args, **kwargs) -> Path:
        """长路径必须在文件系统解析前被拒绝。

        Args:
            path: 待解析路径。
            args: Path.resolve 位置参数。
            kwargs: Path.resolve 命名参数。
        """
        if path == soffice:
            raise AssertionError("长路径在校验前进入了文件系统解析")
        return original_resolve(path, *args, **kwargs)

    monkeypatch.setattr(xls, "sys", SimpleNamespace(platform="win32"), raising=False)
    monkeypatch.setattr(Path, "resolve", reject_long_path_resolution)

    with pytest.raises(RuntimeError, match="Windows conversion path is too long"):
        xls.configure_conversion_runtime(soffice, tmp_path / "profile")
