"""验证 1012 重庆同步模板匹配规则。"""

from __future__ import annotations

import copy
import json
import shutil
import subprocess
import sys
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook

from risk_audit.checks.registry import CheckContext, build_registry
from risk_audit.business_identity import identify_business
from risk_audit.configuration.loader import load_pack
from risk_audit.engine import run_engine
from risk_audit.models import Entity, FieldValue, FileRecord, ParsedSheet, Record
from risk_audit.runner import CONFIRMED_RULE_VERSIONS, PREPROCESSING_RULE_VERSIONS
from risk_audit.pipeline_versions import SYSTEM_TYPE_PREPROCESSING_RULE_VERSIONS
from risk_audit.readers.excel import parse_workbook
from risk_audit.system_type_preprocessing import classify_system_owner
from risk_audit.util import sha256_file
from tools.publish_rules_1012_chongqing import (
    BASE_VERSION,
    BUSINESS_MAPPING,
    EXPECTED_BUSINESS_NAMES,
    SOURCE_NAME,
    extract_key_points,
)


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "规则来源" / SOURCE_NAME
RELEASE = ROOT / "rulepacks/releases/1.9.22"
SPLIT_LISTS_RELEASE = ROOT / "rulepacks/releases/1.9.23"
CURRENT_RELEASE = ROOT / "rulepacks/releases/1.9.27"
UPDATED_SOURCE = ROOT / "规则来源/风控矩阵1012-重庆.docx"
UPDATED_RELEASE = ROOT / "rulepacks/releases/1.9.26"
REVISED_SOURCE = ROOT / "规则来源/风控矩阵及三清单审核规则 审核确认稿（重庆公司版）.docx"


def make_record(record_type: str, row: int, business_id: str = "financial_management", **values: str) -> Record:
    """构造重庆专项规则测试记录。

    Args:
        record_type: 审核记录类型。
        row: 原始工作表行号。
        business_id: 稳定业务标识。
        values: 按逻辑字段名给出的字段值。
    """

    fields = {name: FieldValue(value, value, f"A{row}") for name, value in values.items()}
    return Record(
        record_type,
        "CQ001",
        "02",
        "default",
        "财务管理.xlsx",
        record_type,
        row,
        fields,
        f"{record_type}-{row}",
        business_id=business_id,
    )


def make_context(
    records: list[Record],
    *,
    entity_name: str = "国网重庆市电力公司某供电分公司",
    files: list[FileRecord] | None = None,
    business_id: str = "financial_management",
    baselines: dict[tuple[str, str], dict] | None = None,
) -> CheckContext:
    """构造包含重庆主体与全部记录的审核上下文。

    Args:
        records: 当前审核单元内的全部记录。
        entity_name: 当前主体名称。
        files: 可选的原始材料文件。
        business_id: 当前审核单元的业务标识。
        baselines: 按业务和变体索引的重庆冻结模板数据。
    """

    return CheckContext(
        [],
        files or [],
        records,
        {"CQ001": Entity("CQ001", entity_name)},
        baselines or {},
        {},
        ["CQ001"],
        [{"business_id": business_id, "business_code": "02", "variant_id": "default"}],
    )


def test_responsibility_alignment_reports_format_and_cross_list_mismatches() -> None:
    """责任主体缺公司、使用各级单位或与岗位清单不一致时应分别提示。"""

    capability = build_registry().get("chongqing_responsibility_alignment", 1)
    assert capability is not None
    generic = make_record(
        "matrix",
        3,
        measure_id="财务管理-1.会计主体完整性管控-控制措施01",
        responsibility="各级单位-财务资产部-总账报表管理",
    )
    missing_company = make_record(
        "matrix",
        4,
        measure_id="财务管理-2.业务完整性管控-控制措施01",
        responsibility="财务资产部-预算管理",
    )
    mismatched = make_record(
        "matrix",
        5,
        measure_id="财务管理-3.资金管理-控制措施01",
        responsibility="重庆供电公司-财务资产部-资金管理",
    )
    duty = make_record(
        "position_duty",
        8,
        measure_id="财务管理-3.资金管理-控制措施01",
        department="综合管理部",
        position="综合管理",
    )

    issues = capability.runner(
        make_context([generic, missing_company, mismatched, duty]),
        {},
    )

    assert {issue["evidence"]["issue_type"] for issue in issues} == {
        "responsibility_generic_company",
        "responsibility_company_missing",
        "responsibility_department_mismatch",
        "responsibility_position_mismatch",
    }
    assert {issue["evidence"]["opinion"] for issue in issues} == {
        "请将责任主体修改为“公司名称-部门名称-岗位名称”的形式",
        "需在矩阵责任主体中补充公司名称",
        "请核实岗位责任清单中的部门名称与矩阵责任主体中的部门名称是否一致",
        "请核实岗位责任清单中的岗位名称与矩阵责任主体中的岗位名称是否一致",
    }


def test_responsibility_alignment_accepts_exact_department_and_position() -> None:
    """同一措施的公司、部门和岗位结构完整且清单名称一致时应通过。"""

    capability = build_registry().get("chongqing_responsibility_alignment", 1)
    assert capability is not None
    matrix = make_record(
        "matrix",
        3,
        measure_id="财务管理-1.会计主体完整性管控-控制措施01",
        responsibility="重庆供电公司-财务资产部-总账报表管理",
    )
    duty = make_record(
        "position_duty",
        8,
        measure_id="财务管理-1.会计主体完整性管控-控制措施01",
        department="财务资产部",
        position="总账报表管理",
    )

    assert capability.runner(make_context([matrix, duty]), {}) == []


def test_responsibility_alignment_ignores_separators_inside_parentheses() -> None:
    """括号内组织说明的连字符不得覆盖顶层公司、部门和岗位。"""

    capability = build_registry().get("chongqing_responsibility_alignment", 1)
    assert capability is not None
    position = "业务人员（营销服务中心-95598业务管理部-业务人员）"
    matrix = make_record(
        "matrix",
        3,
        measure_id="营销售电-1.客户服务-控制措施01",
        responsibility=f"市公司-营销部-{position}",
    )
    duty = make_record(
        "position_duty",
        8,
        measure_id="营销售电-1.客户服务-控制措施01",
        department="营销部",
        position=position,
    )

    assert capability.runner(make_context([matrix, duty]), {}) == []


def test_responsibility_alignment_keeps_item_separators_inside_parentheses() -> None:
    """括号内岗位说明的顿号不得被解析为另一个责任主体。"""

    capability = build_registry().get("chongqing_responsibility_alignment", 1)
    assert capability is not None
    position = "膳食管理人员（验收人员、库管人员）"
    matrix = make_record(
        "matrix",
        3,
        measure_id="物业后勤业务-1.膳食管理-控制措施01",
        responsibility=f"各级单位-膳食部门-{position}",
    )
    duty = make_record(
        "position_duty",
        8,
        measure_id="物业后勤业务-1.膳食管理-控制措施01",
        department="膳食部门",
        position=position,
    )

    issues = capability.runner(make_context([matrix, duty]), {})

    assert [issue["evidence"]["issue_type"] for issue in issues] == [
        "responsibility_generic_company",
    ]


def test_control_measure_completeness_reports_shortage_by_key_point() -> None:
    """报送关键控制点的去重措施数少于重庆模板最低数时应提示。"""

    capability = build_registry().get("chongqing_control_measure_completeness", 1)
    assert capability is not None
    first = make_record(
        "matrix",
        3,
        measure_id="财务管理-1.会计主体完整性管控-控制措施01",
        key_control_point="1.会计主体完整性管控",
    )
    duplicate = make_record(
        "matrix",
        4,
        measure_id="财务管理-1.会计主体完整性管控-控制措施01",
        key_control_point="1.会计主体完整性管控",
    )
    baselines = {
        ("financial_management", "default"): {
            "available": True,
            "matrix_rows": [
                {
                    "measure_id": "财务管理-1.会计主体完整性管控-控制措施01",
                    "key_control_point": "1.会计主体完整性管控",
                    "responsibility": "各级单位-财务资产部-总账报表管理",
                },
                {
                    "measure_id": "财务管理-1.会计主体完整性管控-控制措施02",
                    "key_control_point": "1.会计主体完整性管控",
                    "responsibility": "各级单位-财务资产部-资产管理",
                },
                {
                    "measure_id": "财务管理-1.会计主体完整性管控-控制措施03",
                    "key_control_point": "1.会计主体完整性管控",
                    "responsibility": "市公司-财务资产部-主任",
                },
            ],
        },
    }

    issues = capability.runner(
        make_context([first, duplicate], baselines=baselines),
        {},
    )

    assert len(issues) == 1
    assert issues[0]["record"] is first
    assert issues[0]["evidence"] == {
        "issue_type": "control_measure_incomplete",
        "opinion": "请核实该关键控制点下控制措施是否完整",
        "business_id": "financial_management",
        "key_point": [1],
        "expected_count": 2,
        "actual_count": 1,
    }


def test_control_measure_completeness_accepts_expected_measure_count() -> None:
    """报送措施数达到模板中各级单位最低数量时应通过。"""

    capability = build_registry().get("chongqing_control_measure_completeness", 1)
    assert capability is not None
    rows = [
        make_record(
            "matrix",
            row_number,
            measure_id=f"财务管理-1.会计主体完整性管控-控制措施0{index}",
            key_control_point="1.会计主体完整性管控",
        )
        for row_number, index in ((3, 1), (4, 2))
    ]
    baselines = {
        ("financial_management", "default"): {
            "available": True,
            "matrix_rows": [
                {
                    "measure_id": f"财务管理-1.会计主体完整性管控-控制措施0{index}",
                    "key_control_point": "1.会计主体完整性管控",
                    "responsibility": "各级单位-财务资产部-总账报表管理",
                }
                for index in (1, 2)
            ],
        },
    }

    assert capability.runner(make_context(rows, baselines=baselines), {}) == []


def test_control_measure_completeness_counts_each_entity_independently() -> None:
    """多个主体同时审核时，一个主体的完整措施不得掩盖另一个主体的缺项。"""

    capability = build_registry().get("chongqing_control_measure_completeness", 1)
    assert capability is not None
    rows = [
        make_record(
            "matrix",
            row_number,
            measure_id=f"财务管理-1.会计主体完整性管控-控制措施0{index}",
            key_control_point="1.会计主体完整性管控",
        )
        for row_number, index in ((3, 1), (4, 2))
    ]
    incomplete = make_record(
        "matrix",
        5,
        measure_id="财务管理-1.会计主体完整性管控-控制措施01",
        key_control_point="1.会计主体完整性管控",
    )
    incomplete.entity_code = "CQ002"
    baselines = {
        ("financial_management", "default"): {
            "available": True,
            "matrix_rows": [
                {
                    "measure_id": f"财务管理-1.会计主体完整性管控-控制措施0{index}",
                    "key_control_point": "1.会计主体完整性管控",
                    "responsibility": "各级单位-财务资产部-总账报表管理",
                }
                for index in (1, 2)
            ],
        },
    }
    context = make_context([*rows, incomplete], baselines=baselines)
    context.scope_entity_codes = ["CQ001", "CQ002"]

    issues = capability.runner(context, {})

    assert len(issues) == 1
    assert issues[0]["record"] is incomplete
    assert issues[0]["evidence"]["actual_count"] == 1


def test_control_measure_completeness_isolates_business_and_variant() -> None:
    """其他业务或变体的控制措施不得补足当前业务变体的措施数量。"""

    capability = build_registry().get("chongqing_control_measure_completeness", 1)
    assert capability is not None
    current = make_record(
        "matrix",
        3,
        measure_id="财务管理-1.会计主体完整性管控-控制措施01",
        key_control_point="1.会计主体完整性管控",
    )
    other_business = make_record(
        "matrix",
        4,
        business_id="contract_management",
        measure_id="财务管理-1.会计主体完整性管控-控制措施02",
        key_control_point="1.会计主体完整性管控",
    )
    other_variant = make_record(
        "matrix",
        5,
        measure_id="财务管理-1.会计主体完整性管控-控制措施02",
        key_control_point="1.会计主体完整性管控",
    )
    other_variant.variant_id = "special"
    baselines = {
        ("financial_management", "default"): {
            "available": True,
            "matrix_rows": [
                {
                    "measure_id": f"财务管理-1.会计主体完整性管控-控制措施0{index}",
                    "key_control_point": "1.会计主体完整性管控",
                    "responsibility": "各级单位-财务资产部-总账报表管理",
                }
                for index in (1, 2)
            ],
        },
    }

    issues = capability.runner(
        make_context([current, other_business, other_variant], baselines=baselines),
        {},
    )

    assert len(issues) == 1
    assert issues[0]["record"] is current
    assert issues[0]["evidence"]["actual_count"] == 1


def test_system_control_consistency_only_reports_missing_numeric_measure_key() -> None:
    """参数说明：数字措施键存在时忽略建设状态和系统名称，仅报告缺失键。"""

    capability = build_registry().get("system_control_consistency", 1)
    assert capability is not None
    matrix_missing = make_record(
        "matrix", 3, measure_id="1.1", control_mode="系统已建", control_system="ERP",
    )
    matrix_mismatch = make_record(
        "matrix", 4, measure_id="2.1", control_mode="系统已建", control_system="ERP",
    )
    system_mismatch = make_record(
        "system_rule", 8, measure_id="2.1", build_status="系统待建", system_name="PMS",
    )

    issues = capability.runner(make_context([matrix_missing, matrix_mismatch, system_mismatch]), {})

    assert [issue["evidence"]["issue_type"] for issue in issues] == [
        "system_rule_missing",
    ]


def test_system_control_consistency_matches_measure_ids_by_numbers_only() -> None:
    """参数说明：数字键相同即建立对应关系，控制方式和其他编号文字不参与比较。"""

    capability = build_registry().get("system_control_consistency", 1)
    assert capability is not None
    matrix = make_record(
        "matrix",
        3,
        measure_id="营销售电业务-2.客户用电容量确定-控制措施02",
        control_mode="线下管控",
        control_system="营销2.0系统",
    )
    system = make_record(
        "system_rule",
        8,
        measure_id="其他编号文字-2.不同名称-控制措施02",
        build_status="系统待建",
        system_name="营销2.0",
    )

    assert capability.runner(make_context([matrix, system]), {}) == []


def test_system_control_consistency_accepts_matching_numeric_measure_key() -> None:
    """参数说明：矩阵与系统规则清单存在相同数字措施键时不应报错。"""

    capability = build_registry().get("system_control_consistency", 1)
    matrix = make_record(
        "matrix", 3, measure_id="1.1", control_mode="系统已建", control_system="ERP、PMS",
    )
    system = make_record(
        "system_rule", 8, measure_id="1.1", build_status="系统已建", system_name="PMS,ERP",
    )

    assert capability.runner(make_context([matrix, system]), {}) == []


def test_system_control_consistency_accepts_multiple_rows_with_same_numeric_key() -> None:
    """参数说明：系统清单同一数字措施键存在多行时均视为已建立对应关系。"""

    capability = build_registry().get("system_control_consistency", 1)
    matrix = make_record(
        "matrix", 3, measure_id="1.1", control_mode="系统已建", control_system="ERP、PMS",
    )
    system_erp = make_record(
        "system_rule", 8, measure_id="1.1", build_status="系统已建", system_name="ERP",
    )
    system_pms = make_record(
        "system_rule", 9, measure_id="1.1", build_status="系统已建", system_name="PMS",
    )

    assert capability.runner(make_context([matrix, system_erp, system_pms]), {}) == []


def test_branch_key_points_reports_missing_and_extra_points() -> None:
    """供电分公司的关键控制点缺失或多报时应分别提示。"""

    capability = build_registry().get("power_supply_branch_key_points", 1)
    assert capability is not None
    actual = [
        make_record("matrix", 3, measure_id="29.1"),
        make_record("matrix", 4, measure_id="99.1"),
    ]
    params = {
        "expected_key_points": {"financial_management": [[29], [45]]},
        "business_aliases": {"financial_management": ["财务管理"]},
    }

    issues = capability.runner(make_context(actual), params)

    assert {issue["evidence"]["issue_type"] for issue in issues} == {
        "missing_key_points",
        "extra_key_point",
    }
    missing = next(issue for issue in issues if issue["evidence"]["issue_type"] == "missing_key_points")
    assert missing["evidence"]["missing_key_points"] == [[45]]


def test_branch_key_points_only_applies_to_power_supply_branch() -> None:
    """非供电分公司不执行该专项模板规则。"""

    capability = build_registry().get("power_supply_branch_key_points", 1)
    params = {
        "expected_key_points": {"financial_management": [[29]]},
        "business_aliases": {"financial_management": ["财务管理"]},
    }

    issues = capability.runner(
        make_context([], entity_name="国网重庆市电力公司本部"),
        params,
    )

    assert issues == []


def test_branch_key_points_uses_directory_when_unmatched_filename_conflicts() -> None:
    """未匹配模板时只能按业务目录归一，文件名不得造成冲突。"""

    capability = build_registry().get("power_supply_branch_key_points", 1)
    record = make_record("matrix", 3, business_id="unmatched_1234", measure_id="29.1")
    record.file_path = "11.财务管理矩阵及三清单/30物业后勤风控矩阵.xlsx"
    params = {
        "expected_key_points": {"financial_management": [[29]]},
        "business_aliases": {
            "financial_management": ["财务管理"],
            "property_logistics": ["物业后勤"],
        },
    }

    context = make_context([record], business_id="unmatched_1234")
    context.resources.update({
        "_entity_business_ids": ["unmatched_1234"],
        "_entity_business_files": {
            "unmatched_1234": ["11.财务管理矩阵及三清单/30物业后勤风控矩阵.xlsx"],
        },
        "_current_business_id": "unmatched_1234",
    })

    assert capability.runner(context, params) == []


def test_branch_key_points_normalizes_unmatched_source_business() -> None:
    """来源新业务未进入基准时应通过目录别名归一，不得同时报缺失和多报。"""

    capability = build_registry().get("power_supply_branch_key_points", 1)
    record = make_record("matrix", 3, business_id="unmatched_abcd", measure_id="5.1")
    record.file_path = "13.特高压管理矩阵及三清单/财务管理风控矩阵.xlsx"
    context = make_context([record], business_id="unmatched_abcd")
    context.resources.update({
        "_entity_business_ids": ["unmatched_abcd"],
        "_entity_business_files": {
            "unmatched_abcd": ["13.特高压管理矩阵及三清单/财务管理风控矩阵.xlsx"],
        },
        "_current_business_id": "unmatched_abcd",
    })
    params = {
        "expected_key_points": {"ultra_high_voltage_construction": [[5]]},
        "business_aliases": {"ultra_high_voltage_construction": ["特高压管理"]},
    }

    assert capability.runner(context, params) == []


def test_branch_key_points_keeps_files_in_same_directory_as_one_business() -> None:
    """同一业务目录中的文件名不得把审核单元拆成其他业务。"""

    capability = build_registry().get("power_supply_branch_key_points", 1)
    ultra_high_voltage = make_record(
        "matrix", 3, business_id="unmatched_same", measure_id="5.1",
    )
    ultra_high_voltage.file_path = "13.特高压管理矩阵及三清单/特高压建设风控矩阵.xlsx"
    second_matrix = make_record(
        "matrix", 4, business_id="unmatched_same", measure_id="7.1",
    )
    second_matrix.file_path = "13.特高压管理矩阵及三清单/数科平台业务风控矩阵.xlsx"
    context = make_context(
        [ultra_high_voltage, second_matrix],
        business_id="unmatched_same",
    )
    context.resources.update({
        "_entity_business_ids": ["unmatched_same"],
        "_entity_business_files": {
            "unmatched_same": [
                "13.特高压管理矩阵及三清单/特高压建设风控矩阵.xlsx",
                "13.特高压管理矩阵及三清单/数科平台业务风控矩阵.xlsx",
            ],
        },
        "_current_business_id": "unmatched_same",
    })
    params = {
        "expected_key_points": {
            "ultra_high_voltage_construction": [[5], [7]],
        },
        "business_aliases": {
            "ultra_high_voltage_construction": ["特高压管理"],
            "digital_platform_business": ["数科平台业务"],
        },
    }

    assert capability.runner(context, params) == []


def test_branch_key_points_reports_all_businesses_for_entity_without_business_files() -> None:
    """供电分公司只有说明材料时应报告全部来源业务缺失。"""

    capability = build_registry().get("power_supply_branch_key_points", 1)
    context = make_context([], business_id="__entity_completeness__")
    context.resources.update({
        "_entity_business_ids": [],
        "_entity_business_files": {},
        "_current_business_id": "__entity_completeness__",
        "_entity_completeness_pass": True,
    })
    params = {
        "expected_key_points": {
            "financial_management": [[29]],
            "marketing_management": [[1], [2]],
        },
        "business_aliases": {
            "financial_management": ["财务管理"],
            "marketing_management": ["营销管理"],
        },
    }

    issues = capability.runner(context, params)

    assert {issue["evidence"]["business_id"] for issue in issues} == {
        "financial_management",
        "marketing_management",
    }


def test_branch_key_points_only_checks_current_business_unit() -> None:
    """单业务审核单元不应重复报其他业务的缺失。"""

    capability = build_registry().get("power_supply_branch_key_points", 1)
    params = {
        "expected_key_points": {
            "financial_management": [[29]],
            "marketing_management": [[1]],
        },
        "business_aliases": {
            "financial_management": ["财务管理"],
            "marketing_management": ["营销管理"],
        },
    }

    issues = capability.runner(make_context([]), params)

    assert len(issues) == 1
    assert issues[0]["evidence"]["business_id"] == "financial_management"


def test_branch_key_points_reports_wholly_missing_business_once_per_entity() -> None:
    """供电分公司整项未报业务时应在主体首个单元提示完整性。"""

    capability = build_registry().get("power_supply_branch_key_points", 1)
    actual = [make_record("matrix", 3, measure_id="29.1")]
    context = make_context(actual)
    context.resources.update({
        "_entity_business_ids": ["financial_management"],
        "_current_business_id": "financial_management",
        "_entity_completeness_pass": True,
    })
    params = {
        "expected_key_points": {
            "financial_management": [[29]],
            "marketing_management": [[1], [2]],
        },
        "business_aliases": {
            "financial_management": ["财务管理"],
            "marketing_management": ["营销管理"],
        },
    }

    issues = capability.runner(context, params)

    assert len(issues) == 1
    assert issues[0]["evidence"]["business_id"] == "marketing_management"
    assert issues[0]["evidence"]["missing_key_points"] == [[1], [2]]


def test_branch_key_points_reports_extra_business_matrix() -> None:
    """供电分公司提交不在来源 17 项中的整张矩阵时应提示适用性。"""

    capability = build_registry().get("power_supply_branch_key_points", 1)
    record = make_record("matrix", 3, business_id="education_training", measure_id="1.1")
    context = make_context([record], business_id="education_training")
    context.resources.update({
        "_entity_business_ids": ["education_training"],
        "_current_business_id": "education_training",
        "_entity_completeness_pass": True,
    })
    params = {
        "expected_key_points": {"financial_management": [[29]]},
        "business_aliases": {"financial_management": ["财务管理"]},
    }

    issues = capability.runner(context, params)

    assert {issue["evidence"]["issue_type"] for issue in issues} == {
        "extra_business",
        "missing_key_points",
    }
    extra = next(issue for issue in issues if issue["evidence"]["issue_type"] == "extra_business")
    assert extra["record"] is record


def test_branch_key_points_does_not_repeat_entity_missing_on_second_variant() -> None:
    """同一业务的第二个变体不得重复输出主体级整项缺报。"""

    capability = build_registry().get("power_supply_branch_key_points", 1)
    record = make_record("matrix", 3, business_id="grid_construction", measure_id="1.1")
    context = make_context([record], business_id="grid_construction")
    context.resources.update({
        "_entity_business_ids": ["grid_construction"],
        "_entity_business_files": {"grid_construction": ["电网基建风控矩阵.xlsx"]},
        "_current_business_id": "grid_construction",
        "_entity_completeness_pass": False,
    })
    params = {
        "expected_key_points": {
            "grid_construction": [[1]],
            "marketing_management": [[2]],
        },
        "business_aliases": {
            "grid_construction": ["电网基建"],
            "marketing_management": ["营销管理"],
        },
    }

    assert capability.runner(context, params) == []


def _modified_document(tmp_path: Path) -> Path:
    """复制来源文档以便构造发布失败样例。

    Args:
        tmp_path: pytest 临时目录。
    """

    target = tmp_path / SOURCE_NAME
    shutil.copyfile(SOURCE, target)
    return target


def test_source_extracts_all_confirmed_businesses_and_points() -> None:
    """来源文档必须解析出 17 个已确认业务和 368 个关键控制点。"""

    points = extract_key_points(SOURCE, BUSINESS_MAPPING)

    assert set(points) == set(BUSINESS_MAPPING.values())
    assert sum(len(items) for items in points.values()) == 368


def test_source_rejects_unknown_business(tmp_path: Path) -> None:
    """来源文档出现未确认业务时必须终止发布。"""

    source = _modified_document(tmp_path)
    document = Document(source)
    cells = document.tables[0].add_row().cells
    cells[0].text = "未知业务"
    cells[1].text = "1.测试"
    document.save(source)

    with pytest.raises(RuntimeError, match="未知业务"):
        extract_key_points(source, BUSINESS_MAPPING)


def test_source_rejects_missing_business(tmp_path: Path) -> None:
    """来源文档缺少任一已确认业务时必须终止发布。"""

    source = _modified_document(tmp_path)
    document = Document(source)
    table = document.tables[0]
    for row in list(table.rows)[2:]:
        if row.cells[0].text.strip() == EXPECTED_BUSINESS_NAMES[0]:
            table._tbl.remove(row._tr)
    document.save(source)

    with pytest.raises(RuntimeError, match="缺少业务"):
        extract_key_points(source, BUSINESS_MAPPING)


def test_source_rejects_multiple_tables(tmp_path: Path) -> None:
    """来源文档出现多张规则表时必须终止发布。"""

    source = _modified_document(tmp_path)
    document = Document(source)
    document.add_table(rows=1, cols=2)
    document.save(source)

    with pytest.raises(RuntimeError, match="1 张"):
        extract_key_points(source, BUSINESS_MAPPING)


def test_source_rejects_illegal_key_point_number(tmp_path: Path) -> None:
    """关键控制点编号无法唯一解析时必须终止发布。"""

    source = _modified_document(tmp_path)
    document = Document(source)
    document.tables[0].rows[2].cells[1].text = "没有编号"
    document.save(source)

    with pytest.raises(RuntimeError, match="关键控制点编号"):
        extract_key_points(source, BUSINESS_MAPPING)


def test_chongqing_system_owner_is_second_level() -> None:
    """重庆市公司系统规则管理主体应归为二级部署系统。"""

    assert classify_system_owner("国网重庆市电力有限公司") == "二级部署系统"
    assert classify_system_owner("国网重庆市电力公司") == "二级部署系统"


def test_exact_chongqing_city_company_is_second_level() -> None:
    """省级别名“重庆市公司”不应被“市公司”关键词排除。"""

    assert classify_system_owner("重庆市公司") == "二级部署系统"
    assert classify_system_owner("国网重庆市电力公司某供电分公司") is None


def test_release_is_chongqing_specific_and_excludes_hunan_rule_17() -> None:
    """1.9.22 应标识重庆地区，且不得继承湖南财务管理第 17 条。"""

    manifest = json.loads((RELEASE / "manifest.json").read_text(encoding="utf-8"))
    rules = [json.loads(path.read_text(encoding="utf-8")) for path in (RELEASE / "rules").glob("*.json")]

    assert BASE_VERSION == "1.9.20"
    assert manifest["version"] == "1.9.22"
    assert manifest["region_code"] == "chongqing"
    assert manifest["region_name"] == "重庆"
    assert {rule["rule_id"] for rule in rules}.issuperset({
        "chongqing.system_control_consistency",
        "chongqing.power_supply_branch_key_points",
    })
    assert "financial_management.key_measure_applicability" not in {rule["rule_id"] for rule in rules}


def test_release_accepts_combined_and_split_three_list_names() -> None:
    """重庆命名规则应同时接受合并三清单和三个独立清单的模糊名称。"""
    rules = [json.loads(path.read_text(encoding="utf-8")) for path in (SPLIT_LISTS_RELEASE / "rules").glob("*.json")]
    completeness = next(rule for rule in rules if rule["rule_id"] == "documents.completeness")
    naming_check = next(check for check in completeness["checks"] if check["check_id"] == "file_naming")

    assert naming_check["params"]["naming_keywords"] == {
        "matrix": ["风控矩阵", "矩阵", "取证"],
        "three_lists": ["三清单", "系统规则", "岗位内控", "不相容"],
    }


def test_current_release_covers_actual_chongqing_business_directories() -> None:
    """当前重庆规则包必须用实际业务目录名完成基准和专项规则映射。"""
    pack = load_pack(CURRENT_RELEASE)
    branch_rule = next(
        rule
        for rule in pack["rules"]
        if rule["rule_id"] == "chongqing.power_supply_branch_key_points"
    )
    aliases = branch_rule["checks"][0]["params"]["business_aliases"]
    sales = next(
        business
        for business in pack["baseline_registry"]["businesses"]
        if business["business_id"] == "sales_electricity"
    )

    assert "营销与售电" in sales["aliases"]
    assert "营销与售电" in aliases["sales_electricity"]
    assert "特高压管理" in aliases["ultra_high_voltage_construction"]
    assert "检修运维成本项目管理" in aliases["maintenance_cost_project_management"]


def test_updated_release_freezes_new_rules_and_all_chongqing_templates() -> None:
    """1.9.26 必须冻结新来源、新增两条规则并登记全部重庆模板。"""

    pack = load_pack(UPDATED_RELEASE)
    rule_ids = {rule["rule_id"] for rule in pack["rules"]}
    templates = {
        variant["template"]["path"]: variant["template"]["sha256"]
        for business in pack["baseline_registry"]["businesses"]
        for variant in business["variants"]
    }
    template_root = ROOT / "baselines/chongqing/templates"
    actual_templates = {
        str(path.relative_to(template_root))
        for path in template_root.rglob("*")
        if path.suffix.lower() in {".xls", ".xlsx"}
    }

    assert pack["manifest"]["version"] == "1.9.26"
    assert pack["manifest"]["source_document_sha256"] == sha256_file(UPDATED_SOURCE)
    assert rule_ids.issuperset({
        "chongqing.responsibility_alignment",
        "chongqing.control_measure_completeness",
    })
    assert set(templates) == actual_templates
    assert len(templates) == 29
    assert all(sha256_file(template_root / path) == digest for path, digest in templates.items())


def test_updated_release_identifies_every_chongqing_business_directory_uniquely() -> None:
    """29 个重庆业务目录必须分别匹配唯一稳定业务标识。"""

    pack = load_pack(UPDATED_RELEASE)
    for business in pack["baseline_registry"]["businesses"]:
        template_path = business["variants"][0]["template"]["path"]
        identity = identify_business(Path(template_path), pack["baseline_registry"])

        assert identity.business_id == business["business_id"]


def test_current_release_uses_confirmed_scan_and_preprocessing_pipeline() -> None:
    """当前重庆发布包必须进入新版扫描和基准预处理流程。"""
    assert "1.9.27" in CONFIRMED_RULE_VERSIONS
    assert "1.9.27" in PREPROCESSING_RULE_VERSIONS
    assert "1.9.27" in SYSTEM_TYPE_PREPROCESSING_RULE_VERSIONS


@pytest.mark.parametrize(("headers", "values"), [
    (
        ["控制措施编号", "控制措施", "控制系统", "控制载体", "责任主体", "业务备注"],
        ["1.1", "测试措施", "ERP", "制度", "重庆公司-财务部-会计", "正常"],
    ),
    (
        [
            "控制措施编号",
            "控制措施",
            "控制系统",
            "控制载体",
            "责任主体",
            "是否适用及原因",
            "业务备注",
        ],
        ["1.1", "测试措施", "ERP", "制度", "重庆公司-财务部-会计", "/", "正常"],
    ),
])
def test_current_release_does_not_enforce_deleted_applicability_requirements(
    tmp_path: Path,
    headers: list[str],
    values: list[str],
) -> None:
    """重庆新版第3条不得再检查适用性字段存在、内容或末列位置。

    Args:
        tmp_path: pytest 提供的临时目录。
        headers: 待审核矩阵的业务表头。
        values: 与业务表头对应的一行测试数据。
    """

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "风控矩阵"
    sheet.append(["风控矩阵"])
    sheet.append(headers)
    sheet.append(values)
    source = tmp_path / "02财务管理矩阵.xlsx"
    workbook.save(source)

    pack = copy.deepcopy(load_pack(CURRENT_RELEASE))
    pack["rules"] = [rule for rule in pack["rules"] if rule["display_code"] == "R03"]
    pack["submission_scope"] = {
        "batch": "second",
        "entity_codes": ["CQ001"],
        "businesses": [{
            "business_id": "financial_management",
            "business_code": "02",
            "variant_id": "default",
            "required": True,
        }],
    }
    file = FileRecord(
        source,
        Path(source.name),
        sha256_file(source),
        "xlsx",
        "CQ001",
        [],
        False,
        "02",
        "default",
        "matrix",
        business_id="financial_management",
        preservation={"template_match": {"status": "matched"}},
    )
    file.sheets = parse_workbook(file, source, pack["field_aliases"])

    findings, _ = run_engine(
        pack,
        build_registry(),
        [file],
        {"CQ001": Entity("CQ001", "国网重庆市电力公司")},
        {},
        run_id="chongqing-applicability-revision-test",
    )

    assert findings == []


def test_current_release_freezes_revised_source_and_removes_applicability_rule() -> None:
    """重庆 1.9.27 必须冻结新稿，并移除已删除的适用性填写子规则。"""

    pack = load_pack(CURRENT_RELEASE)
    rules = {rule["rule_id"]: rule for rule in pack["rules"]}
    schema_params = rules["schema.applicability"]["checks"][0]["params"]

    assert pack["manifest"]["version"] == "1.9.27"
    assert pack["manifest"]["source_document_sha256"] == sha256_file(REVISED_SOURCE)
    assert "applicability.values" not in rules
    assert "applicability" not in schema_params["required_fields"]
    assert "check_applicability_tail" not in schema_params
    assert schema_params["compare_baseline"] is True
    assert "matrices.measure_applicability" in rules


def test_workbook_title_fallback_satisfies_chongqing_material_naming() -> None:
    """工作簿标题已确认材料类型时，重庆不应再报文件名缺少关键字。"""
    rules = [json.loads(path.read_text(encoding="utf-8")) for path in (SPLIT_LISTS_RELEASE / "rules").glob("*.json")]
    completeness = next(rule for rule in rules if rule["rule_id"] == "documents.completeness")
    naming_check = next(check for check in completeness["checks"] if check["check_id"] == "file_naming")
    file = FileRecord(
        Path("北碚-违章罚款收入管理.xlsx"),
        Path("北碚-违章罚款收入管理.xlsx"),
        "digest",
        "xlsx",
        "CQ001",
        [],
        False,
        "17",
        "default",
        "matrix",
        business_id="violation_fine_income_management",
    )
    file.preservation["material_match"] = {
        "source": "workbook",
        "material_type": "matrix",
    }
    context = CheckContext(
        [],
        [file],
        [],
        {"CQ001": Entity("CQ001", "测试单位")},
        {},
        {"entity_aliases": {}},
        ["CQ001"],
        [{"business_id": "violation_fine_income_management", "business_code": "17", "variant_id": "default", "required": True}],
    )
    capability = build_registry().get("required_documents", 4)

    assert capability is not None
    assert capability.runner(context, naming_check["params"]) == []


@pytest.mark.parametrize(
    "release",
    [RELEASE, SPLIT_LISTS_RELEASE, UPDATED_RELEASE, CURRENT_RELEASE],
)
def test_release_contains_no_hunan_specific_resources(release: Path) -> None:
    """重庆发布包不得包含湖南主体、地名或冻结样本引用。

    Args:
        release: 待检查的重庆历史或当前发布包。
    """

    markers = ("国网湖南", "湖南省电力", "湘西", "星通公司")
    leaks = {
        str(path.relative_to(release)): marker
        for path in release.rglob("*.json")
        for marker in markers
        if marker in path.read_text(encoding="utf-8")
    }

    assert leaks == {}


def test_repository_region_index_only_pins_chongqing_release() -> None:
    """重庆分支地区索引只能固定重庆发布包。"""

    index = json.loads((ROOT / "regions/index.json").read_text(encoding="utf-8"))
    assert [item["region_code"] for item in index["regions"]] == ["chongqing"]
    profile = index["regions"][0]
    manifest = json.loads(
        (ROOT / "rulepacks/releases/1.9.29/manifest.json").read_text(encoding="utf-8")
    )
    assert profile["rulepack_version"] == "1.9.29"
    assert profile["rulepack_content_hash"] == manifest["content_hash"]


@pytest.mark.parametrize(("rule_id", "rule_number", "record_values"), [
    (
        "chongqing.system_control_consistency",
        16,
        {"measure_id": "29.1", "control_mode": "系统已建"},
    ),
    (
        "chongqing.power_supply_branch_key_points",
        17,
        {"measure_id": "29.1"},
    ),
    (
        "chongqing.responsibility_alignment",
        18,
        {
            "measure_id": "财务管理-1.会计主体完整性管控-控制措施01",
            "responsibility": "各级单位-财务资产部-总账报表管理",
        },
    ),
    (
        "chongqing.control_measure_completeness",
        19,
        {
            "measure_id": "财务管理-1.会计主体完整性管控-控制措施01",
            "key_control_point": "1.会计主体完整性管控",
        },
    ),
])
def test_current_chongqing_special_opinion_includes_rule_number(
    rule_id: str,
    rule_number: int,
    record_values: dict[str, str],
) -> None:
    """当前重庆专项规则命中时，最终审核意见必须显示规则编号。

    Args:
        rule_id: 单独执行的重庆专项规则标识。
        rule_number: 意见前必须显示的规则条款编号。
        record_values: 用于稳定触发当前规则的矩阵字段。
    """

    index = json.loads((ROOT / "regions/index.json").read_text(encoding="utf-8"))
    chongqing = next(
        item for item in index["regions"] if item["region_code"] == "chongqing"
    )
    release = ROOT / "rulepacks/releases" / chongqing["rulepack_version"]
    pack = copy.deepcopy(load_pack(release))
    pack["rules"] = [
        rule
        for rule in pack["rules"]
        if rule["rule_id"] == rule_id
    ]
    pack["submission_scope"] = {
        "batch": "second",
        "entity_codes": ["CQ001"],
        "businesses": [{
            "business_id": "financial_management",
            "business_code": "02",
            "variant_id": "default",
            "required": True,
        }],
    }
    record = make_record("matrix", 3, **record_values)
    sheet = ParsedSheet("风控矩阵", "matrix", [2], {}, {}, 8, 3, [record])
    file = FileRecord(
        Path("财务管理.xlsx"),
        Path("财务管理.xlsx"),
        "hash",
        "xlsx",
        "CQ001",
        [],
        False,
        "02",
        "default",
        "matrix",
        [sheet],
        business_id="financial_management",
    )

    baselines = {}
    if rule_id == "chongqing.control_measure_completeness":
        baselines = {
            ("financial_management", "default"): {
                "available": True,
                "matrix_rows": [
                    {
                        "measure_id": f"财务管理-1.会计主体完整性管控-控制措施0{index}",
                        "key_control_point": "1.会计主体完整性管控",
                        "responsibility": "各级单位-财务资产部-总账报表管理",
                    }
                    for index in (1, 2)
                ],
            },
        }

    findings, _ = run_engine(
        pack,
        build_registry(),
        [file],
        {"CQ001": Entity("CQ001", "国网重庆市电力公司某供电分公司")},
        baselines,
        run_id="chongqing-rule-17-number-test",
    )

    assert findings
    expected_prefix = f"【第{rule_number}条】"
    assert all(finding.message.startswith(expected_prefix) for finding in findings)


def test_verify_only_is_deterministic_and_does_not_activate_release() -> None:
    """重庆发布包必须可重复构建，验证时不得改动全局激活指针。"""

    active_path = ROOT / "rulepacks/active.json"
    before = active_path.read_bytes()
    result = subprocess.run(
        [sys.executable, str(ROOT / "tools/publish_rules_1012_chongqing.py"), "--verify-only"],
        cwd=ROOT,
        text=True,
        capture_output=True,
        check=False,
    )

    assert result.returncode == 0, result.stderr
    assert json.loads(result.stdout)["verified"] is True
    assert active_path.read_bytes() == before


def test_updated_release_is_deterministic_and_does_not_activate_release() -> None:
    """重庆 1.9.26 必须可重复构建，复验时不得改动全局激活指针。"""

    active_path = ROOT / "rulepacks/active.json"
    before = active_path.read_bytes()
    result = subprocess.run(
        [
            sys.executable,
            str(ROOT / "tools/publish_rules_1012_chongqing_1926.py"),
            "--verify-only",
        ],
        cwd=ROOT,
        text=True,
        capture_output=True,
        check=False,
    )

    assert result.returncode == 0, result.stderr
    assert json.loads(result.stdout)["verified"] is True
    assert active_path.read_bytes() == before


def test_revised_release_is_deterministic_and_does_not_activate_release() -> None:
    """重庆 1.9.27 必须可重复构建，复验时不得改动全局激活指针。"""

    active_path = ROOT / "rulepacks/active.json"
    before = active_path.read_bytes()
    result = subprocess.run(
        [
            sys.executable,
            str(ROOT / "tools/publish_rules_1012_chongqing_1927.py"),
            "--verify-only",
        ],
        cwd=ROOT,
        text=True,
        capture_output=True,
        check=False,
    )

    assert result.returncode == 0, result.stderr
    assert json.loads(result.stdout)["verified"] is True
    assert active_path.read_bytes() == before


def test_split_lists_release_is_deterministic_and_does_not_activate_release() -> None:
    """分拆三清单规则包必须可重建，复验时不得改动全局激活指针。"""

    active_path = ROOT / "rulepacks/active.json"
    before = active_path.read_bytes()
    result = subprocess.run(
        [sys.executable, str(ROOT / "tools/publish_chongqing_split_lists_1923.py"), "--verify-only"],
        cwd=ROOT,
        text=True,
        capture_output=True,
        check=False,
    )

    assert result.returncode == 0, result.stderr
    assert json.loads(result.stdout)["verified"] is True
    assert active_path.read_bytes() == before


def test_rule_number_release_is_deterministic_and_does_not_activate_release() -> None:
    """规则编号修复包必须可重建，复验时不得改动全局激活指针。"""

    active_path = ROOT / "rulepacks/active.json"
    before = active_path.read_bytes()
    result = subprocess.run(
        [
            sys.executable,
            str(ROOT / "tools/publish_chongqing_rule_numbers_1924.py"),
            "--verify-only",
        ],
        cwd=ROOT,
        text=True,
        capture_output=True,
        check=False,
    )

    assert result.returncode == 0, result.stderr
    assert json.loads(result.stdout)["verified"] is True
    assert active_path.read_bytes() == before
