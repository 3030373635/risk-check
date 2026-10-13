"""验证 1013 重庆规则同步。"""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

from risk_audit.checks.registry import CheckContext, build_registry
from risk_audit.configuration.loader import load_pack
from risk_audit.models import Entity, FieldValue, FileRecord, Record
from risk_audit.output_0916 import EXISTING_HEADER, enrich_matrices
from risk_audit.pipeline_versions import SYSTEM_TYPE_PREPROCESSING_RULE_VERSIONS
from risk_audit.readers.confirmed_v180 import parse_workbook_v180
from risk_audit.runner import CONFIRMED_RULE_VERSIONS, PREPROCESSING_RULE_VERSIONS
from risk_audit.util import sha256_file


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "规则来源/风控矩阵1013-重庆.docx"
RELEASE = ROOT / "rulepacks/releases/1.9.29"
UPDATED_SOURCE = ROOT / "规则来源/风控矩阵-1013-重庆.docx"
UPDATED_RELEASE = ROOT / "rulepacks/releases/1.9.30"


def make_record(record_type: str, row: int, **values: str) -> Record:
    """构造编号引用测试记录。

    Args:
        record_type: 审核记录类型。
        row: 原始工作表行号。
        values: 按逻辑字段名给出的字段值。

    Returns:
        使用同一主体、业务和变体边界的审核记录。
    """

    fields = {
        name: FieldValue(value, value, f"A{row}")
        for name, value in values.items()
    }
    return Record(
        record_type,
        "CQ001",
        "02",
        "default",
        "财务管理.xlsx",
        record_type,
        row,
        fields,
        f"{record_type}-{row}",
        business_id="financial_management",
    )


def make_context(records: list[Record]) -> CheckContext:
    """构造重庆岗位责任清单审核上下文。

    Args:
        records: 本次规则检查使用的岗位责任清单记录。

    Returns:
        包含固定重庆主体和财务管理业务范围的审核上下文。
    """

    return CheckContext(
        records,
        [],
        records,
        {"CQ001": Entity("CQ001", "国网重庆市电力公司")},
        {},
        {},
        ["CQ001"],
        [{
            "business_id": "financial_management",
            "business_code": "02",
            "variant_id": "default",
        }],
    )


def updated_person_name_check() -> tuple[dict, object]:
    """读取重庆 1.9.30 第20条及其审核能力。

    Returns:
        第20条检查配置和已注册的审核能力。
    """

    pack = load_pack(UPDATED_RELEASE)
    rules = {rule["rule_id"]: rule for rule in pack["rules"]}
    check = rules["chongqing.person_name_single"]["checks"][0]
    capability = build_registry().get(check["operator"], check["operator_version"])
    assert capability is not None
    return check, capability


def test_person_name_single_rule_rejects_name_separators() -> None:
    """第20条应将顿号、中英文逗号识别为同一行填写多个人员。"""

    check, capability = updated_person_name_check()

    for row, person_names in enumerate(("张三、李四", "张三，李四", "张三,李四"), start=8):
        record = make_record("position_duty", row, person_names=person_names)
        issues = capability.runner(make_context([record]), check["params"])

        assert len(issues) == 1
        assert issues[0]["kind"] == "violation"
        assert issues[0]["record"] is record

    expected_message = "【第20条】每行岗位名称中只能有一个对应人员，若有多人，则应分多行填写。"
    assert check["message"] == {
        "review": expected_message,
        "violation": expected_message,
    }


def test_person_name_single_rule_accepts_one_person() -> None:
    """第20条不得对单个人员姓名产生审核意见。"""

    check, capability = updated_person_name_check()
    record = make_record("position_duty", 8, person_names="张三")

    assert capability.runner(make_context([record]), check["params"]) == []


def test_reference_exists_v3_ignores_matrix_applicability() -> None:
    """岗位编号已在当前矩阵中存在时，不得再依据适用性产生第9条意见。"""

    capability = build_registry().get("reference_exists", 3)
    assert capability is not None
    duty = make_record("position_duty", 8, measure_id="财务管理-1-控制措施01")
    matrix = make_record(
        "matrix",
        3,
        measure_id="财务管理-1-控制措施01",
        applicability="不适用",
    )
    context = CheckContext(
        [duty],
        [],
        [duty, matrix],
        {"CQ001": Entity("CQ001", "国网重庆市电力公司")},
        {},
        {},
        ["CQ001"],
        [{
            "business_id": "financial_management",
            "business_code": "02",
            "variant_id": "default",
        }],
    )

    issues = capability.runner(
        context,
        {
            "source_field": "measure_id",
            "target_field": "measure_id",
            "target_record_type": "matrix",
        },
    )

    assert issues == []


def test_matrix_enrichment_removes_existing_measure_column(tmp_path: Path) -> None:
    """1013 输出不得新增或保留岗位清单已有措施编号列。

    Args:
        tmp_path: pytest 提供的临时目录。
    """

    path = tmp_path / "02财务管理矩阵.xlsx"
    workbook = Workbook()
    matrix = workbook.active
    matrix.title = "风控矩阵"
    matrix.append([
        "控制措施编号",
        "控制措施",
        "责任主体",
        "是否适用",
        "省公司版本责任主体（核对后删除）",
        EXISTING_HEADER,
        "审核意见",
    ])
    matrix.append([
        "财务管理-1-控制措施01",
        "复核资料",
        "重庆公司-财务部-会计",
        "不适用",
        "旧责任主体",
        "旧措施编号",
        "人工意见",
    ])
    workbook.save(path)

    file = FileRecord(
        path,
        Path(path.name),
        sha256_file(path),
        "xlsx",
        "CQ001",
        [],
        False,
        "02",
        "default",
        "matrix",
        business_id="financial_management",
        preservation={"template_match": {"status": "matched"}},
    )
    aliases = load_pack(ROOT / "rulepacks/releases/1.9.27")["field_aliases"]
    file._parser_policy = {"version": 3}
    file.sheets = parse_workbook_v180(file, path, aliases)

    enrich_matrices(
        file,
        path,
        {},
        [],
        [file],
        include_existing_measure_column=False,
    )

    output = load_workbook(path).active
    headers = [output.cell(1, column).value for column in range(1, output.max_column + 1)]
    assert EXISTING_HEADER not in headers
    assert "旧措施编号" not in [
        output.cell(2, column).value
        for column in range(1, output.max_column + 1)
    ]
    assert "人工意见" in [
        output.cell(2, column).value
        for column in range(1, output.max_column + 1)
    ]


def test_1013_release_is_rebuildable() -> None:
    """重庆 1.9.29 应冻结 1013 来源、重排删除项后的编号并可确定性复验。"""

    assert SOURCE.is_file()
    pack = load_pack(RELEASE)
    rules = {rule["rule_id"]: rule for rule in pack["rules"]}
    reference_check = rules["references.current_matrix"]["checks"][0]
    responsibility_rule = rules["chongqing.responsibility_alignment"]
    completeness_rule = rules["chongqing.control_measure_completeness"]

    assert pack["manifest"]["version"] == "1.9.29"
    assert pack["manifest"]["source_document_sha256"] == sha256_file(SOURCE)
    assert "matrices.measure_applicability" not in rules
    assert reference_check["operator_version"] == 3
    assert "适用性" not in reference_check["message"]["violation"]
    assert responsibility_rule["display_code"] == "R18"
    assert responsibility_rule["checks"][0]["message"]["violation"] == "【第18条】{opinion}"
    assert completeness_rule["display_code"] == "R19"
    assert completeness_rule["checks"][0]["message"]["violation"] == "【第19条】{opinion}"
    result = subprocess.run(
        [
            sys.executable,
            str(ROOT / "tools/publish_chongqing_rule_numbers_1929.py"),
            "--verify-only",
        ],
        cwd=ROOT,
        text=True,
        capture_output=True,
        check=False,
    )

    assert result.returncode == 0, result.stderr
    assert json.loads(result.stdout)["verified"] is True


def test_1013_updated_release_is_active_and_rebuildable() -> None:
    """重庆 1.9.30 应冻结新版来源、启用第20条并可确定性复验。"""

    assert UPDATED_SOURCE.is_file()
    pack = load_pack(UPDATED_RELEASE)
    rules = {rule["rule_id"]: rule for rule in pack["rules"]}
    person_name_rule = rules["chongqing.person_name_single"]
    profile = json.loads((ROOT / "regions/index.json").read_text(encoding="utf-8"))["regions"][0]

    assert pack["manifest"]["version"] == "1.9.30"
    assert pack["manifest"]["source_document_sha256"] == sha256_file(UPDATED_SOURCE)
    assert person_name_rule["display_code"] == "R20"
    assert person_name_rule["scope"]["record_type"] == "position_duty"
    assert "1.9.30" in CONFIRMED_RULE_VERSIONS
    assert "1.9.30" in PREPROCESSING_RULE_VERSIONS
    assert "1.9.30" in SYSTEM_TYPE_PREPROCESSING_RULE_VERSIONS
    assert profile["rulepack_version"] == "1.9.30"
    assert profile["rulepack_content_hash"] == pack["manifest"]["content_hash"]

    result = subprocess.run(
        [
            sys.executable,
            str(ROOT / "tools/publish_rules_1013_chongqing_1930.py"),
            "--verify-only",
        ],
        cwd=ROOT,
        text=True,
        capture_output=True,
        check=False,
    )

    assert result.returncode == 0, result.stderr
    assert json.loads(result.stdout)["verified"] is True
