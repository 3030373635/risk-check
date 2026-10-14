"""验证03合同管理在单业务和全量批次中的审核结果一致。"""

from __future__ import annotations

import json
import shutil
from collections import Counter
from pathlib import Path

from openpyxl import Workbook

from risk_audit import __version__
from risk_audit.runner import audit


ENTITY_NAME = "国网湖南省电力有限公司电力科学研究院本部"


def _write_contract_materials(contract_directory: Path) -> None:
    """写入合同管理矩阵和三清单；contract_directory 为03合同管理目录。"""

    contract_directory.mkdir(parents=True)

    matrix = Workbook()
    matrix_sheet = matrix.active
    matrix_sheet.title = "合同管理"
    matrix_sheet.append([
        "控制措施编号",
        "控制措施",
        "控制载体",
        "责任主体",
        "是否适用",
        "不适用理由",
    ])
    matrix_sheet.append(["合同管理-1-控制措施01", "复核合同", "审批单", "法律部", "适用", ""])
    matrix.save(contract_directory / "03风控矩阵-合同管理-电科院.xlsx")

    checklist = Workbook()
    duty_sheet = checklist.active
    duty_sheet.title = "岗位职责清单"
    duty_sheet.append(["部门", "岗位名称", "人员姓名", "岗位职责编号", "角色", "岗位职责", "控制措施编号"])
    duplicate_row = [
        "法律部",
        "合同管理专责",
        "张三",
        "合同管理-岗位职责01",
        "审核",
        "对合同资料完整性负审核责任。",
        "合同管理-1-控制措施01",
    ]
    duty_sheet.append(duplicate_row)
    duty_sheet.append(duplicate_row)
    helper_sheet = checklist.create_sheet("Sheet1")
    helper_sheet.sheet_state = "hidden"
    # 复现电科院辅助表里的跨业务文字，防止其再次污染合同管理业务身份。
    helper_sheet["A6"] = "供电服务指挥中心（配网调控中心）"
    checklist.save(contract_directory / "“三清单”-合同管理-电科院.xlsx")


def _write_other_business(full_root: Path) -> None:
    """写入全量批次中的无关业务；full_root 为全量审核输入根目录。"""

    other_directory = full_root / ENTITY_NAME / "01供电服务"
    other_directory.mkdir(parents=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "供电服务"
    worksheet.append(["控制措施编号", "控制措施", "责任主体", "是否适用"])
    worksheet.append(["供电服务-1-控制措施01", "受理服务申请", "服务中心", "适用"])
    workbook.save(other_directory / "01风控矩阵-供电服务.xlsx")


def _normalized_contract_findings(run_directory: Path) -> list[dict]:
    """读取稳定的合同管理意见；run_directory 为单次审核报告目录。"""

    findings = json.loads((run_directory / "findings.json").read_text(encoding="utf-8"))
    return [
        {key: value for key, value in finding.items() if key != "finding_id"}
        for finding in findings
        if finding.get("business_id") == "contract_management"
    ]


def test_contract_findings_are_same_in_single_and_full_audit(tmp_path: Path) -> None:
    """真实审核链路必须保持03意见一致；tmp_path 为输入、输出和运行记录的隔离目录。"""

    project_root = Path(__file__).resolve().parents[2]
    active = json.loads(
        (project_root / "risk-audit/rulepacks/active.json").read_text(encoding="utf-8")
    )
    assert active["version"] == __version__

    single_root = tmp_path / "single-input"
    full_root = tmp_path / "full-input"
    single_contract = single_root / ENTITY_NAME / "03合同管理"
    full_contract = full_root / ENTITY_NAME / "03合同管理"
    _write_contract_materials(single_contract)
    full_contract.mkdir(parents=True)
    for source in single_contract.iterdir():
        shutil.copy2(source, full_contract / source.name)
    _write_other_business(full_root)

    rulepack = project_root / "risk-audit/rulepacks/releases" / active["version"]
    entity_file = project_root / "reference-data/会计主体清单20260907.xlsx"
    config_file = project_root / "audit-config.json"
    runs_root = tmp_path / "runs"
    single_result = audit(
        single_root,
        tmp_path / "single-output",
        rulepack,
        entity_file,
        project_root,
        runs_root,
        write=False,
        run_id="single",
        config_file=config_file,
    )
    full_result = audit(
        full_root,
        tmp_path / "full-output",
        rulepack,
        entity_file,
        project_root,
        runs_root,
        write=False,
        run_id="full",
        config_file=config_file,
    )

    single_business = next(item for item in single_result["business_results"] if item["business_code"] == "03")
    full_business = next(item for item in full_result["business_results"] if item["business_code"] == "03")
    single_findings = _normalized_contract_findings(runs_root / "single")
    full_findings = _normalized_contract_findings(runs_root / "full")
    single_rule_counts = Counter(item["display_code"] for item in single_findings)
    full_rule_counts = Counter(item["display_code"] for item in full_findings)

    assert single_business["business_id"] == full_business["business_id"] == "contract_management"
    assert single_business["input_files"] == full_business["input_files"] == 2
    assert single_business["findings"] == full_business["findings"] == len(single_findings)
    assert single_findings
    assert full_findings == single_findings
    # 两条重复职责意见证明三清单已进入合同管理审核，而非只比较矩阵结果。
    assert single_rule_counts["R15"] == 2
    assert full_rule_counts == single_rule_counts
