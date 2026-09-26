import shutil

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

from risk_audit.configuration.loader import load_pack
from risk_audit.inventory import OLE, true_format
from risk_audit.inventory_v180 import scan_package_v180
from risk_audit.models import Entity
from risk_audit.readers.excel import parse_files
from risk_audit.util import sha256_file
from risk_audit.writer import write_outputs


def test_current_scanner_includes_biff_et_file(tmp_path, project_root):
    """当前扫描器应接收 BIFF ET；tmp_path 为测试目录，project_root 为项目根目录。"""
    fixture = (
        project_root
        / 'templates/第一批省公司通用版本整合0818/第一批省公司通用版本整合9.4'
        / '09 职工福利保障与薪酬管理'
        / '09风控矩阵-职工福利保障与薪酬管理-省公司（审定）8.13.xls'
    )
    business_directory = tmp_path / '测试单位有限公司/09 职工福利保障与薪酬管理'
    business_directory.mkdir(parents=True)
    source = business_directory / '09风控矩阵-测试单位有限公司.et'
    shutil.copyfile(fixture, source)

    files = scan_package_v180(tmp_path, {'E1': Entity('E1', '测试单位有限公司')}, {})

    assert len(files) == 1
    assert files[0].true_format == 'xls'
    assert files[0].business_code == '09'
    assert not files[0].parse_errors


def test_current_scanner_includes_ooxml_et_file(tmp_path):
    """当前扫描器应接收 OOXML ET；tmp_path 为测试目录。"""
    business_directory = tmp_path / '测试单位有限公司/06 设备（资产）管理'
    business_directory.mkdir(parents=True)
    source = business_directory / '06风控矩阵-测试单位有限公司.ET'
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(['控制措施编号', '控制措施'])
    worksheet.append(['M1', '核对报表'])
    workbook.save(source)

    files = scan_package_v180(tmp_path, {'E1': Entity('E1', '测试单位有限公司')}, {})

    assert len(files) == 1
    assert files[0].true_format == 'xlsx'
    assert files[0].business_code == '06'
    assert not files[0].parse_errors


@pytest.mark.parametrize('body', [b'unsupported-native-et', OLE + b'not-a-biff-workbook'])
def test_unrecognized_et_is_reported_in_inventory(tmp_path, body):
    """未知 ET 应保留解析错误；tmp_path 为测试目录，body 为文件内容。"""
    source = tmp_path / '06风控矩阵-测试单位有限公司.et'
    source.write_bytes(body)

    files = scan_package_v180(tmp_path, {'E1': Entity('E1', '测试单位有限公司')}, {})

    assert len(files) == 1
    assert files[0].true_format == 'unknown'
    assert '文件签名不是受支持的 Excel、Word 或 PDF 格式' in files[0].parse_errors
    assert true_format(source) == 'unknown'


def test_ooxml_et_reads_rich_text_and_writes_xlsx_copy(tmp_path, project_root):
    """当前链路应读取 OOXML ET 富文本并写 XLSX；参数为测试目录和项目根目录。"""
    pack = load_pack(project_root / 'risk-audit/rulepacks/releases/1.9.14')
    input_root = tmp_path / 'input'
    business_directory = input_root / '测试单位有限公司/06 设备（资产）管理'
    business_directory.mkdir(parents=True)
    source = business_directory / '06三清单-测试单位有限公司.ET'
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = '岗位职责清单'
    worksheet.append(['岗位职责清单'])
    worksheet.append(['部门', '岗位名称', '岗位职责', '控制措施编号'])
    worksheet.append([
        '财务部',
        '财务专责',
        CellRichText([
            TextBlock(InlineFont(strike=True), '旧职责'),
            TextBlock(InlineFont(), '对报表准确性负主体责任'),
        ]),
        'M1',
    ])
    workbook.save(source)
    source_hash = sha256_file(source)

    files = scan_package_v180(input_root, {'E1': Entity('E1', '测试单位有限公司')}, {})
    assert files[0].true_format == 'xlsx'
    files[0]._rules_version = pack['manifest']['version']
    parse_files(files, pack['field_aliases'], tmp_path / 'run')

    assert not files[0].parse_errors
    duty = files[0].sheets[0].records[0].fields['duty']
    assert duty.current == '对报表准确性负主体责任'
    assert duty.deleted_spans[0]['text'] == '旧职责'

    output_root = tmp_path / 'output'
    warnings, _ = write_outputs(files, [], output_root, metadata_dir=tmp_path / 'metadata')
    result = output_root / source.relative_to(input_root).with_suffix('.xlsx')

    assert not warnings
    assert result.exists()
    assert not (output_root / source.relative_to(input_root)).exists()
    with result.open('rb') as stream:
        result_workbook = load_workbook(stream, rich_text=True)
    assert str(result_workbook['岗位职责清单']['C3'].value) == '旧职责对报表准确性负主体责任'
    assert sha256_file(source) == source_hash
