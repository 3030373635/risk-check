"""确认稿读取器的业务区域及原始坐标回归测试。"""

import json
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font

from risk_audit.models import FileRecord
from risk_audit.readers import excel
from risk_audit.readers.excel import parse_files


@pytest.fixture
def aliases():
    """加载正式字段别名；无参数。"""
    path = Path(__file__).parents[1] / "rulepacks/releases/1.7.0/field_aliases.json"
    return json.loads(path.read_text())


def read_book(tmp_path, wb, aliases):
    """保存并读取工作簿；参数分别为临时目录、工作簿、字段别名。"""
    path = tmp_path / "资料.xlsx"
    wb.save(path)
    file = FileRecord(path, Path(path.name), "digest", "xlsx", "E1", [], False, "B1", "default", "three_lists")
    # 新接口不存在时调用成熟读取器，使失败落在实际业务断言上。
    try:
        from risk_audit.readers.confirmed_v180 import parse_workbook_v180
    except ImportError:
        parse_workbook_v180 = excel.parse_workbook
    return parse_workbook_v180(file, path, aliases), file, path


def parse_three_lists_file(tmp_path, wb, aliases, *, include_hidden=False):
    """通过正式文件解析入口读取三清单；参数为测试目录、工作簿、字段别名和隐藏表开关。"""
    path = tmp_path / '三清单.xlsx'
    wb.save(path)
    file = FileRecord(path, Path(path.name), 'digest', 'xlsx', 'E1', [], False,
                      'B1', 'default', 'three_lists')
    file._rules_version = '1.9.7'
    parse_files([file], aliases, tmp_path / 'work', include_hidden=include_hidden)
    return file


def add_three_lists_sheet(book, title, headers):
    """向测试工作簿添加一张业务表；参数 book 为工作簿，title 为表名，headers 为表头。"""
    sheet = book.create_sheet(title)
    sheet.append(headers)
    return sheet


def complete_three_lists_book():
    """构造结构合法的三清单工作簿；无参数。"""
    book = Workbook()
    position = book.active
    position.title = '岗位内控责任清单'
    position.append(['控制措施编号', '部门', '岗位名称', '岗位职责'])
    add_three_lists_sheet(book, '不相容岗位清单', ['不相容业务角色', '岗位A', '岗位B'])
    add_three_lists_sheet(book, '系统控制规则清单', ['控制措施编号', '规则名称', '规则内容'])
    return book


def test_three_lists_missing_optional_sheet_still_parses_matched_sheets(tmp_path, aliases):
    """缺少某类清单时继续解析已匹配表；参数为目录和字段别名。"""
    book = complete_three_lists_book()
    book.remove(book['系统控制规则清单'])

    file = parse_three_lists_file(tmp_path, book, aliases)

    assert [(sheet.title, sheet.sheet_type) for sheet in file.sheets] == [
        ('岗位内控责任清单', 'position_duty'),
        ('不相容岗位清单', 'incompatible_position'),
    ]
    assert file.parse_errors == []


def test_three_lists_with_extra_unrecognized_visible_sheet_is_parsed(tmp_path, aliases):
    """额外未识别可见表不审核且不阻断已匹配表；参数为目录和字段别名。"""
    book = complete_three_lists_book()
    add_three_lists_sheet(book, 'Sheet2', ['说明'])

    file = parse_three_lists_file(tmp_path, book, aliases)

    assert [(sheet.title, sheet.sheet_type) for sheet in file.sheets] == [
        ('岗位内控责任清单', 'position_duty'),
        ('不相容岗位清单', 'incompatible_position'),
        ('系统控制规则清单', 'system_rule'),
    ]
    assert all(sheet.title != 'Sheet2' for sheet in file.sheets)
    assert file.parse_errors == []


def test_three_lists_with_duplicate_matched_type_parses_every_sheet(tmp_path, aliases):
    """同类清单匹配多张时全部解析；参数为目录和字段别名。"""
    book = complete_three_lists_book()
    book.remove(book['系统控制规则清单'])
    book['岗位内控责任清单'].append(['M1', '财务部', '会计', '复核报表'])
    backup = add_three_lists_sheet(book, '岗位职责清单-备份',
                                   ['控制措施编号', '部门', '岗位名称', '岗位职责'])
    backup.append(['M2', '审计部', '审计', '检查报表'])

    file = parse_three_lists_file(tmp_path, book, aliases)

    assert [(sheet.title, sheet.sheet_type) for sheet in file.sheets] == [
        ('岗位内控责任清单', 'position_duty'),
        ('不相容岗位清单', 'incompatible_position'),
        ('岗位职责清单-备份', 'position_duty'),
    ]
    assert file.parse_errors == []
    assert {
        (sheet.title, record.value('measure_id'))
        for sheet in file.sheets if sheet.sheet_type == 'position_duty'
        for record in sheet.records
    } == {('岗位内控责任清单', 'M1'), ('岗位职责清单-备份', 'M2')}


def test_complete_three_lists_is_parsed(tmp_path, aliases):
    """三类清单全部匹配时正常读取；参数为目录和字段别名。"""
    file = parse_three_lists_file(tmp_path, complete_three_lists_book(), aliases)

    assert {sheet.sheet_type for sheet in file.sheets} == {
        'position_duty', 'incompatible_position', 'system_rule',
    }
    assert file.parse_errors == []


def test_row_only_sort_state_is_parsed_without_modifying_source(tmp_path, aliases):
    """整行排序范围不应阻断业务解析；参数为目录和字段别名。"""
    book = complete_three_lists_book()
    position = book['岗位内控责任清单']
    position.append(['M1', '财务部', '会计', '复核'])
    position.auto_filter.ref = 'A1:D2'
    position.auto_filter.add_sort_condition('A2:A2')
    path = tmp_path / '整行排序范围.xlsx'
    book.save(path)

    with ZipFile(path) as archive:
        parts = {info.filename: (info, archive.read(info.filename)) for info in archive.infolist()}
    sheet_info, sheet_xml = parts['xl/worksheets/sheet1.xml']
    invalid_xml = sheet_xml.replace(b'<sortState ref="A1:D2">', b'<sortState ref="1:2">')
    assert invalid_xml != sheet_xml
    parts['xl/worksheets/sheet1.xml'] = (sheet_info, invalid_xml)
    with ZipFile(path, 'w', ZIP_DEFLATED) as archive:
        for info, content in parts.values():
            archive.writestr(info, content)

    source_bytes = path.read_bytes()
    file = FileRecord(path, Path(path.name), 'digest', 'xlsx', 'E1', [], False,
                      'B1', 'default', 'three_lists')
    file._rules_version = '1.9.10'
    parse_files([file], aliases, tmp_path / 'work')

    assert file.parse_errors == []
    assert {sheet.sheet_type for sheet in file.sheets} == {
        'position_duty', 'incompatible_position', 'system_rule',
    }
    assert path.read_bytes() == source_bytes


@pytest.mark.parametrize('sheet_state', ['hidden', 'veryHidden'])
def test_hidden_auxiliary_sheet_is_not_audited(tmp_path, aliases, sheet_state):
    """隐藏附加表不参与审核；参数为目录、字段别名和隐藏状态。"""
    book = complete_three_lists_book()
    extra = add_three_lists_sheet(book, '历史说明', ['说明'])
    extra.sheet_state = sheet_state

    file = parse_three_lists_file(tmp_path, book, aliases)

    assert {sheet.sheet_type for sheet in file.sheets} == {
        'position_duty', 'incompatible_position', 'system_rule',
    }
    assert file.parse_errors == []
    assert file.preservation['hidden_sheets'] == [
        {'title': '历史说明', 'state': sheet_state, 'action': 'skipped_hidden'},
    ]


@pytest.mark.parametrize('sheet_state', ['hidden', 'veryHidden'])
def test_hidden_matched_sheet_is_skipped_without_blocking_visible_sheets(tmp_path, aliases, sheet_state):
    """隐藏的已匹配清单不审核且不阻断可见表；参数为目录、字段别名和隐藏状态。"""
    book = complete_three_lists_book()
    book['不相容岗位清单'].sheet_state = sheet_state

    file = parse_three_lists_file(tmp_path, book, aliases)

    assert [(sheet.title, sheet.sheet_type) for sheet in file.sheets] == [
        ('岗位内控责任清单', 'position_duty'),
        ('系统控制规则清单', 'system_rule'),
    ]
    assert file.parse_errors == []
    assert file.preservation['hidden_sheets'] == [
        {'title': '不相容岗位清单', 'state': sheet_state, 'action': 'skipped_hidden'},
    ]


def test_material_title_in_carrier_is_business_text_not_region(tmp_path, aliases):
    """参数为目录与别名；明细中的完整清单名称不能截断矩阵或误划业务区域。"""
    book = Workbook()
    sheet = book.active
    sheet.title = '风控矩阵'
    sheet.append(['控制措施编号', '控制措施', '控制载体', '责任主体', '是否适用及原因'])
    sheet.append(['M1', '编制岗位责任清单', '岗位内控责任清单', '财务部', '适用'])
    sheet.append(['M2', '复核岗位责任清单', '业务单据', '财务部', '适用'])
    parsed, _, _ = read_book(tmp_path, book, aliases)
    assert [part.sheet_type for part in parsed] == ['matrix']
    assert [row.value('measure_id') for row in parsed[0].records] == ['M1', 'M2']


def test_region_title_allows_explicit_reporting_metadata(tmp_path, aliases):
    """参数为目录与别名；标题行的填报单位元信息不应使合法区域消失。"""
    book = Workbook()
    sheet = book.active
    sheet.title = '综合填报表'
    sheet.append(['风控矩阵', None, None, None, None, '填报单位：信通公司'])
    sheet.append(['控制措施编号', '控制措施', '控制载体', '责任主体', '是否适用及原因'])
    sheet.append(['M1', '核对报表', '报表', '财务部', '适用'])
    sheet.append([])
    sheet.append(['岗位内控责任清单', None, None, None, None, '填报单位：信通公司'])
    sheet.append(['控制措施编号', '部门', '岗位名称', '姓名', '岗位职责', '角色'])
    sheet.append(['M1', '财务部', '会计', '张三', '负主体责任', '经办'])
    parsed, _, _ = read_book(tmp_path, book, aliases)
    assert [part.sheet_type for part in parsed] == ['matrix', 'position_duty']
    assert [row.row for part in parsed for row in part.records] == [3, 7]


def test_mixed_workbook_detects_matrix_despite_three_lists_material(tmp_path, aliases):
    """三清单分类的文件仍包含矩阵；参数为临时目录和字段别名。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "风控矩阵"
    ws.append(["控制措施编号", "控制措施", "责任主体"])
    ws.append(["M1", "复核", "部门"])
    wb.create_sheet("不相容岗位清单")
    parsed, file, path = read_book(tmp_path, wb, aliases)
    assert {sheet.sheet_type for sheet in parsed} == {"matrix", "incompatible_position"}


@pytest.mark.parametrize("message", [None, "无", "不涉及"])
def test_empty_incompatible_list_is_present(tmp_path, aliases, message):
    """空表和无事项声明均属于已提交清单；参数为目录、别名、声明。"""
    wb = Workbook()
    wb.active.title = "不相容岗位清单"
    wb.active.cell(1, 1, message)
    parsed, _, _ = read_book(tmp_path, wb, aliases)
    assert len(parsed) == 1
    assert parsed[0].sheet_type == "incompatible_position"
    assert parsed[0].records == []


@pytest.mark.parametrize("missing", ["部门", "控制措施编号", "岗位名称", "岗位职责", "姓名"])
def test_incomplete_position_headers_are_kept_for_column_findings(tmp_path, aliases, missing):
    """缺失一个字段时仍读取岗位清单；参数为目录、别名、缺失列名。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "岗位内控责任清单"
    values = {"部门": "财务", "岗位名称": "会计", "岗位职责": "复核", "控制措施编号": "M1", "姓名": "张三"}
    headers = [header for header in values if header != missing]
    ws.append(headers)
    ws.append([values[header] for header in headers])
    parsed, _, _ = read_book(tmp_path, wb, aliases)
    assert len(parsed) == 1
    assert parsed[0].sheet_type == "position_duty"
    assert len(parsed[0].records) == 1


def test_vertical_regions_keep_physical_rows_and_do_not_fill_between_regions(tmp_path, aliases):
    """四个区域独立读取，保留坐标和删除线；参数为目录、别名。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "综合填报"
    ws.append(["风控矩阵"])
    ws.append(["控制措施编号", "控制措施", "责任主体", "审核意见"])
    ws.append(["M1", "删除措施", "财务"])
    ws["B3"].font = Font(strike=True)
    ws.merge_cells("C3:C40")  # 合并范围触及下一区域标题，区域视图禁止跨界解引用。
    ws.cell(40, 1, "岗位内控责任清单")
    for col, value in enumerate(["部门", "岗位名称", "岗位职责", "控制措施编号", "姓名"], 1):
        ws.cell(41, col, value)
    for col, value in enumerate(["财务", "会计", None, "M1", "张三"], 1):
        if value:
            ws.cell(42, col, value)
    ws.cell(80, 1, "不相容岗位清单")
    for col, value in enumerate(["不相容业务角色", "岗位A", "岗位B"], 1):
        ws.cell(81, col, value)
    for col, value in enumerate(["付款与审批", "会计", "经理"], 1):
        ws.cell(82, col, value)
    ws.cell(120, 1, "系统控制规则清单")
    for col, value in enumerate(["控制措施编号", "系统规则名称", "规则内容"], 1):
        ws.cell(121, col, value)
    for col, value in enumerate(["M1", "权限校验", "禁止越权"], 1):
        ws.cell(122, col, value)
    parsed, file, path = read_book(tmp_path, wb, aliases)
    assert [(sheet.sheet_type, [r.row for r in sheet.records]) for sheet in parsed] == [
        ("matrix", [3]), ("position_duty", [42]), ("incompatible_position", [82]), ("system_rule", [122])]
    assert {sheet.title for sheet in parsed} == {"综合填报"}
    # D列在岗位区域是业务列，复用矩阵区域D列意见会损坏业务资料。
    assert {sheet.output_column for sheet in parsed} == {6}
    assert all(sheet.audit_columns == {"审核意见": 4} for sheet in parsed)
    assert parsed[0].records[0].fields["control_measure"].deleted_spans[0]["text"] == "删除措施"
    assert parsed[1].records[0].fields["duty"].coordinate == "C42"
    assert parsed[1].records[0].value("duty") == ""
    assert parsed[3].header_rows == [121]
    assert file.preservation["audit_column_conflicts"][0]["sheet"] == "综合填报"
    assert file.preservation["audit_column_conflicts"][0]["original_columns"] == [4]
    assert file.preservation["audit_column_conflicts"][0]["output_column"] == 6
    from openpyxl import load_workbook
    original = load_workbook(path)
    assert original["综合填报"]["D41"].value == "控制措施编号"
    assert original["综合填报"]["D42"].value == "M1"
    original.close()
    from risk_audit.writer import patch_ooxml
    destination = tmp_path / "审核后.xlsx"
    patch_ooxml(path, destination, {"综合填报": {"column": parsed[0].output_column,
                 "header_row": 2, "header_rows": [2, 41, 81, 121], "values": {42: "请核实岗位职责"}}})
    audited = load_workbook(destination)
    assert audited["综合填报"]["D41"].value == "控制措施编号"
    assert audited["综合填报"]["D42"].value == "M1"
    assert audited["综合填报"]["D2"].value == "审核意见"
    assert audited["综合填报"]["F42"].value == "请核实岗位职责"
    audited.close()


def test_late_header_and_formula_cache_are_preserved(tmp_path, aliases):
    """超过30行的标题及公式缓存仍可读取；参数为目录、别名。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "岗位内控责任清单"
    for col, value in enumerate(["部门", "岗位名称", "岗位职责", "控制措施编号"], 1):
        ws.cell(60, col, value)
    for col, value in enumerate(["财务", "会计", "复核", '= "M1"'], 1):
        ws.cell(61, col, value)
    _, file, path = read_book(tmp_path, wb, aliases)
    with ZipFile(path) as archive:
        content = {name: archive.read(name) for name in archive.namelist()}
    content["xl/worksheets/sheet1.xml"] = content["xl/worksheets/sheet1.xml"].replace(b'<v></v>', b'<v>7</v>')
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, data in content.items():
            archive.writestr(name, data)
    try:
        from risk_audit.readers.confirmed_v180 import parse_workbook_v180
    except ImportError:
        parse_workbook_v180 = excel.parse_workbook
    parsed = parse_workbook_v180(file, path, aliases)
    field = parsed[0].records[0].fields["measure_id"]
    assert parsed[0].records[0].row == 61
    assert field.coordinate == "D61"
    assert field.cached == 7
    assert field.state == "formula_cached"


def test_auxiliary_sheet_is_excluded(tmp_path, aliases):
    """填报说明包含示例字段也不作为业务区域；参数为目录、别名。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "填报说明"
    ws.append(["部门", "岗位名称", "岗位职责", "控制措施编号"])
    ws.append(["财务", "会计", "示例", "M1"])
    parsed, _, _ = read_book(tmp_path, wb, aliases)
    assert parsed == []


def test_side_by_side_regions_keep_physical_columns(tmp_path, aliases):
    """并排区域保持各自列号，矩阵不读取岗位字段；参数为目录、别名。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "综合填报"
    ws.cell(1, 1, "风控矩阵")
    ws.cell(1, 6, "岗位内控责任清单")
    for col, value in enumerate(["控制措施编号", "控制措施", "责任主体"], 1):
        ws.cell(2, col, value)
    for col, value in enumerate(["M1", "复核", "财务"], 1):
        ws.cell(3, col, value)
    for col, value in enumerate(["部门", "岗位名称", "岗位职责", "控制措施编号"], 6):
        ws.cell(2, col, value)
    for col, value in enumerate(["财务", "会计", "核查", "M2"], 6):
        ws.cell(3, col, value)
    parsed, _, _ = read_book(tmp_path, wb, aliases)
    assert len(parsed) == 2
    assert parsed[0].records[0].value("measure_id") == "M1"
    assert parsed[1].records[0].value("measure_id") == "M2"
    assert parsed[1].records[0].fields["department"].coordinate == "F3"
    assert parsed[0].records[0].record_id != parsed[1].records[0].record_id
    assert {sheet.output_column for sheet in parsed} == {10}


def test_hidden_empty_incompatible_sheet_preserves_visibility_scope(tmp_path, aliases):
    """隐藏空清单只在明确包含隐藏表时存在；参数为目录、别名。"""
    wb = Workbook()
    wb.active.title = "填报说明"
    wb.create_sheet("不相容岗位清单").sheet_state = "hidden"
    parsed, file, path = read_book(tmp_path, wb, aliases)
    assert parsed == []
    assert file.preservation["hidden_sheets"] == [{"title": "不相容岗位清单", "state": "hidden", "action": "skipped_hidden"}]
    from risk_audit.readers.confirmed_v180 import parse_workbook_v180
    parsed = parse_workbook_v180(file, path, aliases, include_hidden=True)
    assert len(parsed) == 1
    assert file.preservation["hidden_sheets"][0]["action"] == "audited"


def test_struck_business_title_is_not_an_active_region(tmp_path, aliases):
    """删除线标题不划分有效区域；参数为目录、别名。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "综合填报"
    ws.cell(1, 1, "岗位内控责任清单").font = Font(strike=True)
    parsed, _, _ = read_book(tmp_path, wb, aliases)
    assert parsed == []


def test_generic_sheet_with_late_matrix_headers_is_detected(tmp_path, aliases):
    """通用表名依据晚出现的明确业务字段识别；参数为目录、别名。"""
    wb = Workbook()
    ws = wb.active
    for col, value in enumerate(["控制措施编号", "控制措施", "责任主体"], 1):
        ws.cell(60, col, value)
    for col, value in enumerate(["M1", "复核", "财务"], 1):
        ws.cell(61, col, value)
    parsed, _, _ = read_book(tmp_path, wb, aliases)
    assert parsed[0].sheet_type == "matrix"
    assert parsed[0].header_rows == [60]


def test_late_headers_preserve_physical_rows_in_owner_resolution_evidence(tmp_path, aliases):
    """主体归属证据内的行号也应恢复；参数为目录、别名。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "风控矩阵"
    for col, value in enumerate(["控制措施编号", "控制措施", "市公司责任主体", "市公司是否适用"], 1):
        ws.cell(60, col, value)
    for row in (61, 62):
        for col, value in enumerate([f"M{row}", "复核", "财务公司-财务部", "适用"], 1):
            ws.cell(row, col, value)
    _, file, path = read_book(tmp_path, wb, aliases)
    file._parser_policy = {"version": 3}
    file._confirmed_unit_names = ["财务公司"]
    file.entity_evidence = ["标准全称:财务公司"]
    from risk_audit.readers.confirmed_v180 import parse_workbook_v180
    parse_workbook_v180(file, path, aliases)
    selection = next(item for item in file.preservation["column_selections"] if item["field"] == "responsibility")
    assert selection["candidates"][0]["owner_resolution"]["current_unit_body_rows"] == [61, 62]


def test_business_named_sheet_uses_matrix_title_with_missing_responsibility(tmp_path, aliases):
    """表名仅含业务名时仍按表内标题读取；tmp_path 为目录，aliases 为字段别名。"""
    book = Workbook()
    sheet = book.active
    sheet.title = '营销售电'
    sheet.append(['风控矩阵-营销售电业务'])
    sheet.append(['控制措施编号', '控制措施', '尚未确认的责任列'])
    sheet.append(['M1', '核对客户资料', '业务监控部-业扩专责'])
    parsed, _, _ = read_book(tmp_path, book, aliases)
    assert len(parsed) == 1
    assert parsed[0].sheet_type == 'matrix'
    assert parsed[0].records[0].row == 3
    assert parsed[0].records[0].value('measure_id') == 'M1'
    assert 'responsibility' not in parsed[0].columns


@pytest.mark.parametrize('owner,selected', [('供电服务中心', True), ('外单位服务中心', False)])
def test_center_responsibility_header_keeps_unit_ownership(tmp_path, aliases, owner, selected):
    """中心表头须核对主体；参数为目录、别名、表头单位及是否允许选中。"""
    book = Workbook()
    sheet = book.active
    sheet.title = '营销售电'
    sheet.append(['风控矩阵-营销售电业务'])
    sheet.append(['控制措施编号', '控制措施', owner + '\n责任主体'])
    sheet.append(['M1', '核对客户资料', '供电服务中心-业务监控部-业扩专责'])
    _, file, path = read_book(tmp_path, book, aliases)
    file.entity_code = '20JQ'
    file.entity_evidence = ['标准全称:国网湖南省电力有限公司供电服务中心（计量中心）本部']
    file._parser_policy = {'version': 3}
    from risk_audit.readers.confirmed_v180 import parse_workbook_v180
    parsed = parse_workbook_v180(file, path, aliases)
    assert len(parsed) == 1
    assert ('responsibility' in parsed[0].columns) == selected
    if selected:
        assert parsed[0].columns['responsibility'] == 3
        assert parsed[0].records[0].fields['responsibility'].coordinate == 'C3'
    else:
        selection = next(item for item in file.preservation['column_selections'] if item['field'] == 'responsibility')
        assert selection['selected_column'] is None
        assert selection['candidates'][0]['owner'] == '外单位服务中心'


def test_a1_instruction_mentioning_matrix_does_not_override_duty_headers(tmp_path, aliases):
    """说明文字引用矩阵不能改变岗位表类型；tmp_path 为目录，aliases 为字段别名。"""
    book = Workbook()
    sheet = book.active
    sheet.title = '岗位资料'
    sheet.append(['依据风控矩阵整理岗位分工'])
    sheet.append(['控制措施编号', '部门', '岗位名称', '岗位职责', '控制措施'])
    sheet.append(['M1', '财务部', '核算专责', '负主体责任', '核对报表'])
    parsed, _, _ = read_book(tmp_path, book, aliases)
    assert [part.sheet_type for part in parsed] == ['position_duty']
    assert parsed[0].records[0].value('duty') == '负主体责任'


def test_foreign_center_in_body_blocks_template_owner_correction(tmp_path, aliases):
    """明确外单位中心不能因本单位占九成而忽略；参数为目录和字段别名。"""
    book = Workbook()
    sheet = book.active
    sheet.title = '风控矩阵'
    sheet.append(['控制措施编号', '控制措施', '外单位服务中心责任主体'])
    for index in range(10):
        owner = '供电服务中心' if index < 9 else '外单位服务中心'
        sheet.append([f'M{index}', '核对资料', owner + '-财务部-核算专责'])
    _, file, path = read_book(tmp_path, book, aliases)
    file.entity_code = '20JQ'
    file.entity_evidence = ['标准全称:国网湖南省电力有限公司供电服务中心（计量中心）本部']
    file._parser_policy = {'version': 3}
    from risk_audit.readers.confirmed_v180 import parse_workbook_v180
    parsed = parse_workbook_v180(file, path, aliases)
    assert 'responsibility' not in parsed[0].columns
    selection = next(item for item in file.preservation['column_selections'] if item['field'] == 'responsibility')
    proof = selection['candidates'][0]['owner_resolution']
    assert proof['foreign_unit_body_rows'] == [{'row': 11, 'units': ['外单位服务中心']}]


@pytest.mark.parametrize('owner,selected', [
    ('湖南朗晟电力产业发展有限公司计量分公司', True),
    ('湖南朗晟电力产业发展有限公司其他分公司', False),
])
def test_responsibility_header_with_parenthesized_unit_preserves_owner(tmp_path, aliases, owner, selected):
    """括号内全称须核对归属；参数为目录、别名、表头单位及是否允许选中。"""
    book = Workbook()
    sheet = book.active
    sheet.title = '职工福利保障与薪酬管理（企管部）'
    sheet.append(['风控矩阵-职工福利保障与薪酬管理'])
    sheet.append(['控制措施编号', '控制措施', f'责任主体（{owner}）'])
    sheet.append(['M1', '核对薪酬', '湖南朗晟计量分公司-综合管理部-员工管理专责'])
    _, file, path = read_book(tmp_path, book, aliases)
    file.entity_code = '209Z'
    file.entity_evidence = ['标准全称:湖南朗晟电力产业发展有限公司计量分公司']
    file._parser_policy = {'version': 3}
    from risk_audit.readers.confirmed_v180 import parse_workbook_v180
    parsed = parse_workbook_v180(file, path, aliases)
    assert len(parsed) == 1
    assert ('responsibility' in parsed[0].columns) == selected
    if selected:
        assert parsed[0].records[0].fields['responsibility'].coordinate == 'C3'
    else:
        selection = next(item for item in file.preservation['column_selections'] if item['field'] == 'responsibility')
        assert selection['candidates'][0]['owner'] == owner
