"""验证主体只按单位目录识别，文件名称不影响归属及清单外标记。"""
import pytest
from openpyxl import Workbook

from risk_audit.entity_alerts import build_entity_alerts
from risk_audit.entity_groups import prepare_entity_groups
from risk_audit.inventory_v180 import scan_package_v180
from risk_audit.models import Entity, FileRecord


@pytest.mark.parametrize('filename', [
    '07风控矩阵-随意简称.xlsx',
    '07风控矩阵-其他服务中心.xlsx',
    '07风控矩阵-另一个已登记有限公司.xlsx',
])
def test_registered_directory_ignores_filename(tmp_path, filename):
    """tmp_path 为测试根目录，filename 为任意报送名称；目录全称决定归属。"""
    name = '湖南朗晟电力产业发展有限公司计量分公司'
    directory = tmp_path / name / '07物资（服务）采购与实施-随意简称'
    directory.mkdir(parents=True)
    Workbook().save(directory / filename)
    entities = {'209Z': Entity('209Z', name), 'OTHER': Entity('OTHER', '另一个已登记有限公司')}
    files = scan_package_v180(tmp_path, entities, {})
    groups = prepare_entity_groups(files, tmp_path, entities)
    assert list(groups) == ['209Z']
    assert files[0].entity_code == '209Z'
    assert not files[0].entity_conflict


def test_unlisted_directory_remains_one_marked_subject(tmp_path):
    """tmp_path 为测试根目录；同目录的不同文件名不能拆主体或借用已登记代码。"""
    name = '清单外样例有限公司'
    directory = tmp_path / name / '07物资（服务）采购与实施'
    directory.mkdir(parents=True)
    for filename in ['07风控矩阵-另一个已登记有限公司.xlsx', '07三清单-其他服务中心.xlsx']:
        Workbook().save(directory / filename)
    entities = {'OTHER': Entity('OTHER', '另一个已登记有限公司')}
    files = scan_package_v180(tmp_path, entities, {})
    groups = prepare_entity_groups(files, tmp_path, entities)
    assert len(groups) == 1
    identity = next(iter(groups.values()))
    assert identity['entity_name'] == name
    assert identity['entity_code'] == ''
    assert identity['entity_registry_status'] == 'not_listed'
    assert identity['entity_registry_message'] == '主体不在会计主体清单中'
    assert len({file.entity_code for file in files}) == 1
    assert all(not file.entity_conflict for file in files)


def test_parent_directory_mode_uses_directory_as_confirmed_subject(tmp_path):
    """目录识别模式直接使用业务上层目录，不生成名册待确认提示。"""
    name = '重庆样例供电分公司'
    directory = tmp_path / name / '07物资（服务）采购与实施'
    directory.mkdir(parents=True)
    Workbook().save(directory / '07风控矩阵.xlsx')
    files = scan_package_v180(tmp_path, {}, {})

    groups = prepare_entity_groups(files, tmp_path, {}, entity_mode='parent_directory')

    identity = next(iter(groups.values()))
    assert identity['entity_name'] == name
    assert identity['entity_code'] == ''
    assert identity['entity_registry_status'] == 'directory_matched'
    assert identity['entity_registry_message'] == '主体已按业务上层目录识别'
    assert files[0].entity_code.startswith('@目录主体:')
    assert build_entity_alerts(files, {}) == []


def test_parent_directory_mode_preserves_exact_subject_directory_name(tmp_path):
    """主体显示名必须保留业务上级目录原名，不得删除目录序号。"""
    name = '01.重庆样例供电分公司'
    directory = tmp_path / name / '07物资（服务）采购与实施'
    directory.mkdir(parents=True)
    Workbook().save(directory / '01风控矩阵-其他业务.xlsx')
    files = scan_package_v180(tmp_path, {}, {})

    identity = next(iter(prepare_entity_groups(
        files,
        tmp_path,
        {},
        entity_mode='parent_directory',
    ).values()))

    assert identity['entity_name'] == name


def test_parent_directory_mode_uses_parent_of_unknown_business_directory(tmp_path):
    """未配置业务也必须按固定层级取主体，不能把业务目录本身当作主体。"""
    subject_name = '重庆新业务报送主体'
    business_name = '17.全新专项矩阵及“三清单”'
    directory = tmp_path / subject_name / business_name
    directory.mkdir(parents=True)
    Workbook().save(directory / '17风控矩阵-文件名称不得参与归属.xlsx')

    files = scan_package_v180(tmp_path, {}, {})
    identity = next(iter(prepare_entity_groups(
        files,
        tmp_path,
        {},
        entity_mode='parent_directory',
    ).values()))

    assert identity['entity_name'] == subject_name


def test_parent_directory_mode_uses_input_root_as_exact_subject_name(tmp_path):
    """输入根本身是业务上级目录时，主题名称必须完整保留根目录原名。"""
    root = tmp_path / '重庆 新建文件-20261012-194522'
    directory = root / '11.财务管理矩阵及“三清单”'
    directory.mkdir(parents=True)
    Workbook().save(directory / '02风控矩阵-文件名称不得参与归属.xlsx')

    files = scan_package_v180(root, {}, {})
    identity = next(iter(prepare_entity_groups(
        files,
        root,
        {},
        entity_mode='parent_directory',
    ).values()))

    assert identity['entity_name'] == root.name


def test_selected_business_root_uses_outer_directory_as_subject(tmp_path):
    """直接选择业务目录时，主体必须是业务根的外层目录而非业务目录自身。"""
    root = tmp_path / '11.财务管理矩阵及“三清单”'
    root.mkdir()
    Workbook().save(root / '02风控矩阵-文件名称不得参与归属.xlsx')

    files = scan_package_v180(root, {}, {})
    groups = prepare_entity_groups(files, root, {}, entity_mode='parent_directory')
    identity = next(iter(groups.values()))

    assert identity['entity_name'] == tmp_path.name
    assert identity['group_key'] == f'@目录主体:../{tmp_path.name}'


def test_same_unlisted_directory_name_merges_across_all_business_directories(tmp_path):
    """tmp_path 为测试根目录；同一主体下不同业务目录的材料必须合并为一个主体。"""
    name = '清单外样例有限公司'
    files = []
    for business_directory, business_code, material_type in [
        ('历史业务目录', '06', 'matrix'),
        ('未来新增业务目录', '42', 'three_lists'),
    ]:
        path = tmp_path / name / business_directory / f'{business_code}{material_type}.xlsx'
        files.append(FileRecord(path, path.relative_to(tmp_path), 'hash', 'xlsx', None, [], False,
                                business_code, 'default', material_type))

    groups = prepare_entity_groups(files, tmp_path, {})

    assert list(groups) == [f'@目录:{name}']
    assert {file.entity_code for file in files} == {f'@目录:{name}'}


def test_same_subject_name_in_different_paths_stays_separate(tmp_path):
    """主体目录同名但完整路径不同时不得合并，目录路径才是主体身份。"""
    files = []
    for parent in ('集团甲', '集团乙'):
        path = tmp_path / parent / '同名报送主体' / '17.专项业务' / '风控矩阵.xlsx'
        files.append(FileRecord(path, path.relative_to(tmp_path), 'hash', 'xlsx', None, [], False,
                                None, 'default', 'matrix', business_id='unmatched'))

    groups = prepare_entity_groups(files, tmp_path, {}, entity_mode='parent_directory')

    assert set(groups) == {
        '@目录主体:集团甲/同名报送主体',
        '@目录主体:集团乙/同名报送主体',
    }


def test_unlisted_child_directory_does_not_inherit_parent(tmp_path):
    """tmp_path 为测试根目录；清单外子单位目录不能沿父目录寻找正式代码。"""
    parent = tmp_path / '已登记母公司有限公司'
    directory = parent / '清单外子公司有限公司' / '07物资（服务）采购与实施'
    directory.mkdir(parents=True)
    Workbook().save(directory / '07风控矩阵.xlsx')
    entities = {'PARENT': Entity('PARENT', parent.name)}
    files = scan_package_v180(tmp_path, entities, {})
    groups = prepare_entity_groups(files, tmp_path, entities)
    identity = next(iter(groups.values()))
    assert identity['entity_code'] == ''
    assert identity['entity_name'] == '清单外子公司有限公司'
    assert identity['entity_registry_status'] == 'not_listed'
