"""验证模板业务唯一标识、别名匹配和配置约束。"""
from pathlib import Path

import pytest
from openpyxl import Workbook

from risk_audit.business_identity import (
    BusinessIdentityError,
    identify_business,
    normalize_business_text,
    validate_business_registry,
)
from risk_audit.inventory import classify_material
from risk_audit.inventory_v180 import scan_package_v180
from risk_audit.models import Entity
from risk_audit.runner import material_naming_keywords
from risk_audit.util import read_json, sha256_file


@pytest.fixture
def business_registry() -> dict:
    """返回包含重复展示编号和电压变体的最小业务配置。"""
    return {
        "schema_version": "2.0",
        "businesses": [
            {
                "business_id": "sales_electricity",
                "business_code": "01",
                "business_name": "营销售电",
                "aliases": ["营销售电"],
                "variants": [{"variant_id": "default", "aliases": [], "template": {
                    "path": "first-01.xlsx", "sha256": "a" * 64,
                }}],
            },
            {
                "business_id": "equity_management",
                "business_code": "01",
                "business_name": "股权（产权）管理",
                "aliases": ["股权（产权）管理", "股权产权管理", "股权管理"],
                "variants": [{"variant_id": "default", "aliases": [], "template": {
                    "path": "second-01.xlsx", "sha256": "b" * 64,
                }}],
            },
            {
                "business_id": "grid_construction",
                "business_code": "03",
                "business_name": "电网基建",
                "aliases": ["电网基建"],
                "variants": [
                    {"variant_id": "35-220kv", "aliases": ["35kv-220kv", "35千伏-220千伏"],
                     "template": {"path": "grid-low.xlsx", "sha256": "c" * 64}},
                    {"variant_id": "500-750kv", "aliases": ["500kv-750kv", "500千伏-750千伏"],
                     "template": {"path": "grid-high.xlsx", "sha256": "d" * 64}},
                ],
            },
            {
                "business_id": "contract_management",
                "business_code": "03",
                "business_name": "合同管理",
                "aliases": ["合同管理"],
                "variants": [{"variant_id": "default", "aliases": [], "template": {
                    "path": "contract.xlsx", "sha256": "e" * 64,
                }}],
            },
            {
                "business_id": "education_training",
                "business_code": "28",
                "business_name": "教育培训",
                "aliases": ["教育培训"],
                "variants": [{"variant_id": "default", "aliases": [], "template": {
                    "path": "education.xlsx", "sha256": "f" * 64,
                }}],
            },
            {
                "business_id": "distribution_grid_project_management",
                "business_code": "05",
                "business_name": "配网工程管理",
                "aliases": ["配网工程管理", "配网工程"],
                "variants": [{"variant_id": "default", "aliases": [], "template": {
                    "path": "distribution-grid.xlsx", "sha256": "2" * 64,
                }}],
            },
        ],
    }


def test_normalize_business_text_ignores_spacing_parentheses_hyphens_and_case():
    """不同括号、分隔符、空白和大小写不得改变业务匹配文本。"""
    assert normalize_business_text(" 股权（ 产权 ）管理-A ") == normalize_business_text("股权(产权)管理a")


@pytest.mark.parametrize(
    ("relative_path", "business_id", "business_code", "variant_id"),
    [
        ("第一批/01 营销售电/01风控矩阵.xlsx", "sales_electricity", "01", "default"),
        ("第二批/01 股权(产权)管理-省公司版/01三清单.xlsx", "equity_management", "01", "default"),
        ("第二批/28 教育 培训—省公司版/28风控矩阵.xlsx", "education_training", "28", "default"),
        ("第一批/03 电网基建 35KV—220KV/03风控矩阵.xlsx", "grid_construction", "03", "35-220kv"),
        ("第一批/03 电网基建 500千伏-750千伏/03风控矩阵.xlsx", "grid_construction", "03", "500-750kv"),
        ("第二批/03 合同管理/03风控矩阵.xlsx", "contract_management", "03", "default"),
    ],
)
def test_identify_business_uses_name_not_display_code(
    business_registry: dict,
    relative_path: str,
    business_id: str,
    business_code: str,
    variant_id: str,
):
    """relative_path 为报送路径，其余参数为期望身份；重复编号必须按名称隔离。"""
    result = identify_business(Path(relative_path), business_registry)
    assert (result.business_id, result.business_code, result.variant_id) == (
        business_id,
        business_code,
        variant_id,
    )


def test_longest_alias_resolves_containment(business_registry: dict):
    """路径同时命中“股权管理”和完整名称时，必须选择最长别名。"""
    result = identify_business(Path("01 股权（产权）管理/股权管理矩阵.xlsx"), business_registry)
    assert result.business_id == "equity_management"
    assert result.matched_alias == "股权（产权）管理"


def test_identify_business_uses_direct_parent_directory_code(business_registry: dict):
    """父目录只有业务编号和简称时，必须通过唯一编号识别业务。"""
    result = identify_business(
        Path("本部/05配网/附件：三清单模板-本部927.xlsx"),
        business_registry,
    )

    assert (result.business_id, result.business_code, result.business_name) == (
        "distribution_grid_project_management",
        "05",
        "配网工程管理",
    )


def test_identify_business_uses_selected_root_for_direct_package_variant(business_registry: dict):
    """业务包根目录包含模板变体时，必须同时识别业务及变体。"""
    result = identify_business(
        Path("03风控矩阵.xlsx"),
        business_registry,
        "测试主体-03 电网基建 35千伏-220千伏",
    )

    assert (result.business_id, result.business_code, result.variant_id) == (
        "grid_construction",
        "03",
        "35-220kv",
    )


def test_identify_business_uses_directory_when_filename_conflicts(
    business_registry: dict,
):
    """文件名指向其他业务时，业务目录仍是唯一归属依据。"""
    result = identify_business(
        Path("01 营销售电/股权产权管理矩阵.xlsx"),
        business_registry,
    )

    assert (result.business_id, result.business_code) == ("sales_electricity", "01")


def test_identify_business_rejects_duplicate_directory_code(business_registry: dict):
    """父目录只命中重复业务编号时不得猜测具体业务。"""
    with pytest.raises(BusinessIdentityError, match="业务编号映射冲突"):
        identify_business(Path("01业务/清单.xlsx"), business_registry)


def test_identify_business_rejects_filename_only_match_when_directory_is_unknown(
    business_registry: dict,
):
    """父目录无法识别时，文件名不得替代业务目录推定归属。"""
    with pytest.raises(BusinessIdentityError, match="未找到业务模板映射"):
        identify_business(Path("未知目录/营销售电矩阵.xlsx"), business_registry)


def test_identify_business_rejects_workbook_only_match_when_directory_is_unknown(
    business_registry: dict,
):
    """工作表名和表内标题不得替代业务目录推定归属。"""
    with pytest.raises(BusinessIdentityError, match="未找到业务模板映射"):
        identify_business(
            Path("未知目录/报送资料.xlsx"),
            business_registry,
            workbook_evidence=["股权（产权）管理矩阵", "岗位清单"],
        )


def test_identify_business_uses_directory_for_variant_when_file_conflicts(
    business_registry: dict,
) -> None:
    """文件名与业务目录的变体冲突时，必须使用目录中的变体。"""
    result = identify_business(
        Path("03 电网基建 35千伏-220千伏/500千伏-750千伏矩阵.xlsx"),
        business_registry,
        workbook_evidence=["500千伏-750千伏电网基建"],
    )

    assert result.variant_id == "35-220kv"


def test_workbook_business_evidence_excludes_body_rows(tmp_path, business_registry: dict) -> None:
    """工作簿证据读取不得把正文误当标题信息。"""

    from risk_audit.business_identity import workbook_business_evidence

    path = tmp_path / "股权产权管理矩阵.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "股权产权管理"
    sheet.append(["控制措施编号", "控制措施", "责任主体"])
    sheet.append(["股权管理-1-控制措施01", "执行合同管理制度", "财务部"])
    workbook.save(path)

    evidence = workbook_business_evidence(path, "xlsx")

    assert "股权产权管理" in evidence
    assert all("合同管理制度" not in item for item in evidence)


def test_workbook_business_evidence_keeps_title_with_header_keyword(tmp_path) -> None:
    """包含“岗位职责”的业务标题仍应进入模板匹配证据。"""

    from risk_audit.business_identity import workbook_business_evidence

    path = tmp_path / "报送清单.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "岗位清单"
    sheet.append(["财务管理岗位职责清单"])
    sheet.append(["控制措施编号", "岗位职责"])
    workbook.save(path)

    evidence = workbook_business_evidence(path, "xlsx")

    assert "财务管理岗位职责清单" in evidence


def test_unknown_nested_child_does_not_inherit_registered_grandparent(
    tmp_path,
    business_registry: dict,
) -> None:
    """未知业务下的未登记子单位不得继承上级正式主体代码。"""

    from risk_audit.inventory_v180 import scan_package_v180
    from risk_audit.models import Entity

    path = tmp_path / "已登记父公司有限公司" / "未登记子单位" / "未知业务" / "风控矩阵.xlsx"
    path.parent.mkdir(parents=True)
    workbook = Workbook()
    workbook.active.append(["控制措施编号", "控制措施", "责任主体"])
    workbook.save(path)

    files = scan_package_v180(
        tmp_path,
        {"P001": Entity("P001", "已登记父公司有限公司")},
        {},
        business_registry,
    )

    assert len(files) == 1
    assert files[0].entity_code is None


def test_direct_business_directory_uses_registered_parent_as_subject(
    tmp_path,
    business_registry: dict,
) -> None:
    """材料直接父目录按业务处理，其上级登记目录必须作为主体。"""

    from risk_audit.inventory_v180 import scan_package_v180
    from risk_audit.models import Entity

    path = tmp_path / "已登记父公司有限公司" / "未登记子单位" / "财务管理矩阵.xlsx"
    path.parent.mkdir(parents=True)
    workbook = Workbook()
    workbook.active.title = "财务管理"
    workbook.active.append(["控制措施编号", "控制措施", "责任主体"])
    workbook.save(path)

    files = scan_package_v180(
        tmp_path,
        {"P001": Entity("P001", "已登记父公司有限公司")},
        {},
        business_registry,
    )

    assert len(files) == 1
    assert files[0].entity_code == "P001"


def test_scan_package_does_not_use_sheet_name_or_title_for_template_match(
    tmp_path: Path,
    business_registry: dict,
):
    """目录无法识别时，工作表名和标题不得为业务补充归属。"""
    entity_name = "测试供电分公司"
    directory = tmp_path / entity_name / "未知业务"
    directory.mkdir(parents=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "股权（产权）管理矩阵"
    worksheet["A1"] = "股权（产权）管理"
    workbook.save(directory / "风控矩阵.xlsx")

    files = scan_package_v180(
        tmp_path,
        {"E001": Entity("E001", entity_name)},
        {},
        business_registry,
    )

    assert files[0].business_id.startswith("unmatched_")
    assert files[0].business_code is None
    assert files[0].preservation["template_match"]["status"] == "unmatched"
    assert files[0].parse_errors == []


def test_scan_package_keeps_unmatched_workbook_auditable(
    tmp_path: Path,
    business_registry: dict,
):
    """没有任何业务证据时，文件应保留审核资格而不记为解析失败。"""
    entity_name = "测试供电分公司"
    directory = tmp_path / entity_name / "未知业务"
    directory.mkdir(parents=True)
    workbook = Workbook()
    workbook.active.title = "矩阵"
    workbook.active["A1"] = "风控矩阵"
    workbook.save(directory / "风控矩阵.xlsx")

    files = scan_package_v180(
        tmp_path,
        {"E001": Entity("E001", entity_name)},
        {},
        business_registry,
    )

    assert files[0].business_id.startswith("unmatched_")
    assert files[0].business_code is None
    assert files[0].preservation["template_match"]["status"] == "unmatched"
    assert files[0].parse_errors == []


@pytest.mark.parametrize(
    ("relative_path", "expected_business_id"),
    [
        ("本部/02营销购电矩阵/附件清单.xlsx", "power_purchase_trading"),
        ("本部/08 数字化与研发/附件矩阵.xlsx", "digitalization_rd_investment"),
    ],
)
def test_active_registry_recognizes_submission_directory_aliases(
    relative_path: str,
    expected_business_id: str,
):
    """relative_path 为实际报送目录；活动规则包必须覆盖业务目录简称。"""
    registry_path = Path(__file__).resolve().parents[1] / "rulepacks/releases/1.9.20/baseline_registry.json"
    registry = read_json(registry_path)

    result = identify_business(Path(relative_path), registry)

    assert result.business_id == expected_business_id


def test_equal_length_business_alias_conflict_is_rejected(business_registry: dict):
    """两个业务的最长命中别名同长度时不得猜测。"""
    business_registry["businesses"].append({
        "business_id": "other_training",
        "business_code": "99",
        "business_name": "培训教育",
        "aliases": ["培训教育"],
        "variants": [{"variant_id": "default", "aliases": [], "template": {
            "path": "other.xlsx", "sha256": "1" * 64,
        }}],
    })
    with pytest.raises(BusinessIdentityError, match="业务模板映射冲突.*baseline_registry.json"):
        identify_business(Path("教育培训与培训教育/矩阵.xlsx"), business_registry)


def test_missing_business_and_variant_are_rejected(business_registry: dict):
    """未知业务或电网基建缺少电压信息时必须停止映射。"""
    with pytest.raises(BusinessIdentityError, match="未找到业务模板映射"):
        identify_business(Path("未知业务/01风控矩阵.xlsx"), business_registry)
    with pytest.raises(BusinessIdentityError, match="模板变体无法确定"):
        identify_business(Path("03 电网基建/03风控矩阵.xlsx"), business_registry)


def test_multiple_non_default_variants_are_rejected_even_when_alias_lengths_differ(business_registry: dict):
    """同一业务命中两个非默认变体时不得用最长别名替用户猜测。"""
    grid = next(item for item in business_registry["businesses"] if item["business_id"] == "grid_construction")
    grid["variants"][0]["aliases"].append("35kv")

    with pytest.raises(BusinessIdentityError, match="模板变体映射冲突"):
        identify_business(Path("03 电网基建 35kv与500kv-750kv/矩阵.xlsx"), business_registry)


@pytest.mark.parametrize(
    ("filename", "expected"),
    [
        ("附件清单.xlsx", "three_lists"),
        ("附件三清单.xlsx", "three_lists"),
        ("内部控制矩阵.xlsx", "matrix"),
        ("风控矩阵.xlsx", "matrix"),
        ("矩阵及三清单.xlsx", "three_lists"),
        ("普通附件.xlsx", None),
    ],
)
def test_classify_material_uses_filename_keywords(filename: str, expected: str | None):
    """表格材料类型必须只由文件名中的清单或矩阵关键字决定。"""
    assert classify_material(Path(filename)) == expected


@pytest.mark.parametrize(
    "filename",
    [
        "系统规则.xlsx",
        "岗位内控-交易与购电.xlsx",
        "不相容-北碚公司.xlsx",
    ],
)
def test_classify_material_uses_configured_chongqing_list_keywords(filename: str):
    """重庆独立清单省略“清单”后缀时，应按地区配置识别为三清单。"""
    naming_keywords = {
        "matrix": ["风控矩阵", "矩阵", "取证"],
        "three_lists": ["三清单", "系统规则", "岗位内控", "不相容"],
    }

    assert classify_material(Path(filename), naming_keywords=naming_keywords) == "three_lists"


def test_default_material_keywords_do_not_change_hunan_matching() -> None:
    """未传入地区配置时不得扩大原有湖南材料识别范围。"""
    assert classify_material(Path("系统规则.xlsx")) is None


def test_hunan_rulepack_does_not_enable_material_keyword_expansion() -> None:
    """湖南 1.9.21 即使有命名关键字，也不得用于扩大材料扫描范围。"""
    project_root = Path(__file__).resolve().parents[1]
    hunan_pack = read_json(project_root / "rulepacks/releases/1.9.21/manifest.json")
    pack = {
        "manifest": hunan_pack,
        "rules": [{
            "checks": [{
                "operator": "required_documents",
                "params": {
                    "naming_mode": True,
                    "naming_keywords": {"matrix": ["矩阵", "取证"], "three_lists": ["三清单"]},
                },
            }],
        }],
    }

    assert material_naming_keywords(pack) is None


def test_classify_material_uses_chongqing_workbook_title_fallback() -> None:
    """重庆文件名无材料关键字时，应使用工作簿标题匹配矩阵。"""
    naming_keywords = {
        "matrix": ["风控矩阵", "矩阵", "取证"],
        "three_lists": ["三清单", "系统规则", "岗位内控", "不相容"],
    }

    assert classify_material(
        Path("北碚-违章罚款收入管理.xlsx"),
        naming_keywords=naming_keywords,
        workbook_evidence=[
            "违章罚款收入管理",
            "风控矩阵-财务管理-违章罚款收入管理",
            "7.会计不相容岗位管控",
        ],
    ) == "matrix"


@pytest.mark.parametrize("worksheet_title", ["系统规则", "岗位内控-交易与购电", "不相容-交易与购电"])
def test_chongqing_scanner_uses_fuzzy_workbook_titles(tmp_path: Path, worksheet_title: str) -> None:
    """文件名未命中时，重庆应使用工作表短标题识别独立清单。

    Args:
        tmp_path: pytest 隔离目录。
        worksheet_title: 重庆独立清单的模糊工作表名。
    """

    naming_keywords = {
        "matrix": ["风控矩阵", "矩阵", "取证"],
        "three_lists": ["三清单", "系统规则", "岗位内控", "不相容"],
    }
    workbook = Workbook()
    workbook.active.title = worksheet_title
    workbook.save(tmp_path / "附件.xlsx")

    files = scan_package_v180(
        tmp_path,
        {},
        {},
        naming_keywords=naming_keywords,
        use_workbook_evidence=True,
    )

    assert len(files) == 1
    assert files[0].material_type == "three_lists"


def test_matrix_incompatible_position_header_is_not_classified_as_three_lists(tmp_path: Path) -> None:
    """矩阵表头中的“不相容岗位”不得被当成不相容岗位清单标题。

    Args:
        tmp_path: pytest 隔离目录。
    """

    naming_keywords = {
        "matrix": ["风控矩阵", "矩阵", "取证"],
        "three_lists": ["三清单", "系统规则", "岗位内控", "不相容"],
    }
    workbook = Workbook()
    workbook.active.title = "交易与购电"
    workbook.active.append(["控制措施编号", "控制措施", "不相容岗位"])
    workbook.save(tmp_path / "附件.xlsx")

    files = scan_package_v180(
        tmp_path,
        {},
        {},
        naming_keywords=naming_keywords,
        use_workbook_evidence=True,
    )

    assert files == []


def test_validate_business_registry_rejects_unstable_or_ambiguous_configuration(business_registry: dict):
    """空身份、纯数字别名、重复变体和重复模板路径都必须被配置校验拒绝。"""
    duplicate = business_registry["businesses"][0].copy()
    duplicate["aliases"] = ["01"]
    duplicate["variants"] = [
        {"variant_id": "default", "aliases": [], "template": {"path": "same.xlsx", "sha256": "a" * 64}},
        {"variant_id": "default", "aliases": [], "template": {"path": "same.xlsx", "sha256": "a" * 64}},
    ]
    invalid = {"schema_version": "2.0", "businesses": [business_registry["businesses"][0], duplicate]}
    errors = validate_business_registry(invalid)
    assert any("business_id 重复" in error for error in errors)
    assert any("不允许纯数字" in error for error in errors)
    assert any("variant_id 重复" in error for error in errors)
    assert any("模板路径重复" in error for error in errors)


@pytest.mark.parametrize("scanner_name", ["current", "confirmed"])
def test_scanner_uses_filename_only_for_spreadsheet_material(
    tmp_path,
    business_registry: dict,
    scanner_name: str,
):
    """scanner_name 为扫描器版本；无关键字表格即使内容像矩阵也必须忽略。"""
    from risk_audit.inventory import scan_package
    from risk_audit.inventory_v180 import scan_package_v180
    from risk_audit.models import Entity

    business_directory = tmp_path / "测试主体有限公司" / "05配网"
    business_directory.mkdir(parents=True)

    three_lists = Workbook()
    three_lists.active.append(["部门", "岗位名称", "岗位职责"])
    three_lists.save(business_directory / "附件清单.xlsx")

    matrix = Workbook()
    matrix.active.append(["控制措施编号", "控制措施"])
    matrix.save(business_directory / "业务矩阵.xlsx")
    matrix.save(business_directory / "普通附件.xlsx")

    scanner = scan_package if scanner_name == "current" else scan_package_v180
    files = scanner(
        tmp_path,
        {"A001": Entity("A001", "测试主体有限公司")},
        {},
        business_registry,
    )

    assert {(file.relative_path.name, file.material_type) for file in files} == {
        ("附件清单.xlsx", "three_lists"),
        ("业务矩阵.xlsx", "matrix"),
    }
    assert all(file.business_id == "distribution_grid_project_management" for file in files)


def test_confirmed_scanner_carries_unique_business_identity(tmp_path, business_registry: dict):
    """tmp_path 为报送根目录；同为 01 的两个业务必须生成不同业务 ID。"""
    from risk_audit.inventory_v180 import scan_package_v180
    from risk_audit.models import Entity

    unit = tmp_path / "测试主体有限公司"
    paths = [
        unit / "01 营销售电" / "01风控矩阵-测试主体有限公司.xlsx",
        unit / "01 股权（产权）管理" / "01风控矩阵-测试主体有限公司.xlsx",
        unit / "28 教育培训" / "28风控矩阵-测试主体有限公司.xlsx",
    ]
    for path in paths:
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        workbook.active.append(["控制措施编号", "控制措施", "责任主体"])
        workbook.save(path)

    files = scan_package_v180(
        tmp_path,
        {"A001": Entity("A001", "测试主体有限公司")},
        {},
        business_registry,
    )

    identities = {
        file.relative_path.parent.name: (file.business_id, file.business_code, file.variant_id)
        for file in files
    }
    assert identities == {
        "01 营销售电": ("sales_electricity", "01", "default"),
        "01 股权（产权）管理": ("equity_management", "01", "default"),
        "28 教育培训": ("education_training", "28", "default"),
    }
    assert all(not any("业务模板" in error or "业务编号" in error for error in file.parse_errors) for file in files)


@pytest.mark.parametrize("scanner_name", ["current", "confirmed"])
def test_selected_business_root_does_not_guess_subject_from_business_name(tmp_path, scanner_name: str):
    """直接选择业务目录时可识别业务，但不得从同一目录名称猜测主体。"""
    from risk_audit.inventory import scan_package
    from risk_audit.inventory_v180 import scan_package_v180
    from risk_audit.models import Entity

    project_root = Path(__file__).resolve().parents[2]
    registry = read_json(project_root / "risk-audit/rulepacks/releases/1.9.20/baseline_registry.json")
    entity_name = "湖南星通电力信息通信有限公司"
    package_root = tmp_path / f"{entity_name}第二批-02 财务管理风控矩阵（含三清单）"
    package_root.mkdir()

    matrix = Workbook()
    matrix.active.append(["控制措施编号", "控制措施", "责任主体"])
    matrix.save(package_root / "02风控矩阵-财务管理--星通公司.xlsx")

    three_lists = Workbook()
    three_lists.active.append(["部门", "岗位名称", "岗位职责"])
    three_lists.save(package_root / "“三清单”--02 财务管理-星通公司.xlsx")

    scanner = scan_package if scanner_name == "current" else scan_package_v180
    files = scanner(
        package_root,
        {"A001": Entity("A001", entity_name)},
        {},
        registry,
    )

    assert len(files) == 2
    assert {(file.business_id, file.business_code, file.variant_id) for file in files} == {
        ("financial_management", "02", "default"),
    }
    assert {file.entity_code for file in files} == {None}
    assert all(not any("业务模板" in error for error in file.parse_errors) for file in files)


@pytest.mark.parametrize("scanner_name", ["current", "confirmed"])
def test_selected_business_root_ignores_embedded_entity_names(tmp_path, scanner_name: str):
    """业务目录名称中的多个单位名称不得作为主体证据或制造冲突。"""
    from risk_audit.inventory import scan_package
    from risk_audit.inventory_v180 import scan_package_v180
    from risk_audit.models import Entity

    project_root = Path(__file__).resolve().parents[2]
    registry = read_json(project_root / "risk-audit/rulepacks/releases/1.9.20/baseline_registry.json")
    package_root = tmp_path / "已登记主体甲有限公司与已登记主体乙有限公司-02 财务管理"
    package_root.mkdir()
    workbook = Workbook()
    workbook.active.append(["控制措施编号", "控制措施", "责任主体"])
    workbook.save(package_root / "02风控矩阵.xlsx")

    scanner = scan_package if scanner_name == "current" else scan_package_v180
    files = scanner(
        package_root,
        {
            "A001": Entity("A001", "已登记主体甲有限公司"),
            "B001": Entity("B001", "已登记主体乙有限公司"),
        },
        {},
        registry,
    )

    assert len(files) == 1
    assert files[0].entity_code is None
    assert not files[0].entity_conflict
    assert "无法确定会计主体" in files[0].parse_errors


def test_load_baselines_isolates_duplicate_display_codes(tmp_path):
    """tmp_path 为项目根目录；两个 01 模板必须按业务 ID 同时加载。"""
    from risk_audit.baselines import load_baselines

    templates = []
    for filename, responsibility in [
        ("sales.xlsx", "省公司-营销部-营销专责"),
        ("equity.xlsx", "省公司-财务部-产权专责"),
    ]:
        path = tmp_path / filename
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "风控矩阵"
        sheet.append(["控制措施编号", "控制措施", "控制载体", "责任主体"])
        sheet.append(["M1", "核对资料", "审批单", responsibility])
        workbook.save(path)
        templates.append((path, responsibility))
    registry = {
        "schema_version": "2.0",
        "businesses": [
            {
                "business_id": business_id,
                "business_code": "01",
                "business_name": business_name,
                "aliases": [business_name],
                "variants": [{"variant_id": "default", "aliases": [], "template": {
                    "path": path.name,
                    "sha256": sha256_file(path),
                }}],
            }
            for (business_id, business_name), (path, _) in zip(
                [("sales_electricity", "营销售电"), ("equity_management", "股权产权管理")],
                templates,
            )
        ],
    }
    aliases = read_json(Path(__file__).resolve().parents[1] / "rulepacks/releases/1.9.18/field_aliases.json")

    baselines = load_baselines(tmp_path, registry, aliases, tmp_path / "work")

    assert set(baselines) == {
        ("sales_electricity", "default"),
        ("equity_management", "default"),
    }
    assert baselines[("sales_electricity", "default")]["business_code"] == "01"
    assert baselines[("equity_management", "default")]["business_code"] == "01"
    assert baselines[("sales_electricity", "default")]["responsibilities"][0]["measure_id"] == "M1"
    assert baselines[("equity_management", "default")]["responsibilities"][0]["measure_id"] == "M1"
    assert baselines[("sales_electricity", "default")]["matrix_rows"] == [{
        "measure_id": "M1",
        "key_control_point": "",
        "responsibility": "省公司-营销部-营销专责",
    }]


def test_unreadable_baseline_degrades_to_unavailable_entry(tmp_path) -> None:
    """冻结模板不存在时应记录不可用状态，并允许其他规则继续审核。"""

    from risk_audit.baselines import load_baselines

    registry = {
        "schema_version": "2.0",
        "businesses": [{
            "business_id": "financial_management",
            "business_code": "02",
            "business_name": "财务管理",
            "aliases": ["财务管理"],
            "variants": [{
                "variant_id": "default",
                "aliases": [],
                "template": {
                    "path": "不存在的模板.xlsx",
                    "sha256": "a" * 64,
                },
            }],
        }],
    }
    aliases = read_json(
        Path(__file__).resolve().parents[1]
        / "rulepacks/releases/1.9.18/field_aliases.json"
    )

    baselines = load_baselines(tmp_path, registry, aliases, tmp_path / "work")

    baseline = baselines[("financial_management", "default")]
    assert baseline["available"] is False
    assert "冻结基准不存在" in baseline["error"]
