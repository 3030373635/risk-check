"""按单位、业务、规则编号和意见内容生成审核统计表。"""
from __future__ import annotations

from collections import defaultdict
import logging
import os
from pathlib import Path
import re
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from risk_audit.audit_statistics_html import render_audit_statistics_html
from risk_audit.models import FileRecord, Finding
from risk_audit.util import natural_key
from risk_audit.writer import _output_path


RULE_LABEL_PATTERN = re.compile(r'【第(\d+)条】')
PREPROCESSING_LABEL = '【预处理】'
MATERIAL_LABEL = '【资料级】'
COUNT_SUFFIX_PATTERN = re.compile(r'\s+\d+\s*条[.。]?\s*$')
BRANCH_KEY_POINT_SUMMARIES = {
    'missing_key_points': '【第17条】缺少供电分公司适用清单要求的关键控制点',
    'extra_key_point': '【第17条】存在供电分公司适用清单外的关键控制点',
}


def _safe_excel_text(value: str) -> str:
    """将外部文本安全写入 Excel；value 为目录名或审核意见文本。"""
    return f"'{value}" if value.startswith(('=', '+', '-', '@', '\t', '\r')) else value


def _unit_and_business(
    file: FileRecord,
    package_name: str,
    root_subject_name: str,
) -> tuple[str, str] | None:
    """按主体、业务、文件层级计算展示位置。

    Args:
        file: 当前业务文件。
        package_name: 审核包根目录名。
        root_subject_name: 输入根本身是业务目录时的外层主体目录名。
    """

    business_directory = file.business_directory or file.relative_path.parent
    if business_directory == Path('.'):
        business_name = str(file.preservation.get('business_name') or package_name)
        return root_subject_name or '未识别主题', business_name
    # 扫描阶段已排除“三清单”等附件子目录，统计不再自行猜测目录层级。
    business_name = business_directory.name
    unit_directories = business_directory.parts[:-1]
    unit_name = '/'.join(part for part in (package_name, *unit_directories) if part)
    return unit_name, business_name


def _business_group_key(file: FileRecord, business_name: str) -> tuple[str, str]:
    """生成稳定业务分组键。

    Args:
        file: 当前业务文件。
        business_name: 当前文件解析出的业务显示名称。
    """

    del business_name
    # 统计与审核运行器共用标准业务身份，物理目录只用于展示。
    business_identity = file.business_id or file.business_code or ''
    identity = f'business:{file.entity_code}:{business_identity}'
    return identity, file.variant_id or 'default'


def _register_unit_group(
    grouped: dict,
    file: FileRecord,
    package_name: str,
    root_subject_name: str,
) -> tuple[str, dict] | None:
    """按稳定主体身份登记统计外层分组。

    Args:
        grouped: 按主体和业务归集的统计容器。
        file: 当前业务文件。
        package_name: 审核包名称。
        root_subject_name: 输入根本身是业务目录时的外层主体名。

    Returns:
        稳定主体键及其统计容器；无法定位时返回 None。
    """

    location = _unit_and_business(file, package_name, root_subject_name)
    if location is None:
        return None
    unit_name, _ = location
    unit_identity = f'entity:{file.entity_code}' if file.entity_code else f'path:{unit_name}'
    unit = grouped.setdefault(
        unit_identity,
        {'name': unit_name, 'source_names': set(), 'businesses': {}},
    )
    unit['source_names'].add(unit_name)
    unit['name'] = '、'.join(sorted(unit['source_names'], key=natural_key))
    return unit_identity, unit


def _register_grouped_business(
    grouped: dict,
    file: FileRecord | None,
    package_name: str,
    root_subject_name: str,
) -> tuple[str, dict] | None:
    """登记文件所属业务并返回主体及业务统计容器。

    Args:
        grouped: 按主体和稳定业务身份归集的统计容器。
        file: 当前业务文件；无法定位文件时为 None。
        package_name: 审核包名称。
        root_subject_name: 输入根本身是业务目录时的外层主体目录名。
    """

    if file is None:
        return None
    location = _unit_and_business(file, package_name, root_subject_name)
    unit_registration = _register_unit_group(grouped, file, package_name, root_subject_name)
    if location is None or unit_registration is None:
        return None
    _, business_name = location
    unit_key, unit = unit_registration
    business_key = _business_group_key(file, business_name)
    business = unit['businesses'].setdefault(
        business_key,
        {'name': business_name, 'source_names': set(), 'items': {}},
    )
    business['source_names'].add(business_name)
    business['name'] = '、'.join(sorted(business['source_names'], key=natural_key))
    return unit_key, business


def _register_material_finding_business(
    grouped: dict,
    files: list[FileRecord],
    host_file: FileRecord,
    finding: Finding,
    package_name: str,
    root_subject_name: str,
) -> tuple[str, dict] | None:
    """按资料级意见的真实业务登记统计分组。

    Args:
        grouped: 按主体和业务归集的统计容器。
        files: 本次审核扫描到的全部文件。
        host_file: 实际承载资料级意见的文件。
        finding: 与当前意见对应的资料级审核结果。
        package_name: 审核包名称。
        root_subject_name: 输入根本身是业务目录时的外层主体目录名。
    """

    if finding.check_id != 'power_supply_branch_key_points' or not finding.business_id:
        return None
    business_file_path = str(finding.evidence.get('business_file_path') or '')
    # unmatched_* 业务已在规则中按目录归一，统计时用原文件路径找回现有分组。
    matching_file = next((
        file for file in files
        if file.entity_code == finding.entity_code
        and (file.variant_id or 'default') == (finding.variant_id or 'default')
        and (
            file.business_id == finding.business_id
            or business_file_path in {
                str(file.relative_path),
                str(_output_path(file)),
            }
        )
    ), None)
    if matching_file is not None:
        return _register_grouped_business(
            grouped,
            matching_file,
            package_name,
            root_subject_name,
        )

    unit_registration = _register_unit_group(
        grouped,
        host_file,
        package_name,
        root_subject_name,
    )
    if unit_registration is None:
        return None
    unit_key, unit = unit_registration
    business_name = str(finding.evidence.get('business_name') or finding.business_id)
    business_key = (f'finding:{finding.business_id}', finding.variant_id or 'default')
    business = unit['businesses'].setdefault(
        business_key,
        {'name': business_name, 'items': {}},
    )
    return unit_key, business


def _finding_for_occurrence(
    message: str,
    state: dict,
    file: FileRecord,
    findings_by_message: dict[str, list[Finding]],
) -> Finding | None:
    """查找与统计表单元格意见对应的结构化审核结果。

    Args:
        message: 当前拆分后的程序意见。
        state: 当前意见单元格归属记录。
        file: 当前意见所在的业务文件。
        findings_by_message: 按完整意见文本索引的审核结果。
    """

    candidates = findings_by_message.get(message, [])
    row_match = re.search(r'\d+', str(state.get('cell') or ''))
    row_number = int(row_match.group()) if row_match else None
    exact = next((
        finding for finding in candidates
        if finding.file_path == str(file.relative_path)
        and (not finding.sheet or finding.sheet == state.get('sheet', ''))
        and (finding.row is None or finding.row == row_number)
    ), None)
    if exact is not None:
        return exact
    return next((
        finding for finding in candidates
        if (finding.location_policy == 'material' or not finding.file_path)
        and (not finding.entity_code or finding.entity_code == file.entity_code)
    ), None)


def _format_finding_key_point(point: list[int]) -> str:
    """格式化统计明细中的关键控制点编号。

    Args:
        point: 一个关键控制点的分级编号。
    """

    return '.'.join(str(value) for value in point)


def _branch_key_point_details(finding: Finding) -> list[str]:
    """生成第17条统计汇总下的关键控制点明细。

    Args:
        finding: 第17条供电分公司适用清单审核结果。
    """

    issue_type = finding.evidence.get('issue_type')
    if issue_type == 'missing_key_points':
        return [
            _format_finding_key_point(point)
            for point in finding.evidence.get('missing_key_points') or []
        ]
    if issue_type == 'extra_key_point':
        point_name = str(finding.evidence.get('key_point_name') or '').strip()
        if not point_name:
            point_name = _format_finding_key_point(finding.evidence.get('key_point') or [])
        if finding.sheet and finding.row:
            return [f'{point_name}（{finding.sheet}第{finding.row}行）']
        return [point_name]
    return []


def _rule_numbers(message: str) -> tuple[int, ...]:
    """提取并规范规则序号组合；message 为已写入审核意见单元格的一条意见。"""
    numbers = [int(value) for value in RULE_LABEL_PATTERN.findall(message)]
    return tuple(sorted(set(numbers)))


def _opinion_group(message: str) -> tuple[str, tuple[int, ...]] | None:
    """返回统计分组。

    Args:
        message: 一条程序生成的审核意见。
    """

    rules = _rule_numbers(message)
    if rules:
        return "rule", rules
    if message.startswith(PREPROCESSING_LABEL):
        return "preprocessing", ()
    return None


def _normalized_message(
    message: str,
    group: tuple[str, tuple[int, ...]],
) -> str:
    """统一审核意见展示文本。

    Args:
        message: 原始程序意见。
        group: 意见类别及排序后的规则编号组合。
    """

    category, rules = group
    is_material = message.startswith(MATERIAL_LABEL)
    body = RULE_LABEL_PATTERN.sub('', message).strip()
    body = body.removeprefix(MATERIAL_LABEL).strip()
    body = body.removeprefix(PREPROCESSING_LABEL).strip()
    body = COUNT_SUFFIX_PATTERN.sub('', body).strip()
    prefix = (
        PREPROCESSING_LABEL
        if category == "preprocessing"
        else f'{MATERIAL_LABEL if is_material else ""}{"".join(f"【第{number}条】" for number in rules)}'
    )
    return f'{prefix}{body}'


def _program_opinions(program_text: str) -> list[str]:
    """拆分单元格内的程序意见；program_text 为写回器记录的完整程序文本。"""
    text = program_text.strip()
    if not text:
        return []
    starts = [
        match.start()
        for match in re.finditer(r'(?m)^(?:【资料级】【第\d+条】|【第\d+条】|【预处理】)', text)
    ]
    if not starts:
        return [text]
    return [text[start:starts[index + 1] if index + 1 < len(starts) else len(text)].strip()
            for index, start in enumerate(starts)]


def _collect_statistics(
    files: list[FileRecord],
    ownership: dict[str, list[dict]],
    package_name: str,
    business_failures: list[dict] | None = None,
    root_subject_name: str = '',
    findings: list[Finding] | None = None,
) -> dict:
    """归集审核统计。

    Args:
        files: 本次审核扫描到的文件。
        ownership: 实际写回审核意见的单元格归属。
        package_name: 审核包名称。
        business_failures: 已跳过业务的结构化失败信息。
        root_subject_name: 输入根本身是业务目录时的外层主体目录名。
        findings: 本次结构化审核结果，用于还原资料级意见的真实业务归属。
    """
    logger = logging.getLogger(__name__)
    grouped = defaultdict(dict)
    file_by_output = {
        str(_output_path(file)): file
        # 晚表头或非标准文件名可能在扫描时仍是 unclassified，实际写回记录仍应纳入。
        for file in files if file.material_type != 'explanation'
    }
    findings_by_message: dict[str, list[Finding]] = defaultdict(list)
    for finding in findings or []:
        findings_by_message[finding.message].append(finding)
    for file in files:
        if file.material_type == 'explanation' or not (file.business_id or file.business_code):
            continue
        # 先登记目录业务；未匹配标准规则的目录即使没有程序意见也须明确展示。
        _register_grouped_business(grouped, file, package_name, root_subject_name)
    for output_path, states in ownership.items():
        file = file_by_output.get(output_path)
        registration = _register_grouped_business(grouped, file, package_name, root_subject_name)
        if registration is None:
            logger.warning('审核统计表未能定位意见所属业务目录：文件=%s', output_path)
            continue
        _, business = registration
        for state in states:
            for message in _program_opinions(state.get('program_text', '')):
                finding = _finding_for_occurrence(
                    message,
                    state,
                    file,
                    findings_by_message,
                )
                routed_registration = (
                    _register_material_finding_business(
                        grouped,
                        files,
                        file,
                        finding,
                        package_name,
                        root_subject_name,
                    )
                    if finding is not None
                    and (finding.location_policy == 'material' or not finding.file_path)
                    else None
                )
                # 资料级意见按实际缺失业务展示，文件只承担写回载体职责。
                target_business = routed_registration[1] if routed_registration else business
                group = _opinion_group(message)
                if group is None:
                    logger.warning('审核统计表未能识别意见规则序号：文件=%s，意见=%s',
                                   output_path, message)
                    continue
                normalized_message = _normalized_message(message, group)
                issue_type = finding.evidence.get('issue_type') if finding is not None else None
                if (
                    finding is not None
                    and finding.check_id == 'power_supply_branch_key_points'
                    and issue_type in BRANCH_KEY_POINT_SUMMARIES
                ):
                    statistics_key = ('branch_key_points', issue_type)
                    item = target_business['items'].setdefault(
                        statistics_key,
                        {
                            'message': BRANCH_KEY_POINT_SUMMARIES[issue_type],
                            'occurrences': set(),
                            'details': set(),
                        },
                    )
                    item['details'].update(_branch_key_point_details(finding))
                    occurrence = (
                        output_path,
                        state.get('sheet', ''),
                        state.get('cell', ''),
                        message,
                    )
                    item['occurrences'].add(occurrence)
                    continue
                # 规则编号相同但整改要求不同的意见必须分别统计，不能仅按规则编号合并。
                statistics_key = (group[0], group[1], normalized_message)
                item = target_business['items'].setdefault(
                    statistics_key,
                    {'message': normalized_message, 'occurrences': set()},
                )
                occurrence = (output_path, state.get('sheet', ''), state.get('cell', ''), message)
                # ownership 是最终写回结果，相同单元格内的完全相同意见只计一次。
                item['occurrences'].add(occurrence)
    file_by_relative_path = {str(file.relative_path): file for file in files}
    for index, failure in enumerate(business_failures or []):
        failure_files = failure.get('files') or []
        failed_file = next(
            (file_by_relative_path[path] for path in failure_files if path in file_by_relative_path),
            None,
        )
        registration = _register_grouped_business(grouped, failed_file, package_name, root_subject_name)
        if registration is None:
            logger.warning('审核统计表未能定位跳过业务：%s', failure)
            continue
        _, business = registration
        # 序号保证同一业务出现多条失败记录时不会互相覆盖。
        business['items'][('failure', index)] = {'failure': failure}
    return grouped


def _business_records(items: dict) -> list[str]:
    """生成一个业务的展示记录；items 为该业务归集后的意见和失败信息。"""
    records: list[str] = []
    for group, item in sorted(items.items()):
        if group[0] == 'failure':
            failure = item['failure']
            records.extend([
                '【审核状态】已跳过，未完成审核',
                f'【失败阶段】{failure["stage"]}',
                f'【涉及文件】{"；".join(failure.get("files") or ["未定位具体文件"])}',
                f'【异常类型】{failure["error_type"]}',
                f'【跳过原因】{failure["reason"]}',
                f'【详细日志】{failure["log_reference"]}',
            ])
            continue
        if item.get('details'):
            details = sorted(item['details'], key=natural_key)
            records.append(
                f'{item["message"]} {len(details)}项：\n' + '\n'.join(details)
            )
            continue
        records.append(f'{item["message"]} {len(item["occurrences"])}条。')
    return records or ['无']


def _statistics_source_data(grouped: dict) -> dict[str, list[dict[str, object]]]:
    """转换为两种报告共用的数据结构；grouped 为按主体和业务归集的数据。"""
    units: list[dict[str, object]] = []
    for unit in sorted(grouped.values(), key=lambda item: natural_key(item['name'])):
        businesses = [
            {
                'name': business['name'],
                'records': _business_records(business['items']),
            }
            for business in sorted(
                unit['businesses'].values(),
                key=lambda item: natural_key(item['name']),
            )
        ]
        units.append({'path': unit['name'], 'businesses': businesses})
    return {'units': units}


def write_audit_statistics(
    path: Path,
    files: list[FileRecord],
    ownership: dict[str, list[dict]],
    package_name: str,
    business_failures: list[dict] | None = None,
    root_subject_name: str = '',
    findings: list[Finding] | None = None,
) -> None:
    """写入审核统计表。

    Args:
        path: 审核统计表输出路径。
        files: 本次审核扫描到的文件。
        ownership: 实际写回审核意见的单元格归属。
        package_name: 审核包名称。
        business_failures: 已跳过业务的结构化失败信息。
        root_subject_name: 输入根本身是业务目录时的外层主体目录名。
        findings: 本次结构化审核结果，用于还原资料级意见的真实业务归属。
    """
    grouped = _collect_statistics(
        files,
        ownership,
        package_name,
        business_failures,
        root_subject_name,
        findings,
    )
    source_data = _statistics_source_data(grouped)
    book = Workbook()
    sheet = book.active
    sheet.title = '审核统计表'
    sheet.append(['单位', '备注'])

    for unit in source_data['units']:
        for business in unit['businesses']:
            # 每个业务独占一行，避免主体全部内容超过 Excel 单元格字符上限。
            business_section = f'{business["name"]}：\n' + '\n'.join(business['records'])
            sheet.append([
                _safe_excel_text(unit['path']),
                _safe_excel_text(business_section),
            ])

    header_fill = PatternFill('solid', fgColor='305496')
    header_font = Font(bold=True, color='FFFFFF')
    header_border = Border(
        left=Side(style='thin', color='D9E2F3'),
        right=Side(style='thin', color='D9E2F3'),
        top=Side(style='thin', color='D9E2F3'),
        bottom=Side(style='thin', color='D9E2F3'),
    )
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.border = header_border
        cell.alignment = Alignment(horizontal='center', vertical='center')
    for row in sheet.iter_rows(min_row=2):
        row[0].alignment = Alignment(vertical='top')
        row[1].alignment = Alignment(wrap_text=True, vertical='top')
        line_count = max(2, str(row[1].value or '').count('\n') + 1)
        sheet.row_dimensions[row[0].row].height = min(409, line_count * 20)
    sheet.column_dimensions['A'].width = 36
    sheet.column_dimensions['B'].width = 100
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix='.audit-statistics-', suffix='.xlsx', dir=path.parent,
    )
    os.close(descriptor)
    temporary_path = Path(temporary_name)
    try:
        # 先写同目录临时文件，成功后原子替换，避免中途失败留下损坏工作簿。
        book.save(temporary_path)
        temporary_path.replace(path)
    finally:
        temporary_path.unlink(missing_ok=True)


def write_audit_statistics_html(
    path: Path,
    files: list[FileRecord],
    ownership: dict[str, list[dict]],
    package_name: str,
    business_failures: list[dict] | None = None,
    root_subject_name: str = '',
    findings: list[Finding] | None = None,
) -> None:
    """写入审核统计 HTML。

    Args:
        path: HTML 报告输出路径。
        files: 本次审核扫描到的文件。
        ownership: 实际写回审核意见的单元格归属。
        package_name: 审核包名称。
        business_failures: 已跳过业务的结构化失败信息。
        root_subject_name: 输入根本身是业务目录时的外层主体目录名。
        findings: 本次结构化审核结果，用于还原资料级意见的真实业务归属。
    """
    grouped = _collect_statistics(
        files,
        ownership,
        package_name,
        business_failures,
        root_subject_name,
        findings,
    )
    document = render_audit_statistics_html(_statistics_source_data(grouped))
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix='.audit-statistics-', suffix='.html', dir=path.parent,
    )
    os.close(descriptor)
    temporary_path = Path(temporary_name)
    try:
        # HTML 同样先写入同目录临时文件，避免异常中断留下半份报告。
        temporary_path.write_text(document, encoding='utf-8')
        temporary_path.replace(path)
    finally:
        temporary_path.unlink(missing_ok=True)
