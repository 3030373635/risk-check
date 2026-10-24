"""冻结标准库的公司性质匹配、措施核对和结果表输出。"""

from __future__ import annotations

from collections import defaultdict
import hashlib
from pathlib import Path
import re
from typing import Any, Iterable
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from risk_audit.models import FileRecord, Record
from risk_audit.util import measure_id_key


MISSING_OPINION = "需核实标准库中适用的控制措施是否完整"
COMPANY_OPINION = "请核实公司名称与单位性质，无法唯一匹配附件2，无法完成标准库比对"
REPORT_HEADERS = ("矩阵名称", "矩阵类型", "控制措施编码", "重庆控制措施编号", "对应的单位性质", "标记为“是”的单元格", "提示内容")


def normalize_company_name(name: str) -> str:
    """标准化公司名称；name 为全称或带目录序号的报送名称。"""
    value = unicodedata.normalize("NFKC", name)
    value = re.sub(r"\s+", "", value)
    return re.sub(r"^\d+[.、_\-]+", "", value)


def company_aliases(name: str) -> set[str]:
    """生成可唯一核对的简称；name 为附件2中的公司全称，不使用模糊相似度。"""
    full_name = normalize_company_name(name)
    aliases = {full_name}
    short_name = re.sub(r"^(?:国网重庆市电力有限公司|国网重庆市电力公司|重庆市|重庆)", "", full_name)
    aliases.add(short_name)
    # 本部、分公司等层级信息保留；不把子公司简称归并为上级公司。
    if not re.search(r"\(本部\)|分公司", short_name):
        aliases.add(re.sub(r"(?:工程有限责任公司|工程有限公司|有限责任公司|有限公司)$", "", short_name))
    return {alias for alias in aliases if len(alias) >= 2}


def match_company(name: str, companies: list[dict[str, Any]]) -> dict[str, str]:
    """匹配公司性质；name 为报送名称，companies 为冻结的公司名称与性质记录。"""
    normalized = normalize_company_name(name)
    exact = [company for company in companies if normalize_company_name(company["company_name"]) == normalized]
    candidates = exact or [company for company in companies if normalized in company_aliases(company["company_name"])]
    if len(candidates) != 1:
        return {"match_status": "ambiguous" if candidates else "unmatched", "company_name": name, "unit_nature": "", "opinion": COMPANY_OPINION}
    return {"match_status": "matched", "company_name": name, "matched_company_name": candidates[0]["company_name"], "unit_nature": candidates[0]["unit_nature"], "opinion": ""}


def bind_company_groups(files: list[FileRecord], groups: dict[str, dict], library: dict) -> dict[str, dict]:
    """合并附件明确指向同一公司的报送目录。

    Args:
        files: 已按目录识别的全部材料，更新其主体分组及追溯信息。
        groups: 原始主体分组元数据。
        library: 冻结的公司名称与单位性质资源；未匹配或歧义主体保留原目录隔离。
    """
    merged = {}
    for file in files:
        identity = file.preservation.get('entity_identity', {})
        if identity.get('entity_registry_status') != 'directory_matched':
            merged.setdefault(file.entity_code, groups[file.entity_code])
            continue
        match = match_company(identity['entity_name'], library['companies'])
        key = '@重庆公司:' + match['matched_company_name'] if match['match_status'] == 'matched' else file.entity_code
        # 只有附件唯一确认的公司全称才能合并，不能按短名相似度跨主体借用资料。
        shared = merged.setdefault(key, {**identity, 'group_key': key, 'company_match': match})
        file.entity_code = key
        file.preservation['entity_identity'] = shared
    return merged


def compare_standard_library(
    library: dict[str, Any],
    company_name: str,
    records: Iterable[Record],
    *,
    entity_code: str,
    blocked_business_ids: Iterable[str] = (),
) -> dict[str, Any]:
    """按主体、业务及编号数字核对适用措施。

    Args:
        library: 包含 companies、measures 的冻结标准库。
        company_name: 本次报送主体名称。
        records: 同批已解析的矩阵和清单记录，仅矩阵参与措施存在性判断。
        entity_code: 当前主体分组标识，防止借用其他主体的措施。
        blocked_business_ids: 解析或审核失败的业务，不将其缺失记录认定为正常或缺项。

    Returns:
        公司匹配状态和全部适用措施结果；正常记录仅用于独立报告。
    """
    result: dict[str, Any] = {**match_company(company_name, library.get("companies", [])), "rows": []}
    if result["match_status"] != "matched":
        return result
    blocked = set(blocked_business_ids)
    present: dict[tuple[str, tuple], list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        key = measure_id_key(record.value("measure_id"))
        if record.record_type == "matrix" and record.entity_code == entity_code and key and all(isinstance(part, int) for part in key):
            present[(record.business_id or record.business_code, key)].append({"file": record.file_path, "sheet": record.sheet, "row": record.row})
    for measure in library.get("measures", []):
        cell = measure["applicable_cells"].get(result["unit_nature"])
        if not cell:
            continue
        locations = present.get((measure["business_id"], measure_id_key(measure["measure_id"])), [])
        unavailable = measure["business_id"] in blocked
        # 业务读取失败时，即使部分记录存在也不能宣称整项比对正常。
        status = "unavailable" if unavailable else "present" if locations else "missing"
        opinion = "该业务材料未完成读取或审核，无法完成比对" if unavailable else "无异常" if locations else MISSING_OPINION
        result["rows"].append({**measure, "unit_nature": result["unit_nature"], "applicable_cell": cell, "status": status, "opinion": opinion, "locations": locations})
    return result


def write_standard_library_report(output_root: Path, comparison: dict[str, Any], *, suffix: str = "") -> Path:
    """输出独立标准库结果表。

    Args:
        output_root: 本次审核副本根目录。
        comparison: compare_standard_library 返回的结果。
        suffix: 同名主体目录的稳定消歧后缀，默认无后缀。

    Returns:
        最终 Excel 路径；同一主体重跑会覆盖本次程序报告。
    """
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", comparison["company_name"]).strip(" .") or "未识别公司"
    if len(name) > 80:
        name = name[:70] + "-" + hashlib.sha256(name.encode()).hexdigest()[:8]
    path = output_root / f"标准库比对结果-{name}{suffix}.xlsx"
    if path.exists() and not path.is_file() and not path.is_symlink():
        raise ValueError(f"标准库报告路径不是文件：{path}")
    # 不跟随旧输出中的符号链接写入其他目录。
    if path.is_symlink():
        path.unlink()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "标准库比对结果"
    sheet.append(REPORT_HEADERS)
    for row in comparison["rows"]:
        sheet.append([row[key] for key in ("matrix_name", "matrix_type", "standard_code", "measure_id", "unit_nature", "applicable_cell", "opinion")])
    if comparison["match_status"] != "matched":
        sheet.append(["", "", "", "", "", "", comparison["opinion"]])
    widths = (28, 14, 23, 62, 20, 25, 56)
    side = Side(style="thin", color="D9D9D9")
    border = Border(left=side, right=side, top=side, bottom=side)
    for column, width in zip("ABCDEFG", widths):
        sheet.column_dimensions[column].width = width
    for row in sheet:
        for cell in row:
            # 附件文字始终按文本保存，不把以等号开头的内容解释为公式。
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = border
            cell.font = Font(name="微软雅黑", size=11)
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor="D9E2F3")
                cell.font = Font(name="微软雅黑", size=11, bold=True)
        # 长编号和提示显式折行；随内容增加高度，不截断正文。
        lines = max((sum(max(1, (len(line) * 2 + widths[cell.column - 1] - 1) // widths[cell.column - 1]) for line in str(cell.value or "").split("\n")) for cell in row), default=1)
        sheet.row_dimensions[row[0].row].height = max(32, 18 * lines + 12)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.print_title_rows = "1:1"
    output_root.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return path
