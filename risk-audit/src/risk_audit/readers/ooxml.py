"""openpyxl 读取 OOXML 工作簿时的最小兼容层。"""

from __future__ import annotations

from io import BytesIO
import logging
import os
from pathlib import Path
import re
import shutil
import tempfile
from typing import Any
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries


ROW_ONLY_SORT_STATE = re.compile(
    rb'(?P<prefix><(?:[A-Za-z_][\w.-]*:)?sortState\b[^>]*?\bref=")'
    rb'\$?(?P<first_row>\d+):\$?(?P<last_row>\d+)(?P<suffix>")'
)
MAX_WORKSHEET_ROW = 1_048_576
OPENPYXL_EXTENSIONS = {'.xlsx', '.xlsm', '.xltx', '.xltm'}
WORKSHEET_XML_SIZE_LIMIT = 64 * 1024 * 1024
WORKSHEET_DIMENSION_CELL_LIMIT = 1_000_000
SAFE_CONTENT_COLUMN_LIMIT = 256
STREAM_CHUNK_SIZE = 1024 * 1024
MAIN = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
DOCUMENT_REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PACKAGE_REL = 'http://schemas.openxmlformats.org/package/2006/relationships'
CELL_TAG = re.compile(rb'<c\b[^>]*\br="([A-Z]{1,3})(\d+)"[^>]*>')
ROW_TAG = re.compile(rb'<row\b[^>]*\br="(\d+)"[^>]*>')
DIMENSION_TAG = re.compile(rb'(<dimension\b[^>]*\bref=")([^"]+)("[^>]*/?>)')
ROW_SPANS = re.compile(rb'(\bspans=")\d+:\d+(\")')
MERGE_CELL_TAG = re.compile(
    rb'<(?:[A-Za-z_][\w.-]*:)?mergeCell\b[^>]*\bref="([^"]+)"[^>]*/>'
)


class OversizedWorksheetError(ValueError):
    """表示异常膨胀工作表无法在不猜测业务边界的前提下安全净化。"""


def _dimension_exceeds_cell_limit(stream: Any, cell_limit: int) -> bool:
    """判断工作表声明范围是否超过安全单元格数。

    Args:
        stream: 工作表 XML 二进制流。
        cell_limit: 允许直接交给 openpyxl 的最大声明单元格数。
    """

    for is_row, fragment in _row_fragments(stream):
        if is_row:
            return False
        match = DIMENSION_TAG.search(fragment)
        if match is None:
            continue
        try:
            minimum_column, minimum_row, maximum_column, maximum_row = range_boundaries(
                match.group(2).decode('ascii')
            )
        except (TypeError, ValueError):
            return False
        if None in {minimum_column, minimum_row, maximum_column, maximum_row}:
            return False
        declared_cells = (
            (maximum_column - minimum_column + 1)
            * (maximum_row - minimum_row + 1)
        )
        return declared_cells > cell_limit
    return False


def _worksheet_paths(archive: ZipFile) -> dict[str, str]:
    """读取工作表名称和部件路径；archive 为已打开的 OOXML 压缩包。"""
    from xml.etree import ElementTree

    workbook = ElementTree.fromstring(archive.read('xl/workbook.xml'))
    relationships = ElementTree.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
    targets = {
        item.get('Id'): item.get('Target', '')
        for item in relationships.findall(f'{{{PACKAGE_REL}}}Relationship')
    }
    paths = {}
    for sheet in workbook.findall(f'.//{{{MAIN}}}sheet'):
        target = targets[sheet.get(f'{{{DOCUMENT_REL}}}id')]
        normalized = target.lstrip('/')
        if not normalized.startswith('xl/'):
            normalized = f'xl/{normalized}'
        paths[sheet.get('name', normalized)] = normalized
    return paths


def _row_fragments(stream: Any):
    """逐段产出工作表 XML；stream 为二进制输入流，返回值标明是否为完整 row 节点。"""
    buffer = b''
    while True:
        chunk = stream.read(STREAM_CHUNK_SIZE)
        if chunk:
            buffer += chunk
        while True:
            row_start = buffer.find(b'<row')
            if row_start < 0:
                if chunk and len(buffer) > STREAM_CHUNK_SIZE:
                    # 工作表元数据标签远小于一个分块；保留尾部以承接跨块标签。
                    yield False, buffer[:-4096]
                    buffer = buffer[-4096:]
                break
            if row_start:
                yield False, buffer[:row_start]
                buffer = buffer[row_start:]
            row_end = buffer.find(b'</row>')
            if row_end < 0:
                break
            row_end += len(b'</row>')
            yield True, buffer[:row_end]
            buffer = buffer[row_end:]
        if not chunk:
            if buffer:
                yield False, buffer
            return


def _worksheet_profile(stream: Any, sheet_name: str, safe_content_column_limit: int) -> dict[str, Any]:
    """扫描异常工作表的真实内容边界。

    stream 为工作表 XML 流，sheet_name 用于错误定位，safe_content_column_limit
    为允许自动净化的最右真实内容列；返回维度及真实内容边界。
    """
    content_max_column = 0
    content_max_row = 0
    original_dimension = ''
    for is_row, fragment in _row_fragments(stream):
        if not is_row:
            if not original_dimension and (match := DIMENSION_TAG.search(fragment)):
                original_dimension = match.group(2).decode('ascii')
            continue
        previous_column = 0
        for match in CELL_TAG.finditer(fragment):
            column = column_index_from_string(match.group(1).decode('ascii'))
            if column < previous_column:
                raise OversizedWorksheetError(f'{sheet_name} 单元格顺序异常，无法安全净化。')
            previous_column = column
            if not match.group(0).rstrip().endswith(b'/>'):
                content_max_column = max(content_max_column, column)
                content_max_row = max(content_max_row, int(match.group(2)))
    content_max_column = max(content_max_column, 1)
    content_max_row = max(content_max_row, 1)
    if content_max_column > safe_content_column_limit:
        coordinate = get_column_letter(content_max_column)
        raise OversizedWorksheetError(
            f'{sheet_name} 的 {coordinate} 列存在真实内容，超过自动净化安全边界 '
            f'{get_column_letter(safe_content_column_limit)}，请人工确认工作表范围。'
        )
    if not original_dimension:
        raise OversizedWorksheetError(f'{sheet_name} 缺少工作表维度，无法安全净化。')
    first, separator, last = original_dimension.partition(':')
    last_coordinate = last if separator else first
    if re.search(r'\d+$', last_coordinate) is None:
        raise OversizedWorksheetError(f'{sheet_name} 的工作表维度无效：{original_dimension}')
    sanitized_dimension = f'{first}:{get_column_letter(content_max_column)}{content_max_row}'
    return {
        'sheet': sheet_name,
        'original_dimension': original_dimension,
        'sanitized_dimension': sanitized_dimension,
        'content_max_column': content_max_column,
        'content_max_row': content_max_row,
    }


def _trim_row(fragment: bytes, content_max_column: int) -> bytes:
    """删除真实内容右侧的纯样式单元格；fragment 为完整 row XML，content_max_column 为边界。"""
    for match in CELL_TAG.finditer(fragment):
        column = column_index_from_string(match.group(1).decode('ascii'))
        if column <= content_max_column:
            continue
        if not match.group(0).rstrip().endswith(b'/>'):
            raise OversizedWorksheetError(
                f'{get_column_letter(column)}{match.group(2).decode()} 存在真实内容，无法安全净化。'
            )
        extension = fragment.find(b'<extLst', match.start())
        suffix = extension if extension >= 0 else fragment.rfind(b'</row>')
        return fragment[:match.start()] + fragment[suffix:]
    return fragment


def _trim_out_of_bounds_merges(
    fragment: bytes,
    content_max_column: int,
    content_max_row: int,
) -> bytes:
    """删除完全位于真实内容边界外的合并区域。

    Args:
        fragment: 工作表 XML 的非行片段。
        content_max_column: 真实内容的最右列。
        content_max_row: 真实内容的最后一行。
    """

    def replace(match: re.Match[bytes]) -> bytes:
        """保留业务区合并，删除与真实区域完全不相交的合并。

        Args:
            match: 当前 mergeCell 标签的正则匹配结果。
        """

        try:
            minimum_column, minimum_row, _, _ = range_boundaries(
                match.group(1).decode("ascii")
            )
        except (TypeError, ValueError):
            return match.group(0)
        if minimum_column > content_max_column or minimum_row > content_max_row:
            return b""
        return match.group(0)

    return MERGE_CELL_TAG.sub(replace, fragment)


def _sanitize_worksheet(source: Any, target: Any, profile: dict[str, Any]) -> None:
    """流式写出净化后的工作表；source/target 为压缩包成员流，profile 为预检结果。"""
    content_max_column = profile['content_max_column']
    content_max_row = profile['content_max_row']
    original = profile['original_dimension'].encode('ascii')
    sanitized = profile['sanitized_dimension'].encode('ascii')
    for is_row, fragment in _row_fragments(source):
        if is_row:
            row_match = ROW_TAG.search(fragment)
            if row_match is not None and int(row_match.group(1)) > content_max_row:
                # 真实内容边界之后的行只含样式，整行丢弃才能防止写回阶段遍历百万空行。
                continue
            fragment = _trim_row(fragment, content_max_column)
            fragment = ROW_SPANS.sub(
                lambda match: match.group(1) + b'1:' + str(content_max_column).encode('ascii') + match.group(2),
                fragment,
            )
        else:
            fragment = DIMENSION_TAG.sub(
                lambda match: match.group(1) + sanitized + match.group(3)
                if match.group(2) == original else match.group(0),
                fragment,
            )
            # WPS 可能在真实表格外写入数千个远端合并，必须在 openpyxl 加载前剔除。
            fragment = _trim_out_of_bounds_merges(
                fragment,
                content_max_column,
                content_max_row,
            )
        target.write(fragment)


def prepare_oversized_workbook(
        workbook_path: Path,
        destination: Path,
        *,
        worksheet_size_limit: int | None = None,
        worksheet_dimension_cell_limit: int = WORKSHEET_DIMENSION_CELL_LIMIT,
        safe_content_column_limit: int = SAFE_CONTENT_COLUMN_LIMIT,
) -> tuple[Path, list[dict[str, Any]]]:
    """为异常膨胀工作簿生成安全临时副本。

    workbook_path 为源文件，destination 为临时副本路径，worksheet_size_limit
    为触发净化的工作表 XML 字节阈值，worksheet_dimension_cell_limit 为声明区域的
    单元格数阈值，safe_content_column_limit 为允许自动判断的最右真实内容列。
    无需净化时返回源路径和空报告，且绝不修改源文件。
    """
    limit = WORKSHEET_XML_SIZE_LIMIT if worksheet_size_limit is None else worksheet_size_limit
    with ZipFile(workbook_path) as archive:
        paths = _worksheet_paths(archive)
        names = {path: title for title, path in paths.items()}
        oversized = {}
        for info in archive.infolist():
            if info.filename not in names:
                continue
            if info.file_size > limit:
                oversized[info.filename] = info
                continue
            with archive.open(info) as worksheet:
                exceeds_dimension_limit = _dimension_exceeds_cell_limit(
                    worksheet,
                    worksheet_dimension_cell_limit,
                )
            # XML 体积可能很小，但虚假最大维度仍会让 iter_rows 实例化数十亿空单元格。
            if exceeds_dimension_limit:
                oversized[info.filename] = info
        if not oversized:
            return workbook_path, []
        logging.getLogger(__name__).warning(
            '检测到异常膨胀工作表，流式净化开始：文件=%s，工作表=%s',
            workbook_path,
            [names[path] for path in oversized],
        )
        profiles = {}
        for path in oversized:
            with archive.open(path) as worksheet:
                profiles[path] = _worksheet_profile(
                    worksheet,
                    names[path],
                    safe_content_column_limit,
                )

    destination.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix='.risk-audit-sanitized-',
        suffix=destination.suffix or '.xlsx',
        dir=destination.parent,
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        with ZipFile(workbook_path) as source_archive, ZipFile(temporary, 'w') as target_archive:
            for info in source_archive.infolist():
                with source_archive.open(info) as source_part, target_archive.open(info, 'w') as target_part:
                    if info.filename in profiles:
                        _sanitize_worksheet(source_part, target_part, profiles[info.filename])
                    else:
                        shutil.copyfileobj(source_part, target_part, length=STREAM_CHUNK_SIZE)
        temporary.replace(destination)
    finally:
        temporary.unlink(missing_ok=True)
    logging.getLogger(__name__).warning(
        '异常膨胀工作表流式净化完成：文件=%s，工作表=%s',
        workbook_path,
        [profiles[path]['sheet'] for path in oversized],
    )
    return destination, [profiles[path] for path in oversized]


def _normalize_row_only_sort_state(worksheet_xml: bytes) -> bytes:
    """将整行排序范围转为 openpyxl 可读的单元格范围。

    参数 worksheet_xml 为单个工作表的 OOXML 字节；返回仅规范化
    sortState ref 后的字节，不改变业务单元格或筛选范围。
    """

    def replace(match: re.Match[bytes]) -> bytes:
        """生成等价的整表宽度范围；参数 match 为整行 ref 匹配结果。"""
        first_row = int(match.group('first_row'))
        last_row = int(match.group('last_row'))
        if not 1 <= first_row <= last_row <= MAX_WORKSHEET_ROW:
            # 仅兼容 Excel 行边界内的正常范围，其他错误仍交给 openpyxl 拒绝。
            return match.group(0)
        return (
            match.group('prefix')
            + b'A' + match.group('first_row')
            + b':XFD' + match.group('last_row')
            + match.group('suffix')
        )

    return ROW_ONLY_SORT_STATE.sub(replace, worksheet_xml)


def _build_compatible_package(workbook_path: Path) -> BytesIO | None:
    """构造仅供读取的兼容包。

    参数 workbook_path 为源工作簿路径；存在需要规范化的排序范围时
    返回内存副本，否则返回 None。
    """
    replacements: dict[str, bytes] = {}
    with ZipFile(workbook_path) as source:
        for info in source.infolist():
            if not info.filename.startswith('xl/worksheets/') or not info.filename.endswith('.xml'):
                continue
            original = source.read(info.filename)
            normalized = _normalize_row_only_sort_state(original)
            if normalized != original:
                replacements[info.filename] = normalized
        if not replacements:
            return None

        package = BytesIO()
        with ZipFile(package, 'w') as target:
            for info in source.infolist():
                # 复用 ZipInfo 以保留原包的压缩方式、时间和文件属性。
                target.writestr(info, replacements.get(info.filename, source.read(info.filename)))
    package.seek(0)
    return package


def load_compatible_workbook(workbook_path: Path, **options: Any) -> Any:
    """通过 openpyxl 读取兼容 Excel/WPS 整行排序范围的工作簿。

    参数 workbook_path 为源文件路径，options 为透传给 openpyxl.load_workbook
    的选项；仅修正内存副本，不修改源文件。
    """
    package = _build_compatible_package(workbook_path)
    if package is None and workbook_path.suffix.lower() not in OPENPYXL_EXTENSIONS:
        # OOXML 内容可能使用 WPS .et 扩展名，二进制流可跳过后缀拒绝。
        package = BytesIO(workbook_path.read_bytes())
    source = package if package is not None else workbook_path
    return load_workbook(source, **options)
