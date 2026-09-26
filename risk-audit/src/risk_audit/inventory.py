from __future__ import annotations

import re
from pathlib import Path

from risk_audit.business_identity import BusinessIdentityError, identify_business as identify_business_v2
from risk_audit.entities import identify_material_entity
from risk_audit.models import Entity, FileRecord
from risk_audit.util import sha256_file


OOXML = b"PK\x03\x04"
OLE = bytes.fromhex("D0CF11E0A1B11AE1")


def true_format(path: Path) -> str:
    sig = path.read_bytes()[:8]
    if sig[:4] == OOXML:
        try:
            import zipfile
            with zipfile.ZipFile(path) as zf: content = zf.read("[Content_Types].xml")
            if b"wordprocessingml.document" in content: return "docx"
            if b"spreadsheetml.sheet" in content: return "xlsx"
        except Exception: return "unknown"
        return "unknown"
    if sig == OLE:
        if path.suffix.lower() == '.et':
            # WPS's extension alone does not identify its internal format.
            # Only route an OLE .et through the XLS reader if BIFF is readable.
            try:
                import xlrd
                book = xlrd.open_workbook(path, on_demand=True)
                book.release_resources()
            except Exception:
                return 'unknown'
        return "doc" if path.suffix.lower() == ".doc" else "xls"
    return "unknown"


def classify_material(path: Path) -> str | None:
    name = path.name
    auxiliary = ("适用性匹配", "意见建议", "成果清单", "应用清单", "统计表")
    if any(x in name for x in auxiliary): return None
    if path.suffix.lower() in {".doc", ".docx"}: return "explanation" if "说明" in name else None
    if "三清单" in name: return "three_lists"
    if "风控矩阵" in name: return "matrix"
    return None


def is_auxiliary(path: Path) -> bool:
    return any(x in path.name for x in ("适用性匹配", "意见建议", "成果清单", "应用清单", "统计表"))


def probe_material(path: Path) -> str | None:
    if path.suffix.lower() == ".docx":
        try:
            from docx import Document
            doc = Document(path)
            text = "\n".join([p.text for p in doc.paragraphs[:30]] + [str(cell.text) for table in doc.tables[:5] for row in table.rows[:10] for cell in row.cells])
            return "explanation" if "说明" in text else None
        except Exception:
            return None
    if path.suffix.lower() != ".xlsx": return None
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=False, data_only=False)
        for ws in wb.worksheets:
            values = {str(c.value).strip().replace(" ", "") for c in ws._cells.values() if c.row <= 8 and c.value not in (None, "")}
            if {"部门", "岗位名称", "岗位职责"} <= values: return "three_lists"
            if "控制措施" in values and ("控制措施编号" in values or "控制措施/编号" in values): return "matrix"
    except Exception:
        return None
    return None


def identify_business(relative: Path) -> tuple[str | None, str]:
    candidates = []
    business_hints = ("矩阵", "清单", "取证", "营销售电", "交易与购电", "交易购电", "基建", "迁改", "配网", "设备", "物资", "数字化", "职工福利", "员工报账")
    for index, part in enumerate(relative.parts):
        # A package may number companies 1..17. In particular "10某有限公司"
        # is a company folder, not business 10 (employee expenses).
        if index < len(relative.parts) - 1 and re.search(r"(?:公司|分公司)(?:本部)?$", part) and not any(hint in part for hint in business_hints):
            continue
        match = re.match(r"^\D*(0[1-9]|10)(?:\D|$)", part)
        if match: candidates.append(match.group(1))
    business = candidates[-1] if candidates and len(set(candidates)) == 1 else (candidates[0] if len(set(candidates)) == 1 else None)
    text = str(relative)
    variant = "default"
    if business == "03":
        if re.search(r"35\s*(?:k?v|千伏).*220", text, re.I): variant = "35-220kv"
        elif re.search(r"500\s*(?:k?v|千伏).*750", text, re.I): variant = "500-750kv"
        else: variant = "unknown"
    return business, variant


def scan_package(
    root: str | Path,
    entities: dict[str, Entity],
    aliases: dict,
    business_registry: dict | None = None,
) -> list[FileRecord]:
    """扫描报送材料；root 为输入根目录，entities/aliases 为主体配置，business_registry 为模板业务目录。"""
    base = Path(root).resolve(); out = []
    for path in sorted(base.rglob("*")):
        # 点开头文件和 Office 锁文件属于系统或编辑临时文件，不进入报送材料清单。
        if (not path.is_file() or path.name.startswith((".", "~$"))
                or path.suffix.lower() not in {".xlsx", ".xls", ".et", ".doc", ".docx"}):
            continue
        if is_auxiliary(path): continue
        material = classify_material(path) or probe_material(path)
        if not material: continue
        rel = path.relative_to(base)
        business_id = None
        business_error = ""
        if material == "explanation":
            business, variant = None, "default"
        elif business_registry and business_registry.get("schema_version") == "2.0":
            try:
                identity = identify_business_v2(rel, business_registry)
                business_id, business, variant = identity.business_id, identity.business_code, identity.variant_id
            except BusinessIdentityError as error:
                business, variant, business_error = None, "default", str(error)
        else:
            # 历史规则包仍按原编号读取，仅用于旧运行结果复核；当前发布包必须提供 v2 目录。
            business, variant = identify_business(rel)
            business_id = business
        identity_texts = [base.name, *rel.parts]
        if material == "explanation" and path.suffix.lower() == ".docx":
            try:
                from docx import Document
                doc = Document(path); identity_texts.extend(p.text for p in doc.paragraphs[:30])
            except Exception: pass
        code, evidence, conflict = identify_material_entity(
            path, base, entities, aliases, identity_texts, business_registry, business_id, business,
        )
        fmt = true_format(path)
        errors = []
        if fmt == "unknown": errors.append("文件签名不是受支持的OOXML或OLE格式")
        if business is None and material != "explanation":
            errors.append(business_error or "文件路径中的业务编号不唯一或无法识别")
        if code is None:
            errors.append("主体证据冲突" if conflict else "无法确定会计主体")
        record = FileRecord(path, rel, sha256_file(path), fmt, code, evidence, conflict,
                            business, variant, material, parse_errors=errors, business_id=business_id)
        if (material != "explanation" and business_registry
                and business_registry.get("schema_version") == "2.0" and business_id):
            record.preservation["business_name"] = identity.business_name
        out.append(record)
    return out
