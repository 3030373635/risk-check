"""按 0923-2 口径规范系统控制规则清单的系统类型列。"""

from __future__ import annotations

import re
from collections import OrderedDict
from copy import copy
from pathlib import Path
from typing import Any

from openpyxl.formula import Tokenizer
from openpyxl.worksheet.cell_range import MultiCellRange
from openpyxl.worksheet.table import TableColumn
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

from risk_audit.models import FileRecord, ParsedSheet
from risk_audit.readers.ooxml import load_compatible_workbook
from risk_audit.util import norm_text


SYSTEM_TYPE_HEADER = "系统类型"
SYSTEM_TYPE_FIRST = "一级部署系统"
SYSTEM_TYPE_SECOND = "二级部署系统"
SYSTEM_TYPE_THIRD = "三级部署系统"
SYSTEM_TYPE_VALUES = {SYSTEM_TYPE_FIRST, SYSTEM_TYPE_SECOND, SYSTEM_TYPE_THIRD}

_HUNAN_PROVINCE_NAMES = {
    "湖南",
    "省公司",
    "湖南电力",
    "湖南省",
    "国网湖南省电力有限公司",
    "湖南省电力公司",
    "湖南电力公司",
    "湖南省公司",
    "湖南省级单位",
}
_SUBORDINATE_MARKERS = ("市公司", "县公司", "供电公司", "供电分公司", "分公司", "子公司")
_HEADQUARTERS_NAMES = ("总部", "国网总部", "国家电网总部", "国家电网有限公司总部")
_HEADQUARTERS_NEGATIONS = ("非总部", "不属于总部", "不是总部")
_HUNAN_PROVINCE_MARKERS = (
    "国网湖南省电力有限公司",
    "湖南省电力公司",
    "湖南电力公司",
    "湖南省公司",
    "湖南电力",
)
_FORMULA_COLUMN_RE = re.compile(r"(?<![A-Za-z0-9_.])(\$?)([A-Z]{1,3})(?=(?:\$?\d+|:|$))")


def classify_system_owner(owner: str) -> str | None:
    """根据系统规则管理主体返回应填系统类型。

    Args:
        owner: 系统控制规则清单中的系统规则管理主体。
    """

    value = norm_text(owner)
    if not value:
        return None
    # 所属市县或分子公司不能仅因全称包含“湖南”而误判为省级单位。
    if any(marker in value for marker in _SUBORDINATE_MARKERS):
        return None
    if value in _HUNAN_PROVINCE_NAMES or any(marker in value for marker in _HUNAN_PROVINCE_MARKERS):
        return SYSTEM_TYPE_SECOND
    if not any(marker in value for marker in _HEADQUARTERS_NEGATIONS) and any(
        marker in value for marker in _HEADQUARTERS_NAMES
    ):
        return SYSTEM_TYPE_FIRST
    return None


def _header_row(sheet: ParsedSheet) -> int:
    """返回系统规则逻辑区的末级表头行。

    Args:
        sheet: 已解析的系统规则逻辑工作表。
    """

    return max(sheet.header_rows, default=max((sheet.first_data_row or 2) - 1, 1))


def _region_columns(sheet: ParsedSheet, worksheet: Any) -> range:
    """返回逻辑区允许扫描的物理列范围。

    Args:
        sheet: 已解析的系统规则逻辑工作表。
        worksheet: 对应的 openpyxl 物理工作表。
    """

    if sheet.region_bounds:
        return range(sheet.region_bounds[2], sheet.region_bounds[3] + 1)
    return range(1, worksheet.max_column + 1)


def _exact_system_type_header(sheet: ParsedSheet, worksheet: Any) -> tuple[int, int] | None:
    """查找原始表头中精确命名的系统类型单元格。

    Args:
        sheet: 已解析的系统规则逻辑工作表。
        worksheet: 对应的 openpyxl 物理工作表。
    """

    for row in sheet.header_rows:
        for column in _region_columns(sheet, worksheet):
            if norm_text(worksheet.cell(row, column).value) == norm_text(SYSTEM_TYPE_HEADER):
                return row, column
    return None


def _enum_system_type_column(sheet: ParsedSheet, worksheet: Any) -> int | None:
    """从明细枚举值定位无表头的系统类型列。

    Args:
        sheet: 已解析的系统规则逻辑工作表。
        worksheet: 对应的 openpyxl 物理工作表。
    """

    data_rows = {record.row for record in sheet.records}
    candidates = {
        column
        for column in _region_columns(sheet, worksheet)
        if any(norm_text(worksheet.cell(row, column).value) in SYSTEM_TYPE_VALUES for row in data_rows)
    }
    if len(candidates) == 1:
        return next(iter(candidates))
    logical_column = sheet.columns.get("system_type")
    return logical_column if logical_column in candidates else None


def _merged_anchor(worksheet: Any, row: int, column: int) -> tuple[int, int, str] | None:
    """返回目标单元格所属合并范围的锚点及范围。

    Args:
        worksheet: 当前 openpyxl 工作表。
        row: 拟写入的表头行。
        column: 拟写入的表头列。
    """

    for cell_range in worksheet.merged_cells.ranges:
        minimum_column, minimum_row, maximum_column, maximum_row = range_boundaries(str(cell_range))
        if minimum_row <= row <= maximum_row and minimum_column <= column <= maximum_column:
            return minimum_row, minimum_column, str(cell_range)
    return None


def _copy_cell_style(source: Any, target: Any) -> None:
    """复制新增系统类型单元格所需格式。

    Args:
        source: 提供原表格式的相邻单元格。
        target: 新增或补写的系统类型单元格。
    """

    if source.has_style:
        target._style = copy(source._style)
    if source.number_format:
        target.number_format = source.number_format
    target.alignment = copy(source.alignment)
    target.protection = copy(source.protection)


def _adjusted_column(column: int, inserted_columns: list[int]) -> int:
    """计算全列插入后的原列位置。

    Args:
        column: 插入前列号。
        inserted_columns: 按原坐标记录的插入列号。
    """

    return column + sum(inserted <= column for inserted in inserted_columns)


def _inserted_column(column: int, inserted_columns: list[int]) -> int:
    """计算某个新增列在全部插入完成后的列号。

    Args:
        column: 新列插入前的目标列号。
        inserted_columns: 按原坐标记录的插入列号。
    """

    return column + sum(inserted < column for inserted in inserted_columns)


def _adjusted_range(cell_range: str, inserted_columns: list[int]) -> str:
    """计算全列插入后的单元格范围。

    Args:
        cell_range: 插入前的 A1 范围。
        inserted_columns: 按原坐标记录的插入列号。
    """

    # 直接平移原范围的列端点，同时支持 C:C 和 1:1 等整列、整行表达式。
    return _shift_formula_range(cell_range, inserted_columns)


def _adjusted_reference(reference: str, inserted_columns: list[int]) -> str:
    """调整一个或多个空格分隔的 A1 范围。

    Args:
        reference: 工作表中的 A1 范围表达式。
        inserted_columns: 按原坐标记录的插入列号。
    """

    return " ".join(_adjusted_range(part, inserted_columns) for part in reference.split())


def _shift_formula_range(value: str, inserted_columns: list[int]) -> str:
    """将公式范围中受插列影响的列号右移。

    Args:
        value: Tokenizer 识别的单个范围操作数。
        inserted_columns: 按原坐标记录的插入列号。
    """

    if not re.search(r"\d|:", value):
        return value

    def replace(match: re.Match[str]) -> str:
        column = column_index_from_string(match.group(2))
        adjusted = _adjusted_column(column, inserted_columns)
        return f"{match.group(1)}{get_column_letter(adjusted)}"

    return _FORMULA_COLUMN_RE.sub(replace, value)


def _adjust_formula(
    formula: str,
    formula_sheet: str,
    target_sheet: str,
    inserted_columns: list[int],
) -> str:
    """调整公式中指向插列工作表的 A1 引用。

    Args:
        formula: 原始 Excel 公式。
        formula_sheet: 公式所在工作表。
        target_sheet: 发生插列的工作表。
        inserted_columns: 按原坐标记录的插入列号。
    """

    has_equals = formula.startswith("=")
    # 条件格式公式通常不含等号，补齐后才能按公式词法解析。
    tokenizer = Tokenizer(formula if has_equals else f"={formula}")
    for token in tokenizer.items:
        if token.type != "OPERAND" or token.subtype != "RANGE":
            continue
        sheet_prefix = None
        range_value = token.value
        if "!" in range_value:
            sheet_prefix, range_value = range_value.rsplit("!", 1)
            normalized_sheet = sheet_prefix.strip("'").replace("''", "'")
            if normalized_sheet != target_sheet:
                continue
        elif formula_sheet != target_sheet:
            continue
        shifted = _shift_formula_range(range_value, inserted_columns)
        token.value = f"{sheet_prefix}!{shifted}" if sheet_prefix is not None else shifted
    rendered = tokenizer.render()
    return rendered if has_equals else rendered.removeprefix("=")


def _adjust_workbook_references(
    workbook: Any,
    target_sheet: str,
    inserted_columns: list[int],
) -> None:
    """同步插列影响的公式及工作表范围。

    Args:
        workbook: 当前 openpyxl 工作簿。
        target_sheet: 发生插列的工作表名。
        inserted_columns: 按原坐标记录的插入列号。
    """

    for sheet in workbook.worksheets:
        # iter_rows 会按 max_row/max_column 展开整个矩形；只遍历已实例化单元格才能保持稀疏表。
        for cell in tuple(sheet._cells.values()):
            if cell.data_type == "f" and isinstance(cell.value, str):
                cell.value = _adjust_formula(
                    cell.value,
                    sheet.title,
                    target_sheet,
                    inserted_columns,
                )

    worksheet = workbook[target_sheet]
    for validation in worksheet.data_validations.dataValidation:
        validation.sqref = MultiCellRange(
            _adjusted_reference(str(validation.sqref), inserted_columns)
        )
        for attribute in ("formula1", "formula2"):
            formula = getattr(validation, attribute)
            if isinstance(formula, str) and formula:
                setattr(
                    validation,
                    attribute,
                    _adjust_formula(formula, target_sheet, target_sheet, inserted_columns),
                )
    if worksheet.auto_filter.ref:
        worksheet.auto_filter.ref = _adjusted_range(worksheet.auto_filter.ref, inserted_columns)
    for table in worksheet.tables.values():
        minimum_column, _, maximum_column, _ = range_boundaries(table.ref)
        table_insertions = [
            column for column in inserted_columns
            if minimum_column < column <= maximum_column
        ]
        for offset, column in enumerate(table_insertions):
            insertion_index = column - minimum_column + offset
            table.tableColumns.insert(
                insertion_index,
                TableColumn(id=insertion_index + 1, name=SYSTEM_TYPE_HEADER),
            )
        # Excel 表格列 ID 必须与插入后的表头次序一致。
        for identifier, table_column in enumerate(table.tableColumns, 1):
            table_column.id = identifier
        table.ref = _adjusted_range(table.ref, inserted_columns)
        if table.autoFilter and table.autoFilter.ref:
            table.autoFilter.ref = _adjusted_range(table.autoFilter.ref, inserted_columns)

    adjusted_rules = OrderedDict()
    for conditional_format, rules in worksheet.conditional_formatting._cf_rules.items():
        conditional_format.sqref = MultiCellRange(
            _adjusted_reference(str(conditional_format.sqref), inserted_columns)
        )
        for rule in rules:
            rule.formula = [
                _adjust_formula(formula, target_sheet, target_sheet, inserted_columns)
                for formula in (rule.formula or [])
            ]
        adjusted_rules[conditional_format] = rules
    worksheet.conditional_formatting._cf_rules = adjusted_rules

    for defined_name in workbook.defined_names.values():
        formula_sheet = ""
        if defined_name.localSheetId is not None:
            formula_sheet = workbook.worksheets[defined_name.localSheetId].title
        if isinstance(defined_name.attr_text, str) and defined_name.attr_text:
            defined_name.attr_text = _adjust_formula(
                defined_name.attr_text,
                formula_sheet,
                target_sheet,
                inserted_columns,
            )


def preprocess_system_types(file: FileRecord, work_dir: Path, aliases: dict[str, Any]) -> None:
    """规范系统类型列并重新解析预处理副本。

    Args:
        file: 已完成首次解析的业务材料。
        work_dir: 当前业务运行目录。
        aliases: 当前规则包的字段别名。
    """

    system_sheets = [sheet for sheet in file.sheets if sheet.sheet_type == "system_rule"]
    if not system_sheets:
        return

    source = Path(getattr(file, "_preprocessed_path", getattr(file, "_converted_path", file.source)))
    workbook = load_compatible_workbook(source, data_only=False, rich_text=True, keep_links=True)
    reports: list[dict[str, Any]] = []
    try:
        for title in dict.fromkeys(sheet.title for sheet in system_sheets):
            worksheet = workbook[title]
            parts = [sheet for sheet in system_sheets if sheet.title == title]
            plans: list[dict[str, Any]] = []
            for sheet in parts:
                exact_header = _exact_system_type_header(sheet, worksheet)
                exact_column = exact_header[1] if exact_header is not None else None
                enum_column = None if exact_header is not None else _enum_system_type_column(sheet, worksheet)
                owner_column = sheet.columns.get("system_owner")
                system_name_column = sheet.columns.get("system_name")
                insert_column = None
                if exact_column is None and enum_column is None:
                    insert_column = owner_column or (system_name_column + 1 if system_name_column else None)
                header_row = exact_header[0] if exact_header is not None else _header_row(sheet)
                split_header_merge = None
                if exact_header is None and enum_column is not None:
                    merged_header = _merged_anchor(worksheet, header_row, enum_column)
                    if merged_header is not None:
                        header_row, anchor_column, merged_range = merged_header
                        # 横向合并无法为枚举列单独命名，需解除后在实际枚举列写表头。
                        if anchor_column != enum_column:
                            split_header_merge = merged_range
                plans.append({
                    "sheet": sheet,
                    "type_column": exact_column or enum_column,
                    "insert_column": insert_column,
                    "renamed": exact_column is None and enum_column is not None,
                    "header_row": header_row,
                    "split_header_merge": split_header_merge,
                })

            inserted_columns = sorted({
                plan["insert_column"] for plan in plans if plan["insert_column"] is not None
            })
            merged_ranges = [str(cell_range) for cell_range in worksheet.merged_cells.ranges]
            split_header_merges = {
                plan["split_header_merge"]
                for plan in plans
                if plan["split_header_merge"] is not None
            }
            for cell_range in split_header_merges:
                worksheet.unmerge_cells(cell_range)
            # 先解除合并再移动单元格，避免 openpyxl 将移动后的锚点留在旧合并范围内。
            if inserted_columns:
                for cell_range in merged_ranges:
                    if cell_range in split_header_merges:
                        continue
                    worksheet.unmerge_cells(cell_range)
            # 从右向左插列，保证每个目标仍使用首次解析得到的原坐标。
            for column in reversed(inserted_columns):
                worksheet.insert_cols(column, 1)
            if inserted_columns:
                _adjust_workbook_references(workbook, title, inserted_columns)
                for cell_range in merged_ranges:
                    if cell_range in split_header_merges:
                        continue
                    worksheet.merge_cells(_adjusted_range(cell_range, inserted_columns))

            for plan in plans:
                sheet = plan["sheet"]
                if plan["insert_column"] is not None:
                    type_column = _inserted_column(plan["insert_column"], inserted_columns)
                elif plan["type_column"] is not None:
                    type_column = _adjusted_column(plan["type_column"], inserted_columns)
                else:
                    # 缺少系统名称和管理主体时无法确定新增位置，由字段完整性检查处理。
                    continue

                original_owner_column = sheet.columns.get("system_owner")
                owner_column = (
                    _adjusted_column(original_owner_column, inserted_columns)
                    if original_owner_column is not None
                    else None
                )
                header_row = plan["header_row"]
                style_source_column = owner_column or max(type_column - 1, 1)
                header_cell = worksheet.cell(header_row, type_column)
                _copy_cell_style(worksheet.cell(header_row, style_source_column), header_cell)
                header_cell.value = SYSTEM_TYPE_HEADER

                unresolved_rows = []
                for record in sheet.records:
                    type_cell = worksheet.cell(record.row, type_column)
                    if norm_text(type_cell.value) == SYSTEM_TYPE_THIRD:
                        continue
                    owner = worksheet.cell(record.row, owner_column).value if owner_column else ""
                    expected = classify_system_owner(str(owner or ""))
                    if expected is None:
                        # 一、二级必须按管理主体重算，未决行不保留旧分类。
                        type_cell.value = None
                        unresolved_rows.append(record.row)
                        continue
                    _copy_cell_style(worksheet.cell(record.row, style_source_column), type_cell)
                    type_cell.value = expected

                reports.append({
                    "sheet": title,
                    "header_row": header_row,
                    "column": type_column,
                    "inserted": plan["insert_column"] is not None,
                    "renamed": plan["renamed"],
                    "unresolved_rows": unresolved_rows,
                })
                column_letter = get_column_letter(type_column)
                if not worksheet.column_dimensions[column_letter].width:
                    worksheet.column_dimensions[column_letter].width = 16

        destination = work_dir / "preprocessed" / "system_types" / file.relative_path.with_suffix(".xlsx")
        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(destination)
    finally:
        workbook.close()

    file._preprocessed_path = str(destination)
    from risk_audit.readers.confirmed_v180 import parse_workbook_v180

    file.sheets = parse_workbook_v180(
        file,
        destination,
        aliases,
        include_hidden=getattr(file, "_include_hidden", False),
    )
    file.preservation["system_type_preprocessing"] = reports
